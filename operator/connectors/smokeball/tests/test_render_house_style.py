"""The drafting classes held to an attorney's house style, enforced in-class.

The standard (an engaged firm's written drafting instructions, 2026-07-29 and
2026-10-07): Times New Roman 12 on plain Letter paper, never pleading paper; a
mediation brief's roman-numeral headings centered and bold ONLY, its lettered
subsections indented, bold and underlined, its body double-spaced; a discovery
item led by an all-caps bold underlined label, the request a double-spaced
first-line-indented paragraph, double-spaced between items and never more; a
deposition outline numbered at the bottom of every page.

"In-class" is the point: a firm template's own style (an underlined Heading 1,
an Arial body, a space-after) must not move a served document off the standard,
so each test runs on the starter AND on a firm base that tries to.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

from smokeball_connector import docx_classes
from smokeball_connector.docx_format import CLASS_RULES, render_document

HOUSE = ("discovery_set", "discovery_response", "mediation_brief", "memo", "depo_outline")

BRIEF = """| PLAINTIFF ALPHA EXAMPLE, | Case No. {{FILL: case number}} |
| v. |  |
| BETA EXAMPLE |  |

# I. INTRODUCTION

This brief is submitted on behalf of plaintiff.

## A. Past Medical Expenses

Body text of the subsection.

# II. LIABILITY

More body.
"""

DISCOVERY = """# PLAINTIFF'S SPECIAL INTERROGATORIES, SET ONE

## DEFINITIONS

"INCIDENT" means the collision described in the complaint.

**SPECIAL INTERROGATORY NO. 1:**

Identify each person who witnessed the INCIDENT.

**SPECIAL INTERROGATORY NO. 2:**

State all facts supporting your contention.
"""


def _hostile_base(*, footer: str = "", page_field: bool = False) -> bytes:
    """A firm base that tries every way to move the standard: A4 paper, line
    numbering (pleading paper), an underlined Arial Heading 1 and SMD Heading 1,
    an Arial 14 body with space after."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)  # A4
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    sec._sectPr.append(ln)
    for name in ("Heading 1", "SMD Heading 1"):
        st = doc.styles[name] if name == "Heading 1" else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.font.underline = True
        st.font.name = "Arial"
    body = doc.styles.add_style("SMD Body", WD_STYLE_TYPE.PARAGRAPH)
    body.font.name, body.font.size = "Arial", Pt(14)
    body.paragraph_format.space_after = Pt(18)
    if footer:
        sec.footer.paragraphs[0].text = footer
    if page_field:
        r = sec.footer.paragraphs[0].add_run()
        for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
            el = OxmlElement("w:fldChar" if kind else "w:instrText")
            if kind:
                el.set(qn("w:fldCharType"), kind)
            else:
                el.text = text
            r._r.append(el)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


BASES = {"starter": None, "hostile firm base": _hostile_base()}


def _doc(blob: bytes) -> Document:
    return Document(io.BytesIO(blob))


def _effective_underline(run, para) -> bool:
    if run.underline is not None:
        return bool(run.underline)
    style = para.style
    while style is not None:
        if style.font.underline is not None:
            return bool(style.font.underline)
        style = style.base_style
    return False


def level1_headings_underlined(blob: bytes) -> bool:
    """True when any level-1 (roman numeral) heading renders underlined."""
    for p in _doc(blob).paragraphs:
        if p.text.split(".")[0].strip() in ("I", "II", "III", "IV"):
            if any(_effective_underline(r, p) for r in p.runs if r.text.strip()):
                return True
    return False


# ---- font, paper -----------------------------------------------------------------


@pytest.mark.parametrize("cls", HOUSE)
@pytest.mark.parametrize("base", list(BASES))
def test_every_run_is_times_new_roman_12_in_class(cls: str, base: str) -> None:
    blob, _ = render_document(BRIEF + DISCOVERY + "- a bullet\n1. a numbered item\n", cls, BASES[base])
    doc = _doc(blob)
    runs = [r for p in doc.paragraphs for r in p.runs if r.text.strip()]
    runs += [r for t in doc.tables for c in t._cells for p in c.paragraphs for r in p.runs if r.text.strip()]
    assert runs
    for r in runs:
        fonts = r._r.rPr.find(qn("w:rFonts"))
        assert fonts is not None and {fonts.get(qn(a)) for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs")} == {
            "Times New Roman"
        }, r.text
        assert r.font.size == Pt(12), r.text


@pytest.mark.parametrize("cls", HOUSE)
def test_plain_letter_paper_never_pleading_paper(cls: str) -> None:
    blob, report = render_document(BRIEF, cls, BASES["hostile firm base"])
    sec = _doc(blob).sections[0]
    assert (sec.page_width, sec.page_height) == (Inches(8.5), Inches(11))
    assert sec._sectPr.find(qn("w:lnNumType")) is None
    notes = report.to_dict()["notes"]
    assert any("US Letter" in n for n in notes) and any("never pleading paper" in n for n in notes)


def test_letter_classes_keep_the_firms_own_typography_and_paper() -> None:
    blob, _ = render_document("Body.", "letter", BASES["hostile firm base"])
    sec = _doc(blob).sections[0]
    assert sec.page_width == Twips(11906)


# ---- mediation brief headings ---------------------------------------------------------


@pytest.mark.parametrize("base", list(BASES))
def test_mediation_level1_headings_are_centered_bold_and_never_underlined(base: str) -> None:
    blob, _ = render_document(BRIEF, "mediation_brief", BASES[base])
    doc = _doc(blob)
    h1 = [p for p in doc.paragraphs if p.text in ("I. INTRODUCTION", "II. LIABILITY")]
    assert len(h1) == 2
    for p in h1:
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert all(r.bold for r in p.runs)
    assert not level1_headings_underlined(blob)


def test_the_underline_check_can_fail(monkeypatch: pytest.MonkeyPatch) -> None:
    """Falsifier: a mediation rule that underlines level 1 is caught."""
    import dataclasses

    bad = dataclasses.replace(CLASS_RULES["mediation_brief"], heading_underline=(True, True, False))
    monkeypatch.setitem(CLASS_RULES, "mediation_brief", bad)
    blob, _ = render_document(BRIEF, "mediation_brief", None)
    assert level1_headings_underlined(blob)
    assert docx_classes.CLASS_RULES is CLASS_RULES  # one table, re-exported


@pytest.mark.parametrize("base", list(BASES))
def test_mediation_level2_is_indented_bold_and_underlined(base: str) -> None:
    blob, _ = render_document(BRIEF, "mediation_brief", BASES[base])
    p = next(p for p in _doc(blob).paragraphs if p.text == "A. Past Medical Expenses")
    assert p.paragraph_format.left_indent == Inches(0.5)
    assert all(r.bold and r.underline for r in p.runs)


@pytest.mark.parametrize("base", list(BASES))
def test_mediation_body_double_spaced_and_caption_first(base: str) -> None:
    blob, _ = render_document(BRIEF, "mediation_brief", BASES[base])
    doc = _doc(blob)
    body = next(p for p in doc.paragraphs if p.text.startswith("This brief"))
    assert body.paragraph_format.line_spacing == 2.0
    assert body.paragraph_format.space_after == Pt(0)
    first = doc.element.body[0]
    assert first.tag == qn("w:tbl")  # the caption table leads the document


# ---- discovery --------------------------------------------------------------------


@pytest.mark.parametrize("cls", ("discovery_set", "discovery_response"))
@pytest.mark.parametrize("base", list(BASES))
def test_discovery_items_double_spaced_between_items_only(cls: str, base: str) -> None:
    blob, _ = render_document(DISCOVERY, cls, BASES[base])
    doc = _doc(blob)
    labels = [p for p in doc.paragraphs if p.text.startswith("SPECIAL INTERROGATORY NO.")]
    items = [p for p in doc.paragraphs if p.text.startswith(("Identify", "State all"))]
    assert len(labels) == 2 and len(items) == 2
    for p in labels:
        assert all(r.bold and r.underline and r.font.all_caps for r in p.runs if r.text.strip())
        pf = p.paragraph_format
        assert (pf.line_spacing, pf.space_before, pf.space_after) == (2.0, Pt(0), Pt(0))
    for p in items:
        pf = p.paragraph_format
        assert pf.first_line_indent == Inches(0.5)
        assert (pf.line_spacing, pf.space_before, pf.space_after) == (2.0, Pt(0), Pt(0))


def test_discovery_spacing_check_can_fail(monkeypatch: pytest.MonkeyPatch) -> None:
    import dataclasses

    bad = dataclasses.replace(CLASS_RULES["discovery_set"], item_space_after_pt=12.0)
    monkeypatch.setitem(CLASS_RULES, "discovery_set", bad)
    blob, _ = render_document(DISCOVERY, "discovery_set", None)
    item = next(p for p in _doc(blob).paragraphs if p.text.startswith("Identify"))
    assert item.paragraph_format.space_after == Pt(12)


def test_discovery_definitions_section_is_supported() -> None:
    blob, report = render_document(DISCOVERY, "discovery_set", None)
    doc = _doc(blob)
    assert any(p.text == "DEFINITIONS" for p in doc.paragraphs)
    d = next(p for p in doc.paragraphs if p.text.startswith('"INCIDENT"'))
    assert d.paragraph_format.line_spacing == 2.0


# ---- deposition outline ---------------------------------------------------------------


def _footer_xml(blob: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(blob))
    return "".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))


@pytest.mark.parametrize(
    "base",
    [None, _hostile_base(), _hostile_base(footer="123 Example Street")],
    ids=["starter", "firm base", "firm base with footer text"],
)
def test_depo_outline_footer_always_carries_a_centered_page_field(base: bytes | None) -> None:
    blob, _ = render_document("# Witness background\n\n1. Name and address.\n", "depo_outline", base)
    xml = _footer_xml(blob)
    assert "PAGE" in xml and 'w:fldCharType="begin"' in xml
    footer = _doc(blob).sections[0].footer
    page_para = next(p for p in footer.paragraphs if "PAGE" in p._p.xml)
    assert page_para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    if base is not None and b"123 Example" in base:
        assert any("123 Example Street" in p.text for p in footer.paragraphs)


def test_depo_outline_does_not_add_a_second_page_field() -> None:
    blob, report = render_document("Body.", "depo_outline", _hostile_base(page_field=True))
    assert _footer_xml(blob).count("PAGE") == 1
    assert "page number field added to the footer" not in report.to_dict()["notes"]


def test_memo_is_times_new_roman_12() -> None:
    blob, _ = render_document("# To the file\n\nA memo body.", "memo", BASES["hostile firm base"])
    p = next(p for p in _doc(blob).paragraphs if p.text == "A memo body.")
    assert all(r.font.name == "Times New Roman" and r.font.size == Pt(12) for r in p.runs)
