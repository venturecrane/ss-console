"""Records orders built from a matter's filed gap audit (gap_orders.py, gap_locate.py).

Synthetic throughout: a gap audit docx in the demand job's rendered shape, a
fake Smokeball that serves it and one cited bill, and a vendor directory
stand-in. Nothing here is a real matter, provider address or client."""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from smokeball_connector import gap_locate, gap_orders
from smokeball_connector.records_orders import _CUSTODIAN_ID

MATTER = "8d7c2a4e-1f3b-4c5d-9e6f-0a1b2c3d4e5f"
CLIENT = "c0ffee00-0000-4000-8000-000000000001"
TODAY = date(2026, 10, 7)
HEAD = [
    "Item",
    "Provider",
    "What's missing",
    "Where the file points to it",
    "Basis",
    "Suggested request type",
    "Priority",
]
ROWS = {
    "A. Referral and order trail": [
        [
            "1",
            "Example Imaging Center (ordered by Dr. A)",
            "Missing: MRI report and films for 3/10/2026",
            "Example Imaging bill, p.1",
            "Referenced in record",
            "Radiology report, images and bill request",
            "Blocks demand",
        ],
        [
            "2",
            "Physical therapy (provider not named)",
            "Missing: PT notes",
            "ER record, p.2",
            "Referenced in record",
            "Ask client which PT clinic",
            "Strengthens demand",
        ],
    ],
    "B. Provider-by-provider completeness": [
        [
            "3",
            "Northfield Spine Clinic (Dr. B)",
            "Missing prior records, same spine: injections 12/2/2023",
            "Northfield notes, p.6",
            "Referenced in record",
            "Prior records request, 5-year look-back",
            "Blocks demand",
        ],
        [
            "4",
            "Medicare",
            "Missing: conditional payment letter",
            "CMS letter, p.1",
            "Referenced in record",
            "Conditional payment letter request",
            "Strengthens demand",
        ],
        [
            "5",
            "Example Imaging Center",
            "Missing: lien and balance",
            "Example Imaging bill, p.1",
            "Referenced in record",
            "Lien and balance request",
            "Housekeeping",
        ],
        [
            "6",
            "Northfield Spine Clinic",
            "Visit 2/20/2026 billed with no itemization",
            "Northfield bill, p.1",
            "Billing mismatch",
            "Itemized billing ledger with CPT detail",
            "Strengthens demand",
        ],
        [
            "7",
            "Westside Medical Group",
            "Missing: records from 2/12/2026",
            "Westside note, p.1",
            "Referenced in record",
            "Records and itemized billing request",
            "Strengthens demand",
        ],
        [
            "8",
            "Unlisted Clinic",
            "Missing: records from 2/14/2026",
            "ER record, p.3",
            "Referenced in record",
            "Records request",
            "Strengthens demand",
        ],
    ],
    "C. Billing-to-record reconciliation": [
        [
            "9",
            "Example Imaging Center",
            "Tab shows a charge with no date",
            "preflight (Medicals tab vs bills)",
            "Billing mismatch",
            "Update Medicals tab",
            "Housekeeping",
        ],
    ],
    "D. Treatment timeline": [
        [
            "10",
            "Example Imaging Center",
            "3/10/2026 to 5/6/2026, 57 days",
            "Example Imaging bill, p.1",
            "Referenced in record",
            "None (timeline fact)",
            "Housekeeping",
        ],
    ],
}


def gap_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Records and Billing Gap Audit", 1)
    for section, rows in ROWS.items():
        d.add_heading(section, 2)
        t = d.add_table(rows=1, cols=7)
        for i, h in enumerate(HEAD):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = v
    d.add_heading("Possible, needs paralegal check", 2)
    p = d.add_table(rows=2, cols=4)
    for i, v in enumerate(["Provider", "What's missing", "Where the file points to it", "Suggested request type"]):
        p.rows[0].cells[i].text = v
    for i, v in enumerate(["Guess Clinic", "maybe records", "nowhere", "Records request"]):
        p.rows[1].cells[i].text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def bill_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Westside Medical Group")
    d.add_paragraph("200 Oak Ave, Springfield, CA 95811   (916) 555-0100")
    d.add_paragraph("Statement of charges")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class FakeSmokeball:
    def __init__(self, ssn: str = "123456789") -> None:
        self.ssn = ssn
        self.downloads: list[str] = []

    files = [
        {
            "id": "f-gap",
            "name": "Gap Audit - 900101 - 10-07-26",
            "fileExtension": ".docx",
            "dateCreated": "2026-10-07T09:00",
        },
        {
            "id": "f-old",
            "name": "Gap Audit - 900101 - 10-01-26",
            "fileExtension": ".docx",
            "dateCreated": "2026-10-01T09:00",
        },
        {"id": "f-west", "name": "Westside note", "fileExtension": ".docx", "dateCreated": "2026-03-01"},
    ]

    def get(self, path: str, **params: Any) -> Any:
        if path == f"/matters/{MATTER}/documents/files":
            return {"value": self.files if params.get("Offset", 0) == 0 else []}
        if path == f"/matters/{MATTER}":
            return {"id": MATTER, "number": "900101", "clientIds": [CLIENT]}
        if path == f"/contacts/{CLIENT}":
            return {
                "id": CLIENT,
                "person": {
                    "firstName": "Pat",
                    "lastName": "Example",
                    "identificationNumber": self.ssn,
                    "birthDate": "1985-06-15T00:00:00",
                    "residentialAddress": {
                        "addressLine1": "1 Main St",
                        "city": "Springfield",
                        "state": "CA",
                        "zipCode": "95811",
                    },
                    "email": "pat@example.com",
                },
            }
        raise LookupError(path)  # the Medicals tab is not served: no hints, never a crash

    def download_file(self, matter_id: str, file_id: str) -> tuple[dict[str, Any], bytes]:
        self.downloads.append(file_id)
        return {}, {"f-gap": gap_docx(), "f-west": bill_docx()}[file_id]


def loc(i: str, name: str, street: str, city: str, zip_: str) -> dict[str, Any]:
    return {"id": i, "value": name, "street": street, "city": city, "state": "CA", "postalcode": zip_}


class Directory:
    def __init__(self) -> None:
        self.asked: list[tuple[str, str | None]] = []

    def get_locations(self, term: str, zip_code: str | None = None) -> list[dict[str, Any]]:
        self.asked.append((term, zip_code))
        book = {
            "Example Imaging Center": [loc("neo_101", "Example Imaging Center", "5 Elm St", "Springfield", "95811")],
            "Northfield Spine Clinic": [
                loc("npi_2020202020", "Northfield Spine Clinic", "9 Pine Rd", "Shelbyville", "95822")
            ],
            "Westside Medical Group": [
                loc("neo_301", "Westside Medical Group", "200 Oak Ave", "Springfield", "95811"),
                loc("neo_302", "Westside Medical Group", "8 Bay Rd", "Capital City", "95899"),
                loc("neo_303", "Westside Medical Group - Pediatrics", "200 Oak Ave", "Springfield", "95811"),
            ],
        }
        rows = book.get(term, [])
        return [r for r in rows if not zip_code or r["postalcode"] == zip_code]


@pytest.fixture(autouse=True)
def _seat(tmp_path: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "customer.yaml"
    path.write_text("records_orders:\n  pre_approved_custodian_fee: 100\n  authorization: e_auth\n", encoding="utf-8")
    monkeypatch.setenv("SMD_CUSTOMER_YAML_PATH", str(path))


def _build(sb: FakeSmokeball | None = None, d: Directory | None = None) -> dict[str, Any]:
    return gap_orders.build(
        sb or FakeSmokeball(), d or Directory(), MATTER, "paralegal@firm.example", TODAY, "the vendor"
    )


def test_only_orderable_rows_addressed_to_a_medical_provider_are_kept() -> None:
    rows = gap_orders.audit_rows(gap_docx())
    assert [r["item"] for r in rows] == [str(i) for i in range(1, 11)]  # the possible list is never read
    why = {r["item"]: gap_orders.classify(r) for r in rows}
    assert why["1"] is None and why["3"] is None and why["6"] is None and why["7"] is None
    assert "does not name" in why["2"]
    assert "payer" in why["4"]
    assert "lien" in why["5"]
    assert "question or a confirmation" in why["9"]
    assert "timeline" in why["10"]


def test_rows_group_by_custodian_with_types_and_dates() -> None:
    rows = [r for r in gap_orders.audit_rows(gap_docx()) if gap_orders.classify(r) is None]
    groups = {g["name"]: g for g in gap_orders.group(rows, TODAY, gap_orders.first_date(rows, TODAY))}
    north = groups["Northfield Spine Clinic"]
    assert north["items"] == [3, 6] and set(north["types"]) == {"Medical", "Billing"}
    assert north["service_start"] == "2021-02-12"  # a 5-year look-back from the first treatment date the audit names
    imaging = groups["Example Imaging Center"]
    assert set(imaging["types"]) == {"Radiology Record", "Radiology Image", "Billing"}
    assert imaging["service_start"] == "2026-03-10" and imaging["service_end"] == "2026-10-07"


def test_the_file_settles_an_ambiguous_provider_and_the_order_is_prepared() -> None:
    sb, d = FakeSmokeball(), Directory()
    out = _build(sb, d)
    assert out["status"] == "ready", out
    assert out["counts"] == {"matched": 3, "ambiguous": 0, "no_match": 1}
    order = out["orders"][0]["order"]
    ids = [x["custodian_id"] for x in order["locations"]]
    # Westside: three directory entries; the cited note's letterhead (200 Oak Ave,
    # 95811) settles it, and the Pediatrics department is not the custodian
    assert ids == ["neo_101", "npi_2020202020", "neo_301"]
    assert "f-west" in sb.downloads
    assert order["authorization"] == "e_auth"
    assert [w["custodian"] for w in out["would_order"]] == [
        "Example Imaging Center",
        "Northfield Spine Clinic",
        "Westside Medical Group",
    ]


def test_what_the_file_cannot_settle_is_one_question_never_a_list() -> None:
    out = _build()
    assert len(out["questions"]) == 1
    q = out["question"]
    assert q.startswith("One thing before these are ordered") and "Unlisted Clinic" in q and q.count("?") <= 1


def test_a_missing_client_fact_still_shows_what_would_be_ordered() -> None:
    out = _build(FakeSmokeball(ssn=""))
    assert out["orders"][0]["status"] == "missing_client_facts"
    assert len(out["would_order"]) == 3  # she sees the value while the file need is stated


def test_a_name_that_only_shares_an_address_or_a_generic_word_is_not_the_provider() -> None:
    assert gap_locate.names_match("Roseville Family Healthcare", "Roseville Family Healthcare - Main")
    assert not gap_locate.names_match("Roseville Family Healthcare", "Oakridge Healthcare Center")


def test_directory_duplicates_at_one_address_are_one_place() -> None:
    a = {"custodian_id": "npi_1", "name": "Example Lab LLC", "address": "4 Way, Town, CA 95811"}
    b = {"custodian_id": "neo_2", "name": "Example Lab", "address": "4 Way, Town, CA 95811"}
    assert gap_locate.same_place([a, b]) == [b]


def test_the_vendors_custodian_id_shapes_are_accepted() -> None:
    for good in ("552211", "neo_77456", "npi_1881861417", "611d698f-8ea9-11ea-989c-42010a473806"):
        assert _CUSTODIAN_ID.match(good)
    for bad in ("", "../x", "neo 1", "a" * 80):
        assert not _CUSTODIAN_ID.match(bad)


def test_two_providers_on_one_location_are_one_facility() -> None:
    f = [
        {
            "custodian_id": "neo_1",
            "record_types": ["Medical"],
            "service_start": "2026-01-01",
            "service_end": "2026-10-07",
        },
        {
            "custodian_id": "neo_1",
            "record_types": ["Billing"],
            "service_start": "2021-01-01",
            "service_end": "2026-10-07",
        },
    ]
    assert gap_orders.merge_same_custodian(f) == [
        {
            "custodian_id": "neo_1",
            "record_types": ["Medical", "Billing"],
            "service_start": "2021-01-01",
            "service_end": "2026-10-07",
        }
    ]


def test_no_filed_gap_audit_is_said_plainly() -> None:
    sb = FakeSmokeball()
    sb.files = [f for f in sb.files if not f["name"].startswith("Gap Audit")]
    assert _build(sb)["status"] == "refused"
