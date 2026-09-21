"""The firm's letterhead: authored once in customer.yaml, printed by code.

Three precedence states, each pinned: the firm's own template file wins and is
untouched; else the authored ``firm_identity`` is printed on the starter's
first-page header for the letter classes; else nothing is printed, no marker is
added, and the report says so. The values in these tests are fictitious.
"""

from __future__ import annotations

import io
import textwrap

import httpx
import pytest
from docx import Document

from smokeball_connector import server
from smokeball_connector.docx_format import FormatReport, render_document
from smokeball_connector.letterhead import LETTERHEAD_CLASSES, FirmIdentity, load_firm_identity
from smokeball_connector.library import CUSTOMER_YAML_ENV

from .test_library_resolution import _handler, _mock_client, _put_bytes, _stub_record_check
from .test_render_document import make_firm_template

_FULL = FirmIdentity(
    authored=True,
    name="ACME LAW, LLP",
    street="100 Example Way",
    city_state_zip="Springfield, CA 90000",
    phone="(555) 010-0000",
    fax="(555) 010-0001",
    website="www.acme-law.example",
    source="test",
)

_YAML = """\
firm_identity:
  name: 'ACME LAW, LLP'
  street: '100 Example Way'
  city_state_zip: 'Springfield, CA 90000'
  phone: '(555) 010-0000'
  fax: '(555) 010-0001'
  website: 'www.acme-law.example'
"""


def _write(tmp_path, body: str) -> str:
    p = tmp_path / "customer.yaml"
    p.write_text(textwrap.dedent(body))
    return str(p)


def _first_page_header_lines(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    section = doc.sections[0]
    if not section.different_first_page_header_footer:
        return []
    return [p.text for p in section.first_page_header.paragraphs if p.text.strip()]


def _all_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            parts.extend(p.text for p in part.paragraphs)
    return "\n".join(parts)


# ---- config ---------------------------------------------------------------------


def test_authored_block_loads_every_field(tmp_path) -> None:
    fi = load_firm_identity(_write(tmp_path, _YAML))
    assert fi.authored is True
    assert fi.lines() == [
        "ACME LAW, LLP",
        "100 Example Way",
        "Springfield, CA 90000",
        "Telephone (555) 010-0000  |  Facsimile (555) 010-0001",
        "www.acme-law.example",
    ]


def test_absent_fields_are_omitted_never_invented(tmp_path) -> None:
    fi = load_firm_identity(_write(tmp_path, "firm_identity:\n  name: 'ACME LAW, LLP'\n  phone: '(555) 010-0000'\n"))
    assert fi.lines() == ["ACME LAW, LLP", "Telephone (555) 010-0000"]


@pytest.mark.parametrize(
    ("body", "reason"),
    [
        ("customer_id: x\n", "not authored"),
        ("firm_identity:\n  street: '100 Example Way'\n", "without a name"),
        ("firm_identity: 'ACME'\n", "not authored"),
    ],
)
def test_unauthored_or_nameless_block_is_not_authored_with_a_reason(tmp_path, body: str, reason: str) -> None:
    fi = load_firm_identity(_write(tmp_path, body))
    assert fi.authored is False and fi.lines() == []
    assert reason in fi.source


def test_missing_file_is_not_authored(tmp_path) -> None:
    fi = load_firm_identity(str(tmp_path / "absent.yaml"))
    assert fi.authored is False and "not readable" in fi.source


# ---- renderer ---------------------------------------------------------------------


@pytest.mark.parametrize("cls", sorted(LETTERHEAD_CLASSES))
def test_starter_prints_the_authored_letterhead_on_the_first_page(cls: str) -> None:
    data, report = render_document("Dear Counsel:\n\nBody.", cls, None, firm_identity=_FULL)
    assert _first_page_header_lines(data) == _FULL.lines()
    assert report.letterhead == {"source": "firm_identity", "lines": _FULL.lines()}
    # First page only: the continuation pages carry no letterhead.
    doc = Document(io.BytesIO(data))
    assert not "".join(p.text for p in doc.sections[0].header.paragraphs).strip()
    # The body is untouched by the letterhead.
    assert [p.text for p in doc.paragraphs if p.text.strip()] == ["Dear Counsel:", "Body."]
    assert report.to_dict()["letterhead"]["source"] == "firm_identity"


def test_letterhead_border_sits_in_schema_order() -> None:
    """Word calls a w:pBdr after w:spacing/w:jc unreadable content."""
    from docx.oxml.ns import qn

    data, _ = render_document("Body.", "letter", None, firm_identity=_FULL)
    last = Document(io.BytesIO(data)).sections[0].first_page_header.paragraphs[-1]
    tags = [child.tag for child in last._p.pPr]
    assert qn("w:pBdr") in tags
    for later in ("w:spacing", "w:jc"):
        if qn(later) in tags:
            assert tags.index(qn("w:pBdr")) < tags.index(qn(later))


@pytest.mark.parametrize("identity", [None, FirmIdentity(authored=False, source="firm_identity not authored")])
def test_absent_identity_prints_nothing_adds_no_marker_and_says_so(identity) -> None:
    data, report = render_document("Body.", "letter", None, firm_identity=identity)
    assert _first_page_header_lines(data) == []
    assert "{{" not in _all_text(data)
    assert report.letterhead["source"] == "none" and report.letterhead["lines"] == []
    assert any(n.startswith("no letterhead") for n in report.notes)


@pytest.mark.parametrize("cls", ["memo", "discovery_set", "discovery_response", "mediation_brief"])
def test_non_letter_classes_carry_no_letterhead(cls: str) -> None:
    data, report = render_document("Body.", cls, None, firm_identity=_FULL)
    assert _first_page_header_lines(data) == []
    assert "ACME LAW" not in _all_text(data)
    assert report.letterhead is None


def test_firm_template_base_is_untouched_by_the_authored_identity() -> None:
    template = make_firm_template(header_text="FIRM OWN LETTERHEAD", header_image=False)
    data, report = render_document("Body.", "letter", template, FormatReport("letter"), firm_identity=_FULL)
    doc = Document(io.BytesIO(data))
    assert doc.sections[0].header.paragraphs[0].text == "FIRM OWN LETTERHEAD"
    assert doc.sections[0].different_first_page_header_footer is False
    assert "ACME LAW" not in _all_text(data)
    assert report.letterhead == {"source": "firm_template", "lines": ["FIRM OWN LETTERHEAD"]}


def test_a_template_filed_from_the_starter_carries_its_letterhead_into_later_drafts() -> None:
    """The reason the letterhead lives in the header: a filed template is the
    FORMAT BASE for every later draft and its body is cleared, so body text
    would be lost. The header survives."""
    template, _ = render_document("# SKELETON\n\n{{FILL: body | matter record}}", "letter", None, firm_identity=_FULL)
    draft, report = render_document("Draft body.", "letter", template, FormatReport("letter"), firm_identity=None)
    assert _first_page_header_lines(draft) == _FULL.lines()
    assert report.letterhead["source"] == "firm_template"


# ---- the tools ----------------------------------------------------------------------


def test_template_tool_prints_the_authored_letterhead_on_the_starter(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv(CUSTOMER_YAML_ENV, _write(tmp_path, _YAML))
    captured: list[httpx.Request] = []
    monkeypatch.setattr(server, "_get_client", lambda: _mock_client(_handler(captured)))
    skeleton = "# SKELETON\n\n{{FILL: salutation | matter contacts}}\n"
    out = server.render_docx_template("m-ops", "Template - Letter", skeleton, document_class="letter")
    assert out["refusals"] == [] and out["fileId"] == "file-88"
    assert out["formatApplied"]["letterhead"]["source"] == "firm_identity"
    assert _first_page_header_lines(_put_bytes(captured))[0] == "ACME LAW, LLP"


def test_draft_tool_without_authored_identity_reports_none(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv(CUSTOMER_YAML_ENV, str(tmp_path / "absent.yaml"))
    captured: list[httpx.Request] = []
    monkeypatch.setattr(server, "_get_client", lambda: _mock_client(_handler(captured)))
    _stub_record_check(monkeypatch)
    out = server.render_docx_draft("m-1", "Draft", "Body.", document_class="demand_letter")
    assert out["formatApplied"]["letterhead"]["source"] == "none"
    assert _first_page_header_lines(_put_bytes(captured)) == []
