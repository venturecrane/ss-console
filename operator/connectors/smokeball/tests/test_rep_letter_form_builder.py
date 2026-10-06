"""``tools/build_rep_letter_forms.py --third-party-task-form``, on a synthetic
letter shaped like the one the firm's 3rd party TASK makes: the insurer, its
address, the client, the date of loss and the attorney merged by Smokeball
AUTOMATIONFIELD codes; the delivery line, the claim number and the title
typed by the person filling the task. The real task-form letter is client data
and is never committed; this one carries invented values only.

What it defends: each merged value is found by its FIELD, not its text; the
typed values by their position; field codes and the last address line's
paragraph are gone; the build refuses a letter missing a field.
"""

from __future__ import annotations

import importlib.util
import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from smokeball_connector.form_docx import DOCUMENT_PART, FormError, document_paragraphs, placeholders_in

_TOOL = Path(__file__).parents[1] / "tools" / "build_rep_letter_forms.py"
_spec = importlib.util.spec_from_file_location("build_rep_letter_forms", _TOOL)
assert _spec and _spec.loader
builder = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(builder)

_PATH = "AUTOMATIONFIELD customtables_data_source\\*abc/def\\*\\*"


def _run(paragraph: Any, **kind: str) -> None:
    run = OxmlElement("w:r")
    paragraph._p.append(run)
    if "char" in kind:
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind["char"])
        run.append(char)
    elif "instr" in kind:
        instr = OxmlElement("w:instrText")
        instr.text = kind["instr"]
        run.append(instr)
    else:
        text = OxmlElement("w:t")
        text.text = kind["text"]
        run.append(text)


def _field(paragraph: Any, path: str, shown: str) -> None:
    _run(paragraph, char="begin")
    _run(paragraph, instr=_PATH + path)
    _run(paragraph, char="separate")
    if shown:
        _run(paragraph, text=shown)
    _run(paragraph, char="end")


def _task_letter(*, with_claim: bool = True) -> bytes:
    doc = Document()
    doc.add_paragraph("\tVia Fax: 555-000-1111")
    attn = doc.add_paragraph("Attn:  Claims")
    _field(attn, "Defendant/Insurance Policy & Vehicle Details/Adjuster/Full Name", "")
    for path, shown in (
        ("Defendant 1/Insurance Policy Details 1/Insurer/Full Name", "Invented Mutual"),
        ("Defendant 1/Insurance Policy Details 1/Insurer/Street Address - Lines 1 & 2 (formatted across)", "1 Way"),
        ("Defendant 1/Insurance Policy Details 1/Insurer/Street Address - Last Line", "Town, ST 00000"),
    ):
        _field(doc.add_paragraph(), path, shown)
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 1).paragraphs[0].add_run("Our Client:")
    _field(table.cell(0, 2).paragraphs[0], "Plaintiff 1/Full Name(s) (All Parties)", "Pat Invented")
    if with_claim:
        table.cell(1, 1).paragraphs[0].add_run("Claim#:")
        table.cell(1, 2).paragraphs[0].add_run("CLM-0001")
    loss = doc.add_paragraph()
    _field(loss, "Case Details/Incident Details/Date", "01/02/2026")
    doc.add_paragraph("Cordially,")
    _field(doc.add_paragraph(), "Info/Attorney Responsible/Full Name", "Lee Lawyer")
    doc.add_paragraph("\tCounsel")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_builds_the_third_party_form_from_fields_and_positions() -> None:
    form = builder.build_third_party_task_form(_task_letter())
    assert tuple(placeholders_in(form)) == builder.THIRD_PARTY_FIELDS
    texts = [t for t in document_paragraphs(form) if t.strip()]
    assert texts == [
        "\t{{delivery_lines}}",
        "Attn:  Claims",
        "{{carrier_name}}",
        "{{carrier_address}}",
        "Our Client:",
        "{{client_name}}",
        "Claim#:",
        "{{claim_number}}",
        "{{date_of_loss}}",
        "Cordially,",
        "{{signer_name}}",
        "\t{{signer_title}}",
    ]
    body = zipfile.ZipFile(io.BytesIO(form)).read(DOCUMENT_PART)
    for leaked in (b"AUTOMATIONFIELD", b"<w:fldChar", b"Invented", b"Town, ST", b"CLM-0001", b"555-000"):
        assert leaked not in body


def test_refuses_a_letter_missing_a_value_it_must_place() -> None:
    with pytest.raises(FormError, match="Claim#"):
        builder.build_third_party_task_form(_task_letter(with_claim=False))
