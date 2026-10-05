"""The firm's own rep-letter forms, filled (``render_firm_form_letter``).

What these defend, each written so that removing the line it defends fails it:

* the filled letter is the FORM with values in it: every package part except
  the body is byte for byte the form's, and the body's fixed text is the
  firm's, line for line (synthetic facts; the real-value comparison against a
  letter the firm made is run locally and never committed);
* a placeholder split across runs (a person edited the form in Word) fills;
* a fact the record does not hold prints as ``[Not in the file: ...]`` and is
  listed in ``unfilled``; nothing is defaulted;
* the adjuster supplies the carrier's fax and email, never its name;
* the signer is the authored signer map's, never the staff record's name;
* a form that does not resolve refuses and files nothing;
* the committed forms carry no source matter's binding.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from smokeball_connector import form_letters as fl
from smokeball_connector.form_docx import (
    CUSTOM_PROPS_PART,
    DOCUMENT_PART,
    build_form,
    document_paragraphs,
    fill_form,
    placeholders_in,
)
from smokeball_connector.library import NotResolved, ResolvedTemplate

FORMS_DIR = Path(__file__).parent / "fixtures" / "forms"
FIRST = (FORMS_DIR / "Form - 1st Party Rep Letter.docx").read_bytes()
THIRD = (FORMS_DIR / "Form - 3rd Party Rep Letter.docx").read_bytes()

MATTER = "m-1"
CLIENT = "c-client"
ADJUSTER = "c-adj"
INSURER_3P = "c-ins3"
STAFF = "s-1"
ASSIST = "s-2"


def _layout(**overrides: Any) -> dict[str, Any]:
    values = {
        "Matter/Plaintiffs/InsurancePolicy/Adjuster": "rel-adj",
        "Matter/Plaintiffs/InsurancePolicy/Claims/Number": "111222333",
        "Matter/Defendants/InsurancePolicy/Insurer": "rel-ins3",
        "Matter/Defendants/InsurancePolicy/Claims/Number": "44-5556667",
        "Matter/CaseDetails/AccidentDetails/AccidentDate": "2026-03-04",
    }
    values.update(overrides)
    return values


class _Record:
    """The read side of a matter shaped like the live one (roles carrying
    relationship ids, layouts carrying those ids, company contacts)."""

    def __init__(self, layout: dict[str, Any] | None = None, contacts: dict[str, Any] | None = None) -> None:
        self.layout = _layout() if layout is None else layout
        self.contacts = {
            CLIENT: {"id": CLIENT, "person": {"firstName": "Dana", "lastName": "Example"}},
            ADJUSTER: {
                "id": ADJUSTER,
                "company": {
                    "name": "Pat Adjuster",
                    "fax": {"areaCode": "800", "number": "555-0101"},
                    "email": "claims@carrier.example",
                    "businessAddress": {"state": "CA", "country": "United States"},
                },
            },
            INSURER_3P: {
                "id": INSURER_3P,
                "company": {
                    "name": "Other Side Mutual",
                    "businessAddress": {
                        "addressLine1": "PO Box 1",
                        "city": "Townville",
                        "state": "OH",
                        "zipCode": "44000",
                    },
                },
            },
        }
        self.contacts.update(contacts or {})
        self.uploads: list[tuple[str, str, bytes]] = []
        self.files: list[dict[str, Any]] = []

    def get(self, path: str, **_params: Any) -> Any:
        if path == f"/matters/{MATTER}":
            return {
                "id": MATTER,
                "clientIds": [CLIENT],
                "personResponsibleStaffId": STAFF,
                "personAssistingStaffId": ASSIST,
            }
        if path == f"/matters/{MATTER}/roles":
            return {
                "roles": [
                    {
                        "name": "Plaintiff",
                        "isClient": True,
                        "contactId": CLIENT,
                        "relationships": [
                            {"id": "rel-atty", "name": "Attorney", "contactId": "MyFirm"},
                            {"id": "rel-adj", "name": "Adjuster", "contactId": ADJUSTER},
                        ],
                    },
                    {
                        "name": "Defendant",
                        "isOtherSide": True,
                        "contactId": "c-def",
                        "relationships": [{"id": "rel-ins3", "name": "Insurer", "contactId": INSURER_3P}],
                    },
                ]
            }
        if path == f"/matters/{MATTER}/layouts":
            return {"value": [{"id": "lay-1"}]}
        if path == f"/matters/{MATTER}/layouts/lay-1":
            return {"values": [{"key": k, "value": v} for k, v in self.layout.items()]}
        if path.startswith("/contacts/"):
            return self.contacts[path.rsplit("/", 1)[1]]
        if path == f"/staff/{STAFF}":
            return {"id": STAFF, "firstName": "Sam", "lastName": "Signer", "role": "Attorney"}
        if path == f"/staff/{ASSIST}":
            return {"id": ASSIST, "firstName": "Alex", "lastName": "Barnes", "role": None}
        if path == f"/matters/{MATTER}/documents/files":
            return {"value": self.files}
        if path.startswith(f"/matters/{MATTER}/documents/files/"):
            return {"id": path.rsplit("/", 1)[1], "name": "1st party letter"}
        raise AssertionError(f"unscripted GET {path}")

    def add_file(self, matter_id: str, file_name: str, data: bytes, *, folder_id: str | None = None) -> dict[str, Any]:
        self.uploads.append((matter_id, file_name, data))
        return {"fileId": "f-new", "matterId": matter_id, "fileName": file_name, "uploaded": True}


@pytest.fixture
def signers(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    path = tmp_path / "customer.yaml"
    path.write_text(
        "form_letters:\n  signers:\n    Sam Signer: {name: 'Samuel Q. Signer', title: 'Attorney at Law', initials: 'SQS'}\n",
        encoding="utf-8",
    )
    monkeypatch.setenv("SMD_CUSTOMER_YAML_PATH", str(path))
    return path


def _values(record: _Record, form: str) -> dict[str, str]:
    found = fl.gather_facts(record, MATTER, fl.FORMS[form], date(2026, 10, 5))
    return {k: f.value if f.value is not None else fl.MARKER.format(f.missing) for k, f in found.items()}


def _texts(blob: bytes) -> list[str]:
    return [p for p in document_paragraphs(blob) if p.strip()]


def test_the_gap_marker_carries_no_all_capitals_word() -> None:
    """The reply to the requester quotes every gap marker, and the outbound
    checklist holds a staff reply carrying an all-capitals word as emphasis
    (overlay shared/output_checklist.py caps_emphasis). On 2026-10-05 the first
    live request's reply was held for "[NOT IN THE FILE: ...]" and never sent."""
    import re

    rendered = fl.MARKER.format("3rd party insurer name")
    assert not re.search(r"\b[A-Z]{2,}\b", rendered), rendered


# ---- the form is the firm's, byte for byte outside the body ----------------


def test_the_committed_forms_carry_their_fields() -> None:
    assert placeholders_in(FIRST) == [
        "date",
        "delivery_lines",
        "carrier_name",
        "carrier_address",
        "client_name",
        "date_of_loss",
        "claim_number",
        "signer_name",
        "signer_title",
    ]
    assert placeholders_in(THIRD) == [
        "date",
        "delivery_lines",
        "carrier_name",
        "client_name",
        "claim_number",
        "date_of_loss",
        "signer_name",
        "signer_initials",
        "signer_title",
    ]


def test_the_committed_forms_carry_no_source_matter_binding() -> None:
    for blob in (FIRST, THIRD):
        with zipfile.ZipFile(io.BytesIO(blob)) as zf:
            assert b"MatterId" not in zf.read(CUSTOM_PROPS_PART)
            assert b"<w:docVars>" not in zf.read("word/settings.xml")
            assert b"AUTOMATIONFIELD" not in zf.read(DOCUMENT_PART)
            assert b"<w:fldChar" not in zf.read(DOCUMENT_PART)


def test_only_the_body_changes_when_filled(signers: Path) -> None:
    filled = fill_form(FIRST, _values(_Record(), "first_party_rep"), fl.PARAGRAPH_FIELDS).data
    with zipfile.ZipFile(io.BytesIO(FIRST)) as a, zipfile.ZipFile(io.BytesIO(filled)) as b:
        assert a.namelist() == b.namelist()
        assert [n for n in a.namelist() if a.read(n) != b.read(n)] == [DOCUMENT_PART]


def test_first_party_letter_reads_as_the_firms_letter(signers: Path) -> None:
    filled = fill_form(FIRST, _values(_Record(), "first_party_rep"), fl.PARAGRAPH_FIELDS)
    assert filled.unknown == []
    assert _texts(filled.data) == [
        "October 5, 2026",
        "VIA FAX: (800) 555-0101",
        "Email: claims@carrier.example",
        "Attn: Claims",
        # No 1st-party insurer on the record: the adjuster is the carrier's
        # claims desk for fax and email, never the carrier's NAME.
        "[Not in the file: 1st party insurer name]\n[Not in the file: 1st party insurer address]",
        "RE:",
        "Our Client /Your Insured:",
        "Dana Example",
        "Date of Loss:",
        "03/04/2026",
        "Claim Number:",
        "111222333",
        "To Whom It May Concern,",
        "ASHTON & PRICE has been retained to represent the above-named claimant with respect to a personal "
        "injury claim stemming from an incident on the indicated date, therefore please direct all future "
        "communication to this office. Please also include a declarations page with your letter of acknowledgement. ",
        "Please be advised that any and all authorizations signed by our client releasing any information are "
        "hereby revoked. Also, request is hereby made that you forward to us a copy of any statement obtained from "
        "our client, any accident report, preservation and copy of vehicle event data recorder, and photographs, "
        "audio, video, and digital recording you may have relating to this accident and any information regarding "
        "witnesses known to you.",
        "Thank you for your anticipated courtesy and cooperation in this matter.  If you have any questions, "
        "please feel free to contact our office at (916) 786-7787. ",
        "Cordially,.",
        "Samuel Q. Signer",
        "Attorney at Law",
    ]


def test_third_party_letter_fills_its_own_layout(signers: Path) -> None:
    filled = fill_form(THIRD, _values(_Record(), "third_party_rep"), fl.PARAGRAPH_FIELDS)
    texts = _texts(filled.data)
    assert texts[0] == "October 5, 2026"
    assert texts[1] == "\t\t\t\t\t[Not in the file: 3rd party insurer fax or email]"
    assert "Other Side Mutual" in texts
    assert "RE:\tOur client:\t\tDana Example" in texts
    assert "\t\tClaim#:\t\t44-5556667" in texts
    assert "Date of Loss:\t\t03/04/2026" in texts
    assert texts[-2:] == ["\t\t\t\t\t\t\tSamuel Q. Signer", "SQS/ab\t\t\t\t\t\t\tAttorney at Law"]


def test_the_reference_line_pairs_the_attorney_with_the_staff_on_the_file(signers: Path) -> None:
    """EAS/cr, CAP/ic: the attorney's authored initials, then the assisting
    staff member's. A file with no assisting staff marks the second half."""
    assert _values(_Record(), "third_party_rep")["signer_initials"] == "SQS/ab"

    class _NoAssist(_Record):
        def get(self, path: str, **params: Any) -> Any:
            if path == f"/matters/{MATTER}":
                return {"id": MATTER, "clientIds": [CLIENT], "personResponsibleStaffId": STAFF}
            return super().get(path, **params)

    assert _values(_NoAssist(), "third_party_rep")["signer_initials"] == (
        "SQS/[Not in the file: preparer initials (no assisting staff on the matter)]"
    )


# ---- the facts ------------------------------------------------------------


def test_an_insurer_on_the_record_is_the_carrier_and_its_channels_win(signers: Path) -> None:
    insurer = {
        "id": "c-ins1",
        "company": {
            "name": "First Carrier Co",
            "email": "first@carrier.example",
            "businessAddress": {"addressLine1": "PO Box 9", "city": "Dallas", "state": "TX", "zipCode": "75266"},
        },
    }
    record = _Record(
        layout=_layout(**{"Matter/Plaintiffs/InsurancePolicy/Insurer": "rel-ins1"}), contacts={"c-ins1": insurer}
    )
    roles = record.get(f"/matters/{MATTER}/roles")
    roles["roles"][0]["relationships"].append({"id": "rel-ins1", "name": "Insurer", "contactId": "c-ins1"})
    record.get = _patched_roles(record, roles)
    values = _values(record, "first_party_rep")
    assert values["carrier_name"] == "First Carrier Co"
    assert values["carrier_address"] == "PO Box 9\nDallas, TX 75266"
    # Fax from the adjuster (the insurer has none), email from the insurer.
    assert values["delivery_lines"] == "VIA FAX: (800) 555-0101\nEmail: first@carrier.example"


def _patched_roles(record: _Record, roles: dict[str, Any]) -> Any:
    original = record.get

    def get(path: str, **params: Any) -> Any:
        return roles if path == f"/matters/{MATTER}/roles" else original(path, **params)

    return get


def test_missing_facts_print_markers_and_are_listed(signers: Path) -> None:
    record = _Record(layout={})
    values = _values(record, "first_party_rep")
    assert values["claim_number"] == "[Not in the file: 1st party claim number]"
    assert values["date_of_loss"] == "[Not in the file: date of loss]"
    assert values["delivery_lines"] == "VIA FAX: (800) 555-0101\nEmail: claims@carrier.example", (
        "the adjuster is still found by its role relationship when the layout holds nothing"
    )


def test_an_unmapped_signer_prints_markers_not_the_staff_record(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "customer.yaml"
    path.write_text("form_letters:\n  signers: {}\n", encoding="utf-8")
    monkeypatch.setenv("SMD_CUSTOMER_YAML_PATH", str(path))
    values = _values(_Record(), "third_party_rep")
    assert values["signer_name"] == "[Not in the file: how Sam Signer signs: name]"
    assert "Sam Signer" not in (values["signer_title"].replace("how Sam Signer signs", ""))


def test_a_signer_is_found_by_staff_id_too(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "customer.yaml"
    path.write_text(
        f"form_letters:\n  signers:\n    {STAFF}: {{name: 'S. Signer', title: 'Partner'}}\n", encoding="utf-8"
    )
    monkeypatch.setenv("SMD_CUSTOMER_YAML_PATH", str(path))
    values = _values(_Record(), "first_party_rep")
    assert (values["signer_name"], values["signer_title"]) == ("S. Signer", "Partner")


# ---- the docx engine ------------------------------------------------------


def _letter(paragraph_runs: list[list[str]]) -> bytes:
    from docx import Document

    doc = Document()
    for runs in paragraph_runs:
        p = doc.add_paragraph()
        for text in runs:
            p.add_run(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_a_placeholder_split_across_runs_fills() -> None:
    blob = _letter([["Dear {{cli", "ent_na", "me}}, hello"], ["{{date}}"]])
    filled = fill_form(blob, {"client_name": "Dana Example", "date": "May 1, 2026"})
    assert _texts(filled.data) == ["Dear Dana Example, hello", "May 1, 2026"]
    assert filled.unknown == []


def test_an_unknown_placeholder_is_reported_and_left_visible() -> None:
    filled = fill_form(_letter([["{{nickname}} and {{date}}"]]), {"date": "May 1, 2026"})
    assert filled.unknown == ["nickname"]
    assert _texts(filled.data) == ["{{nickname}} and May 1, 2026"]


def test_build_form_puts_each_value_in_one_run_and_unwraps_nothing_else() -> None:
    source = _letter([["Our client: ", "Da", "na Ex", "ample"], ["Claim ", "1234"], ["second line"]])
    form = build_form(source, [("Dana Example", "{{client_name}}"), ("1234", "{{claim_number}}")], ["second line"])
    assert _texts(form) == ["Our client: {{client_name}}", "Claim {{claim_number}}"]
    xml = zipfile.ZipFile(io.BytesIO(form)).read(DOCUMENT_PART).decode()
    assert ">{{client_name}}</w:t>" in xml and ">{{claim_number}}</w:t>" in xml
    with zipfile.ZipFile(io.BytesIO(source)) as a, zipfile.ZipFile(io.BytesIO(form)) as b:
        assert [n for n in a.namelist() if a.read(n) != b.read(n)] == [DOCUMENT_PART]


def test_build_form_refuses_a_value_it_cannot_find_exactly_once() -> None:
    from smokeball_connector.form_docx import FormError

    with pytest.raises(FormError):
        build_form(_letter([["a b a"]]), [("a", "{{x}}")], [])


# ---- the tool --------------------------------------------------------------


@pytest.fixture
def record(monkeypatch: pytest.MonkeyPatch, signers: Path) -> _Record:
    r = _Record()
    monkeypatch.setattr(fl, "_client", lambda: r)
    monkeypatch.setattr(fl, "SLEEP", lambda _s: None)
    return r


def _resolves(monkeypatch: pytest.MonkeyPatch, blob: bytes) -> None:
    hit = ResolvedTemplate(bytes=blob, name="Form", file_id="f-form", matter_id="lib", folder_id=None)
    monkeypatch.setattr(fl, "_resolve_form", lambda _client, _spec: hit)


def test_a_form_that_does_not_resolve_files_nothing(record: _Record, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        fl, "_resolve_form", lambda _c, _s: NotResolved("no file named 'Form - 1st Party Rep Letter.docx'")
    )
    out = fl.render_firm_form_letter(MATTER, "first_party_rep")
    assert out["status"] == "refused"
    assert "Form - 1st Party Rep Letter.docx" in out["reason"]
    assert record.uploads == []


def test_files_under_the_firms_name_and_reports_gaps(record: _Record, monkeypatch: pytest.MonkeyPatch) -> None:
    _resolves(monkeypatch, FIRST)
    record.files = [{"id": "f-old", "name": "1st party letter", "fileExtension": ".docx"}]
    out = fl.render_firm_form_letter(MATTER, "first_party_rep", "2026-10-05")
    assert out["status"] == "filed"
    assert record.uploads[0][1] == "1st party letter.docx"
    assert out["unfilled"] == [
        "[Not in the file: 1st party insurer name]",
        "[Not in the file: 1st party insurer address]",
    ]
    assert out["same_name_on_matter"] == ["f-old"]
    assert "signer_initials" not in out["facts_used"], "a field the form does not carry is not reported as used"
    assert out["facts_used"]["claim_number"] == "layout Matter/Plaintiffs/InsurancePolicy/Claims/Number"


def test_an_unknown_form_or_bad_date_refuses(record: _Record) -> None:
    assert fl.render_firm_form_letter(MATTER, "demand")["status"] == "refused"
    assert fl.render_firm_form_letter(MATTER, "third_party_rep", "10/05/2026")["status"] == "refused"
    assert record.uploads == []


def test_a_record_that_cannot_be_read_files_nothing(record: _Record, monkeypatch: pytest.MonkeyPatch) -> None:
    _resolves(monkeypatch, THIRD)
    record.contacts.pop(INSURER_3P)
    out = fl.render_firm_form_letter(MATTER, "third_party_rep")
    assert out["status"] == "refused"
    assert "could not be read" in out["reason"]
    assert record.uploads == []


def test_a_file_that_is_not_a_form_is_refused(record: _Record, monkeypatch: pytest.MonkeyPatch) -> None:
    _resolves(monkeypatch, _letter([["A finished letter, no fields."]]))
    out = fl.render_firm_form_letter(MATTER, "first_party_rep")
    assert out["status"] == "refused"
    assert record.uploads == []
