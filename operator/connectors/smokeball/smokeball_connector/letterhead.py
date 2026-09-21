"""The firm's letterhead on the starter base, printed from AUTHORED config.

WHY THIS EXISTS (observed 2026-09-21). The document library renders a class
template onto the starter when the firm has no Word file of its own for that
class, and the starter had no letterhead. The establishment turn tried to type
the firm's letterhead into the template as text; the template content gate
refused the street number, the zip and the phone digits (a bare run of five or
more digits outside a marker), so the whole letterhead was markered as a FILL.
Even had it passed, the letterhead would have been lost: a filed template is
opened as the FORMAT BASE for every later draft and its body is cleared, so
only what lives in the header survives into a draft.

So the letterhead is not model text at all. The firm authors its identity once,
in ``customer.yaml`` (``firm_identity``), and tool code prints it into the
FIRST-PAGE HEADER of the starter for the letter classes. It never passes through
the content gate because the model never types it: it is config, reviewed in a
PR, like every other authored value.

Precedence, and each step is a real authoring decision:

1. **The firm's own Word file wins.** When the class resolves to a file in the
   firm's library, the renderer opens it as the base and its header, whatever
   the firm built, is kept exactly. Nothing here runs.
2. **Else the authored identity**, printed on the starter's first page header.
   A template filed from that starter carries the letterhead in its header, so
   it survives into every later draft rendered into that template.
3. **Else nothing**, reported. Never invented, never a marker: an unauthored
   field is an omitted line, and an unauthored block is a sentence in the
   delivery note.

WHICH CLASSES. ``letter`` and ``demand_letter``: correspondence, where the
firm's name and address head the first page. Not ``memo`` (internal). Not the
pleadings (``discovery_set``, ``discovery_response``, ``mediation_brief``): on
pleading paper the firm block sits in the BODY beneath the signing attorney's
name and State Bar number, which vary per document, and a body is cleared every
time a template is used as a base. Printing the firm lines in a header would
put them above the attorney line and reverse the pleading's order, so the
pleading caption block stays content the drafter writes, as the skeletons show.
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Any

from .library import CUSTOMER_YAML_ENV, DEFAULT_CUSTOMER_YAML

#: The classes whose starter carries the authored letterhead. See the module
#: docstring for why the pleadings and the memo are not here.
LETTERHEAD_CLASSES = frozenset({"letter", "demand_letter"})

#: The authored fields, in print order. ``name`` is required; every other field
#: is optional and an absent one is an omitted line.
FIELDS = ("name", "street", "city_state_zip", "phone", "fax", "website")

_NAME_SIZE_PT = 14
_LINE_SIZE_PT = 10


@dataclass(frozen=True)
class FirmIdentity:
    authored: bool
    name: str | None = None
    street: str | None = None
    city_state_zip: str | None = None
    phone: str | None = None
    fax: str | None = None
    website: str | None = None
    source: str = ""

    def lines(self) -> list[str]:
        """The letterhead's lines, exactly as authored. The only words code adds
        are the two labels on the telephone line; the values are the firm's."""
        if not self.authored or not self.name:
            return []
        out = [self.name]
        out.extend(v for v in (self.street, self.city_state_zip) if v)
        contact = []
        if self.phone:
            contact.append(f"Telephone {self.phone}")
        if self.fax:
            contact.append(f"Facsimile {self.fax}")
        if contact:
            out.append("  |  ".join(contact))
        if self.website:
            out.append(self.website)
        return out


def _scalar(value: Any) -> str | None:
    if value is None or isinstance(value, (dict, list)):
        return None
    text = str(value).strip()
    return text or None


def load_firm_identity(path: str | None = None) -> FirmIdentity:
    """Read ``firm_identity`` from the seat's live customer.yaml. A missing
    file, a missing block, an unparseable file, or a block without a ``name``
    all mean "not authored", with the reason carried in ``source``."""
    path = path or os.environ.get(CUSTOMER_YAML_ENV) or DEFAULT_CUSTOMER_YAML
    try:
        with open(path, encoding="utf-8") as fh:
            raw = fh.read()
    except OSError as exc:
        return FirmIdentity(authored=False, source=f"customer.yaml not readable at {path}: {exc.__class__.__name__}")
    try:
        import yaml

        data = yaml.safe_load(raw) or {}
    except Exception as exc:  # noqa: BLE001 - an unparseable config is "not authored", reported
        return FirmIdentity(authored=False, source=f"customer.yaml not parseable: {exc.__class__.__name__}")
    block = data.get("firm_identity") if isinstance(data, dict) else None
    if not isinstance(block, dict):
        return FirmIdentity(authored=False, source="firm_identity not authored in customer.yaml")
    values = {f: _scalar(block.get(f)) for f in FIELDS}
    if not values["name"]:
        return FirmIdentity(authored=False, source="firm_identity authored without a name")
    return FirmIdentity(authored=True, source=path, **values)


def apply_letterhead(doc, identity: FirmIdentity) -> list[str]:
    """Print ``identity`` into the first-page header of ``doc`` (the STARTER,
    never a firm's file). Returns the lines printed; an empty list means
    nothing was printed and the header was not touched."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    lines = identity.lines()
    if not lines:
        return []
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.first_page_header
    header.is_linked_to_previous = False
    for i, text in enumerate(lines):
        para = header.paragraphs[0] if i == 0 and header.paragraphs else header.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(text)
        run.font.size = Pt(_NAME_SIZE_PT if i == 0 else _LINE_SIZE_PT)
        if i == 0:
            run.bold = True
    # A thin rule under the last line, and breathing room before the body.
    last = header.paragraphs[-1]
    last.paragraph_format.space_after = Pt(12)
    ppr = last._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "auto")
    pbdr.append(bottom)
    # Schema order matters: Word treats a w:pBdr placed after w:spacing or w:jc
    # as unreadable content, so it goes before its successors, not at the end.
    ppr.insert_element_before(pbdr, *_PBDR_SUCCESSORS)
    return lines


_PBDR_SUCCESSORS = (
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)


__all__ = ["FIELDS", "LETTERHEAD_CLASSES", "FirmIdentity", "apply_letterhead", "load_firm_identity"]
