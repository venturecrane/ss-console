"""Excel workbooks: detect one, read its cells as text, and build one.

WHY THIS EXISTS. Until 2026-09-25 an ``.xlsx`` on a matter took the DOCX road
in :mod:`extract` (every Office Open XML file carries ``[Content_Types].xml``,
which is what ``_looks_like_docx`` keys on), python-docx refused it, and the
file landed in the drafting record check's ``unextractable`` list, so ONE
spreadsheet on a matter refused every ``render_docx_draft`` there. The Operator
also had no way to produce a workbook at all: the costs-and-advances workbooks
the firm received were built off the seat.

READING is deliberately dumb, like the rest of :mod:`extract`: cell VALUES as
text, one line per row, never a formula evaluated here. A formula cell whose
file carries no saved result says so in words, because a blank would read as
zero and a zero is a claim about money.

BUILDING takes typed columns and does the arithmetic in code (``totals``), so
the model never adds a column of money. A text cell is always a literal string:
a value beginning ``=`` is text in the file, never a formula (spreadsheet
formula injection). Every workbook built here is stamped with
:data:`OPERATOR_CREATOR`, and the drafting record check uses that stamp to
exclude the Operator's own output from "the firm's record".
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from decimal import Decimal, InvalidOperation
from typing import Any

#: Stamped into every workbook built here (``docProps/core.xml`` creator).
OPERATOR_CREATOR = "SMD Operator"

#: OLE compound file: a legacy .xls/.doc, or an encrypted Office file.
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

FORMATS = ("text", "number", "integer", "currency", "date")
_NUMBER_FORMATS = {
    "number": "#,##0.00",
    "integer": "#,##0",
    "currency": '"$"#,##0.00',
    "date": "yyyy-mm-dd",
}
NO_SAVED_VALUE = "[formula, no saved value]"

MAX_SHEETS = 50
MAX_CELLS = 200_000
_BAD_SHEET_CHARS = set("[]:*?/\\")


class WorkbookRefused(ValueError):
    """The input cannot become a workbook. The message names what and where."""


def _zip_names(blob: bytes) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as zf:
            return zf.namelist()
    except (zipfile.BadZipFile, ValueError):
        return []


def is_xlsx(blob: bytes) -> bool:
    """A zip holding ``xl/workbook.xml``. Read from the central directory, not a
    byte window: the part can sit anywhere in the archive."""
    return blob.startswith(b"PK\x03\x04") and "xl/workbook.xml" in _zip_names(blob)


def is_pptx(blob: bytes) -> bool:
    return blob.startswith(b"PK\x03\x04") and "ppt/presentation.xml" in _zip_names(blob)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def xlsx_text(blob: bytes) -> str:
    """Every sheet as ``## Sheet: <name>`` followed by one line per non-empty
    row, cells joined ``" | "``, trailing empty cells dropped."""
    from openpyxl import load_workbook

    try:
        values = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(blob), read_only=True, data_only=False)
    except Exception as exc:  # any parse failure is "unsupported", named by the caller
        raise ValueError(f"XLSX could not be parsed: {exc}") from exc
    parts: list[str] = []
    for ws in values.worksheets:
        parts.append(f"## Sheet: {ws.title}")
        raw_rows = formulas[ws.title].iter_rows(values_only=True)
        for row, raw in zip(ws.iter_rows(values_only=True), raw_rows):
            cells = []
            for value, source in zip(row, raw):
                if value is None and isinstance(source, str) and source.startswith("="):
                    cells.append(NO_SAVED_VALUE)
                else:
                    cells.append(_cell_text(value))
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                parts.append(" | ".join(cells))
    values.close()
    formulas.close()
    return "\n".join(parts)


def xlsx_creator(blob: bytes) -> str:
    """The ``docProps/core.xml`` creator, or ``""``."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(blob), read_only=True)
    except Exception:  # noqa: BLE001 - an unreadable workbook simply has no creator stamp
        return ""
    try:
        return str(wb.properties.creator or "")
    finally:
        wb.close()


# ---- Building ---------------------------------------------------------------


def _coerce(value: Any, fmt: str, where: str) -> Any:
    """One input value as the cell value its column's format asks for."""
    if value is None or value == "":
        return None
    if isinstance(value, (dict, list, tuple)):
        raise WorkbookRefused(f"{where}: a cell must be a single value, not a {type(value).__name__}")
    if fmt == "text":
        return value if isinstance(value, str) else _cell_text(value)
    if fmt == "date":
        if isinstance(value, str):
            try:
                return dt.date.fromisoformat(value.strip()[:10])
            except ValueError:
                pass
        raise WorkbookRefused(f"{where}: {value!r} is not an ISO date (YYYY-MM-DD)")
    if isinstance(value, bool):
        raise WorkbookRefused(f"{where}: {value!r} is not a number")
    try:
        number = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except InvalidOperation:
        raise WorkbookRefused(f"{where}: {value!r} is not a number") from None
    if not number.is_finite():
        raise WorkbookRefused(f"{where}: {value!r} is not a finite number")
    if fmt == "integer":
        if number != number.to_integral_value():
            raise WorkbookRefused(f"{where}: {value!r} is not a whole number")
        return int(number)
    return float(number)


def _check_sheet_name(name: Any, seen: set[str]) -> str:
    if not isinstance(name, str) or not name.strip():
        raise WorkbookRefused("every sheet needs a name")
    name = name.strip()
    if len(name) > 31 or set(name) & _BAD_SHEET_CHARS:
        raise WorkbookRefused(f"sheet name {name!r} must be 31 characters or fewer without []:*?/\\")
    if name.lower() in seen:
        raise WorkbookRefused(f"sheet name {name!r} is used twice")
    seen.add(name.lower())
    return name


def _columns(name: str, columns: Any) -> tuple[list[str], list[str]]:
    if not isinstance(columns, list) or not columns:
        raise WorkbookRefused(f"sheet {name!r}: columns must be a non-empty list")
    headers, formats = [], []
    for col in columns:
        header = col.get("header") if isinstance(col, dict) else None
        fmt = (col.get("format") or "text") if isinstance(col, dict) else None
        if not isinstance(header, str) or not header.strip():
            raise WorkbookRefused(f"sheet {name!r}: every column needs a header")
        if fmt not in FORMATS:
            raise WorkbookRefused(f"sheet {name!r}, column {header!r}: format must be one of {', '.join(FORMATS)}")
        if header.strip() in headers:
            raise WorkbookRefused(f"sheet {name!r}: column {header.strip()!r} is used twice")
        headers.append(header.strip())
        formats.append(fmt)
    return headers, formats


def _rows(name: str, rows: Any, headers: list[str], formats: list[str]) -> list[list[Any]]:
    if not isinstance(rows, list):
        raise WorkbookRefused(f"sheet {name!r}: rows must be a list of rows")
    out = []
    for r, row in enumerate(rows, start=1):
        if not isinstance(row, list) or len(row) > len(headers):
            raise WorkbookRefused(f"sheet {name!r}, row {r}: must be a list of at most {len(headers)} values")
        coerced = [_coerce(v, formats[c], f"sheet {name!r}, row {r}, column {headers[c]!r}") for c, v in enumerate(row)]
        out.append(coerced + [None] * (len(headers) - len(row)))
    return out


def _totals(name: str, wanted: Any, headers: list[str], formats: list[str], rows: list[list[Any]]) -> dict[str, Any]:
    """Column sums computed here, in Decimal, so no total is model arithmetic."""
    if wanted is not None and not isinstance(wanted, list):
        raise WorkbookRefused(f"sheet {name!r}: totals must be a list of column headers")
    totals: dict[str, Any] = {}
    for header in wanted or []:
        if header not in headers:
            raise WorkbookRefused(f"sheet {name!r}: totals names {header!r}, which is not a column")
        c = headers.index(header)
        if formats[c] not in ("number", "integer", "currency"):
            raise WorkbookRefused(f"sheet {name!r}: totals column {header!r} is not numeric")
        if c == 0:
            raise WorkbookRefused(f"sheet {name!r}: the first column holds the Total label and cannot be totalled")
        exact = sum((Decimal(str(row[c])) for row in rows if row[c] is not None), Decimal(0))
        totals[header] = int(exact) if formats[c] == "integer" else float(exact)
    return totals


def normalize_sheets(sheets: Any) -> list[dict[str, Any]]:
    """Validate the input and return it with every cell coerced and every total
    computed. Raises :class:`WorkbookRefused` naming the first problem."""
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= MAX_SHEETS:
        raise WorkbookRefused(f"sheets must be a list of 1 to {MAX_SHEETS} sheets")
    seen: set[str] = set()
    cells = 0
    out = []
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise WorkbookRefused("each sheet must be an object with name, columns and rows")
        name = _check_sheet_name(sheet.get("name"), seen)
        headers, formats = _columns(name, sheet.get("columns"))
        rows = _rows(name, sheet.get("rows") or [], headers, formats)
        cells += len(headers) * (len(rows) + 1)
        if cells > MAX_CELLS:
            raise WorkbookRefused(f"more than {MAX_CELLS} cells")
        totals = _totals(name, sheet.get("totals"), headers, formats, rows)
        out.append({"name": name, "headers": headers, "formats": formats, "rows": rows, "totals": totals})
    return out


def _totals_row(sheet: dict[str, Any]) -> list[Any]:
    return [sheet["totals"].get(h, "Total" if i == 0 else None) for i, h in enumerate(sheet["headers"])]


def build_workbook(sheets: list[dict[str, Any]]) -> bytes:
    """Bytes of an .xlsx for already-normalized sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = OPERATOR_CREATOR
    bold = Font(bold=True)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["name"])
        ws.append(sheet["headers"])
        for cell in ws[1]:
            cell.font = bold
        body = list(sheet["rows"])
        if sheet["totals"]:
            body.append(_totals_row(sheet))
        for r, row in enumerate(body, start=2):
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                cell = ws.cell(row=r, column=c)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"  # literal text, never a formula
                fmt = sheet["formats"][c - 1]
                if fmt in _NUMBER_FORMATS and not isinstance(value, str):
                    cell.number_format = _NUMBER_FORMATS[fmt]
        if sheet["totals"]:
            for cell in ws[len(body) + 1]:
                cell.font = bold
        ws.freeze_panes = "A2"
        for c, header in enumerate(sheet["headers"], start=1):
            width = max([len(header)] + [len(_cell_text(row[c - 1])) for row in body])
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def verify_workbook(blob: bytes, sheets: list[dict[str, Any]]) -> list[str]:
    """Re-open the built bytes and compare EVERY cell to the normalized input.
    Returns the mismatches; empty means the file says exactly what was asked."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(blob), data_only=False)
    problems: list[str] = []
    if wb.sheetnames != [s["name"] for s in sheets]:
        problems.append(f"sheet names {wb.sheetnames} != {[s['name'] for s in sheets]}")
        return problems
    for sheet in sheets:
        ws = wb[sheet["name"]]
        expected = [sheet["headers"], *sheet["rows"]]
        if sheet["totals"]:
            expected.append(_totals_row(sheet))
        width = len(sheet["headers"])
        actual = [list(r) for r in ws.iter_rows(min_row=1, max_row=len(expected), max_col=width, values_only=True)]
        if ws.max_row > len(expected):
            problems.append(f"sheet {sheet['name']!r}: {ws.max_row} rows, expected {len(expected)}")
        for r, (want, got) in enumerate(zip(expected, actual), start=1):
            for c, (w, g) in enumerate(zip(want, got), start=1):
                if isinstance(w, dt.date) and isinstance(g, dt.datetime):
                    g = g.date()
                if w != g:
                    problems.append(f"sheet {sheet['name']!r} row {r} column {c}: wrote {g!r}, expected {w!r}")
    return problems
