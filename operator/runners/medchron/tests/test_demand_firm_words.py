"""Nothing the demand lane files shows the firm this system's own vocabulary,
and the pre-delivery review stops raising two false findings (the 2026-10-15
proof run, held on five MUST-FIX findings):

* the letter carried ``{{NOT IN RECORD: ... removed by the final audit ...}}``
  in a draft written in the firm's square-bracket markers;
* the attorney notes carried a "delivery note" naming the seat's tool and its
  manifest, "Premise gate. G1 passed ...", the check labels (DRIFTS) and
  "Removed or flagged by the final audit";
* the review read the letter's paragraphs and then its tables, so a section
  that IS a table read as "a bare header with no text";
* the review flagged the numbered format's script line plus printed name as
  "the signature repeats the name".

Each test fails on the code before the change. Every name and figure here is
invented."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from demand_testkit import DRAFT, ScriptedClient, make_inputs, make_job, seat_with, standard_docs
from medchron_testkit import PRICING
from medchron.demand import draft, finalpass, format_check, premise, review as rv
from medchron.demand.run import DemandRun

CHECKER = Path(__file__).resolve().parents[3] / "templates" / "drafting" / "drafting_gate_check.py"
LETTER = (
    "## Liability\n\nThe light was red (police report, p. 1). The other driver admitted fault at the scene "
    "(cite). More.\n\n## Damages\n\nThe bill totals $500.00 (bill, p. 1).\n"
)
AUDIT = (
    "## AUDIT: Liability\n\n- The other driver admitted fault at the scene | INVENTED | no cite in the record\n\n"
    "## AUDIT: Damages\n\n- The bill totals $500.00 | ARITHMETIC | the lines sum to $450.00\n"
    "- the light was red | DRIFTS | the report says amber\n"
)


@pytest.fixture(autouse=True)
def _checker(monkeypatch):
    monkeypatch.setenv("SMD_DRAFTING_GATE_CHECK", str(CHECKER))


@pytest.fixture
def pricing(tmp_path: Path) -> Path:
    p = tmp_path / "pricing.json"
    p.write_text(json.dumps(PRICING), encoding="utf-8")
    return p


# ---- 1. the marker code writes into the letter -------------------------------------------
def test_the_brief_decides_the_marker_dialect():
    assert format_check.dialect("Draft the demand. Mark gaps {{NOT IN RECORD: what}}.") == format_check.CURLY
    assert format_check.dialect("Use {{FILL}} for anything the firm supplies.") == format_check.CURLY
    assert format_check.dialect("Draft the demand for the full limits.") == format_check.SQUARE
    assert format_check.dialect("") == format_check.SQUARE


@pytest.mark.parametrize("dialect", ["square", "curly"])
def test_a_settled_marker_is_in_the_drafts_dialect_counted_and_names_no_stage(dialect):
    md, notes = finalpass.settle(LETTER, AUDIT, dialect)
    assert "admitted fault" not in md and "$500.00" not in md
    opens = format_check.OPEN_ITEM.findall(md)
    assert len(opens) == 2  # the format check counts both as open items for the attorney
    if dialect == format_check.SQUARE:
        assert "{{" not in md and "[TO BE SUPPLIED: " in md and "[CONFIRM the arithmetic (" in md
    else:
        assert "[TO BE SUPPLIED" not in md and "{{NOT IN RECORD: " in md and "{{ATTORNEY: verify arithmetic: " in md
    # FALSIFIER: the marker said "removed by the final audit" and the notes "the final audit checked"
    assert rv.firm_banned(md) == [] and "audit" not in md.lower()
    assert all(rv.firm_banned(n) == [] and "audit" not in n.lower() for n in notes)
    assert finalpass.summary(notes) == {"removed": 1, "flagged": 1, "cleared": 0}


def test_a_check_that_states_its_own_figure_correct_is_noted_in_firm_words():
    _md, notes = finalpass.settle(
        "## Damages\n\nOne visit at $1,200.00 (bill).\n",
        "## AUDIT: Damages\n\n- one visit at $1,200.00 | ARITHMETIC | 1 x 1200.00 = 1200.00; correct\n",
    )
    assert finalpass.summary(notes) == {"removed": 0, "flagged": 0, "cleared": 1}
    assert "final audit" not in notes[0] and "left as written" in notes[0]


# ---- 2. the notes lines code writes -----------------------------------------------------
def test_a_partly_supported_finding_reaches_the_notes_without_its_label():
    (line,) = draft.finding_text(AUDIT, ("DRIFTS",))
    assert line == "Damages: the light was red (the report says amber)"
    assert rv.firm_banned(line) == [] and "|" not in line


def test_the_seat_delivery_rule_never_reaches_the_drafter():
    disc = (
        "## Part I\n\n1. ZERO INVENTION. Never invent.\n\n7. PLAIN REGISTER. No flourishes.\n\n"
        "8. READ THE AUTHORED SPEC. Your authored-spec pointer block names the file. Deliver through\n"
        "   `smd_deliver_draft`; the bytes are verified against the root manifest.\n\n"
        "Your output is the draft document only. Deliver it with `smd_deliver_draft`.\n\n## Part II\n\ngates\n"
    )
    part = draft.discipline_part_one(disc)
    assert "1. ZERO INVENTION" in part and "7. PLAIN REGISTER" in part
    assert "smd_" not in part and "manifest" not in part and "8. READ" not in part
    assert "G1 to G4" in draft.FIRM_READS and "never the skeleton" in draft.FIRM_READS


# ---- 3. the banned-term check, which holds -----------------------------------------------
@pytest.mark.parametrize(
    "said",
    [
        "Delivery note. smd_deliver_draft was not available to this turn.",
        "No spec file was read or verified against the root manifest.",
        "Premise gate. G1 passed. G2 passed.",
        "- **G4 passed.** A pre-suit claim.",
        "Taken out by the final audit: a sentence.",
        "Damages: a claim | DRIFTS | the record says less",
        "Skeleton marker in Duty to Accept",
        "General damages (fixed skeleton text)",
        "The drafting gate repaired a quotation.",
    ],
)
def test_each_system_term_is_a_must_fix_that_names_the_words(said):
    found = rv.wording({"attorney_notes": f"# Attorney notes\n\n{said}\n"})
    assert found and all(f.grade == "MUST-FIX" and f.check == "5 system vocabulary" for f in found)
    assert any(f'says "' in f.finding for f in found)


@pytest.mark.parametrize(
    "clean",
    [
        "No acute abnormality of the visualized axial skeleton (MRI report, p. 1).",
        "Her symptoms manifest within days of the collision (Dr. Example, 2/1/26).",
        "Diagnosis: epilepsy, G40.909 (clinic note). Obstetric history G1P0.",
        "The claimant worked as an auditor for an accounting firm (wage letter).",
        "The record supports the treatment (ER record, p. 2).",
        "[TO BE SUPPLIED: the bill; looked in the billing folder]",
    ],
)
def test_record_words_that_resemble_a_term_never_hold(clean):
    assert rv.wording({"demand": clean, "attorney_notes": clean}) == []


def test_a_system_term_in_the_attorney_notes_holds_the_job_and_nothing_is_filed(tmp_path, pricing):
    planted = DRAFT.replace(
        "=== ATTORNEY NOTES ===\n",
        "=== ATTORNEY NOTES ===\n\nDelivery note. smd_deliver_draft was not available to this turn.\n",
        1,
    )
    make_inputs(tmp_path / "inputs")
    make_job(tmp_path / "job")
    seat = seat_with(standard_docs())
    v = DemandRun(
        tmp_path / "job",
        inputs_dir=str(tmp_path / "inputs"),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=ScriptedClient(draft=planted),
        log=lambda m: None,
        readback_pause=0.0,
        vendor_factory=lambda: None,
    ).run()
    assert v.outcome == "failed" and v.reason.startswith("review_held:") and "smd_deliver_draft" in v.reason
    assert seat.sent == [] and seat.created == []


def test_the_coverage_report_names_no_gate():
    decision = {
        "gates": [
            {"gate": "G1 identified carrier", "passed": True, "evidence": ["matter insurer field: Example Mutual"]},
            {"gate": "G2 coverage not denied", "passed": False, "evidence": ["Carrier letter: we deny the claim"]},
            {"gate": "G4 the right instrument (pre-suit, unresolved)", "passed": True, "evidence": []},
        ],
        "premise_facts": ["conditional limits language (a constraint on figures): Letter: if limits apply"],
    }
    md = premise.coverage_report(SimpleNamespace(matter_number="100001"), decision, "October 15, 2026")
    assert rv.wording({"coverage_report": md}) == []  # FALSIFIER: "## The premise gates", "G1 identified carrier"
    assert "Coverage not denied in writing | No | Carrier letter: we deny the claim" in md


# ---- 4. the review reads a table where it stands -----------------------------------------
def test_the_review_reads_a_table_under_its_heading(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("I. LIABILITY")
    d.add_paragraph("The light was red.")
    d.add_paragraph("II. DAMAGES")
    t = d.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "Past medical", (0, 1): "$1,200.00", (1, 0): "General", (1, 1): "TBD"}.items():
        t.cell(r, c).text = text
    d.add_paragraph("III. OBJECTIVE FINDINGS")
    p = tmp_path / "letter.docx"
    d.save(str(p))
    lines = rv.docx_text(p).splitlines()
    damages, findings = lines.index("II. DAMAGES"), lines.index("III. OBJECTIVE FINDINGS")
    # FALSIFIER: the tables were read after every paragraph, so DAMAGES had nothing under it
    assert lines[damages + 1 : findings] == ["| Past medical | $1,200.00 |", "| General | TBD |"]


# ---- 5. the review knows the authored signature block -------------------------------------
def test_the_review_is_told_the_authored_signature_block():
    firm = SimpleNamespace(
        attorney={"signer": "Gamma Q. Counsel", "title": "Attorney at Law", "initials": "GQC:ab"},
        variant={"layout": "numbered"},
    )
    block = rv.signature_block(firm)
    assert "Gamma Q. Counsel" in block and "script signature line" in block and "by design" in block
    user = rv.demand_user("Draft it.", [], [], set(), {"demand": "letter"}, [], signature=block)
    assert "THE SIGNATURE BLOCK, AS THE FIRM AUTHORED IT" in user and "by design" in user
    plain = rv.signature_block(SimpleNamespace(attorney={"signer": "Gamma Q. Counsel", "title": "T"}, variant={}))
    assert "script" not in plain and "Gamma Q. Counsel" in plain
