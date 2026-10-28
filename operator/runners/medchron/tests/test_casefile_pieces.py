"""Three pieces the Case File Manager owes: low limits flagged in code, new
files flagged with the opening tasks they lack, and specialty appointments
proposed for the calendar as numbered lines a person approves.

Everything here is synthetic (Example / Exampletown)."""

from __future__ import annotations

import copy
import datetime as dt
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

import casefile_testkit as kit
from medchron.casefile import apply as apply_mod, firm as firm_mod, flags, own_writes, propose, rank, recheck, state
from medchron.casefile.apply import ApplyRun
from medchron.casefile.book import render_html
from medchron.casefile.run import CasefileRun
from medchron.casefile.validate import check

TODAY = dt.date(2026, 10, 26)
APPLY_JOB = "01KTJ0BX0000000000000000AP"
LOW = {"per_person_usd": 30000, "bills_trigger_usd": 1000, "carrier_patterns": [r"example mutual"]}
NEW = {
    "opened_within_days": 30,
    "opening_tasks": [
        {"name": "Police report request", "pattern": "police report"},
        {"name": "Wage loss letter", "pattern": "wage loss"},
    ],
}


def _firm(tmp: Path, **over) -> firm_mod.CasefileFirm:
    return firm_mod.load(kit.make_inputs(tmp / "inputs", low_limits=LOW, new_matter=NEW, **over))


# ---- the firm's thresholds --------------------------------------------------------------


def test_the_new_firm_keys_are_accepted_and_checked() -> None:
    assert firm_mod.validate(kit.firm_data(low_limits=LOW, new_matter=NEW)) == []
    bad = firm_mod.validate(kit.firm_data(low_limits={"per_person": 30000}))
    assert bad and "low_limits" in bad[0]
    bad = firm_mod.validate(kit.firm_data(new_matter={"opening_tasks": [{"name": "X", "pattern": "("}]}))
    assert any("invalid regex" in b for b in bad)
    bad = firm_mod.validate(kit.firm_data(new_matter={"opening_tasks": [{"name": "X"}]}))
    assert any("{name, pattern}" in b for b in bad)


def test_with_no_thresholds_authored_nothing_is_flagged(tmp_path: Path) -> None:
    firm = firm_mod.load(kit.make_inputs(tmp_path / "inputs"))
    rec = {"insurance": {"third_party": [{"carrier": "Example Mutual", "limits": "15/30"}]}}
    assert flags.low_limits(rec, firm) == []
    assert flags.new_matter({"detail": {"openedDate": "2026-10-20"}}, firm, TODAY) is None


# ---- low limits ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text,want",
    [
        ("30/60", 30000),
        ("15/30/5", 15000),
        ("$15,000/$30,000", 15000),
        ("100k/300k", 100000),
        ("1M CSL", 1000000),
        ("$250,000 per person", 250000),
        # As dec pages print split limits: a label after each figure (the 10-28 pilot proof).
        ("$15,000 each person / $30,000 each accident", 15000),
        ("$15,000 per person / $30,000 per accident", 15000),
        ("BI: $25,000 per person / $50,000 per occurrence", 25000),
        ("Bodily Injury $100,000 each person/$300,000 each accident", 100000),
        ("see the dec page", None),
        ("$15,000 each person and something else", None),
        ("30/60 (claimed by adjuster)", None),
        ("", None),
        (None, None),
    ],
)
def test_a_limits_string_parses_only_when_it_reads_as_limits(text, want) -> None:
    assert flags.per_person(text) == want


def test_low_limits_are_decided_in_code_from_the_firms_thresholds(tmp_path: Path) -> None:
    firm = _firm(tmp_path)
    rec = {
        "medicals": [{"provider": "A", "charge": 600.0}, {"provider": "B", "charge": 500}],
        "insurance": {
            "first_party": [{"carrier": "Own Carrier", "limits": "100/300"}, {"carrier": "Own Two", "limits": "25/50"}],
            "third_party": [
                {"carrier": "Example Mutual Insurance", "limits": None},
                {"carrier": "Other Carrier", "limits": "see the dec page"},
            ],
        },
    }
    got = flags.low_limits(rec, firm)
    assert [(g["side"], g["carrier"]) for g in got] == [
        ("first_party", "Own Two"),
        ("third_party", "Example Mutual Insurance"),
    ]
    assert got[0]["urgent"] is False  # the client's own policy never carries the bills trigger
    assert got[1]["urgent"] is True and "past the firm's $1,000 line" in got[1]["why"]
    assert "lists as writing low limits" in got[1]["why"]


def test_a_low_limit_ranks_this_month_and_this_week_once_the_bills_pass_the_line(tmp_path: Path) -> None:
    firm = _firm(tmp_path)
    rec = {"posture": "pre_demand", "insurance": {"third_party": [{"carrier": "Other", "limits": "15/30"}]}}
    inv: dict = {}
    got = rank.rank(rec, inv, firm, TODAY, set(), 0, flags.flags_for(rec, inv, firm, TODAY))
    assert got["priority"] == 3 and got["reason"].startswith("Low limit, other side (Other)")
    rec["medicals"] = [{"provider": "A", "charge": 5000.0}]
    got = rank.rank(rec, inv, firm, TODAY, set(), 0, flags.flags_for(rec, inv, firm, TODAY))
    assert got["priority"] == 2
    assert rank.rank(rec, inv, firm, TODAY, set(), 0)["priority"] == 4  # the flag is what ranks it


def test_the_file_status_prints_the_low_limit_flag_on_the_policy(tmp_path: Path) -> None:
    firm = _firm(tmp_path)
    rec = {
        "number": "900001",
        "clients": ["Test Client"],
        "headline": "Synthetic.",
        "insurance": {
            "third_party": [{"carrier": "Other", "limits": "15/30"}, {"carrier": "Big Carrier", "limits": "500/1000"}]
        },
    }
    page = render_html(
        rec, "2026-10-26", {}, {"priority": 3, "reason": "x"}, {"flags": flags.flags_for(rec, {}, firm, TODAY)}
    )
    assert "<th>Low limit</th>" in page
    assert page.count('<span class="bad">Low limit</span>') == 1
    assert "per-person limit $15,000, at or under the firm&#x27;s $30,000" in page


# ---- new files ------------------------------------------------------------------------------


def _inv(opened: str, *subjects: str) -> dict:
    return {"detail": {"openedDate": opened}, "tasks_open": [{"subject": s} for s in subjects], "tasks_done": []}


def test_a_new_file_with_none_of_its_opening_tasks_is_flagged_with_them(tmp_path: Path) -> None:
    firm = _firm(tmp_path)
    assert flags.new_matter(_inv("2026-10-20", "Call client"), firm, TODAY) == {
        "opened": "2026-10-20",
        "missing": ["Police report request", "Wage loss letter"],
    }
    assert flags.new_matter(_inv("2026-10-20", "Request police report"), firm, TODAY) is None  # the workflow ran
    assert flags.new_matter(_inv("2026-08-01", "Call client"), firm, TODAY) is None  # not new
    assert flags.new_matter({"detail": {}}, firm, TODAY) is None  # no opened date: never guessed


def test_the_file_status_names_the_opening_tasks_a_new_file_lacks(tmp_path: Path) -> None:
    firm = _firm(tmp_path)
    rec = {"number": "900002", "clients": ["Test Client"], "headline": "Synthetic.", "posture": "pre_demand"}
    fl = flags.flags_for(rec, _inv("2026-10-20"), firm, TODAY)
    page = render_html(rec, "2026-10-26", {}, {"priority": 3, "reason": "x"}, {"flags": fl})
    assert "<h2>New file: opening tasks not set</h2>" in page
    assert "<li>Police report request</li>" in page and "Nothing was created." in page
    assert rank.rank(rec, {}, firm, TODAY, set(), 0, fl)["priority"] == 3


# ---- specialty appointments: the record ------------------------------------------------------


def _meta() -> dict:
    return {kit.DOC_REPORT: {"state": "text 900", "name": "Letter.pdf"}, kit.DOC_SCAN: {"state": "SCAN 1 page"}}


def _appt(**over) -> dict:
    a = {"provider": "Dr Alpha Example", "specialty": "Orthopedic", "date": "2026-11-10", "time": "14:30",
         "time_zone": "America/Denver", "source_doc_ids": [kit.DOC_REPORT], "note": "Consult letter."}  # fmt: skip
    a.update(over)
    return a


def _violations(appts: list) -> list[str]:
    firm = SimpleNamespace(statute_rx=[], at_stake_rx=None)
    rec = {**kit.review_for("100002", ""), "specialty_appointments": appts}
    inv = {"tasks_open": [{"id": kit.TASK2, "subject": "Call client"}]}
    return [v for v in check(rec, "100002", inv, _meta(), firm, lambda _d: "") if v.startswith("appointment")]


def test_the_validator_refuses_a_bad_appointment_and_accepts_a_good_one() -> None:
    assert _violations([_appt()]) == []
    assert (
        _violations([_appt(date=None, time=None, time_zone=None)]) == []
    )  # no date in the file: reported, never proposed
    got = _violations([_appt(date="11/10/26"), _appt(time="2:30 PM"), _appt(time_zone="Mountain"),
                       _appt(source_doc_ids=[kit.DOC_SCAN]), _appt(specialty="")])  # fmt: skip
    assert len(got) == 5
    assert any("is a SCAN" in g or "SCAN" in g for g in got)


def test_a_recheck_patch_upserts_an_appointment_by_provider_and_date() -> None:
    prior = {"specialty_appointments": [_appt(), _appt(provider="Example Imaging", date="2026-11-12")]}
    got = recheck.merge(prior, {"upsert": {"specialty_appointments": [_appt(time="15:00")]}})
    assert [a["time"] for a in got["specialty_appointments"]] == ["15:00", "14:30"]


# ---- specialty appointments: the proposed line ------------------------------------------------


def test_each_appointment_becomes_one_line_or_says_why_not(tmp_path: Path) -> None:
    firm = _firm(tmp_path)
    rec = {
        "number": "100001",
        "specialty_appointments": [
            _appt(),
            _appt(provider="Dr. Alpha Example", source_doc_ids=[kit.DOC_BILL], time=None, time_zone=None),
            _appt(provider="Example Imaging", specialty="MRI", date="2026-11-12", time="09:00", time_zone=None),
            _appt(provider="Dr Gone Example", date="2026-09-01"),
            _appt(provider="Dr Beta Example", date="2026-11-20"),
            _appt(provider="Dr Nodate Example", date=None),
        ],
    }
    inv = {"events": [{"id": "e1", "startTime": "2026-11-20T10:00:00", "subject": "Ortho consult, Dr. Beta Example"}]}
    lines, skipped = propose.appointment_lines(rec, inv, {}, firm, TODAY)
    assert [ln["facts"]["provider"] for ln in lines] == ["Dr Alpha Example", "Example Imaging"]
    alpha, mri = (ln["facts"] for ln in lines)
    assert alpha["source_doc_ids"] == [kit.DOC_REPORT, kit.DOC_BILL]  # two documents, one appointment
    assert (alpha["start"], alpha["end"], alpha["all_day"], alpha["time_zone"]) == (
        "2026-11-10T14:30:00", "2026-11-10T15:30:00", False, "America/Denver",
    )  # fmt: skip
    assert mri["all_day"] is True and "09:00 with no time zone" in mri["how"]  # never a guessed zone
    assert sorted(skipped.values()) == ["Already on the calendar", "Date passed", "No date in the file"]


def _proposed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, pricing_path: Path) -> dict:
    for k, v in kit.environ(tmp_path).items():
        monkeypatch.setenv(k, v)
    seat = kit.CfSeat()
    seat.matters[1]["openedDate"] = "2026-10-20"

    def review(number: str, pack: str) -> dict:
        rec = kit.review_for(number, pack)
        if number == "100001":
            rec["insurance"] = {"third_party": [{"carrier": "Example Mutual", "limits": "15/30", "layout_only": False}]}
            rec["specialty_appointments"] = [
                _appt(source_doc_ids=[kit.cite(pack, kit.DOC_REPORT)]),
                _appt(provider="Dr Gone Example", date="2026-09-01", source_doc_ids=[kit.cite(pack, kit.DOC_REPORT)]),
            ]
        return rec

    jd = kit.make_job(tmp_path / "job")
    inputs = kit.make_inputs(tmp_path / "inputs", low_limits=LOW, new_matter=NEW)
    r = CasefileRun(
        jd, inputs_dir=str(inputs), pricing=str(pricing_path), state_dir=tmp_path / "state",
        seat_factory=lambda: seat, client=kit.Scripted(review_fn=review), log=lambda m: None,
        today=lambda: TODAY, readback_pause=0,
    )  # fmt: skip
    v = json.loads(r.run().to_json())
    assert v["verdict"] == "delivered", v
    from smokeball_connector import readback

    monkeypatch.setattr(readback, "SLEEP", lambda s: None)
    monkeypatch.setattr(apply_mod, "READBACK_WAITS", (0,))
    return {"seat": seat, "lines": v["lines"], "state": tmp_path / "state", "tmp": tmp_path, "job": jd}


@pytest.fixture
def proposed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, pricing_path: Path) -> dict:
    return _proposed(tmp_path, monkeypatch, pricing_path)


def test_a_job_delivers_the_three_pieces_on_the_page_the_workbook_and_the_lines(proposed) -> None:
    kinds = sorted(ln["kind"] for ln in proposed["lines"])
    assert kinds == ["event", "medicals", "task"]
    data = proposed["job"] / "data"
    html = next((data / "out" / kit.M1).glob("*.html")).read_text()
    event = next(ln for ln in proposed["lines"] if ln["kind"] == "event")
    assert f"Proposed for the calendar, change {event['n']}, awaiting approval" in html
    assert "Date passed" in html and '<span class="bad">Low limit</span>' in html
    html2 = next((data / "out" / kit.M2).glob("*.html")).read_text()
    assert "<h2>New file: opening tasks not set</h2>" in html2
    wb = load_workbook(next((data / "out").glob("Case File Manager *.xlsx")))
    assert {"Low limits", "New files", "Proposed changes"} <= set(wb.sheetnames)
    low = [r for r in wb["Low limits"].iter_rows(min_row=2, values_only=True)]
    assert low[0][0] == "100001" and low[0][1] == "2 This week" and low[0][3] == "Example Mutual"
    assert [r[0] for r in wb["New files"].iter_rows(min_row=2, values_only=True)] == ["100002"]
    kinds_shown = [r[2] for r in wb["Proposed changes"].iter_rows(min_row=2, values_only=True)]
    assert "Calendar entry" in kinds_shown
    pri = {r[2]: r[1] for r in wb["Priorities"].iter_rows(min_row=2, values_only=True)}
    assert pri["100001"] == "2 This week"  # the low limit with the bills past the line


# ---- specialty appointments: the apply -------------------------------------------------------


def _apply(p: dict, lines: list[dict]) -> dict:
    jd = p["tmp"] / f"apply-{len(list(p['tmp'].glob('apply-*')))}"
    jd.mkdir()
    doc = {
        "kind": "casefile", "job_id": APPLY_JOB, "trigger": "request", "mode": "apply",
        "requester": "admin@firm.example", "message_ref": "<approve@firm.example>", "request_text": "yes",
        "slug": "example", "month_cents_used": 0,
        "apply": {"source_job_id": kit.JOB, "approver": "admin@firm.example", "thread_ref": "t-1", "lines": lines},
    }  # fmt: skip
    (jd / "job.json").write_text(json.dumps(doc))
    run = ApplyRun(jd, state_dir=p["state"], seat_factory=lambda: p["seat"], log=lambda m: None, sleep=lambda s: None)
    return {"run": run, "verdict": json.loads(run.run().to_json())}


def _event_line(p: dict) -> dict:
    return next(ln for ln in p["lines"] if ln["kind"] == "event")


def _posts(p: dict) -> list:
    return [w for w in p["seat"].client.writes if w[:2] == ("POST", "/events")]


def test_an_approved_calendar_line_is_written_stamped_on_the_attorneys_calendar_and_read_back(proposed) -> None:
    got = _apply(proposed, [_event_line(proposed)])["verdict"]
    assert got["verdict"] == "delivered", got
    assert got["line_outcomes"][0]["outcome"] == "created"
    ((_m, _p, body),) = _posts(proposed)
    assert body["subject"] == "[Operator] Orthopedic appointment: Dr Alpha Example"
    assert body["matterId"] == kit.M1 and body["attendees"] == [kit.STAFF_R] and body["type"] == "Normal"
    assert (body["startTime"], body["endTime"], body["timeZone"]) == (
        "2026-11-10T14:30:00", "2026-11-10T15:30:00", "America/Denver",
    )  # fmt: skip
    eid = proposed["seat"].events[-1]["id"]
    assert got["writes"] == [{"method": "POST", "path": f"/events/{eid}", "ok": True}]
    assert [r for r in own_writes.rows(proposed["state"]) if r["kind"] == "event"][0]["id"] == eid


def test_the_lanes_own_calendar_entry_is_no_change_on_the_next_recheck(proposed) -> None:
    inv = {"events": [], "files": [], "tasks_open": [], "tasks_done": [], "memos": [], "layouts": {}}
    before = state.fingerprint(inv, state.own_ids(proposed["state"]))
    _apply(proposed, [_event_line(proposed)])
    after_inv = copy.deepcopy(inv)
    after_inv["events"] = [dict(proposed["seat"].events[-1])]
    assert state.fingerprint(after_inv, state.own_ids(proposed["state"])) == before


def test_a_calendar_that_already_has_the_appointment_is_not_written_again(proposed) -> None:
    proposed["seat"].events.append(
        {"id": "e-firm", "matter": kit.M1, "startTime": "2026-11-10T14:30:00", "subject": "Dr Alpha Example ortho"}
    )
    got = _apply(proposed, [_event_line(proposed)])["verdict"]
    assert got["line_outcomes"][0]["outcome"] == "already_present" and not _posts(proposed)


def test_a_resumed_apply_finds_its_entry_and_never_posts_twice(proposed) -> None:
    first = _apply(proposed, [_event_line(proposed)])
    again = json.loads(first["run"].run().to_json())
    assert again["line_outcomes"][0]["outcome"] == "created" and len(_posts(proposed)) == 1


def test_a_calendar_line_not_as_proposed_or_in_a_proof_run_writes_nothing(proposed, monkeypatch) -> None:
    got = _apply(proposed, [{**_event_line(proposed), "sha256": "0" * 64}])["verdict"]
    assert got["line_outcomes"][0]["outcome"] == "refused_not_as_proposed"
    monkeypatch.setenv("MEDCHRON_PROOF_RUN", "1")
    got = _apply(proposed, [_event_line(proposed)])["verdict"]
    assert got["line_outcomes"][0]["outcome"] == "would_create" and not _posts(proposed)


def test_an_entry_that_does_not_read_back_holds_the_apply(proposed, monkeypatch) -> None:
    client = proposed["seat"].client
    real = client.get

    def drifted(path, **params):
        got = real(path, **params)
        if path.startswith("/events/"):
            got["startTime"] = "2026-11-10T21:30:00"  # stored in another clock: the defect the read-back catches
        return got

    monkeypatch.setattr(client, "get", drifted)
    got = _apply(proposed, [_event_line(proposed)])["verdict"]
    assert got["verdict"] == "held" and got["reason"].startswith("readback_mismatch:")
