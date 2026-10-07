"""The demand job (medchron/demand/): the DAG on synthetic fixtures, the
privilege wall, the pre-paid estimate hold, the premise gate, the format check
on a rendered fixture, and the drafting gate before filing."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
import yaml

from demand_testkit import (
    LIBRARY as LIBRARY_ID,
    MATTER as MATTER_ID,
    DRAFT,
    FOOTER,
    SIGNATURE,
    ScriptedClient,
    make_inputs,
    make_job,
    make_pdf,
    seat_with,
    standard_docs,
)
from medchron_testkit import PRICING
from medchron.demand import firm as firm_mod, format_check, preflight, pull
from medchron.demand.run import DemandRun

CHECKER = Path(__file__).resolve().parents[3] / "templates" / "drafting" / "drafting_gate_check.py"


@pytest.fixture(autouse=True)
def _checker(monkeypatch):
    monkeypatch.setenv("SMD_DRAFTING_GATE_CHECK", str(CHECKER))


@pytest.fixture
def pricing(tmp_path: Path) -> Path:
    p = tmp_path / "pricing.json"
    p.write_text(json.dumps(PRICING), encoding="utf-8")
    return p


def _run(tmp_path: Path, pricing: Path, seat, client, vendor=None, **job_kw):
    inputs = tmp_path / "inputs"
    if not inputs.is_dir():
        make_inputs(inputs)
    jd = make_job(tmp_path / "job", **job_kw)
    log: list[str] = []
    r = DemandRun(
        jd,
        inputs_dir=str(inputs),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=client,
        log=log.append,
        readback_pause=0.0,
        vendor_factory=(lambda: vendor) if vendor is not None else (lambda: None),
    )
    return r, r.run(), log


# ---- the firm inputs -------------------------------------------------------------
def test_the_firm_inputs_validate_and_every_file_is_pinned(tmp_path):
    root = make_inputs(tmp_path / "in")
    assert firm_mod.load(root).model("compose") == "claude-sonnet-5"
    (root / "voice_profile.md").write_text("edited by hand", encoding="utf-8")
    with pytest.raises(firm_mod.DemandConfigError, match="does not match the pin"):
        firm_mod.load(root)


def test_an_unknown_key_refuses_rather_than_defaulting(tmp_path):
    root = make_inputs(tmp_path / "in")
    data = yaml.safe_load((root / "demand-firm.yaml").read_text())
    data["budget"]["per_job_cap"] = 5
    (root / "demand-firm.yaml").write_text(yaml.safe_dump(data))
    with pytest.raises(firm_mod.DemandConfigError, match="unknown key"):
        firm_mod.load(root)
    assert firm_mod.main([str(root)]) == 1


# ---- the DAG end to end ------------------------------------------------------------
def test_a_demand_runs_pull_to_read_back_and_files_both_deliverables(tmp_path, pricing):
    seat = seat_with(standard_docs())
    client = ScriptedClient()
    r, v, log = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "delivered", (v.reason, log[-5:])
    names = sorted(f["name"] for f in v.files)
    assert names == sorted(
        [
            "Gap Audit - 100001 - " + r.date_stamp + ".docx",
            "Demand.Alpha Example.docx",
            "Demand.Alpha Example - attorney notes.docx",
        ]
    )
    # read back on the matter, in one dated folder this job created
    assert len(seat.created) == 1 and seat.created[0]["name"].startswith("Demand Prep ")
    assert {s["name"] for s in seat.sent} == set(names)
    # the vendor chronology folder was never pulled
    assert ["d4"] not in seat.mints and all("d4" not in m for m in seat.mints)
    # stages in order, every call live and streamed (no batch API exists on the client)
    stages = client.stages()
    assert stages[0] == "DIGEST" and stages.index("GAP") < stages.index("COMPOSE") < stages.index("AUDIT")
    assert all(c["max_tokens"] > 0 for c in client.calls)
    assert v.dollars > 0 and v.coverage_report is False
    # the rendered demand passes the firm's format check as filed
    demand = tmp_path / "job" / "data" / "out" / "demand" / "Demand.Alpha Example.docx"
    res = format_check.check(
        demand,
        signature=SIGNATURE,
        signer_title="Attorney at Law",
        footer_markers=["Settlement Communication", "1119, 1152"],
    )
    assert res.ok, res.fails
    assert any("open item" in n for n in res.notes)  # the NOT IN RECORD marker is counted, not refused
    bills = json.loads((tmp_path / "job" / "data" / "preflight.json").read_text())["bills"]
    assert bills["tab_providers_without_a_bill_in_file"] == []  # matched on the bill's letterhead, not its file name
    state = json.loads((tmp_path / "job" / "data" / "state.json").read_text())
    assert {
        "facts",
        "destination",
        "pull",
        "preflight",
        "premise",
        "transcribe",
        "summarize",
        "gate.json",
        "gate-gap-audit.json",
        "render",
    } <= set(state)


def test_a_resume_after_delivery_pays_for_nothing_twice(tmp_path, pricing):
    seat = seat_with(standard_docs())
    client = ScriptedClient()
    _run(tmp_path, pricing, seat, client)
    n = len(client.calls)
    r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "delivered" and len(client.calls) == n
    assert len(seat.sent) == 3  # nothing re-sent: the delivery record says it went


def test_the_estimate_holds_before_anything_is_paid(tmp_path, pricing):
    make_inputs(tmp_path / "inputs", budget={"per_job_cap_usd": 1.0})
    seat = seat_with(standard_docs())
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "failed" and v.stage == "per_job_cap_usd"
    assert "estimate before anything was paid" in v.reason
    assert client.calls == [] and v.dollars == 0


def test_the_months_demand_spend_holds_the_estimate_too(tmp_path, pricing):
    seat = seat_with(standard_docs())
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat, client, cents=19_900)
    assert v.outcome == "failed" and v.stage == "monthly_budget_usd" and client.calls == []


def test_a_premise_failure_files_a_coverage_report_and_pays_nothing(tmp_path, pricing):
    docs = standard_docs() + [
        ("d5", "Demand Acceptance letter.pdf", make_pdf(["We accept the policy limits demand."]), "f-corr")
    ]
    seat = seat_with(docs)
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "delivered" and v.coverage_report is True
    assert client.calls == [] and v.dollars == 0
    assert [f["name"].startswith("Coverage Posture Report") for f in v.files] == [True]
    report = (tmp_path / "job" / "data" / "coverage-report.md").read_text()
    assert (
        "Demand Acceptance letter (rule: a document name of class acceptance)" in report
        and "No demand was drafted" in report
    )


def test_a_written_denial_fails_g2(tmp_path, pricing):
    docs = standard_docs() + [
        (
            "d6",
            "Carrier letter 2.pdf",
            make_pdf(
                [
                    "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. Coverage is denied."
                ]
            ),
            "f-corr",
        )
    ]
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), ScriptedClient())
    assert v.coverage_report is True
    decision = json.loads((tmp_path / "job" / "data" / "premise.json").read_text())
    assert [g["passed"] for g in decision["gates"]] == [True, False, True]


def test_a_truncated_digest_is_split_and_redone_without_a_hand_step(tmp_path, pricing):
    client = ScriptedClient(truncate_digest_once=True)
    _r, v, _ = _run(tmp_path, pricing, seat_with(standard_docs()), client)
    assert v.outcome == "delivered"
    assert client.stages().count("DIGEST") == 3  # the chunk, then its two halves
    assert list((tmp_path / "job" / "data" / "digest").glob("truncated-*.md"))


def test_the_drafting_gate_refuses_before_anything_is_filed(tmp_path, pricing):
    bad = DRAFT.replace("Demand is hereby made", "This demand fully addresses every claim. Demand is hereby made")
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(draft=bad))
    assert v.outcome == "held" and "drafting gate refused the letter" in v.reason
    assert seat.sent == [] and seat.created == []


def test_an_unfound_quotation_is_repaired_twice_then_held_never_filed(tmp_path, pricing):
    bad = DRAFT.replace(
        "Demand is hereby made",
        'The record says "the patient reported severe pain radiating to both arms daily" and demand is made',
    )
    seat = seat_with(standard_docs())
    client = ScriptedClient(draft=bad)
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "held" and "1 quotation(s) not found" in v.reason and "after 2 repair(s)" in v.reason
    assert client.stages().count("REPAIR") == 2
    assert seat.sent == [] and seat.created == []


def test_invented_facts_left_after_repair_hold_and_drifts_reach_the_attorney_notes(tmp_path, pricing):
    seat = seat_with(standard_docs())
    client = ScriptedClient(
        audit="- claim | INVENTED | no cite in digest\nSUPPORTED=0 DRIFTS=0 INVENTED=1 ARITHMETIC=0"
    )
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "held" and " invented and " in v.reason and seat.sent == []
    assert client.stages().count("REPAIR") == 2
    tmp2 = tmp_path / "drifts"
    seat2 = seat_with(standard_docs())
    client2 = ScriptedClient(
        audit="- the date | DRIFTS | digest says 1/16\nSUPPORTED=3 DRIFTS=1 INVENTED=0 ARITHMETIC=0"
    )
    _r, v2, _ = _run(tmp2, pricing, seat2, client2)
    assert v2.outcome == "delivered", v2.reason
    import docx

    notes = docx.Document(str(tmp2 / "job" / "data" / "out" / "demand" / "Demand.Alpha Example - attorney notes.docx"))
    text = "\n".join(p.text for p in notes.paragraphs)
    assert "DRIFTS" in text and "digest says 1/16" in text


def test_the_short_re_block_and_demand_table_are_audited():
    from medchron.demand import draft

    secs = draft.sections(DRAFT)
    audited = {h for h, _ in draft.auditable(secs)}
    assert {"PREAMBLE", "Demand", "Summary of Injuries"} <= audited
    assert all(len(b) <= 300 for h, b in secs if h == "Demand")  # short: only the figure test brings it in


def test_a_repair_that_drops_a_section_leaves_no_repaired_draft_to_resume_from(tmp_path):
    from medchron.demand import draft as draft_mod

    d = draft_mod.Drafter.__new__(draft_mod.Drafter)
    d.data = tmp_path
    (tmp_path / "draft-v1.md").write_text(DRAFT)
    (tmp_path / "audit-v1.md").write_text("## AUDIT: Liability\n\n- x | INVENTED | none\n")
    d.firm, d.compose_max = None, 100

    class R:
        text = "## Damages\n\nnothing for Liability"
        stop_reason = "end_turn"

    d.compose_system = lambda: "S"
    d.firm = type("F", (), {"text": lambda self, k: "REPAIR-PROMPT", "model": lambda self, k: "claude-opus-5-5"})()
    d._call = lambda *a, **k: R()
    with pytest.raises(draft_mod.DraftError):
        d.repair("digest", 1)
    assert not (tmp_path / "draft-v2.md").exists()


def test_the_re_block_is_cross_checked_against_the_matter_record(tmp_path, pricing):
    wrong = DRAFT.replace("dol: January 15, 2026", "dol: January 16, 2026")
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(draft=wrong))
    assert v.outcome == "held" and "date of loss" in v.reason and seat.sent == []
    from medchron.demand import crosscheck

    other = DRAFT.replace("client: Alpha Example", "client: Gamma Person")
    assert crosscheck.check(other, {"client_name": "Alpha Example", "date_of_loss": "01/15/2026"})["mismatches"]
    left = DRAFT.replace("client: Alpha Example", "client: {{FILL: client}}")
    out = crosscheck.check(left, {"client_name": "Alpha Example", "date_of_loss": "01/15/2026"})
    assert out["mismatches"] == [] and out["unchecked"]


def test_the_request_is_never_an_audit_source(tmp_path):
    from medchron.demand import draft

    block = draft.brief_block("Carrier claim # 99-1234; limits $15,000.", [])
    assert "never a source of facts" in block
    root = make_inputs(tmp_path / "in")
    d = draft.Drafter(
        tmp_path,
        None,
        firm_mod.load(root).select("pre_suit", {"signer": "S", "title": "T", "initials": "I"}),
        "Carrier claim # 99-1234",
        [],
        print,
        matter_fields=["client: A"],
    )
    (tmp_path / "draft-v1.md").write_text(DRAFT)
    seen = {}

    class Door:
        def call(self, stage, **kw):
            seen["system"] = kw["system"]
            return type("R", (), {"text": "SUPPORTED=0 DRIFTS=0 INVENTED=0 ARITHMETIC=0", "stop_reason": "end_turn"})()

    d.doorway = Door()
    d.audit(1, "digest", "corpus")
    assert "99-1234" not in seen["system"] and "client: A" in seen["system"]


def test_an_off_format_demand_is_refused_before_filing(tmp_path, pricing):
    bad = DRAFT.replace("## Liability", "## Fault")
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(draft=bad))
    assert v.outcome == "held" and "format check refused" in v.reason and "Liability" in v.reason
    assert seat.sent == []


# ---- the privilege wall -------------------------------------------------------------
FIRM = ("firm.example",)


@pytest.mark.parametrize(
    "sender,recipients,client,want",
    [
        (
            "client@mail.example",
            ["atty@firm.example"],
            {"client@mail.example"},
            "between the firm and the client or a personal mailbox",
        ),
        (
            "atty@firm.example",
            ["client@mail.example"],
            {"client@mail.example"},
            "between the firm and the client or a personal mailbox",
        ),
        ("atty@firm.example", ["para@firm.example"], {"client@mail.example"}, "firm internal"),
        # a second personal address (a spouse) is walled although the client's address is known
        (
            "spouse@mail.example",
            ["atty@firm.example"],
            {"client@mail.example"},
            "between the firm and the client or a personal mailbox",
        ),
        # Exchange writes an internal sender as a name or X.500 path: the firm's own
        ("", ["atty@firm.example"], {"client@mail.example"}, "firm internal"),
        ("", ["client@mail.example"], {"client@mail.example"}, "between the firm and the client or a personal mailbox"),
        ("", [], {"client@mail.example"}, "sender and recipients unreadable"),
        # ANY business party on the thread keeps it: carrier, provider, vendor (measured 2026-10-06)
        ("adjuster@carrier.example", ["atty@firm.example"], {"client@mail.example"}, None),
        ("atty@firm.example", ["adjuster@carrier.example", "para@firm.example"], {"client@mail.example"}, None),
        ("client@mail.example", ["atty@firm.example", "records@provider.example"], {"client@mail.example"}, None),
        ("", ["atty@firm.example", "claims@administrator.example"], {"client@mail.example"}, None),
    ],
)
def test_the_wall_is_decided_by_addresses(sender, recipients, client, want):
    assert pull.wall_reason(sender, recipients, FIRM, client, ("mail.example",)) == want


def test_x500_addresses_parse_to_nothing_so_the_wall_holds_them():
    assert pull.addresses("/O=EXCHANGELABS/OU=EXCHANGE ADMINISTRATIVE GROUP/CN=RECIPIENTS/CN=ATTY") == []


def test_walled_email_never_reaches_the_corpus_and_kept_mail_brings_its_body(tmp_path, monkeypatch):
    root = make_inputs(tmp_path / "in")
    firm = firm_mod.load(root)
    data = tmp_path / "data"
    (data / "raw").mkdir(parents=True)
    rows = []
    msgs = {
        "m1": {
            "subject": "my injuries",
            "sender": "client@mail.example",
            "recipients": ["atty@firm.example"],
            "date": "",
            "body": "Privileged words the client wrote to her lawyer.",
            "attachments": [("x.pdf", b"%PDF" * 9000)],
        },
        "m3": {
            "subject": "fwd",
            "sender": "",
            "recipients": ["atty@firm.example"],
            "date": "",
            "body": "Internal X500 note about strategy.",
            "attachments": [],
        },
        "m2": {
            "subject": "claim CLM-0001 limits",
            "sender": "adj@carrier.example",
            "recipients": ["atty@firm.example"],
            "date": "",
            "body": "The carrier's position.",
            "attachments": [("dec page.pdf", b"%PDF-dec" * 9000)],
        },
    }
    for mid in msgs:
        p = data / "raw" / f"{mid}.msg"
        p.write_bytes(b"x")
        rows.append(
            {
                "id": mid,
                "name": mid,
                "ext": ".msg",
                "folder": "/Correspondence",
                "ok": True,
                "path": str(p),
                "sha256": mid,
            }
        )
    monkeypatch.setattr(pull, "_open_msg", lambda path: msgs[path.stem])
    out = pull.split_emails(rows, data, firm, {"client@mail.example"})
    assert sorted(w["reason"] for w in out["walled"]) == [
        "between the firm and the client or a personal mailbox",
        "firm internal",
    ]
    assert [k["name"] for k in out["kept"]] == ["Email: claim CLM-0001 limits", "dec page.pdf"]
    # the privileged words are on disk only as held-out text, never in a kept document's TEXT
    kept_text = "".join(Path(k["text_path"]).read_text() for k in out["kept"] if k.get("text_path"))
    assert "carrier's position" in kept_text
    assert "Privileged words" not in kept_text and "strategy" not in kept_text
    assert "Privileged words" in Path(out["walled"][0]["text_path"]).read_text()


def test_a_printed_email_pdf_is_walled_by_its_header_addresses(tmp_path, pricing):
    printed = make_pdf(
        [
            "From: Alpha Example <client@mail.example>\nSent: Monday, March 2, 2026 9:14 AM\nTo: atty@firm.example\nSubject: my neck\nThe privileged body text."
        ]
    )
    seat = seat_with(standard_docs() + [("d8", "scan 3-1-26.pdf", printed, "f-med")])
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "delivered", v.reason
    walled = json.loads((tmp_path / "job" / "data" / "walled.json").read_text())
    assert [w["name"] for w in walled] == ["scan 3-1-26"]
    assert all("privileged body" not in json.dumps(c) for c in client.calls)


def test_memo_intake_and_notes_are_never_pulled(tmp_path, pricing):
    docs = standard_docs() + [
        ("d9", "Attorney memo re value.pdf", make_pdf(["Firm analysis."]), "f-med"),
        ("d10", "Intake sheet.pdf", make_pdf(["Intake."]), "f-med"),
    ]
    seat = seat_with(docs)
    _run(tmp_path, pricing, seat, ScriptedClient())
    assert all("d9" not in m and "d10" not in m for m in seat.mints)


def test_a_client_contact_read_error_holds_before_anything_is_pulled(tmp_path, pricing):
    seat = seat_with(standard_docs())
    seat.facts = {**seat.facts, "errors": ["client contact: HTTPError: 500"]}
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "held" and "client contact" in v.reason
    assert seat.mints == [] and client.calls == []


# ---- preflight ------------------------------------------------------------------------
def test_a_control_character_page_goes_to_transcription_before_any_summary(tmp_path):
    good = "Exampletown ER itemized statement with enough ordinary words to read as English text here."
    assert not preflight.junk_page(good)
    assert preflight.junk_page("\x01\x02\x03\x04\x05\x06 \x07\x08\x0b\x0c\x0e ok")
    assert preflight.junk_page(" ")


def test_the_estimate_counts_characters_pages_and_the_drafting_tail(tmp_path):
    firm = firm_mod.load(make_inputs(tmp_path / "in"))
    rows = [{"chars": 1_000_000, "pages": 10, "transcribe": [2, 3]}, {"chars": 0, "pages": 4, "transcribe": "all"}]
    e = preflight.estimate(rows, firm)
    assert e["transcription_pages"] == 6 and e["pages"] == 14
    assert e["usd"] == round((1_000_000 + 6 * 2500) / 1e6 * 6.0 + 6 * 0.01 + 4.0, 2)


def test_the_bill_reconciliation_flags_both_directions():
    medicals = [{"provider": "Exampletown Emergency", "charges": []}, {"provider": "Northfield Imaging", "charges": []}]
    rows = [{"name": "Exampletown Emergency bill"}, {"name": "Ridgeview Chiropractic ledger"}]
    rec = preflight.reconcile(medicals, rows)
    assert rec["tab_providers_without_a_bill_in_file"] == ["Northfield Imaging"]
    assert rec["bills_whose_provider_is_not_on_the_tab"] == ["Ridgeview Chiropractic ledger"]


# ---- the format check ------------------------------------------------------------------
def test_the_format_check_refuses_a_leftover_template_slot_and_counts_open_items(tmp_path):
    from medchron.demand import house

    ref = tmp_path / "ref.docx"
    from demand_testkit import make_reference

    make_reference(ref)
    kw = dict(signature=SIGNATURE, signer_title="Attorney at Law", author="Example")
    path, _ = house.render(DRAFT, ref, tmp_path / "ok", **kw)
    ok = format_check.check(path, signature=SIGNATURE, signer_title="Attorney at Law", footer_markers=["1119, 1152"])
    assert ok.ok, ok.fails
    path2, _ = house.render(
        DRAFT.replace("Demand is hereby made", "{{CLIENT_NAME}} Demand is hereby made"), ref, tmp_path / "bad", **kw
    )
    bad = format_check.check(path2, signature=SIGNATURE, signer_title="Attorney at Law", footer_markers=["1119, 1152"])
    assert any("curly-brace marker" in f for f in bad.fails)
    assert FOOTER  # the reference's footer carries the markers the check reads


def _scanned_pdf() -> bytes:
    """Page 1 has a text layer; page 2 is an image with no text (a scanned bill)."""
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Exampletown ER itemized statement page one with readable words.", fontsize=11)
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(72, 72, 300, 300), pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return data


def test_an_image_only_page_is_transcribed_and_merged_under_its_own_marker(tmp_path, pricing):
    docs = standard_docs() + [("d7", "Exampletown ER bill page 2.pdf", _scanned_pdf(), "f-med")]
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), client)
    assert v.outcome == "delivered"
    assert client.stages().count("VISION") == 1  # one page, not the file
    assert client.stages().index("VISION") < client.stages().index("DIGEST")
    text = (tmp_path / "job" / "data" / "text" / "d7.txt").read_text()
    assert "[p.2] (machine transcription)\ntranscribed page text" in text and "readable words" in text
    assert json.loads((tmp_path / "job" / "data" / "transcribed.json").read_text()) == ["Exampletown ER bill page 2"]


def test_an_open_item_in_the_expiry_line_is_counted_not_refused(tmp_path):
    """Live smoke 2026-10-06: a brief that asks for {{FILL}} on firm-supplied
    values put one in the banner's expiry date; the shape check still holds."""
    from demand_testkit import make_reference
    from medchron.demand import house

    ref = tmp_path / "ref.docx"
    make_reference(ref)
    md = DRAFT.replace("Friday, November 6, 2026", "{{FILL: weekday and date 30 days after the letter date}}")
    path, _ = house.render(md, ref, tmp_path / "o", signature=SIGNATURE, signer_title="Attorney at Law", author="E")
    res = format_check.check(path, signature=SIGNATURE, signer_title="Attorney at Law", footer_markers=["1119, 1152"])
    assert res.ok, res.fails
    md2 = DRAFT.replace("This Demand Expires at 5:00 P.M. Pacific Time on Friday, November 6, 2026", "Respond soon")
    path2, _ = house.render(md2, ref, tmp_path / "o2", signature=SIGNATURE, signer_title="Attorney at Law", author="E")
    assert any(
        "expiry line" in f
        for f in format_check.check(
            path2, signature=SIGNATURE, signer_title="Attorney at Law", footer_markers=["1119, 1152"]
        ).fails
    )


def test_an_unmapped_responsible_attorney_holds_before_anything_is_paid(tmp_path, pricing):
    seat = seat_with(standard_docs())
    seat.facts = {**seat.facts, "responsible_attorney": "Other Attorney"}
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "held" and "Other Attorney" in v.reason and "no signature is authored" in v.reason
    assert client.calls == [] and v.dollars == 0 and seat.sent == []


def test_the_signature_block_and_initials_come_from_the_attorney_map_not_the_model(tmp_path, pricing):
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient())  # DRAFT's front matter says "Example Attorney"
    assert v.outcome == "delivered", v.reason
    import docx

    doc = docx.Document(str(tmp_path / "job" / "data" / "out" / "demand" / "Demand.Alpha Example.docx"))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    i = texts.index("Cordially,")
    assert texts[i + 2 : i + 5] == ["Example A. Lawyer", "Attorney at Law", "EAL/dm"]
    assert "CERTIFIED MAIL" in texts[:6]


def test_a_filed_action_selects_litigation_and_an_unauthored_litigation_format_holds(tmp_path, pricing):
    docs = standard_docs() + [
        (
            "d12",
            "Pleading.pdf",
            make_pdf(["SUPERIOR COURT OF CALIFORNIA. COMPLAINT FOR DAMAGES. Case No. 24CV0001."]),
            "f-corr",
        )
    ]
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), client)
    assert v.outcome == "held" and "litigation demand" in v.reason and client.calls == []


def test_defense_counsel_of_record_selects_litigation():
    from medchron.demand import facts

    roles = [
        {
            "isOtherSide": True,
            "name": "Defendant",
            "relationships": [{"name": "Attorney", "contactId": "c-1"}, {"name": "Insurer", "contactId": "c-2"}],
        },
        {"isOtherSide": True, "name": "Defendant", "relationships": [{"name": "Attorney"}]},
    ]  # an empty slot is not counsel
    assert facts.defense_counsel(roles) == ["Attorney"]


def test_a_lawsuit_named_only_in_a_subject_is_unclear_and_holds(tmp_path, pricing):
    docs = standard_docs() + [("d13", "Lawsuit threat email.pdf", make_pdf(["We may consider options."]), "f-corr")]
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), client)
    assert v.outcome == "held" and "pre-suit or a litigation" in v.reason and client.calls == []


def test_the_litigation_variant_refuses_a_pre_suit_shaped_file(tmp_path):
    from demand_testkit import make_reference
    from medchron.demand import house

    ref = tmp_path / "ref.docx"
    make_reference(ref)
    path, _ = house.render(
        DRAFT,
        ref,
        tmp_path / "o",
        signature=SIGNATURE,
        signer_title="Attorney at Law",
        author="E",
        mail_line="CERTIFIED MAIL",
    )
    lit = format_check.Variant(
        headings=("Summary of Injuries",),
        re_required=(),
        re_one_of=(),
        mail_line="",
        banner=False,
        forbidden=("CERTIFIED MAIL",),
        re_vs=True,
    )
    fails = format_check.check(
        path, signature=SIGNATURE, signer_title="Attorney at Law", footer_markers=[], variant=lit
    ).fails
    assert any("vs." in f for f in fails) and any("another demand variant" in f for f in fails)


def test_a_gap_audit_only_job_needs_no_attorney_signature(tmp_path, pricing):
    seat = seat_with(standard_docs())
    seat.facts = {**seat.facts, "responsible_attorney": "Other Attorney"}
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(), deliverables=["gap_audit"])
    assert v.outcome == "delivered", v.reason
    assert [f["role"] for f in v.files] == ["gap_audit"]


def test_a_profile_names_its_own_headings_and_the_check_holds_them(tmp_path):
    from demand_testkit import make_reference
    from medchron.demand import house

    ref = tmp_path / "ref.docx"
    make_reference(ref)
    path, _ = house.render(DRAFT, ref, tmp_path / "o", signature=SIGNATURE, signer_title="Attorney at Law", author="E")
    kw = dict(signature=SIGNATURE, signer_title="Attorney at Law", footer_markers=["1119, 1152"])
    assert format_check.check(path, **kw, headings=["Summary of Injuries", "Liability", "Damages", "Demand"]).ok
    other = format_check.check(path, **kw, headings=["Facts", "Liability"])
    assert any("'Facts' missing" in f for f in other.fails)


def test_a_variant_naming_no_input_refuses_to_load(tmp_path):
    from demand_testkit import firm_data

    bad = firm_data()["variants"]["pre_suit"] | {"skeleton": "nope"}
    root = make_inputs(tmp_path / "in", variants={"pre_suit": bad})
    with pytest.raises(firm_mod.DemandConfigError, match="names no entry in inputs"):
        firm_mod.load(root)


class FakeVendor:
    def __init__(self, answers):
        self.answers, self.asked = answers, []

    def get_locations(self, term, zip_code=None):
        self.asked.append(term)
        a = self.answers[term]
        if isinstance(a, Exception):
            raise a
        return a


def test_each_missing_provider_is_resolved_against_the_vendor_directory_once_after_the_paid_stages(tmp_path, pricing):
    v = FakeVendor(
        {
            "Exampletown ER": [
                {
                    "id": "L-1",
                    "value": "Exampletown Regional",
                    "street": "1 Main",
                    "city": "Exampletown",
                    "state": "CA",
                    "postalcode": "90000",
                }
            ],
            "Northfield Imaging": [],
            "Ridgeview PT": [
                {"id": "L-2", "value": "Ridgeview PT North"},
                {"id": "L-3", "value": "Ridgeview PT South"},
            ],
        }
    )
    client = ScriptedClient()
    _r, verdict, _ = _run(tmp_path, pricing, seat_with(standard_docs()), client, vendor=v)
    assert verdict.outcome == "delivered", verdict.reason
    assert v.asked == ["Northfield Imaging", "Ridgeview PT", "Exampletown ER"]  # once each, in item order
    rows = {r["provider"]: r for r in json.loads((tmp_path / "job" / "data" / "vendor.json").read_text())["rows"]}
    assert rows["Exampletown ER"]["custodian_id"] == "L-1"
    assert rows["Northfield Imaging"]["status"] == "no vendor match"
    assert rows["Ridgeview PT"]["status"] == "several matches" and rows["Ridgeview PT"]["count"] == 2
    import docx

    gap = docx.Document(
        str(tmp_path / "job" / "data" / "out" / "demand" / f"Gap Audit - 100001 - {_r.date_stamp}.docx")
    )
    cells = " ".join(c.text for t in gap.tables for row in t.rows for c in row.cells)
    assert "L-1" in cells and "no vendor match" in cells and "several matches (2)" in cells


def test_vendor_lookups_are_capped_and_a_failed_lookup_is_named(tmp_path):
    from medchron.demand import vendor

    v = FakeVendor({"Exampletown ER": RuntimeError("timeout"), "Northfield Imaging": [], "Ridgeview PT": []})
    from demand_testkit import GAP

    out = vendor.run(tmp_path, GAP, lambda: v, 2, lambda m: None)
    assert [r["status"] for r in out["rows"]] == ["lookup failed", "no vendor match", "not looked up (over the cap)"]
    assert v.asked == ["Exampletown ER", "Northfield Imaging"]
    (tmp_path / "x").mkdir()
    none = vendor.run(tmp_path / "x", GAP, lambda: None, 2, lambda m: None)
    assert "not connected" in vendor.section(none)


# ---- the destination ------------------------------------------------------------------
def test_filing_anywhere_but_the_matter_or_an_authored_rehearsal_matter_is_refused(tmp_path, pricing):
    from demand_testkit import job_doc

    seat = seat_with(standard_docs())
    seat.numbers["0f0f0f0f-0000-4000-8000-0000000000ff"] = "200002"
    jd = tmp_path / "job"
    jd.mkdir()
    doc = job_doc()
    doc["file_to"] = {"id": "0f0f0f0f-0000-4000-8000-0000000000ff", "number": "200002"}
    (jd / "job.json").write_text(json.dumps(doc))
    make_inputs(tmp_path / "inputs")
    r = DemandRun(
        jd,
        inputs_dir=str(tmp_path / "inputs"),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=ScriptedClient(),
        log=lambda m: None,
        readback_pause=0.0,
        vendor_factory=lambda: None,
    )
    v = r.run()
    assert v.outcome == "held" and "neither the matter it reads nor an authored rehearsal matter" in v.reason
    assert seat.mints == [] and seat.sent == []


def test_a_matter_id_that_does_not_carry_its_number_is_refused(tmp_path, pricing):
    seat = seat_with(standard_docs())
    seat.numbers[MATTER_ID] = "999999"
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient())
    assert v.outcome == "held" and "does not carry matter number 100001" in v.reason and seat.mints == []


def test_the_rehearsal_matter_is_allowed_by_number_read_from_its_id(tmp_path, pricing):
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(), file_to=True)
    assert v.outcome == "delivered", v.reason
    seat2 = seat_with(standard_docs())
    seat2.numbers.pop(LIBRARY_ID)  # the library id no longer resolves to its number
    _r, v2, _ = _run(tmp_path / "b", pricing, seat2, ScriptedClient(), file_to=True)
    assert v2.outcome == "held" and "destination's id" in v2.reason


# ---- resume never sends a second copy ---------------------------------------------------
def test_a_resume_after_a_short_read_back_files_the_same_bytes_under_the_same_date(tmp_path, pricing):
    seat = seat_with(standard_docs())
    seat.lag = 99  # the vendor's list never shows the files: the read-back stays short
    days = iter([(2026, 10, 6), (2026, 10, 9)])

    def today():
        y, m, d = next(days)
        import time as _t

        return _t.struct_time((y, m, d, 9, 0, 0, 0, 0, -1))

    make_inputs(tmp_path / "inputs")
    jd = make_job(tmp_path / "job")
    kw = dict(
        inputs_dir=str(tmp_path / "inputs"),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=ScriptedClient(),
        log=lambda m: None,
        readback_pause=0.0,
        vendor_factory=lambda: None,
    )
    v1 = DemandRun(jd, today=today, **kw).run()
    assert v1.outcome == "failed" and "read-back is short" in v1.reason
    sent = list(seat.sent)
    manifest = (jd / "data" / "out" / "demand" / "upload_manifest.json").read_text()
    seat.lag = 0
    seat._pending = [(0, row) for _due, row in seat._pending]  # the index catches up
    v2 = DemandRun(jd, today=today, **kw).run()
    assert v2.outcome == "delivered", v2.reason
    assert seat.sent == sent  # nothing sent twice
    assert (jd / "data" / "out" / "demand" / "upload_manifest.json").read_text() == manifest
    assert all("10-06-26" in f["name"] or not f["name"].startswith("Gap") for f in v2.files)


# ---- premise: text, not only names; G1 is a carrier's own document ---------------------
def test_an_acceptance_found_only_in_a_documents_text_fails_g4(tmp_path, pricing):
    docs = standard_docs() + [
        (
            "d11",
            "letter 5-5-26.pdf",
            make_pdf(["Example Mutual. This is our timely acceptance of your demand."]),
            "f-corr",
        )
    ]
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), client)
    assert v.coverage_report is True and client.calls == []


def test_g1_does_not_pass_on_a_medical_bill_or_a_blank_claim_field():
    from medchron.demand import premise

    texts = [
        ({"name": "ER bill 1-15-26"}, "Patient account. Claim number 55512345. Insured: Alpha Example."),
        ({"name": "Retainer agreement"}, "CLAIM NUMBER:\nYOUR INSURED:\n"),
        ({"name": "Notes"}, "your insured was speeding; claim number: n/a"),
    ]
    assert premise.carrier_documents(texts, ["claim number"]) == []
    carrier = [({"name": "Carrier letter"}, "Our insured Beta Driver. Claim number CLM-0001.")]
    assert len(premise.carrier_documents(carrier, ["claim number"])) == 1


# ---- spend ----------------------------------------------------------------------------
def test_the_estimate_adds_condense_when_the_digest_outgrows_the_budget(tmp_path):
    firm = firm_mod.load(make_inputs(tmp_path / "in", levers={"digest_budget_chars": 1_000_000}))
    rows = [{"chars": 2_000_000, "pages": 10, "transcribe": []}]
    assert preflight.estimate(rows, firm)["usd"] == round(2.0 * 6.0 + 4.0 + 2.0 * 3.3, 2)


def test_a_resume_is_checked_against_what_is_left_to_spend(tmp_path, pricing):
    make_inputs(tmp_path / "inputs", budget={"per_job_cap_usd": 5.0})
    seat = seat_with(standard_docs())
    jd = make_job(tmp_path / "job")
    ledger = jd / "data" / "usage-ledger.jsonl"
    ledger.parent.mkdir(parents=True)
    # 1.50 USD already spent by an earlier attempt; the estimate is 4.00, so
    # 1.50 + (4.00 - 1.50) fits the 5.00 cap, where 1.50 + 4.00 would not.
    ledger.write_text(json.dumps({"stage": "digest", "model": "claude-sonnet-5", "in": 750_000, "out": 0}) + "\n")
    r = DemandRun(
        jd,
        inputs_dir=str(tmp_path / "inputs"),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=ScriptedClient(),
        log=lambda m: None,
        readback_pause=0.0,
        vendor_factory=lambda: None,
    )
    r._stage("facts", r._facts)
    r._stage("pull", r._pull)
    r._stage("preflight", lambda: preflight.run(r.data, r.firm, r._json("facts.json"), r.log))
    r._estimate()  # does not raise


# ---- the compose sentinel path and the upload's own reason --------------------------------
def test_a_model_written_coverage_report_is_gated_before_it_is_filed(tmp_path, pricing):
    from medchron.demand import draft as draft_mod

    report = draft_mod.COVERAGE_SENTINEL + "\n# Coverage\n\nThis report fully addresses every claim.\n"
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(draft=report))
    assert v.outcome == "held" and "coverage report" in v.reason and seat.sent == []


def test_an_upload_refusal_carries_the_upload_stages_own_reason(tmp_path, pricing):
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient())
    assert v.outcome == "delivered"
    seat2 = seat_with(standard_docs())
    seat2.folders.append({"id": "f-x", "name": seat.created[0]["name"], "parentId": None, "path": "/x"})
    _r, v2, _ = _run(tmp_path / "b", pricing, seat2, ScriptedClient())
    assert v2.outcome == "held" and "already exists on the matter" in v2.reason


def test_a_records_vendors_case_number_stamp_is_not_a_lawsuit(tmp_path, pricing):
    """Free replay on the 9/24 trial matters: the records vendor stamps the
    firm's matter number as "Case Number:" on every certificate, and a gate
    that read that as a court case called every pre-suit file litigation."""
    cert = make_pdf(["Certificate of records. Case Number: 100001 Requested Date Range: 1/1/2021 - present."])
    docs = standard_docs() + [("d14", "ER Certificate for Medical.pdf", cert, "f-med")]
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), ScriptedClient())
    assert v.outcome == "delivered", v.reason
    assert json.loads((tmp_path / "job" / "data" / "premise.json").read_text())["variant"]["variant"] == "pre_suit"


def test_no_demand_left_this_cycle_fails_before_anything_is_paid(tmp_path, pricing):
    from demand_testkit import job_doc

    seat = seat_with(standard_docs())
    jd = tmp_path / "job"
    jd.mkdir()
    (jd / "job.json").write_text(json.dumps({**job_doc(), "allowance_remaining": 0}))
    make_inputs(tmp_path / "inputs")
    client = ScriptedClient()
    v = DemandRun(
        jd,
        inputs_dir=str(tmp_path / "inputs"),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=client,
        log=lambda m: None,
        readback_pause=0.0,
        vendor_factory=lambda: None,
    ).run()
    assert v.outcome == "failed" and v.stage == "demand_allowance_per_cycle" and client.calls == []


@pytest.mark.parametrize("estimate,outcome", [(40.0, "delivered"), (140.0, "failed")])
def test_the_runaway_guard_holds_only_above_the_cap(tmp_path, pricing, estimate, outcome):
    """Captain, 2026-10-06: delivery must not stop for a few dollars; the
    per-job cap is a runaway guard (100), so a 40-dollar estimate runs and a
    140-dollar one stops before anything is paid."""
    make_inputs(
        tmp_path / "inputs",
        budget={"per_job_cap_usd": 100.0, "monthly_budget_usd": 750.0, "usd_drafting_fixed": estimate},
    )
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat_with(standard_docs()), client)
    assert v.outcome == outcome, v.reason
    assert (client.calls == []) == (outcome == "failed")
    if outcome == "delivered":
        assert v.dollars > 0  # the job's actual spend rides the verdict, the ledger row and the wake


# ---- re-review 2026-10-06 ---------------------------------------------------------------
def test_an_audit_reply_with_no_tally_line_is_retried_then_held(tmp_path, pricing):
    client = ScriptedClient(audit="- claim | SUPPORTED | cite\n(the auditor stopped before its tally)")
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "held" and "did not complete twice" in v.reason
    assert seat.sent == []


def test_a_truncated_audit_reply_is_never_a_pass(tmp_path, pricing):
    class Truncating(ScriptedClient):
        def _msg(self, params):
            m = super()._msg(params)
            if "AUDIT-PROMPT" in json.dumps(params.get("system")):
                m.stop_reason = "max_tokens"
            return m

    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, Truncating())
    assert v.outcome == "held" and "did not complete" in v.reason and seat.sent == []


def test_finding_rows_count_even_when_the_tally_line_under_counts():
    from medchron.demand import draft

    body = "- a | INVENTED | none\n- b | INVENTED | none\nSUPPORTED=4 DRIFTS=0 INVENTED=0 ARITHMETIC=0"
    assert draft.tally(body)["INVENTED"] == 2 and draft.tally(body)["SUPPORTED"] == 4


def test_a_fax_cover_sheet_is_kept_not_walled(tmp_path):
    firm = firm_mod.load(make_inputs(tmp_path / "in"))
    cover = "FAX COVER SHEET\nTO: Exampletown ER Records\nFROM: Example & Example, LLP\nRE: Alpha Example\nPages: 4"
    assert pull.printed_email_wall(cover, firm, set()) is None
    no_addr = "From: Alpha Example\nSent: Monday\nTo: Example Lawyer\nSubject: hi\nbody"
    assert pull.printed_email_wall(no_addr, firm, set()) is None  # Outlook shape but no address: not walled
    real = "From: Alpha <client@mail.example>\nSent: Monday\nTo: atty@firm.example\nSubject: hi\nbody"
    assert pull.printed_email_wall(real, firm, {"client@mail.example"})


@pytest.mark.parametrize(
    "text",
    [
        "We cannot deny coverage at this time.",
        "We are unable to deny coverage until the investigation ends.",
        "The carrier will not deny coverage for this loss.",
    ],
)
def test_a_negated_denial_is_not_a_denial(text):
    from medchron.demand import premise

    assert premise._phrase_hits([({"name": "letter"}, text)], ["deny coverage"], negatable=True) == []
    assert premise._phrase_hits([({"name": "letter"}, "We deny coverage.")], ["deny coverage"], negatable=True)


def test_the_letter_may_name_any_of_the_matters_clients_and_the_file_is_named_for_that_one():
    from medchron.demand import crosscheck

    facts = {
        "client_name": "Alpha Example",
        "client_names": ["Alpha Example", "Beta Second"],
        "date_of_loss": "01/15/2026",
    }
    second = DRAFT.replace("client: Alpha Example", "client: Beta Second")
    assert crosscheck.check(second, facts)["mismatches"] == []
    assert crosscheck.matched_client(second, facts) == "Beta Second"
    stranger = DRAFT.replace("client: Alpha Example", "client: Gamma Person")
    assert crosscheck.check(stranger, facts)["mismatches"]


def test_every_walled_document_is_listed_by_name_in_the_gap_audit_and_the_notes(tmp_path, pricing):
    printed = make_pdf(["From: Alpha <client@mail.example>\nSent: Monday\nTo: atty@firm.example\nSubject: x\nbody"])
    seat = seat_with(standard_docs() + [("d8", "scan 3-1-26.pdf", printed, "f-med")])
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient())
    assert v.outcome == "delivered", v.reason
    import docx

    out = tmp_path / "job" / "data" / "out" / "demand"
    gap = docx.Document(str(next(out.glob("Gap Audit*.docx"))))
    gtext = "\n".join(p.text for p in gap.paragraphs) + " ".join(
        c.text for t in gap.tables for r in t.rows for c in r.cells
    )
    assert "privilege wall: 1 document(s)" in gtext and "scan 3-1-26" in gtext
    notes = docx.Document(str(out / "Demand.Alpha Example - attorney notes.docx"))
    assert any("scan 3-1-26" in p.text and "privilege wall" in p.text for p in notes.paragraphs)


def test_g1_skips_a_medical_record_by_its_content_not_only_its_name():
    from medchron.demand import premise

    bill_by_content = [
        (
            {"name": "scan 3-1-26"},
            "Patient account 12345. Date of service 01/15/2026. "
            "Amount due $1,200. Insured: Alpha Example. Claim number CLM-0009.",
        )
    ]
    assert premise.carrier_documents(bill_by_content, ["claim number"]) == []


def test_a_header_block_without_a_sent_line_is_not_a_printed_email(tmp_path):
    firm = firm_mod.load(make_inputs(tmp_path / "in"))
    no_sent = "From: Alpha <client@mail.example>\nTo: atty@firm.example\nCc: para@firm.example\nbody"
    assert pull.printed_email_wall(no_sent, firm, {"client@mail.example"}) is None
    with_cc = (
        "From: Alpha <client@mail.example>\nSent: Mon\nTo: atty@firm.example\nCc: para@firm.example\nSubject: x\nb"
    )
    assert pull.printed_email_wall(with_cc, firm, {"client@mail.example"})


PREM = {
    "settled_phrases": ["accept the policy limits", "release in full of all"],
    "_firm_signature": "EXAMPLE & EXAMPLE, LLP",
    "_firm_domains": ["firm.example"],
}


def test_the_firms_own_prior_demand_is_not_an_acceptance():
    from medchron.demand import premise

    ours = (
        {"name": "Demand 3-1-26", "kind": "file"},
        "We ask that you accept the policy limits demand.\nCordially,\nEXAMPLE & EXAMPLE, LLP\nExample A. Lawyer",
    )
    ours_email = (
        {"name": "Email: demand", "kind": "email_body"},
        "Email: demand\nFrom: atty@firm.example\nTo: adj@carrier.example\n\nPlease accept the policy limits.",
    )
    theirs = (
        {"name": "Letter 5-5-26", "kind": "file"},
        "Example & Example, LLP\n1 Main\nRe: Alpha. We hereby accept the policy limits demand.\nSincerely,\nPat Adjuster",
    )
    assert premise.firm_authored(ours[0], ours[1], PREM) and premise.firm_authored(ours_email[0], ours_email[1], PREM)
    assert not premise.firm_authored(theirs[0], theirs[1], PREM)  # addressed to the firm, signed by the carrier
    texts = [ours, ours_email, theirs]
    hits = premise._phrase_hits(
        [(r, t) for r, t in texts if not premise.firm_authored(r, t, PREM)], PREM["settled_phrases"], negatable=True
    )
    assert [h["document"] for h in hits] == ["Letter 5-5-26"]


def test_a_carrier_letter_about_dates_of_service_and_charges_is_still_a_carrier_letter():
    from medchron.demand import premise

    letter = [
        (
            {"name": "Letter from carrier"},
            "Example Mutual Insurance Company, claims department. Our insured "
            "Beta Driver. Claim number CLM-0001. We have reviewed the dates of service and charges you sent.",
        )
    ]
    assert len(premise.carrier_documents(letter, ["claim number"])) == 1


@pytest.mark.parametrize(
    "text",
    [
        "We decline to deny coverage at this time.",
        "We are not able to deny coverage yet.",
        "We won't deny coverage before the investigation ends.",
    ],
)
def test_the_remaining_negations(text):
    from medchron.demand import premise

    assert premise._phrase_hits([({"name": "letter"}, text)], ["deny coverage"], negatable=True) == []


def test_a_financial_responsibility_form_is_not_an_acceptance(tmp_path):
    """The engagements config narrows "hereby accept" to accepting a demand."""

    phrases = ["hereby accept your demand", "hereby accepts your demand", "hereby accept your time-limited demand"]
    from medchron.demand import premise

    form = [({"name": "Lien"}, "I hereby accept financial responsibility for all charges.")]
    acceptance = [({"name": "Letter"}, "Example Mutual hereby accepts your demand dated May 1.")]
    assert premise._phrase_hits(form, phrases, negatable=True) == []
    assert premise._phrase_hits(acceptance, phrases, negatable=True)


# ---- live practice run 2026-10-06: a false G2 on our own draft ------------------
def test_g2_never_reads_the_firms_or_the_operators_own_draft_demand(tmp_path, pricing):
    own = make_pdf(
        [
            "Your claim number CLM-0001. We note your letter: coverage is denied pending review. "
            "Demand is hereby made for the full policy limits."
        ]
    )
    docs = standard_docs() + [("d20", "100001 Demand DRAFT 2026-09-24.pdf", own, "f-corr")]
    client = ScriptedClient()
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), client)
    assert v.outcome == "delivered" and v.coverage_report is False, v.reason
    assert client.calls  # the job went on to draft


def test_a_carrier_refusing_to_disclose_limits_is_not_a_denial(tmp_path, pricing):
    letter = make_pdf(
        [
            "Example Mutual. Claim number CLM-0001. Policy limits are protected private financial "
            "information; coverage is denied disclosure until we request our insured's permission."
        ]
    )
    docs = standard_docs() + [("d21", "Carrier letter limits.pdf", letter, "f-corr")]
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), ScriptedClient())
    assert v.coverage_report is False, v.reason
    real = make_pdf(
        [
            "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. Coverage is denied for this loss."
        ]
    )
    docs2 = standard_docs() + [("d22", "Carrier letter 2.pdf", real, "f-corr")]
    _r, v2, _ = _run(tmp_path / "b", pricing, seat_with(docs2), ScriptedClient())
    assert v2.coverage_report is True
    report = (tmp_path / "b" / "job" / "data" / "coverage-report.md").read_text()
    assert "Carrier letter 2 (rule: a denial phrase)" in report


def test_a_reservation_of_rights_is_a_flag_not_a_fail(tmp_path):
    from medchron.demand import premise

    prem = {
        "denial_phrases": ["coverage is denied"],
        "fail_on": [],
        "settled_phrases": ["timely acceptance"],
        "carrier_phrases": ["claim number"],
        "litigation_phrases": ["complaint for damages"],
        "_firm_signature": "X",
        "_firm_domains": [],
    }
    t = tmp_path / "a.txt"
    t.write_text(
        "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. We issue this reservation of rights."
    )
    out = premise.decide([{"name": "ROR letter", "text_path": str(t)}], {"premise_hits": []}, {"insurer": "X"}, prem)
    assert out["gates"][1]["passed"] is True
    assert any("reservation of rights" in f for f in out["premise_facts"])


# ---- practice job 2 (2026-10-06): one near-quote held a whole paid delivery ---------------
NEAR = (
    'The record states the patient was "seen 01/15/2026 for neck pain following a collision" (ER record 1-15-26, p. 1).'
)
FAR = (
    'The record states "the patient did not immediately initiate chiropractic care for weeks" '
    "(ER record 1-15-26, p. 1)."
)


def _gap_with(sentence: str) -> str:
    from demand_testkit import GAP

    row = f"| B. Provider-by-provider completeness | Exampletown ER | {sentence} | ER record 1-15-26, p. 1 | "
    return GAP.replace("END OF ITEM TABLE", row + "Referenced in record | records | Housekeeping |\nEND OF ITEM TABLE")


def test_a_near_quote_in_the_gap_audit_is_normalized_to_the_source_and_filed(tmp_path, pricing):
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(gap=_gap_with(NEAR)))
    assert v.outcome == "delivered", v.reason
    gap = (tmp_path / "job" / "data" / "gap-audit.md").read_text()
    assert '"seen 01/15/2026 for neck pain after a collision"' in gap  # the record's own words
    g = json.loads((tmp_path / "job" / "data" / "gate-gap-audit.json").read_text())
    assert g["passed"] and g["repairs"] and g["repairs"][0].startswith("quote normalized to source")
    import docx

    notes = docx.Document(
        str(tmp_path / "job" / "data" / "out" / "demand" / "Demand.Alpha Example - attorney notes.docx")
    )
    assert any("quote normalized to source" in p.text for p in notes.paragraphs)


def test_a_quote_with_no_close_region_becomes_a_paraphrase_with_its_cite(tmp_path, pricing):
    seat = seat_with(standard_docs())
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(gap=_gap_with(FAR)))
    assert v.outcome == "delivered", v.reason
    gap = (tmp_path / "job" / "data" / "gap-audit.md").read_text()
    assert '"the patient did not' not in gap
    assert "the patient did not immediately initiate chiropractic care for weeks (ER record 1-15-26, p. 1)" in gap
    assert json.loads((tmp_path / "job" / "data" / "gate-gap-audit.json").read_text())["repairs"][0].startswith(
        "quote converted to paraphrase"
    )


def test_a_near_quote_in_the_letter_is_repaired_too(tmp_path, pricing):
    letter = DRAFT.replace(
        "The emergency department diagnosed a cervical strain",
        'The emergency department record says "seen 01/15/2026 for neck pain following a collision" and diagnosed a cervical strain',
    )
    seat = seat_with(standard_docs())
    # the auditor passes it; only the gate's contiguity check would refuse it
    _r, v, _ = _run(tmp_path, pricing, seat, ScriptedClient(draft=letter))
    assert v.outcome in ("delivered", "held"), v.reason
    if v.outcome == "held":  # the free mechanical check catches it first: also correct, nothing filed
        assert "quotation" in v.reason and seat.sent == []


def test_gate_quote_failures_left_after_repair_are_failed_and_resume_repays_nothing(tmp_path, pricing, monkeypatch):
    from medchron.demand import quotefix

    real = quotefix.repair
    monkeypatch.setattr(quotefix, "repair", lambda md, refusals, sources: (md, []))  # a repair that fixes nothing
    seat = seat_with(standard_docs())
    client = ScriptedClient(gap=_gap_with(NEAR))
    _r, v, _ = _run(tmp_path, pricing, seat, client)
    assert v.outcome == "failed" and v.stage == "gate" and seat.sent == []
    paid = len(client.calls)
    monkeypatch.setattr(quotefix, "repair", real)
    _r, v2, _ = _run(tmp_path, pricing, seat, client)
    assert v2.outcome == "delivered", v2.reason
    assert len(client.calls) == paid  # resume re-ran repair, gate, render and file only


# ---- 2026-10-06: a workers' comp policy's boilerplate stopped G2 ----------------------------
def test_g2_never_reads_a_policy_copy_or_its_boilerplate(tmp_path, pricing):
    policy = make_pdf(
        [
            "WORKERS COMPENSATION AND EMPLOYERS LIABILITY POLICY\nPolicy period 2024-2025. Named insured: "
            "Example Employer. You will not deny coverage under this policy and will reimburse us for any "
            "increase in indemnity. Claim number CLM-0001. Our insured."
        ]
    )
    docs = standard_docs() + [("d30", "Policy copy.pdf", policy, "f-corr")]
    _r, v, _ = _run(tmp_path, pricing, seat_with(docs), ScriptedClient())
    assert v.coverage_report is False, v.reason


def test_g2_reads_only_letters_shaped_carrier_correspondence():
    from medchron.demand import premise

    prem = {"carrier_phrases": ["claim number"], "_firm_signature": "X", "_firm_domains": ["firm.example"]}
    letter = (
        {"name": "Letter from Example Mutual"},
        "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. Coverage is denied.",
    )
    booklet = (
        {"name": "Benefit booklet"},
        "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. Coverage is denied.",
    )
    unshaped = ({"name": "scan 4"}, "Example Mutual. Our insured. Claim number CLM-0001. Coverage is denied.")
    provider = (
        {"name": "ER record"},
        "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. Coverage is denied.",
    )
    policy_text = (
        {"name": "doc 7"},
        "Example Mutual Insurance Company\nRe: Alpha Example, claim number CLM-0001\nDear Counsel:\nOur insured Beta Driver. This policy. Named insured. Policy period. We will pay. Coverage is denied.",
    )
    got = [r["name"] for r, _ in premise.carrier_letters([letter, booklet, unshaped, provider, policy_text], prem)]
    assert got == ["Letter from Example Mutual"]


def test_the_wake_carries_no_email_id():
    import inspect

    from medchron import demand_lane

    assert "Request ref" not in inspect.getsource(demand_lane.DemandLane._compose_wake)
