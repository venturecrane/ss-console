"""The delivered chronology: found by the firm's own name for it, rebuilt to
markdown from the delivered document, and PROVEN to be it with the seat's
renderer before an update may carry a single entry (medchron/update/)."""

from __future__ import annotations

from pathlib import Path

import pytest

import update_testkit as K
from medchron.update import body as B, docx_md, locate as L


def _render0(text: str, dst: Path) -> int:
    K.render(text, dst)
    return 0


def test_the_rebuilt_markdown_renders_back_to_the_delivered_document(tmp_path: Path) -> None:
    delivered = K.render(K.base_md(), tmp_path / "delivered.docx")
    md = docx_md.docx_to_md(delivered, K.HEADINGS)
    r = docx_md.prove(md, delivered, _render0, tmp_path / "prove")
    assert r["rows"] > 20 and r["diffs"] == []
    # Every delivered entry comes back byte-identical, provider tag included.
    _, ents = B.entries_of(md)
    assert ents == [K.ENTRY_PRE, K.ENTRY_POST]
    assert "* One record set was reviewed in full." in md


def test_the_proof_fails_on_a_document_the_markdown_does_not_reproduce(tmp_path: Path) -> None:
    """The falsifier: a delivered document edited after delivery (one word of
    one entry) is not reproduced by any markdown rebuilt before the edit."""
    import docx

    delivered = K.render(K.base_md(), tmp_path / "delivered.docx")
    md = docx_md.docx_to_md(delivered, K.HEADINGS)
    d = docx.Document(str(delivered))
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if "6 of 10" in run.text:
                            run.text = run.text.replace("6 of 10", "7 of 10")
    edited = tmp_path / "edited.docx"
    d.save(str(edited))
    r = docx_md.prove(md, edited, _render0, tmp_path / "prove")
    assert r["diffs"], "a proof that cannot fail measured nothing"


def test_summary_lines_are_the_entries_the_reader_sees(tmp_path: Path) -> None:
    delivered = K.render(K.base_md(), tmp_path / "delivered.docx")
    lines = docx_md.summary_lines(delivered)
    assert "Cervical strain after the collision, pain rated 6 of 10 at intake. (Exhibit 1 - p. 2)" in lines
    assert "Prior Medical History" in lines and "4. Medical Summary" not in lines


def test_a_chronology_another_renderer_made_is_proven_by_its_words(tmp_path: Path) -> None:
    """The laptop-rendered chronology: the seat cannot reproduce its styles
    or bold runs, so the exact proof fails (that is the old hold); its words
    re-render paragraph for paragraph, so the words-only proof passes and the
    entries come back as the delivered text."""
    delivered = K.foreign_docx(tmp_path / "foreign.docx", note=K.NOTE)
    exact = docx_md.docx_to_md(delivered, K.HEADINGS)
    assert docx_md.prove(exact, delivered, _render0, tmp_path / "exact")["diffs"], "the fixture must not be seat-shaped"
    md = docx_md.docx_to_md(delivered, K.HEADINGS, plain=True)
    r = docx_md.prove_text(md, delivered, _render0, tmp_path / "text")
    assert r["rows"] > 20 and r["diffs"] == [], r["diffs"][:5]
    pre, ents = B.entries_of(md)
    assert ents == [K.ENTRY_PRE, K.ENTRY_POST]
    assert K.NOTE in pre and "**" not in md
    assert B.delivered_numbers(md) == {1}


def test_the_words_only_proof_still_fails_on_a_paragraph_the_markdown_cannot_place(tmp_path: Path) -> None:
    delivered = K.foreign_docx(tmp_path / "typed.docx", typed="A note the firm typed under the entries.")
    md = docx_md.docx_to_md(delivered, K.HEADINGS, plain=True)
    assert docx_md.prove_text(md, delivered, _render0, tmp_path / "text")["diffs"]


TEMPLATE = "{CLIENT} - Medical Chronology {MM-DD-YY}.docx"
FOLDERS = [
    {
        "id": "fold-chron",
        "name": "MEDICAL CHRONOLOGY - Alpha Example 09-01-26",
        "parentId": None,
        "path": "/MEDICAL CHRONOLOGY - Alpha Example 09-01-26",
    },
    {"id": "fold-old", "name": "OLD", "parentId": None, "path": "/OLD"},
]


def _row(fid: str, name: str, folder: str, ext: str = ".docx", deleted: bool = False) -> dict:
    return {"id": fid, "name": name, "ext": ext, "folderId": folder, "size": 10, "deleted": deleted}


def test_the_newest_chronology_by_its_name_is_the_one_updated() -> None:
    docs = [
        _row("c1", "Alpha Example - Medical Chronology 03-02-26", "fold-old"),
        _row("c2", "Alpha Example - Medical Chronology 09-01-26", "fold-chron"),
        _row("e1", "Exhibit 1 - Example Clinic - 12-01-2025 (Medical Records)", "fold-chron", ".pdf"),
        _row("w1", "Alpha Example - Medical Billing Worksheet 09-01-26", "fold-chron"),
        _row("x9", "Alpha Example - Medical Chronology 10-01-26", "fold-old", deleted=True),
    ]
    got = L.locate(docs, FOLDERS, template=TEMPLATE, client="Alpha Example", file_id=None)
    assert got["chronology"]["id"] == "c2" and got["folder_id"] == "fold-chron"
    assert [e["exhibit"] for e in got["exhibits"]] == [1]
    assert got["others"] == ["Alpha Example - Medical Billing Worksheet 09-01-26.docx"]


def test_a_named_chronology_must_be_this_clients_chronology() -> None:
    docs = [_row("c2", "Alpha Example - Medical Chronology 09-01-26", "fold-chron"), _row("z", "Notes", "fold-chron")]
    assert L.locate(docs, FOLDERS, template=TEMPLATE, client="Alpha Example", file_id="c2")["chronology"]["id"] == "c2"
    with pytest.raises(L.LocateHold):
        L.locate(docs, FOLDERS, template=TEMPLATE, client="Alpha Example", file_id="z")
    with pytest.raises(L.LocateHold):
        L.locate(docs, FOLDERS, template=TEMPLATE, client="Beta Example", file_id=None)


def test_a_versioned_same_day_chronology_is_the_newest() -> None:
    """A same-day update files "<name> (2).docx" rather than over a file of
    the same name; the next update must find that one."""
    docs = [
        _row("c2", "Alpha Example - Medical Chronology 09-01-26", "fold-chron"),
        _row("c3", "Alpha Example - Medical Chronology 09-01-26 (2)", "fold-chron"),
        _row("c1", "Alpha Example - Medical Chronology 03-02-26 (3)", "fold-old"),
    ]
    got = L.locate(docs, FOLDERS, template=TEMPLATE, client="Alpha Example", file_id=None)
    assert got["chronology"]["id"] == "c3"
    assert (
        L.locate(docs[:1], FOLDERS, template=TEMPLATE, client="Alpha Example", file_id=None)["chronology"]["id"] == "c2"
    )


def test_two_chronologies_of_the_newest_date_hold() -> None:
    docs = [
        _row("c2", "Alpha Example - Medical Chronology 09-01-26", "fold-chron"),
        _row("c3", "Alpha Example - Medical Chronology 09-01-26", "fold-old"),
    ]
    with pytest.raises(L.LocateHold, match="newest date"):
        L.locate(docs, FOLDERS, template=TEMPLATE, client="Alpha Example", file_id=None)
