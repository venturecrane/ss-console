"""The 12-section discovery breakdown (``litigation/discovery12/``), on a fake
seat with a scripted model and a fake Batch API. Synthetic data only."""

from __future__ import annotations

import datetime as dt
import io
import json
import re
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from litigation_testkit import JOB, M1, LitSeat, first_text, make_inputs, make_job, matter_docs
from medchron_testkit import make_pdf

from medchron import lane_proof
from medchron.litigation import job as job_mod, manifest, printout
from medchron.litigation.discovery12 import batch, book, sections
from medchron.litigation.discovery12.proof import Discovery12Proof
from medchron.litigation.discovery12.run import PRINTOUT_ROLE, ROLE, Discovery12Run, plan_for

TODAY = dt.date(2026, 10, 29)
FROGS = "Form Interrogatories, Set One, served by mail on March 5, 2026 on defendant Delta."
RESP = "Responses to Form Interrogatories, Set One. Served April 4, 2026."


def disc_docs() -> list[tuple[str, str, bytes, str]]:
    return [
        *matter_docs(),
        ("f-i1", "Form Interrogatories Set One.pdf", make_pdf([FROGS]), "2026-03-05T10:00:00"),
        ("f-i2", "Responses to Form Interrogatories.pdf", make_pdf([RESP]), "2026-04-04T10:00:00"),
    ]


def frogs_result(n: dict[str, int], served: str = "2026-03-05") -> dict[str, Any]:
    entry = {
        "item": "Form Interrogatories, Set One",
        "direction": "propounded by the firm",
        "party": "Delta",
        "served": {"date": served, "doc": n.get("Form Interrogatories Set One.pdf")},
        "due": {"date": None},
        "response": {"value": "yes", "date": "2026-04-04", "doc": n.get("Responses to Form Interrogatories.pdf")},
        "extension": {"date": None, "detail": ""},
        "notes": "",
    }
    return {"sections": {"frogs": [entry]}, "flags": []}


def _usage() -> Any:
    return SimpleNamespace(
        input_tokens=1000, output_tokens=200, cache_read_input_tokens=0, cache_creation_input_tokens=0
    )


class Scripted12:
    """Single-shot answers: the read records ``fn(doc numbers)``; verify and
    audit record ``verdicts`` / ``audit``. Serves the interactive path and,
    through ``batches``, a fake Batch API."""

    def __init__(self, fn: Any = frogs_result, verdicts: list | None = None, audit: list | None = None) -> None:
        self.fn, self.v12, self.a12 = fn, verdicts or [], audit or []
        self.calls: list[dict[str, Any]] = []
        self.submitted: list[list[dict[str, Any]]] = []
        self._results: dict[str, list[Any]] = {}
        self.messages = SimpleNamespace(
            create=self._msg, stream=self._stream, batches=SimpleNamespace(
                create=self._b_create, retrieve=self._b_retrieve, results=self._b_results)
        )  # fmt: skip

    def _msg(self, **params: Any) -> Any:
        self.calls.append(params)
        system = json.dumps(params.get("system"))
        if "YOUR PASS: determine all twelve" in system:
            nums = {name.strip(): int(k) for k, name in re.findall(r"\[doc (\d+)\] ([^|]+?) \|", first_text(params))}
            final = {"type": "tool_use", "id": "rec", "name": "record_discovery", "input": self.fn(nums)}
        else:
            vs = self.a12 if "last check before" in system else self.v12
            final = {"type": "tool_use", "id": "rec", "name": "record_verdicts", "input": {"verdicts": vs}}
        return SimpleNamespace(content=[final], stop_reason="tool_use", usage=_usage())

    def _stream(self, **params: Any) -> Any:
        from demand_testkit import _Stream

        return _Stream(self._msg(**params))

    def _b_create(self, requests: list[dict[str, Any]]) -> Any:
        bid = f"msgbatch_{len(self.submitted)}"
        self.submitted.append(requests)
        self._results[bid] = [
            SimpleNamespace(
                custom_id=r["custom_id"], result=SimpleNamespace(type="succeeded", message=self._msg(**r["params"]))
            )
            for r in requests
        ]
        return SimpleNamespace(id=bid)

    def _b_retrieve(self, bid: str) -> Any:
        return SimpleNamespace(processing_status="ended")

    def _b_results(self, bid: str) -> Any:
        return iter(self._results[bid])


@pytest.fixture(autouse=True)
def _pricing(tmp_path: Path, pricing_path: Path) -> None:
    (tmp_path / "pricing.json").write_text(pricing_path.read_text())


def _kw(tmp_path: Path, seat: LitSeat, client: Any) -> dict[str, Any]:
    return {
        "inputs_dir": str(make_inputs(tmp_path / "inputs")),
        "pricing": str(tmp_path / "pricing.json"),
        "state_dir": tmp_path / "state",
        "seat_factory": lambda: seat,
        "client": client,
        "log": lambda m: None,
        "today": lambda: TODAY,
        "readback_pause": 0.0,
    }


def _run(tmp_path: Path, seat: LitSeat, client: Any, cents: int = 0, **kw: Any) -> Discovery12Run:
    return Discovery12Run(
        make_job(tmp_path / "job", cents=cents, report="discovery12"), **_kw(tmp_path, seat, client), **kw
    )


def _discovery_rows(xlsx: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx), read_only=True)
    try:
        return [list(r) for r in wb[book.DISCOVERY].iter_rows(values_only=True)][1:]
    finally:
        wb.close()


def _ledger(r: Discovery12Run) -> list[dict[str, Any]]:
    p = r.data / "usage-ledger.jsonl"
    return [json.loads(x) for x in p.read_text().splitlines()] if p.is_file() else []


def test_a_batched_run_files_the_workbook_and_the_printout_billed_at_the_batch_rate(tmp_path):
    seat = LitSeat({M1: disc_docs()})
    client = Scripted12()
    r = _run(tmp_path, seat, client)
    v = r.run()
    assert v.verdict == "delivered", v.reason
    assert [f["role"] for f in v.files] == [ROLE, PRINTOUT_ROLE]
    assert len(client.submitted) == 3  # read, verify, audit: one batch each, one request per chunk
    rows = _discovery_rows(r._out().read_bytes())
    frogs = next(x for x in rows if x[3] == sections.LABEL["frogs"])
    assert frogs[7] == "2026-03-05" and "Form Interrogatories Set One.pdf" in frogs[8]
    assert sum(1 for x in rows if x[4] == sections.NONE_IN_FILE) == 11
    text = "\n".join(printout.text_of(r._pdf()))
    flat = "".join(text.split())
    assert "Discovery(12sections)" in flat and "FormInterrogatoriesSetOne.pdf" in flat
    assert text.count(sections.NONE_IN_FILE) == 11 + 1  # the eleven empty sections, and the cover's legend
    ledger = _ledger(r)
    d12 = [x for x in ledger if x["stage"] in batch.STAGES.values()]
    assert len(d12) == 3 and all(x["batch"] for x in d12)
    one = d12[0]
    assert r.pricing.price_row(one) == pytest.approx(0.5 * r.pricing.price_row({**one, "batch": False}))
    assert manifest.load_prior(tmp_path / "state", M1) is None  # the lane's state untouched


def test_the_interactive_path_gives_the_same_result(tmp_path):
    seat = LitSeat({M1: disc_docs()})
    client = Scripted12()
    r = _run(tmp_path, seat, client, use_batch=False)
    assert r.run().verdict == "delivered" and client.submitted == [] and len(client.calls) == 3
    assert all(not x["batch"] for x in _ledger(r))


def test_a_report_whose_dates_are_mostly_unconfirmed_holds(tmp_path):
    seat = LitSeat({M1: disc_docs()})
    r = _run(tmp_path, seat, Scripted12(fn=lambda n: frogs_result(n, served="2026-03-09")))
    v = r.run()
    assert v.verdict == "held" and v.reason.startswith("gate_refused:") and "of dates unconfirmed" in v.reason
    assert seat.sent == []
    g = json.loads((r.data / "gates.json").read_text())
    assert g["unconfirmed_count"] == 1 and "date not in the cited document" in g["unconfirmed"]


def test_one_unconfirmed_date_is_printed_as_unconfirmed_and_never_as_fact(tmp_path, monkeypatch):
    """2026-10-29: one value no longer holds every case; it prints as unconfirmed, with its paper."""
    monkeypatch.setattr(sections, "MAX_UNCONFIRMED_SHARE", 1.0)
    seat = LitSeat({M1: disc_docs()})
    r = _run(tmp_path, seat, Scripted12(fn=lambda n: frogs_result(n, served="2026-03-09")))
    v = r.run()
    assert v.verdict == "delivered", v.reason
    rows = _discovery_rows(r._out().read_bytes())
    frogs = next(x for x in rows if x[3] == sections.LABEL["frogs"])
    assert frogs[7] == sections.UNCONFIRMED and "Form Interrogatories Set One.pdf" in frogs[8]
    text = "".join("\n".join(printout.text_of(r._pdf())).split())
    assert "".join(sections.UNCONFIRMED.split()) in text and "2026-03-09" not in text


def test_a_verify_overturn_replaces_the_section_whole(tmp_path):
    seat = LitSeat({M1: disc_docs()})
    verdicts = [{"section": "frogs", "verdict": "overturned", "corrected": [], "why": "not served"}]
    r = _run(tmp_path, seat, Scripted12(verdicts=verdicts))
    assert r.run().verdict == "delivered"
    rows = _discovery_rows(r._out().read_bytes())
    assert sum(1 for x in rows if x[4] == sections.NONE_IN_FILE) == 12


def test_a_case_with_no_discovery_paper_is_not_read_and_says_so(tmp_path):
    seat = LitSeat({M1: matter_docs()})
    client = Scripted12()
    r = _run(tmp_path, seat, client)
    assert r.run().verdict == "delivered"
    rows = _discovery_rows(r._out().read_bytes())
    assert len(rows) == 1 and rows[0][4] == book.NOT_READ and client.calls == []
    assert printout.D12_NOT_READ in "\n".join(printout.text_of(r._pdf()))


def test_a_batch_projected_over_the_cap_is_never_submitted(tmp_path, monkeypatch):
    seat = LitSeat({M1: disc_docs()})
    client = Scripted12()
    monkeypatch.setattr(batch, "project_usd", lambda *a: 10_000.0)
    v = _run(tmp_path, seat, client).run()
    assert v.verdict == "failed" and v.reason.startswith("limit:") and client.submitted == [] and seat.sent == []


def test_the_projection_prices_the_batch_rate_and_errs_high(pricing_path):
    from medchron.budget import Pricing

    p = Pricing.load(pricing_path)
    rate = p.rate_for("claude-sonnet-5")
    full = batch.project_usd(rate, 400_000, 1, 1.0)
    assert batch.project_usd(rate, 400_000, 1, p.batch_multiplier) == pytest.approx(full * p.batch_multiplier)
    assert full >= 100_000 * rate.input_per_million_cents / 1e8  # at least the input at four chars a token


def test_chunks_keep_papers_whole_and_cut_an_over_long_one():
    docs = [(1, "a" * 300), (2, "b" * 300), (3, "c" * 900)]
    got = batch.chunk(docs, limit=700)
    assert [[n for n, _t in c] for c in got] == [[1, 2], [3]]
    assert len(got[1][0][1]) == 700 and got[1][0][1].endswith(batch.CUT)


def test_the_merge_keeps_each_entry_once():
    e = {"item": "x"}
    got = batch.merge([{"sections": {"frogs": [e]}, "flags": ["f"]}, {"sections": {"frogs": [e]}, "flags": ["f"]}])
    assert got["sections"]["frogs"] == [e] and got["flags"] == ["f"]


def test_the_proof_reads_the_named_case_and_files_nothing(tmp_path):
    seat = LitSeat({M1: disc_docs()})
    client = Scripted12()
    report = tmp_path / "proof.json"
    r = Discovery12Proof(
        make_job(tmp_path / "job", report="discovery12"),
        report_path=report,
        only=["100001"],
        **_kw(tmp_path, seat, client),
    )
    v = r.run()
    assert v.verdict == "held" and v.reason.startswith("rehearsal:") and seat.sent == []
    rep = json.loads(report.read_text())
    assert rep["usd_total"] > 0 and rep["cases"][0]["number"] == "100001" and rep["cases"][0]["read"] is True
    assert set(rep["files_built"]) == {r._out().name, r._pdf().name}


def test_the_book_check_can_fail(tmp_path):
    ms = [{"matter_id": "x", "number": "1", "title": "A v. B", "read": True, "sections": {}, "flags": []}]
    out = tmp_path / "x.xlsx"
    book.build(ms, out, TODAY)
    assert book.check(out.read_bytes(), ms) == []
    assert book.check(out.read_bytes(), ms + [{**ms[0], "number": "2"}])


def test_the_printout_check_catches_a_discovery_date_not_printed(tmp_path):
    m = {"matter_id": "x", "number": "1", "case_name": "A v. B", "title": "A v. B", "responsible": "Lee",
         "discovery12": {"read": True, "sections": {}, "flags": []}}  # fmt: skip
    out = tmp_path / "x.pdf"
    printout.build([m], out, TODAY)
    assert printout.check(out, [m]) == []
    planted = {**m, "discovery12": {"read": True, "flags": [], "sections": {"rfas": [{"item": "RFA",
               "served": {"date": "2026-02-02"}}]}}}  # fmt: skip
    assert any("discovery dates are not printed" in p for p in printout.check(out, [planted]))


def test_the_plan_takes_every_discovery_paper_uncapped(tmp_path):
    from medchron.litigation import firm as firm_mod

    firm = firm_mod.load(str(make_inputs(tmp_path / "inputs")))
    files = [
        {"id": f"d{i}", "name": f"Request for Production Set {i}", "ext": ".pdf", "modified": "2026-01-01"}
        for i in range(60)
    ]
    p = plan_for(files, firm, TODAY)
    assert len(p["candidates"]) == 60 and p["read_groups"] == ["discovery12"]


def test_the_envelope_and_the_proof_key():
    from litigation_testkit import envelope

    env = {**envelope(JOB), "slug": "example", "month_cents_used": 0, "report": "discovery12"}
    j = job_mod.parse(env, Path("/x"))
    assert j.report == "discovery12" and j.workbooks == ()
    with pytest.raises(job_mod.LitigationJobError):
        job_mod.parse({**env, "workbooks": ["deadlines"]}, Path("/x"))
    assert lane_proof.lane_key("litigation", env) == "litigation-status.discovery12"
