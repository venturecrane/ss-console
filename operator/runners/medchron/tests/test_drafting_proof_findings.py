"""The 2026-10-13 proof run, as falsifiers: each class of MUST-FIX the
pre-delivery review held three propounded sets on, and the code that now
settles that class before a review is paid for.

1. A sworn section 2030.050 paragraph 4 left as a placeholder because a memo
   ABOUT discovery read as a possible prior set (``render.proof``).
2. A note on a set of admissions counting "1 special interrogatory" because a
   request cited a served interrogatory response (``render.special_count``).
3. A caption note that said nothing (``notes.attorney_notes``, ``wording``).
4. Findings the audit flagged and nobody fixed, forwarded as "open items"
   (``compose.STRICT_CLASSES``, ``finish._draft``).
5. A request whose premise the record does not bear out (the same: a served
   paper's audit finding is repaired or the job fails, never marked in place).
6. A defined term worded differently from the served response, and
7. the attorney block and a shared term worded differently across the sets of
   one request (``sets.align``, ``sets.authored_head``).
8. Notes that hedge on who the handling attorney is (``finish._request_facts``,
   ``wording.hedges``), and the system's own words in the notes.

Every name, number and sentence is invented (Example / Exampletown).
"""

from __future__ import annotations

import json

import pytest

from medchron.drafting import compose, notes as notes_mod, render, sets, wording
from test_drafting_review import KINDS, Client, _run, one_set

SROG, RFA, RFP = "Special Interrogatories", "Requests for Admission", "Requests for Production"


@pytest.fixture(autouse=True)
def _checker(monkeypatch):
    from test_drafting_review import CHECKER

    monkeypatch.setenv("SMD_DRAFTING_GATE_CHECK", str(CHECKER))


@pytest.fixture
def pricing(tmp_path):
    from medchron_testkit import PRICING

    p = tmp_path / "pricing.json"
    p.write_text(json.dumps(PRICING), encoding="utf-8")
    return p


class _Firm:
    def attachment(self, key: str) -> str:
        from drafting_testkit import DECL, POS

        return {"decl_2030_050": DECL, "pos": POS}[key]


def _srogs(n: int) -> str:
    items = "\n\n".join(f"SPECIAL INTERROGATORY NO. {i}\n\nState a fact." for i in range(1, n + 1))
    return (
        "| | |\n| --- | --- |\n| PROPOUNDING PARTY: | GAMMA EXAMPLE |\n| RESPONDING PARTY: | DELTA EXAMPLE |\n\n"
        f'# DEFINITIONS\n\n1. "INCIDENT" means the fall.\n\n# SPECIAL INTERROGATORIES\n\n{items}\n\nDated: x\n'
    )


def _digest(*lines: str) -> str:
    return "## MEDICAL | x\nnothing\n\n## DISCOVERY-SETS\n" + "\n".join(lines) + "\n\n## FILES-SEEN\n- x\n"


#: A tool memo that NAMES the set type and was indexed as a set with no count.
MEMO_ROW = (
    "Special Interrogatories | not readable | Gamma Example -> Delta Example | special interrogatories | "
    "not readable | undated | FILE: Discovery tool memo.pdf, p. 1"
)
MEMO = {
    "Discovery tool memo.pdf": (
        "MEMO. Uploaded the defense's Special Interrogatories, Set One to the response tool; "
        "a proof of service is attached to the upload."
    )
}


# ---- 1. paragraph 4 is computed from served sets, never a placeholder ------------------------


def test_a_memo_that_names_the_set_type_is_a_mention_and_paragraph_4_is_counted():
    doc, notes = render.attach(_srogs(60), "discovery_set", _Firm(), _digest(MEMO_ROW), MEMO)
    assert "DECLARATION FOR ADDITIONAL DISCOVERY" in doc  # 60 > 35 on this set alone
    assert "{{ATTORNEY" not in doc.split("DECLARATION")[1].split("PROOF OF SERVICE")[0]
    assert "a total of 0 interrogatories" in doc and "contains 60 specially" in doc
    assert any("the file shows none served" in n for n in notes)
    assert any("read as a mention, not a prior set" in n for n in notes)


def test_a_set_on_file_without_proof_of_service_is_named_never_a_blank_in_the_oath():
    row = MEMO_ROW.replace("FILE: Discovery tool memo.pdf", "FILE: SROG one draft.pdf")
    corpus = {
        "SROG one draft.pdf": (
            "PLAINTIFF GAMMA EXAMPLE'S SPECIAL INTERROGATORIES, SET ONE\nPROPOUNDING PARTY: GAMMA EXAMPLE\n"
            "RESPONDING PARTY: DELTA EXAMPLE\nSET NO.: ONE\nSPECIAL INTERROGATORY NO. 1\nState a fact."
        )
    }
    doc, notes = render.attach(_srogs(60), "discovery_set", _Firm(), _digest(row), corpus)
    assert "a total of 0 interrogatories" in doc and "{{ATTORNEY: the file shows" not in doc
    assert any("not counted, the file does not show it served" in n and "SROG one draft.pdf" in n for n in notes)


def test_a_served_set_whose_count_cannot_be_read_still_leaves_paragraph_4_to_the_attorney_and_says_why():
    row = MEMO_ROW.replace("FILE: Discovery tool memo.pdf", "FILE: SROG one.pdf")
    corpus = {
        "SROG one.pdf": (
            "PLAINTIFF GAMMA EXAMPLE'S SPECIAL INTERROGATORIES, SET ONE\nPROPOUNDING PARTY: GAMMA EXAMPLE\n"
            "RESPONDING PARTY: DELTA EXAMPLE\nPROOF OF SERVICE\nServed by mail."
        )
    }
    doc, _ = render.attach(_srogs(60), "discovery_set", _Firm(), _digest(row), corpus)
    assert "{{ATTORNEY: the file shows a prior set served, but its count is not readable" in doc


# ---- 2. only the set's own labels count --------------------------------------------------------


def test_a_set_of_admissions_that_cites_an_interrogatory_counts_no_interrogatories():
    rfa = one_set(RFA).replace(
        "Admit that YOU owned the PREMISES.",
        "Admit that YOUR response to Special Interrogatory No. 41 identified the PREMISES.",
    )
    assert render.special_count(rfa) == 0
    _doc, notes = render.attach(rfa.split("=== ATTORNEY NOTES")[0], "discovery_set", _Firm(), _digest())
    assert not any("2030.050" in n for n in notes)
    assert any(n.startswith("set number: earlier requests for admission") and n.endswith("none") for n in notes)
    assert render.special_count(_srogs(60)) == 60


# ---- 3. a note that says nothing never reaches the attorney -----------------------------------


def test_a_caption_difference_with_nothing_to_report_is_not_a_line(tmp_path):
    cap = {
        "source": "Complaint.pdf",
        "fields": {},
        "discrepancies": [
            {"field": "court", "document_value": "", "record_value": "x", "why": "differs", "quote": "", "source": "p1"}
        ],
        "corrections": [],
        "compared": ["court"],
    }
    (tmp_path / "caption.json").write_text(json.dumps(cap))
    for name, body in (("final-pass.json", {"settled": [], "drifts": []}), ("gate.json", {})):
        (tmp_path / name).write_text(json.dumps(body))
    md = notes_mod.attorney_notes(tmp_path, "Discovery Set", [{"kind": "ATTORNEY", "text": "declarant's name"}], [])
    assert "- court:" not in md and "No discrepancy on those fields" in md
    assert "- For your decision: declarant's name" in md and "{{" not in md
    assert wording.empty_lines('- court: ""\n- court: the paper reads "X".\n') == ['- court: ""']


# ---- 4 and 5. a served paper's audit findings are applied, or the job fails -------------------


def _audit_flagging(word: str):
    """An auditor that reports a DRIFTS row for any section carrying ``word``."""

    def answer(params):
        body = params["messages"][-1]["content"]
        if word in body:
            return (
                f"- {word} | DRIFTS | the served response says otherwise\nSUPPORTED=0 DRIFTS=1 INVENTED=0 ARITHMETIC=0"
            )
        return "- claim | SUPPORTED | cite\nSUPPORTED=1 DRIFTS=0 INVENTED=0 ARITHMETIC=0"

    return answer


class AuditClient(Client):
    def __init__(self, *a, audit_fn=None, repair_fix: tuple[str, str] | None = None, **kw):
        super().__init__(*a, **kw)
        self.audit_fn, self.repair_fix = audit_fn, repair_fix

    def _answer(self, params):
        system = json.dumps(params.get("system"))
        if "AUDIT-PROMPT" in system and self.audit_fn is not None:
            return self.audit_fn(params)
        text = super()._answer(params)
        if "REPAIR-PROMPT" in system and self.repair_fix:
            text = text.replace(*self.repair_fix)
        return text


def _drafts(**over: str) -> dict[str, str]:
    return {label: over.get(label, one_set(label)) for label in KINDS}


def test_a_drift_the_audit_flags_is_repaired_before_anything_is_delivered(tmp_path, pricing):
    bad = one_set(RFA).replace('"PREMISES" means the store.', '"PREMISES" means the WRONGSTORE.')
    client = AuditClient(
        drafts=_drafts(**{RFA: bad}),
        audit_fn=_audit_flagging("WRONGSTORE"),
        repair_fix=("WRONGSTORE", "store"),
    )
    r, v, _ = _run(tmp_path, pricing, client)
    assert v.outcome == "delivered", (v.stage, v.reason)
    assert client.count("REPAIR-PROMPT") >= 1
    final = (r.data / "set-1" / "draft-final.md").read_text()
    assert "WRONGSTORE" not in final
    assert "WRONGSTORE" not in (r.data / "notes-final.md").read_text()


def test_a_drift_the_repairs_never_clear_fails_and_is_never_forwarded_as_an_open_item(tmp_path, pricing):
    bad = one_set(RFA).replace('"PREMISES" means the store.', '"PREMISES" means the WRONGSTORE.')
    client = AuditClient(drafts=_drafts(**{RFA: bad}), audit_fn=_audit_flagging("WRONGSTORE"))
    r, v, seat = _run(tmp_path, pricing, client)
    assert v.outcome == "failed" and v.reason.startswith("audit_unrepaired: "), (v.stage, v.reason)
    assert "WRONGSTORE" not in v.reason  # counts only: a reason reaches SMD's alerts
    assert seat.created == [] and seat.sent == []
    assert client.count("REPAIR-PROMPT") == 1  # a repair that changed nothing is not audited again
    assert client.count("INDEPENDENT PRE-DELIVERY REVIEW") == 0  # nothing paid past the failure


def test_repairs_that_change_the_text_but_never_clear_the_drift_stop_at_three(tmp_path, pricing):
    bad = one_set(RFA).replace('"PREMISES" means the store.', '"PREMISES" means the WRONGSTORE.')
    client = AuditClient(
        drafts=_drafts(**{RFA: bad}),
        audit_fn=_audit_flagging("WRONGSTORE"),
        repair_fix=("WRONGSTORE", "WRONGSTORE X"),
    )
    _r, v, _ = _run(tmp_path, pricing, client)
    assert v.outcome == "failed" and v.reason.startswith("audit_unrepaired: "), (v.stage, v.reason)
    assert client.count("REPAIR-PROMPT") == 3


def test_an_unsupported_premise_in_a_served_paper_is_never_marked_in_place(tmp_path, pricing):
    def invented(params):
        if "PRESERVATIONLETTER" in params["messages"][-1]["content"]:
            return "- received PRESERVATIONLETTER | INVENTED | no cite\nSUPPORTED=0 DRIFTS=0 INVENTED=1 ARITHMETIC=0"
        return "- claim | SUPPORTED | cite\nSUPPORTED=1 DRIFTS=0 INVENTED=0 ARITHMETIC=0"

    bad = one_set(RFA).replace("Admit that YOU owned the PREMISES.", "Admit that YOU received the PRESERVATIONLETTER.")
    client = AuditClient(drafts=_drafts(**{RFA: bad}), audit_fn=invented)
    _r, v, _ = _run(tmp_path, pricing, client)
    assert v.outcome == "failed" and v.reason.startswith("audit_unrepaired: ")


def test_drifts_block_only_the_served_paper_classes():
    res = {"tallies": {"DRIFTS": 2}, "quotes_not_found": 0}
    assert compose.blocking_findings(res, "discovery_set") == 2
    assert compose.blocking_findings(res, "mediation_brief") == 0


# ---- 6 and 7. one front matter across the sets of one request ----------------------------------

HEAD = (
    "Alpha Example (SBN 000001)\nEXAMPLE LAW\nAttorneys for Plaintiff\n\nSUPERIOR COURT OF THE STATE OF CALIFORNIA\n\n"
)


def test_a_later_set_is_aligned_to_the_first_sets_attorney_block_definitions_and_signature():
    first = HEAD + one_set(RFA).replace("Dated: {{ATTORNEY: date}}", "Dated: today\n\nEXAMPLE LAW\nBy: Alpha Example")
    later = (
        "Alpha Example, Esq.\nExample Law LLP\nAttorneys for Plaintiff\n\nSUPERIOR COURT OF THE STATE OF CALIFORNIA\n\n"
        + one_set(SROG)
        .replace('"PREMISES" means the store.', '"PREMISES" means the parking lot.')
        .replace("Dated: {{ATTORNEY: date}}", "Dated: today\n\nExample Law LLP\nBy: A. Example")
    )
    doc, changes = sets.align(
        later.split("=== ATTORNEY")[0], first.split("=== ATTORNEY")[0], "Requests for Admission set"
    )
    assert doc.startswith("Alpha Example (SBN 000001)\nEXAMPLE LAW\n")
    assert '"PREMISES" means the store.' in doc and "parking lot" not in doc
    assert "By: Alpha Example" in doc and "A. Example" not in doc
    assert "SPECIAL INTERROGATORY NO. 1" in doc  # its own requests untouched
    assert len(changes) == 3
    again, none = sets.align(doc, first.split("=== ATTORNEY")[0], "x")
    assert again == doc and none == []


def test_the_attorney_block_is_printed_as_the_firm_authors_it():
    doc = "Alpha Example, Esq.\nAttorneys for Plaintiff\n\nSUPERIOR COURT OF THE STATE OF CALIFORNIA\n\nBody.\n"
    out, notes = sets.authored_head(doc, ["Alpha Example (SBN 000001)", "EXAMPLE LAW", "Email: alpha@firm.example"])
    assert out.startswith("Alpha Example (SBN 000001)\nEXAMPLE LAW\nEmail: alpha@firm.example\nAttorneys for Plaintiff")
    assert (
        notes
        and sets.authored_head(out, ["Alpha Example (SBN 000001)", "EXAMPLE LAW", "Email: alpha@firm.example"])[1] == []
    )


def test_three_sets_with_three_wordings_deliver_with_one(tmp_path, pricing):
    drafts = {
        RFA: one_set(RFA),
        RFP: one_set(RFP).replace('"PREMISES" means the store.', '"PREMISES" means the store and its lot.'),
        SROG: one_set(SROG).replace('"PREMISES" means the store.', '"PREMISES" means the parking lot.'),
    }
    client = Client(drafts=drafts)
    r, v, _ = _run(tmp_path, pricing, client)
    assert v.outcome == "delivered", (v.stage, v.reason)
    defs = [
        sets.definition_texts((r.data / f"set-{i}" / "draft-final.md").read_text())["PREMISES"][2] for i in (1, 2, 3)
    ]
    assert defs == ['"PREMISES" means the store.'] * 3
    # the later composes were told the first set's front matter
    later = [c for c in client.calls if "SET ALREADY DRAFTED FOR THIS REQUEST" in json.dumps(c.get("messages"))]
    assert later


def test_a_model_written_declaration_never_travels_with_the_signature_block():
    canon = (
        one_set(SROG)
        .split("=== ATTORNEY")[0]
        .replace(
            "Dated: {{ATTORNEY: date}}",
            "Dated: today\n\nBy: Alpha Example\n\nDECLARATION FOR ADDITIONAL DISCOVERY\n\nI declare.",
        )
    )
    later = one_set(RFA).split("=== ATTORNEY")[0].replace("Dated: {{ATTORNEY: date}}", "Dated: today\n\nBy: A. Example")
    doc, _ = sets.align(later, canon, "Special Interrogatories set")
    assert "By: Alpha Example" in doc and "DECLARATION" not in doc and "I declare." not in doc
    assert "DECLARATION" not in sets.front_matter(canon)


def test_a_title_above_the_court_line_is_never_taken_for_an_attorney_block():
    doc = "# PLAINTIFF'S REQUESTS FOR ADMISSION\nAttorneys for Plaintiff\n\nSUPERIOR COURT OF THE STATE OF CALIFORNIA\n"
    start, at, lines = sets.head_lines(doc)
    assert lines == ["Attorneys for Plaintiff"] and doc[:start].startswith("# PLAINTIFF'S REQUESTS")
    out, _ = sets.authored_head(doc, ["Alpha Example (SBN 000001)"])
    assert out.startswith("# PLAINTIFF'S REQUESTS FOR ADMISSION\nAlpha Example (SBN 000001)\nAttorneys for Plaintiff")


def test_a_definition_worded_for_one_kind_of_set_is_not_copied_into_another():
    canon = one_set(SROG).replace(
        '"PREMISES" means the store.', '"PREMISES" means the store, in answering these interrogatories.'
    )
    later = one_set(RFA).replace('"PREMISES" means the store.', '"PREMISES" means the lot.')
    doc, changes = sets.align(later, canon, "x")
    assert '"PREMISES" means the lot.' in doc and not changes


def test_a_scanned_prior_set_is_unread_never_zero():
    row = MEMO_ROW.replace("FILE: Discovery tool memo.pdf", "FILE: SROG one scan.pdf")
    doc, _ = render.attach(_srogs(60), "discovery_set", _Firm(), _digest(row), {"SROG one scan.pdf": "  [p.1]  "})
    assert "{{ATTORNEY: the file shows a prior set served, but its count is not readable" in doc


def test_interrogatory_labels_carrying_their_text_still_count():
    inline = _srogs(40).replace(
        "SPECIAL INTERROGATORY NO. 1\n\nState a fact.", "SPECIAL INTERROGATORY NO. 1: State a fact."
    )
    assert render.special_count(inline) == 40
    assert render.special_count("SPECIAL INTERROGATORY NO. 1: State.\n\nSPECIAL INTERROGATORY NO. 2: State.") == 2


def test_the_compound_marker_names_the_variants_own_rule():
    rfp = one_set(RFP).replace(
        "All DOCUMENTS showing the PREMISES.",
        "All DOCUMENTS showing the PREMISES. {{ATTORNEY: compound under section 2030.060(f)?}}",
    )
    out, notes = sets.compound_markers(rfp)
    assert "2030.060(f)" not in out and "2031.030(c)(1)" in out and notes


# ---- the slots and the words the attorney never receives --------------------------------------


@pytest.mark.parametrize(
    "text, found",
    [
        ("REQUEST FOR ADMISSION NO. 1\n\n{{NOT IN RECORD: the date}}", True),
        ("Case No. {{FILL: case number | operative pleading}}", True),
        ("On {{FILL: date of service | at service}}, I served it.", False),
        ("{{ATTORNEY: declarant's full name}}", False),
        ("{{ATTORNEY}}", True),
        ("PLAINTIFF <NAME>'S REQUESTS", True),
        ("=== ATTORNEY NOTES ===", True),
    ],
)
def test_an_unfilled_slot_in_a_served_paper_is_found(text, found):
    assert bool(wording.slots("discovery_set", text)) is found


def test_the_systems_words_and_hedges_are_found_in_notes_but_the_codes_name_is_not():
    assert wording.vocabulary("Open items from the audit") and wording.vocabulary("counted by code")
    assert (
        wording.vocabulary("All DOCUMENTS showing the audit trail of the chart; return to this job; by Code, § 250")
        == []
    )
    assert wording.document_vocabulary("REQUEST FOR PRODUCTION NO. 1 the audit trail, the pipeline") == []
    assert wording.hedges('The ER note reads "possibly a sprain."') == []
    assert wording.empty_lines("- Prior sets:\n  - Set One, served 03/01/2026") == []
    assert (
        wording.vocabulary("under Code of Civil Procedure section 2030.050, as required by Code of Civil Procedure")
        == []
    )
    assert wording.vocabulary("He was injured on the job at the gate.") == []
    assert wording.hedges("It is unclear whether Alpha Example is the handling attorney.")
    assert wording.hedges("Alpha Example is the responsible attorney (matter record).") == []


def test_a_served_paper_with_a_not_in_record_slot_gets_one_repair_then_fails(tmp_path, pricing):
    bad = one_set(RFA).replace("Admit that YOU owned the PREMISES.", "Admit that {{NOT IN RECORD: the owner}}.")
    client = Client(drafts=_drafts(**{RFA: bad}))
    _r, v, seat = _run(tmp_path, pricing, client)
    assert v.outcome == "failed" and v.reason.startswith("format_check: "), (v.stage, v.reason)
    assert "the owner" not in v.reason and seat.created == []


def test_notes_that_hedge_are_rewritten_once_then_fail(tmp_path, pricing):
    client = Client(drafts=_drafts(), notes="## Open items\n\n- It is unclear whether Alpha is the handling attorney.")
    _r, v, seat = _run(tmp_path, pricing, client)
    assert v.outcome == "failed" and v.reason.startswith("notes_unclean: ") and "Alpha" not in v.reason
    assert client.count("THE ATTORNEY NOTES, FROM THE FINAL DOCUMENT") == 2
    assert seat.created == []


def test_the_notes_are_told_who_asked_and_whose_block_it_is_and_say_it(tmp_path, pricing):
    client = Client(drafts=_drafts())
    r, v, _ = _run(tmp_path, pricing, client)
    assert v.outcome == "delivered", (v.stage, v.reason)
    call = next(c for c in client.calls if "THE ATTORNEY NOTES, FROM THE FINAL DOCUMENT" in json.dumps(c.get("system")))
    user = json.dumps(call["messages"])
    assert "the documents were requested by admin@firm.example" in user
    notes = (r.data / "notes-final.md").read_text()
    assert (
        "- Attorney block: the firm has no authored attorney block for the requesting or the responsible attorney"
        in notes
    )
    assert wording.vocabulary(notes) == [] and "{{" not in notes


# ---- what the replay of the held run showed (2026-10-14) ------------------------------------


def test_the_dated_line_carries_the_drafts_date_in_every_set():
    for raw in ("Dated: {{NOT IN RECORD: date of signing}}", "Dated:", "Dated: `{{FILL: date | today}}`"):
        out, notes = sets.dated(f"Body.\n\n{raw}\n\nEXAMPLE LAW\n", "October 14, 2026")
        assert "Dated: October 14, 2026\n" in out and "{{" not in out and notes
    same = "Dated: October 14, 2026\n"
    assert sets.dated(same, "October 14, 2026") == (same, [])


def test_the_declarant_is_the_authored_blocks_attorney():
    assert render.declarant_name(["Alpha Q. Example, Esq. (SBN 000001)", "EXAMPLE LAW"]) == "Alpha Q. Example"
    assert render.declarant_name([]) is None
    from drafting_testkit import DECL

    class Firm(_Firm):
        def attachment(self, key):
            if key == "decl_2030_050":
                return (
                    "# DECLARATION FOR ADDITIONAL DISCOVERY\n\nI, `{{ATTORNEY: declarant's full name}}`, declare:\n"
                    + DECL
                )
            return super().attachment(key)

    doc, _ = render.attach(_srogs(40), "discovery_set", Firm(), _digest(), {}, declarant="Alpha Q. Example")
    assert "declarant's full name" not in doc and "I, Alpha Q. Example, declare" in doc
    bare, _ = render.attach(_srogs(40), "discovery_set", Firm(), _digest(), {})
    assert "{{ATTORNEY: declarant's full name}}" in bare


def test_fictitious_defendants_are_not_a_spelling_difference():
    from medchron.drafting import caption

    fields = {"defendants": [{"value": "DELTA EXAMPLE, LLC and DOES 1-25", "verified": True, "quote": "q"}]}
    record = {"parties": [{"side": "other", "name": "DELTA EXAMPLE, LLC"}]}
    diffs, compared = caption.compare(fields, record, "Complaint.pdf")
    assert diffs == [] and "defendant" in compared
    fields["defendants"][0]["value"] = "DELTA EXAMPEL, LLC and DOES 1-25"
    assert [d["field"] for d in caption.compare(fields, record, "Complaint.pdf")[0]] == ["defendant"]


def test_the_review_reads_a_document_in_body_order(tmp_path):
    import docx

    from medchron.drafting import format_check

    d = docx.Document()
    d.add_paragraph("ATTORNEY BLOCK")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "GAMMA v. DELTA", "Case No. CV-0001"
    d.add_paragraph("PROOF OF SERVICE")
    text = format_check.doc_text(d)
    assert text.index("Case No.") < text.index("PROOF OF SERVICE")


def test_a_line_every_set_carries_is_said_once(tmp_path):
    for k in ("a", "b"):
        (tmp_path / k).mkdir()
        (tmp_path / k / "final-pass.json").write_text(json.dumps({"settled": [], "drifts": []}))
        (tmp_path / k / "gate.json").write_text(json.dumps({"warnings": ["3 of 9 sources are scans"]}))
    (tmp_path / "walled.json").write_text(json.dumps([{"reason": "firm internal", "name": "INTAKE"}]))
    md = notes_mod.attorney_notes(
        tmp_path, "Discovery Set", [], [], "", [("RFA", tmp_path / "a"), ("RFP", tmp_path / "b")]
    )
    assert md.count("3 of 9 sources are scans") == 1
    assert "privilege wall" not in md and "never read" not in md
    assert 'Email kept out of drafting (firm internal): "INTAKE"' in md


def test_a_walled_email_named_like_a_filed_document_stays_out_and_the_document_stays_in(tmp_path, monkeypatch):
    """The 2026-10-13 review: an internal email titled like the signed lien
    filed in the liens folder. The email (body and attachment) never reaches
    the corpus; the filed lien is an ordinary record and does."""
    from pathlib import Path

    from medchron.demand import pull
    from medchron.demand.firm import DemandFirm

    firm = DemandFirm(
        root=tmp_path,
        data={"privilege": {"firm_domains": ["firm.example"], "consumer_domains": ["mail.example"]}},
    )
    data = tmp_path / "data"
    (data / "raw").mkdir(parents=True)
    rows = []
    for rid, ext, folder in (("e1", ".msg", "/EMAILS"), ("f1", ".pdf", "/LIENS")):
        p = data / "raw" / f"{rid}{ext}"
        p.write_bytes(b"x")
        rows.append(
            {
                "id": rid,
                "name": "LIEN - SIGNED",
                "ext": ext,
                "folder": folder,
                "ok": True,
                "path": str(p),
                "sha256": rid,
            }
        )
    msg = {
        "subject": "LIEN - SIGNED",
        "sender": "staff@firm.example",
        "recipients": ["atty@firm.example"],
        "date": "",
        "body": "Internal strategy words about the lien.",
        "attachments": [("LIEN - SIGNED.pdf", b"%PDF-internal-copy" * 9000)],
    }
    monkeypatch.setattr(pull, "_open_msg", lambda path: msg)
    out = pull.split_emails(rows, data, firm, set())
    assert [w["id"] for w in out["walled"]] == ["e1"] and out["kept"] == []
    assert "Internal strategy words" in Path(out["walled"][0]["text_path"]).read_text()
    assert not any(r.get("from_email") for r in out["kept"])  # the email's attachment never became a document
