"""A request that names court-paper columns and no discovery reads no discovery
and shows no Discovery tab (2026-10-23: Chris asked for complaint, service and
answers on every litigation file; the run read discovery on 152 matters)."""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from litigation_testkit import M1, LitSeat, ScriptedLit, make_inputs, make_job, matter_docs
from medchron.litigation import manifest, vocab
from medchron.litigation.firm import load
from medchron.litigation.run import LitigationRun, court_only

TODAY = dt.date(2026, 10, 7)
CHRIS = (
    "For all files that are in litigation and for which a complaint has been filed, please do the following: "
    "Note the date the complaint was filed. Note the date the defendants were served. "
    "Note the date each of the defendants filed an answer or indicate they have not filed an answer."
)


@pytest.fixture(autouse=True)
def _pricing(tmp_path: Path, pricing_path: Path) -> None:
    (tmp_path / "pricing.json").write_text(pricing_path.read_text())


def _rx(tmp_path: Path) -> list[re.Pattern[str]]:
    return load(make_inputs(tmp_path / "rx")).discovery_rx


@pytest.mark.parametrize(
    ("trigger", "text", "expected"),
    [
        ("request", CHRIS, True),
        ("request", CHRIS + " Also the discovery we propounded.", False),
        ("request", CHRIS + " And any interrogatories served on us.", False),
        ("request", "Send me a fresh litigation status list.", False),  # names no columns: everything
        ("scheduled", CHRIS, False),  # the weekday list always reads everything
    ],
)
def test_court_only_is_a_request_naming_court_columns_and_no_discovery(tmp_path, trigger, text, expected):
    assert court_only(trigger, text, _rx(tmp_path)) is expected


def _run(tmp_path: Path, text: str) -> LitigationRun:
    seat = LitSeat({M1: matter_docs()})
    return LitigationRun(
        make_job(tmp_path / "job", request_text=text),
        inputs_dir=str(make_inputs(tmp_path / "inputs")),
        pricing=str(tmp_path / "pricing.json"),
        state_dir=tmp_path / "state",
        seat_factory=lambda: seat,
        client=ScriptedLit(),
        log=lambda m: None,
        today=lambda: TODAY,
        readback_pause=0.0,
    )


def test_a_court_only_run_reads_no_discovery_and_shows_no_discovery_tab(tmp_path):
    r = _run(tmp_path, CHRIS)
    v = r.run()
    assert v.verdict == "delivered", v.reason
    st = manifest.load_prior(tmp_path / "state", M1)
    assert set(st["fields_read"]) == {vocab.GROUP_CASE, vocab.GROUP_DEFENDANTS}  # the next full list reads discovery
    assert "Discovery" not in load_workbook(r._out()).sheetnames


def _discovery_divergence(monkeypatch: pytest.MonkeyPatch) -> None:
    """Every matter's parity carries one unexplained discovery change, the
    shape that held a court-only list on 2026-10-23."""
    from medchron.litigation import parity

    real = parity.compare

    def compare(*a, **kw):
        return [
            *real(*a, **kw),
            {
                "path": "discovery_propounded[frogs|delta].date",
                "label": "x",
                "old": "2026-06-01",
                "new": None,
                "explained": False,
                "why": "changed with no document cited",
            },
        ]

    monkeypatch.setattr(parity, "compare", compare)


def test_a_court_only_list_does_not_hold_on_discovery_it_did_not_read(tmp_path, monkeypatch):
    _discovery_divergence(monkeypatch)
    v = _run(tmp_path, CHRIS).run()
    assert v.verdict == "delivered", v.reason


def test_a_full_list_still_holds_on_an_unexplained_discovery_change(tmp_path, monkeypatch):
    _discovery_divergence(monkeypatch)
    v = _run(tmp_path, "Send me a fresh litigation status list.").run()
    assert v.verdict != "delivered" and "parity_hold" in (v.reason or ""), v.reason


def test_a_full_request_still_reads_discovery_and_shows_its_tab(tmp_path):
    r = _run(tmp_path, "Send me a fresh litigation status list.")
    v = r.run()
    assert v.verdict == "delivered", v.reason
    assert set(manifest.load_prior(tmp_path / "state", M1)["fields_read"]) == set(vocab.GROUPS)
    assert "Discovery" in load_workbook(r._out()).sheetnames
