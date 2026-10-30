"""The one-sheet-per-case copy of the status list (``litigation/percase.py``):
built from the same final values as the main list, every case on one tab,
every defendant row once, printable, leak-scanned, reviewed, and filed after
the main list and never without it.

Every name, number and date is invented (Example / CV-0001)."""

from __future__ import annotations

import collections
import copy
import datetime as dt
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from litigation_testkit import JOB, M1, LitSeat, ScriptedLit, make_inputs, make_job, matter_docs
from medchron.litigation import book, gates, percase, printout, review, vocab
from medchron.litigation.run import PER_CASE_ROLE, PRINTOUT_ROLE, LitigationRun
from medchron_testkit import doc_row
from test_litigation_gates import good

TODAY = dt.date(2026, 10, 7)


def _matters() -> list[dict]:
    """Three cases: two filed (one with two defendants, one held by another
    attorney), one not filed with no defendants."""
    a = good()
    a["defendants"].append({**copy.deepcopy(a["defendants"][0]), "name": "Zeta Example"})
    b = copy.deepcopy(good())
    b.update(matter_id="m2", number="100002", responsible="Beta Example", case_name="Epsilon v. Eta")
    b["defendants"][0].update(name="Eta Example", status=vocab.SERVED_NO_ANSWER, answered={"date": None})
    c = copy.deepcopy(good())
    c.update(matter_id="m3", number="100003", responsible="Alpha Example", case_name="Theta intake", defendants=[])
    c["case_status"] = {"value": vocab.NOT_FILED, "detail": ""}
    c["complaint_filed"] = {"date": None, "source": None}
    c["notes"] = "Demand sent; waiting on the carrier"
    a["case_name"] = "Gamma v. Delta"
    a["notes"] = "Mediation is set"
    return [a, b, c]


def _both(tmp_path: Path, ms: list[dict], tagged=None) -> tuple[Path, Path]:
    main, per = tmp_path / "main.xlsx", tmp_path / "per.xlsx"
    book.build(ms, {}, main, TODAY, tagged=tagged)
    percase.build_per_case(ms, per, TODAY, tagged=tagged)
    return main, per


def _facts(ws) -> dict[str, object]:
    out, r = {}, 3
    while ws.cell(r, 1).value:
        out[ws.cell(r, 1).value] = ws.cell(r, 2).value
        r += 1
    return out


def _def_rows(ws) -> list[tuple]:
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Defendant")
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        vals = tuple(ws.cell(r, j).value for j in range(1, 10))
        if any(v is not None for v in vals) and vals[0] != percase.NO_DEFENDANTS:
            rows.append(vals)
    return rows


# ---- the content: same values, every case once, every row once ---------------------------------
def test_every_case_is_on_exactly_one_tab_in_case_notes_order(tmp_path):
    main, per = _both(tmp_path, _matters())
    notes = [r[1].value for r in load_workbook(main)["Case notes"].iter_rows(min_row=2)]
    wb = load_workbook(per)
    shown = [_facts(ws)["Matter #"] for ws in wb.worksheets]
    assert shown == notes and len(shown) == len(set(shown)) == 3


def test_every_defendant_row_appears_once_with_the_main_lists_values(tmp_path):
    ms = _matters()
    main, per = _both(tmp_path, ms)
    want = collections.Counter()
    for r in load_workbook(main)["All defendants"].iter_rows(min_row=2, values_only=True):
        want[(r[2], *r[7:16])] += 1  # matter #, then the per-defendant columns
    got = collections.Counter()
    for ws in load_workbook(per).worksheets:
        for row in _def_rows(ws):
            got[(_facts(ws)["Matter #"], *row)] += 1
    assert got == want and sum(got.values()) == 3
    assert percase.check(per.read_bytes(), ms) == []


def test_the_facts_carry_the_matters_own_case_values(tmp_path):
    ms = _matters()
    _, per = _both(tmp_path, ms)
    ws = load_workbook(per)["100001 Gamma v. Delta"]
    f = _facts(ws)
    assert ws["A1"].value == "Gamma v. Delta"
    assert f["Complaint filed"] == "2026-03-02" and "Complaint filed (source)" not in f
    assert f["Responsible (Smokeball)"] == "Alpha Example" and f["Case status"] == vocab.ACTIVE
    assert f["Next court date"] == "2026-11-20 (CMC)" and f["Case notes"] == book.clean(book.case_notes_text(ms[0]))
    assert list(f) == [
        "Matter #",
        "Responsible (Smokeball)",
        "Court / case no.",
        "Complaint filed",
        "Case status",
        "Next court date",
        "Case notes",
    ]


def test_a_not_filed_case_says_where_it_stands_and_lists_no_defendants(tmp_path):
    _, per = _both(tmp_path, _matters())
    ws = load_workbook(per)["100003 Theta intake"]
    f = _facts(ws)
    assert f["Complaint filed"] == percase.NO_COMPLAINT and "Complaint filed (source)" not in f
    assert f["Not filed: where it stands"] == "Demand sent; waiting on the carrier"
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Defendant")
    assert ws.cell(hdr + 1, 1).value == percase.NO_DEFENDANTS and _def_rows(ws) == []
    assert ws.cell(hdr - 1, 1).value == "Defendants"


def test_a_filed_case_with_no_complaint_date_says_not_in_file_and_why(tmp_path):
    ms = _matters()
    ms[1]["complaint_filed"] = {"date": None, "source": None}
    _, per = _both(tmp_path, ms)
    f = _facts(load_workbook(per)["100002 Epsilon v. Eta"])
    assert f["Complaint filed"] == book.NOT_IN_FILE
    assert f["Complaint filed (source)"] == percase.NO_COMPLAINT_WHY
    assert "Not filed: where it stands" not in f


def test_a_people_scope_adds_the_role_row_from_the_tagged_tab(tmp_path):
    ms = _matters()
    tagged = (book.tagged_header(False), [["100001", "Gamma v. Delta", "Responsible, Assisting", vocab.ACTIVE, ""]])
    _, per = _both(tmp_path, ms, tagged)
    wb = load_workbook(per)
    assert _facts(wb["100001 Gamma v. Delta"])["Role"] == "Responsible, Assisting"
    assert list(_facts(wb["100003 Theta intake"]))[-2] == "Role"
    two = (book.tagged_header(True), [["Alpha Example", "100002", "Epsilon v. Eta", "Paralegal", "", ""]])
    _, per2 = _both(tmp_path / "two", ms, two)
    assert _facts(load_workbook(per2)["100002 Epsilon v. Eta"])["Role"] == "Alpha Example: Paralegal"
    assert "Role" not in _facts(load_workbook(_both(tmp_path / "none", ms)[1])["100002 Epsilon v. Eta"])


def test_tab_names_are_sanitised_truncated_and_unique():
    names = percase.tab_names(
        [
            "100001 A/B: v. C's [estate]?*\\",
            "100002 " + "x" * 40,
            "100002 " + "X" * 40,
            "100002 " + "x" * 40,
            "",
        ]
    )
    assert names[0] == "100001 AB v. Cs estate"
    assert names[1] == "100002 " + "x" * 24 and len(names[1]) == 31
    assert names[2] == ("100002 " + "X" * 20) + " (2)" and len(names[2]) == 31
    assert names[3] == ("100002 " + "x" * 20) + " (3)"
    assert names[4] == "Case"
    assert len({n.lower() for n in names}) == len(names)


def test_two_cases_with_one_name_get_two_tabs_through_the_build(tmp_path):
    ms = _matters()
    ms[1]["number"] = ms[0]["number"]
    ms[1]["case_name"] = ms[0]["case_name"]
    _, per = _both(tmp_path, ms)
    assert load_workbook(per).sheetnames == [
        "100001 Gamma v. Delta",
        "100003 Theta intake",
        "100001 Gamma v. Delta (2)",
    ]


# ---- the print layout, asserted on the saved file --------------------------------------------
def test_print_setup_fonts_widths_and_heights_on_the_saved_file(tmp_path):
    _, per = _both(tmp_path, _matters())
    for ws in load_workbook(per).worksheets:
        ps = ws.page_setup
        assert ps.orientation == "landscape" and str(ps.paperSize) == str(ws.PAPERSIZE_LETTER)
        assert int(ps.scale) == 100 and not ws.sheet_properties.pageSetUpPr.fitToPage
        assert (ws.page_margins.left, ws.page_margins.right) == (0.3, 0.3)
        assert (ws.page_margins.top, ws.page_margins.bottom) == (0.5, 0.5)
        hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Defendant")
        assert ws.print_title_rows == f"${hdr}:${hdr}"
        assert ws.oddFooter.center.text == percase.FOOTER
        assert [ws.column_dimensions[c].width for c in "ABCDEFGHI"] == percase.WIDTHS and sum(percase.WIDTHS) == 134
        assert ws["A1"].font.name == "Calibri" and ws["A1"].font.sz == 14 and ws["A1"].font.b
        assert "A1:I1" in {str(m) for m in ws.merged_cells.ranges}
        assert [c.value for c in ws[hdr]] == book.DEF_HEADER[7:]
        assert all(c.fill.fgColor.rgb.endswith(book.HEADER_FILL) and c.font.b for c in ws[hdr])
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for c in row:
                if c.value is None:
                    continue
                assert c.font.name == "Calibri", c.coordinate
                assert c.font.sz in (9, 12), c.coordinate
                assert c.alignment.wrap_text and c.alignment.vertical == "top", c.coordinate
        for r in range(1, ws.max_row + 1):
            assert ws.row_dimensions[r].height, f"{ws.title} row {r} has no set height"


def test_a_long_case_prints_at_8pt_and_never_smaller(tmp_path):
    ms = _matters()
    many = copy.deepcopy(ms[0])
    many["defendants"] = [
        {**copy.deepcopy(ms[0]["defendants"][0]), "name": f"Defendant {i} Example", "flags": ["A long note " * 20]}
        for i in range(60)
    ]
    _, per = _both(tmp_path, [many, ms[2]])
    stats = percase.build_per_case([many, ms[2]], tmp_path / "again.xlsx", TODAY)
    wb = load_workbook(per)
    sizes = {ws.title: {c.font.sz for r in ws.iter_rows(min_row=2) for c in r if c.value is not None} for ws in wb}
    assert sizes["100001 Gamma v. Delta"] == {8, 12}  # body 8pt, the Defendants line 12pt
    assert sizes["100003 Theta intake"] == {9, 12}
    assert stats["small_type_tabs"] == 1


def test_row_heights_grow_with_the_text():
    short = percase.lines("Delta Example", 9, percase.col_pt(14))
    long = percase.lines("word " * 200, 9, percase.col_pt(36))
    assert short == 1 and long > 10
    assert percase.height(long, 9) > percase.height(short, 9)
    assert percase.lines("x" * 300, 9, percase.col_pt(14)) > 5  # a run with no spaces still wraps
    assert percase.pages([100.0] * 12, 0) == 3 and percase.pages([10.0] * 3, 0) == 1


def test_the_per_case_name_is_the_main_lists_with_its_suffix():
    p = percase.path_for(Path("/x/out/Litigation status 2026-10-07 (Operator 0000AA).xlsx"))
    assert p.name == "Litigation status 2026-10-07 (Operator 0000AA) one sheet per case.xlsx"


def test_the_check_catches_a_dropped_or_changed_row(tmp_path):
    ms = _matters()
    _, per = _both(tmp_path, ms)
    drifted = copy.deepcopy(ms)
    drifted[0]["defendants"][1]["name"] = "Someone Else"
    assert any("differ" in p for p in percase.check(per.read_bytes(), drifted))
    assert any("tabs" in p for p in percase.check(per.read_bytes(), ms[:2]))


# ---- the run: built, scanned, reviewed, filed after the list and never without it -----------
def _run(tmp_path: Path, seat: LitSeat, job: str = "job", **kw) -> LitigationRun:
    return LitigationRun(
        make_job(tmp_path / job, **kw),
        inputs_dir=str(make_inputs(tmp_path / "inputs")),
        pricing=str(tmp_path / "pricing.json"),
        state_dir=tmp_path / "state",
        seat_factory=lambda: seat,
        client=ScriptedLit(),
        log=lambda m: None,
        today=lambda: TODAY,
        readback_pause=0.0,
    )


@pytest.fixture(autouse=True)
def _pricing(tmp_path: Path, pricing_path: Path) -> None:
    (tmp_path / "pricing.json").write_text(pricing_path.read_text())


def test_a_delivered_run_files_the_per_case_copy_after_the_list_in_its_folder(tmp_path):
    seat = LitSeat({M1: matter_docs()})
    r = _run(tmp_path, seat)
    v = r.run()
    assert v.verdict == "delivered", v.reason
    assert [s["name"] for s in seat.sent] == [
        r._out().name,
        percase.path_for(r._out()).name,
        printout.path_for(r._out()).name,
    ]
    assert seat.sent[0]["folderId"] == seat.sent[1]["folderId"] == seat.sent[2]["folderId"] == v.folder_id
    assert [f["role"] for f in v.files] == ["workbook", PER_CASE_ROLE, PRINTOUT_ROLE]
    assert all(f["sha256"] and f["file_id"] for f in v.files)
    assert [f["name"] for f in v.proof["files"]] == [f["name"] for f in v.files]
    built = json.loads((r.data / "book-percase.json").read_text())
    assert built["cases"] == 1 and built["defendant_rows"] == 1 and built["leaks"] == []


def test_the_leak_gate_scans_the_per_case_copy(tmp_path, monkeypatch):
    seat = LitSeat({M1: matter_docs()})
    real = percase.facts
    monkeypatch.setattr(percase, "facts", lambda m, role: [*real(m, role), ("Leak", "see results.json")])
    r = _run(tmp_path, seat)
    v = r.run()
    assert v.verdict == "failed" and v.reason.startswith("leak_found:"), v.reason
    assert v.stage == "book" and seat.sent == [] and not percase.path_for(r._out()).exists()
    assert gates.leak_scan(r._out()) == []  # the list itself was clean; the copy is what was caught


def test_a_review_hold_on_the_per_case_copy_files_neither(tmp_path, monkeypatch):
    seat = LitSeat({M1: matter_docs()})
    monkeypatch.setattr(percase, "check", lambda xlsx, ms: ["planted"])
    v = _run(tmp_path, seat).run()
    assert v.verdict == "failed" and v.reason.startswith("review_held:"), v.reason
    assert seat.sent == []


def test_a_held_main_list_files_no_per_case_copy(tmp_path, monkeypatch):
    seat = LitSeat({M1: matter_docs()})
    real = review.findings

    def drifted(xlsx, finals, in_scope, tagged=None):
        finals = [{**m, "defendants": [{**d, "name": "Someone Else"} for d in m["defendants"]]} for m in finals]
        return real(xlsx, finals, in_scope, tagged)

    monkeypatch.setattr(review, "findings", drifted)
    v = _run(tmp_path, seat).run()
    assert v.verdict == "failed" and v.reason.startswith("review_held:") and seat.sent == []


def test_a_refused_main_filing_files_no_per_case_copy(tmp_path):
    seat = LitSeat({M1: matter_docs()})
    seat.folders.append({"id": "fold-x", "name": "Litigation Status", "parentId": None, "path": "Litigation Status"})
    name = f"Litigation status {TODAY.isoformat()} (Operator {JOB[-6:]}).xlsx"
    seat.docs.append(doc_row("someone-elses", name, "fold-x", 10))
    v = _run(tmp_path, seat).run()
    assert v.verdict == "held" and v.reason.startswith("filing_refused:"), v.reason
    assert seat.sent == []


@pytest.mark.parametrize("last", ["book", "review"])
def test_a_job_built_before_the_copy_existed_gets_it_on_resume(tmp_path, last):
    """A job in flight across the release: its book (or review) stage ran on
    the old code, so the copy is not on disk when it resumes."""
    seat = LitSeat({M1: matter_docs()})
    r = _run(tmp_path, seat)
    stages = ("inventory", "diff", "fetch", "extract", "screen", "read1", "read2", "read3", "gates", "parity", "book")
    for stage in (*stages, "review") if last == "review" else stages:
        r._stage(stage, getattr(r, f"_{stage}"))
    percase.path_for(r._out()).unlink()  # as a stage that ran on the old code left it
    printout.path_for(r._out()).unlink()
    v = _run(tmp_path, seat).run()
    assert v.verdict == "delivered", v.reason
    assert [f["role"] for f in v.files] == ["workbook", PER_CASE_ROLE, PRINTOUT_ROLE]
    assert [s["name"] for s in seat.sent] == [
        r._out().name,
        percase.path_for(r._out()).name,
        printout.path_for(r._out()).name,
    ]
