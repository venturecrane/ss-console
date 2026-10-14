"""Synthetic fixtures for an attorney's own NUMBERED demand format
(``layout: numbered``, house_numbered.py): a content-free reference built here
with python-docx (never a firm's file), the variant, the attorney row, and a
draft in that format.

Every name, number and sentence is invented, and the draft's facts are the
ones ``demand_testkit.standard_docs`` carries, so the drafting gate passes it.
"""

from __future__ import annotations

import copy
import hashlib
from pathlib import Path
from typing import Any

import yaml

from demand_testkit import make_inputs

# kind -> (alignment, bold, italic, list level or None)
PARA = {
    "date": ("left", False, False, None),
    "mail": ("left", True, False, None),
    "addr": ("left", False, False, None),
    "banner": ("center", True, False, None),
    "banner_sub": ("center", True, False, None),
    "banner_note": ("center", False, True, None),
    "salute": ("left", False, False, None),
    "lead": ("center", True, False, None),
    "section": ("left", True, False, None),
    "subhead": ("left", True, False, None),
    "body": ("justify", False, False, None),
    "callout": ("justify", True, True, None),
    "quote": ("justify", False, True, None),
    "note": ("right", False, True, None),
    "bullet": ("justify", False, False, "bullet"),
    "bullet2": ("justify", False, False, "bullet"),
    "numbered": ("left", False, False, "number"),
    "sign": ("left", False, False, None),
    "signfirm": ("left", False, False, None),
    "signscript": ("left", False, True, None),
    "signname": ("left", False, False, None),
    "exh_head": ("left", True, False, None),
    "exh_intro": ("left", False, True, None),
    "exh_sub": ("left", True, False, None),
    "exh_item": ("left", True, False, None),
}
# kind -> (columns, rows)
TABLES = {
    "re": (2, 1),
    "kv": (2, 1),
    "summary": (2, 3),
    "grid": (4, 3),
    "chronology": (4, 3),
    "specials": (3, 3),
    "future": (3, 3),
    "causation": (3, 3),
    "exhibits": (4, 3),
}
BANNER_LINES = ["Pursuant to Example Code Section 999 et seq.", "Confidential Settlement Communication"]
FOOTER = "Time-Limited Demand (Example Code 999)  |  {name}  |  Claim No. {claim}  |  Page "
SIGNER, TITLE = "Gamma Q. Counsel", "Attorney for {client}"
HEADINGS = [
    "Damages",
    "Introduction",
    "Facts and Liability",
    "Current Medical Specials",
    "Demand",
    "The Duty to Accept",
]
LEADS = ["Amount.", "Release.", "Time for acceptance."]


def _num_id(doc: Any, style: str) -> str:
    """The numId the default template's list style points at."""
    from docx.oxml.ns import qn

    st = doc.styles[style].element
    return st.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))


def _numbered(p: Any, num_id: str, level: int) -> None:
    from docx.oxml.ns import qn

    ppr = p._p.get_or_add_pPr()
    numpr = ppr.makeelement(qn("w:numPr"), {})
    ilvl = numpr.makeelement(qn("w:ilvl"), {qn("w:val"): str(level)})
    nid = numpr.makeelement(qn("w:numId"), {qn("w:val"): num_id})
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)


def make_numbered_reference(path: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    align = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    d = docx.Document()
    lists = {"bullet": _num_id(d, "List Bullet"), "number": _num_id(d, "List Number")}
    for kind, (al, bold, italic, lst) in PARA.items():
        p = d.add_paragraph(f"PROTO:{kind}")
        p.alignment = align[al]
        p.runs[0].bold, p.runs[0].italic = bold or None, italic or None
        if lst:
            _numbered(p, lists[lst], 1 if kind == "bullet2" else 0)
    for kind, (cols, rows) in TABLES.items():
        t = d.add_table(rows=rows, cols=cols)
        t.cell(0, 0).text = f"PROTO:{kind}"
        if kind in ("re", "kv"):
            t.cell(0, 0).paragraphs[0].runs[0].bold = True
    footer = d.sections[0].footer.paragraphs[0]
    footer.text = FOOTER
    d.save(str(path))


def variant(**over: Any) -> dict[str, Any]:
    return {
        "skeleton": "skeleton_numbered",
        "house_reference": "house_reference_numbered",
        "layout": "numbered",
        "headings": HEADINGS,
        "re_required": ["Re:", "Your Insured:", "Our Client:", "Claim No.:", "Date of Loss:"],
        "re_one_of": [],
        "mail_line": "CERTIFIED MAIL, RETURN RECEIPT REQUESTED",
        "banner": True,
        "forbidden": ["Attn:"],
        "re_vs": False,
        "banner_lines": BANNER_LINES,
        "lead_heading": "CASE DISTILLED",
        "footer_markers": ["Time-Limited Demand (Example Code 999)", "Claim No."],
        "required_leads": LEADS,
        **over,
    }


SKELETON = "---\ndate: [INSERT]\n---\n## Damages\n## Introduction\n## Demand\n## Enclosures: Index of Exhibits\n"


def make_numbered_inputs(root: Path, attorney: str = "Gamma Counsel", **row: Any) -> Path:
    """The standard firm inputs plus a ``numbered`` format and an attorney row
    that selects it. The default format and its attorney are untouched."""
    make_inputs(root)
    ref = root / "house-reference-numbered.docx"
    make_numbered_reference(ref)
    sk = root / "skeleton-numbered.md"
    sk.write_text(SKELETON, encoding="utf-8")
    cfg = root / "demand-firm.yaml"
    data = yaml.safe_load(cfg.read_text(encoding="utf-8"))
    data["inputs"]["house_reference_numbered"] = {
        "path": ref.name,
        "sha256": hashlib.sha256(ref.read_bytes()).hexdigest(),
    }
    data["inputs"]["skeleton_numbered"] = {"path": sk.name, "sha256": hashlib.sha256(sk.read_bytes()).hexdigest()}
    data["formats"] = {"numbered": {"pre_suit": variant()}}
    data["attorneys"][attorney] = {"signer": SIGNER, "title": TITLE, "initials": "", "format": "numbered", **row}
    cfg.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True), encoding="utf-8")
    return root


def edit_config(root: Path, fn: Any) -> None:
    cfg = root / "demand-firm.yaml"
    data = yaml.safe_load(cfg.read_text(encoding="utf-8"))
    data = fn(copy.deepcopy(data)) or data
    cfg.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True), encoding="utf-8")


NUMBERED_DRAFT = """---
date: October 6, 2026
mail: VIA CERTIFIED MAIL, RETURN RECEIPT REQUESTED
attn: Pat Adjuster | Example Mutual | 1 Example Way | Exampletown, CA 90000
re: Alpha Example v. Beta Driver
insured: Beta Driver, and all other insureds under the policy
client: Alpha Example
claim: CLM-0001
dol: January 15, 2026
location: {{NOT IN RECORD: the place of the collision; looked in the matter file, see gap audit item 1}}
banner: TIME-LIMITED POLICY LIMIT DEMAND
salutation: Dear Pat Adjuster:
footer_name: Example
---

## A Collision on January 15, 2026

>> Alpha Example was seen at Exampletown ER on January 15, 2026 for neck pain after a collision (ER record 1-15-26, p. 1).

## Damages

|  |  |
|---|---|
| Current medical specials | $1,200.00 |
| General damages | In excess of policy limits |

## Introduction

This firm represents Alpha Example. Your insured, Beta Driver, is named in the carrier's letter under claim CLM-0001 (Carrier letter of acknowledgment, p. 1).

| Item | Summary |
|---|---|
| Current medical specials | $1,200.00 |

## Facts and Liability

The evidence places fault on your insured:

- Beta Driver is named as the insured under claim CLM-0001 (Carrier letter of acknowledgment, p. 1).

## Treatment

1. **January 15, 2026.** The emergency department diagnosed a cervical strain (ER record 1-15-26, p. 1).
  - The itemized statement lists total charges of $1,200.00 (ER bill 1-15-26, p. 1).
2. **Billing.** Exampletown ER, $1,200.00 (ER bill 1-15-26, p. 1).

Between the lists.

1. **A second list.** It restarts at one.

## Current Medical Specials

| Provider | Amount Billed | Net of Adjustments |
|---|---|---|
| Exampletown ER | $1,200.00 | $1,200.00 |
| TOTAL | $1,200.00 | $1,200.00 |

^ Source: ER bill 1-15-26.

## Demand

Alpha Example offers to settle on the following terms:

- **Amount.** Payment of the full available policy limits.
- **Release.** A complete release of your insureds.
- **Time for acceptance.** {{FILL: the acceptance date}}

## The Duty to Accept

> Demand is hereby made for the full available policy limits.

## Enclosures: Index of Exhibits

Exhibit list

| Exh. | Provider / Description | Dates of Service | Contents |
|---|---|---|---|
| 1 | ER record 1-15-26 | 01/15/2026 | Medical Records |

=== ATTORNEY NOTES ===

## HELD OUT PENDING ATTORNEY PRIVILEGE REVIEW

None.
"""
