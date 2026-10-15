"""The deadlines report (``medchron/litigation/deadlines/``): a $0 litigation
job kind that reads the calendar, tasks and statute fields of the job's
matters and files two workbooks to the library matter.

Every name, number and date here is invented. The parity with the hand kit
that delivered the first report is proven on real data offline (the kit is
client material and does not live in this repo); these tests pin the rules
the parity rests on, the falsifier (it must be able to fail), and the job's
shape end to end against a fake seat."""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path
from typing import Any

import pytest

from litigation_testkit import JOB, LIBRARY, M1, M2, STAFF_A, STAFF_B, LitSeat, make_inputs, make_job
from medchron.litigation.deadlines import book, run as run_mod
from medchron.litigation.deadlines.pull import in_scope
from medchron.litigation import job as job_mod

TODAY = dt.date(2026, 3, 10)
M3 = "33333333-3333-4333-8333-333333333333"
PARA = "cccccccc-cccc-4ccc-8ccc-cccccccccccc"
SOL = "Matter/CaseDetails/StandardCaseDetails/StatuteOfLimitationDate"
FILED = "Matter/CaseDetails/PleadingsDetails/FiledDate"


def ev(day: str, subject: str, **over: Any) -> dict[str, Any]:
    return {"startTime": f"{day}T09:30:00", "subject": subject, "attendees": [{"id": STAFF_A}], **over}


#: Matter id -> (list row, layouts values, events, tasks, file names).
WORLD: dict[str, Any] = {
    M1: (
        {
            "id": M1,
            "number": "100001",
            "title": "100001 - Gamma v. Delta",
            "personResponsibleStaffId": STAFF_A,
            "personAssistingStaffId": PARA,
        },
        {},
        [
            ev("2026-03-12", "Mediation with Example Mediator"),
            ev("2026-03-15", "Responses to SROGs set one due"),
            ev("2026-03-16", "HIDDEN rule date"),
            ev("2026-03-20", "Deleted CMC", isDeleted=True),
            ev("2026-03-01", "CMC"),
            ev("2026-06-01", "Trial setting conference", additionalData={"LawToolBox": {}}),
        ],
        [{"subject": "Calendar the CMC", "dueDate": "2026-02-01", "assignees": [{"id": STAFF_A}]}],
        ["Complaint for damages.pdf", "Medical bill.pdf"],
    ),
    M2: (
        {"id": M2, "number": "100002", "title": "100002 - Epsilon intake", "personResponsibleStaffId": STAFF_B},
        {SOL: "2026-04-01T00:00:00"},
        [],
        [],
        ["Intake form.pdf"],
    ),
    M3: (
        {"id": M3, "number": "100003", "title": "Zeta v. Eta", "personResponsibleStaffId": STAFF_A},
        {SOL: "2025-12-01", FILED: "2025-11-20"},
        [ev("2026-04-02", "Expert exchange")],
        [],
        [],
    ),
}


class Client:
    def __init__(self, world: dict[str, Any], pending: tuple[str, ...] = (M3,)) -> None:
        self.world, self.pending, self.calls = world, pending, []

    def get(self, path: str, **params: Any) -> Any:
        self.calls.append((path, params))
        if path == "/matters":
            ids = [m for m in self.world if (m in self.pending) == (params["Status"] == "Pending")]
            return {"value": [self.world[m][0] for m in ids] if not params.get("Offset") else []}
        if path == "/staff":
            return {"value": [{"id": STAFF_A, "firstName": "Alpha", "lastName": "Example"},
                              {"id": STAFF_B, "firstName": "Beta", "lastName": "Example"},
                              {"id": PARA, "firstName": "Para", "lastName": "Example"}]}  # fmt: skip
        mid = params.get("MatterId") or path.split("/")[2]
        _row, layout, events, tasks, names = self.world[mid]
        if path.endswith("/layouts"):
            return {"value": [{"values": [{"key": k, "value": v} for k, v in layout.items()]}]}
        if path == "/events":
            return {"value": events}
        if path == "/tasks":
            return {"value": tasks}
        if path.endswith("/documents/files"):
            return {"value": [{"id": f"{mid[:4]}-{i}", "name": n} for i, n in enumerate(names)]}
        raise AssertionError(path)


class Seat(LitSeat):
    def __init__(self, world: dict[str, Any]) -> None:
        super().__init__({})
        self.client = Client(world)


def deadline_job(tmp_path: Path, **over: Any) -> Path:
    fields = {"report": "deadlines", "workbooks": ["deadlines", "statute"], "folder_name": "Deadlines",
              "request_text": "Please send me the deadlines report and statute report.",
              "scope": {"attorney_staff_ids": [STAFF_A, STAFF_B]}}  # fmt: skip
    fields.update(over)
    return make_job(tmp_path / "job", **fields)


def make_run(tmp_path: Path, world: dict[str, Any] | None = None, **over: Any) -> tuple[run_mod.DeadlinesRun, Seat]:
    seat = Seat(world or WORLD)
    jd = deadline_job(tmp_path, **over)
    r = run_mod.DeadlinesRun(
        jd, inputs_dir=str(make_inputs(tmp_path / "inputs")), seat_factory=lambda: seat, log=lambda m: None,
        today=lambda: TODAY, readback_pause=0.0,
    )  # fmt: skip
    return r, seat


# ---- the envelope ---------------------------------------------------------------------


def test_a_status_envelope_is_unchanged_and_a_deadlines_one_names_its_workbooks(tmp_path: Path) -> None:
    assert job_mod.load(make_job(tmp_path / "a")).report == "status"
    j = job_mod.load(deadline_job(tmp_path, workbooks=["statute"]))
    assert (j.report, j.workbooks) == ("deadlines", ("statute",))
    with pytest.raises(job_mod.LitigationJobError):
        job_mod.load(deadline_job(tmp_path / "b", workbooks=[]))
    with pytest.raises(job_mod.LitigationJobError):
        job_mod.load(make_job(tmp_path / "c", workbooks=["deadlines"]))


def test_scope_takes_attorneys_paralegals_or_both(tmp_path: Path) -> None:
    j = job_mod.load(deadline_job(tmp_path, scope={"attorney_staff_ids": [STAFF_A], "paralegal_staff_ids": [PARA]}))
    rows = {m: WORLD[m][0] for m in WORLD}
    assert [m for m, r in rows.items() if in_scope(r, j)] == [M1]


# ---- the rules ----------------------------------------------------------------------


def _rules(staff: dict[str, str] | None = None) -> book.Rules:
    import re

    return book.Rules(staff or {}, [re.compile(r"\bcomplaint\b", re.I)], 1, TODAY)


def test_litigation_is_a_court_paper_or_a_filed_date_or_a_case_number() -> None:
    r = _rules()
    assert r.is_lit({"fileNames": ["Complaint for damages.pdf"]})
    assert r.is_lit({"filed": "2025-01-01"}) and r.is_lit({"caseNumber": "CV-1"})
    assert not r.is_lit({"fileNames": ["Intake form.pdf"]})


def test_event_times_are_written_as_given_and_categories_follow_the_subject() -> None:
    assert book.ev_time({"startTime": "2026-03-12T13:05:00"}) == "1:05p"
    assert book.ev_time({"startTime": "2026-03-12T00:00:00"}) == ""
    assert book.ev_time({"startTime": "2026-03-12T09:30:00", "allDay": True}) == ""
    assert book.category("Mediation with X") == "MEDIATION"
    assert book.category("Responses to SROGs due") == "DISCOVERY"
    assert book.category("CMC") == "CMC"
    assert book.category("Trial setting conference") == "TRIAL & MSC"
    assert book.category("Call client") == "OTHER"


def test_the_build_drops_deleted_dates_and_sorts_hidden_ones_last(tmp_path: Path) -> None:
    r, _seat = make_run(tmp_path)
    v = r.run()
    assert v.verdict == "delivered", v.reason
    counts = json.loads((r.data / "book.json").read_text())["counts"]
    assert counts["litigation"] == 2  # M1 by its complaint, M3 by its Filed date; M2 is not in litigation
    assert counts["dates_next_14"] == 3 and counts["dates_last_14"] == 1  # the deleted CMC is gone
    assert counts["tasks_overdue"] == 1
    assert counts["statute_upcoming"] == 1  # M2: statute set, nothing filed; M3 has a Filed date
    from openpyxl import load_workbook

    ws = load_workbook(r.data / "out" / book.file_names(TODAY, JOB[-6:])["deadlines"])["Next 14 days"]
    subjects = [row[5] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert subjects[-1] == "HIDDEN rule date"


# ---- the falsifier --------------------------------------------------------------------


def test_the_falsifier_fires_on_a_rule_date_outside_the_litigation_test(tmp_path: Path) -> None:
    world = dict(WORLD)
    row, layout, events, tasks, names = world[M2]
    world[M2] = (row, layout, [ev("2026-05-01", "Hearing", additionalData={"LawToolBox": {"id": 1}})], tasks, names)
    r, seat = make_run(tmp_path, world)
    v = r.run()
    assert v.verdict == "held" and v.reason.startswith("gate_refused: coverage: 1 matters")
    assert seat.sent == []  # nothing filed
    assert json.loads((r.data / "check.json").read_text())["coverage_misses"] == ["100002"]


def test_the_falsifier_is_quiet_on_the_clean_world() -> None:
    import re

    r = book.Rules({}, [re.compile(r"\bcomplaint\b", re.I)], 1, TODAY)
    clean = [{"number": "1", "fileNames": [], "events": [ev("2026-05-01", "x")]}]
    assert book.coverage_misses(clean, r) == []
    past = [{"number": "1", "fileNames": [], "events": [ev("2026-01-01", "x", additionalData={"LawToolBox": {}})]}]
    assert book.coverage_misses(past, r) == []


def test_the_review_holds_a_workbook_whose_bytes_disagree_with_the_build(tmp_path: Path) -> None:
    r, _seat = make_run(tmp_path)
    r._pull()
    r._book()
    built = json.loads((r.data / "book.json").read_text())
    built["sheets"]["deadlines"]["CASES"] += 1
    (r.data / "book.json").write_text(json.dumps(built))
    findings = run_mod.review((r.data / "out" / built["files"]["deadlines"]).read_bytes(), built["sheets"]["deadlines"])
    assert findings == ["sheet 'CASES': 2 rows, built 3"]


# ---- the job ------------------------------------------------------------------------


def test_the_job_files_both_workbooks_reads_them_back_and_costs_nothing(tmp_path: Path) -> None:
    r, seat = make_run(tmp_path)
    v = r.run()
    assert v.verdict == "delivered" and v.cents == 0
    names = book.file_names(TODAY, JOB[-6:])
    assert [s["name"] for s in seat.sent] == [names["deadlines"], names["statute"]]
    assert [c["name"] for c in seat.created] == ["Deadlines"]
    assert {f["role"] for f in v.files} == {"workbook", "statute"}
    assert v.proof["readback"] == "confirmed" and v.proof["review"] == "code_only"
    assert v.counts["matters_read"] == 3 and v.matters_total == 3
    # read-only otherwise: the only writes are the folder and the two files
    assert all(p in ("/matters", "/staff", "/events", "/tasks") or p.endswith(("/layouts", "/documents/files"))
               for p, _ in seat.client.calls)  # fmt: skip


def test_only_the_named_workbook_is_built(tmp_path: Path) -> None:
    r, seat = make_run(tmp_path, workbooks=["statute"])
    v = r.run()
    assert v.verdict == "delivered"
    assert [s["name"] for s in seat.sent] == [book.file_names(TODAY, JOB[-6:])["statute"]]
    assert v.files[0]["role"] == "workbook"


def test_a_proof_run_stops_before_filing(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("MEDCHRON_PROOF_RUN", "1")
    r, seat = make_run(tmp_path)
    v = r.run()
    assert v.verdict == "dry_run" and seat.sent == [] and seat.created == []
    assert len(v.files) == 2


def test_a_matter_that_cannot_be_read_fails_resumably_and_files_nothing(tmp_path: Path) -> None:
    r, seat = make_run(tmp_path)
    real = seat.client.get

    def flaky(path: str, **params: Any) -> Any:
        if path == "/events" and params.get("MatterId") == M3:
            raise RuntimeError("timeout")
        return real(path, **params)

    seat.client.get = flaky  # type: ignore[method-assign]
    v = r.run()
    assert v.verdict == "failed" and v.reason.startswith("read_incomplete: 1 of 3")
    assert seat.sent == []
    seat.client.get = real  # type: ignore[method-assign]
    before = len(seat.client.calls)
    v = r.run()  # the resume reads only the one matter that is missing
    assert v.verdict == "delivered"
    again = [prm.get("MatterId") for p, prm in seat.client.calls[before:] if p == "/events"]
    assert again == [M3]


def test_no_matter_in_scope_is_held_as_nothing_in_scope(tmp_path: Path) -> None:
    r, seat = make_run(tmp_path, scope={"attorney_staff_ids": ["dddddddd-dddd-4ddd-8ddd-dddddddddddd"]})
    v = r.run()
    assert v.verdict == "held" and v.reason.startswith("scope_empty:")


def test_the_library_matter_is_the_only_place_written(tmp_path: Path) -> None:
    r, seat = make_run(tmp_path)
    r.run()
    assert {s["folderId"] for s in seat.sent} == {seat.created[0]["id"]}
    assert r.job.file_to_id == LIBRARY


# ---- the lane: its proof key and its wake --------------------------------------------

DEADLINES_RUNNER = """
import json, sys, pathlib
jd = pathlib.Path(sys.argv[1])
names = ["Litigation deadlines 10-15-26 (Operator 0000AA).xlsx", "Statute report 10-15-26 (Operator 0000AA).xlsx"]
v = {"verdict": "delivered", "stage": "report", "cents": 0, "matters_total": 601, "matters_reread": 176, "flags_new": 0,
     "files": [{"role": "workbook", "name": names[0], "size": 9, "sha256": "a" * 64, "file_id": "up-1"},
               {"role": "statute", "name": names[1], "size": 7, "sha256": "b" * 64, "file_id": "up-2"}],
     "folder_id": "folder-7", "counts": {"matters_read": 601, "litigation": 176, "dates_next_14": 42},
     "proof": {"readback": "confirmed",
               "files": [{"name": n, "file_id": "f", "size": 1, "sha256": "a" * 64} for n in names],
               "review": "code_only", "review_cents": 0}}
(jd / 'verdict.json').write_text(json.dumps(v))
print(json.dumps(v))
"""


def test_the_deadlines_report_is_its_own_lane_key() -> None:
    from medchron.lane_proof import lane_key

    assert lane_key("litigation", {"report": "deadlines"}) == "litigation-status.deadlines"
    assert lane_key("litigation", {}) == "litigation-status"


def _deadlines_lane(tmp_path: Path, script: str) -> Any:
    from test_litigation_lane import _lane, _submit

    lane, client = _lane(tmp_path, script=script)
    _submit(lane, client, report="deadlines", workbooks=["deadlines", "statute"], folder_name="Deadlines")
    return lane, client


def test_the_deadlines_wake_carries_counts_and_the_composed_file_names_only(tmp_path: Path) -> None:
    lane, client = _deadlines_lane(tmp_path, DEADLINES_RUNNER)
    assert lane.tick() == "delivered"
    assert client.records[-1][2]["counts"] == {"matters_read": 601, "litigation": 176, "dates_next_14": 42}
    task = lane._daemon_state(JOB)["wake"]["task"]
    assert task.splitlines() == [
        f"Run the litigation-status skill's DELIVER mode for litigation job {JOB}.",
        "Kind: litigation. Report: deadlines.",
        "Trigger: request.",
        "Outcome: delivered.",
        "Counts: dates_next_14 42; litigation 176; matters_read 601.",
        "Folder: Deadlines.",
        "Files: Litigation deadlines 10-15-26 (Operator 0000AA).xlsx; Statute report 10-15-26 (Operator 0000AA).xlsx.",
        "Requested by: admin@firm.example.",
    ]


def test_a_file_name_the_code_did_not_compose_never_reaches_the_wake(tmp_path: Path) -> None:
    lane, _client = _deadlines_lane(tmp_path, DEADLINES_RUNNER.replace("Statute report 10-15-26", "Gamma v Delta"))
    lane.tick()
    task = lane._daemon_state(JOB)["wake"]["task"]
    assert "Gamma" not in task
