"""Parity against the last list, and the workbook: its tabs, its responsible
column, clean(), and a leak scan that passes on what book builds."""

from __future__ import annotations

import copy
import datetime as dt

from openpyxl import load_workbook

from medchron.litigation import book, gates, parity, vocab
from test_litigation_gates import good

TODAY = dt.date(2026, 10, 7)


def test_an_overturn_that_cites_a_document_is_explained():
    prior = good()
    new = copy.deepcopy(prior)
    new["defendants"][0]["served"]["date"] = "2026-04-09"
    ov = [{"path": "defendants[Delta Example].served", "source": {"file_id": "f1"}}]
    cs = parity.compare(prior, new, moved_files=set(), overturns=ov, today=TODAY)
    assert cs and not parity.unexplained(cs)


def test_the_same_papers_with_a_different_answer_hold():
    prior = good()
    new = copy.deepcopy(prior)
    new["defendants"][0]["status"] = vocab.SERVED_NO_ANSWER
    new["defendants"][0]["answered"] = {"date": None, "source": None}
    bad = parity.unexplained(parity.compare(prior, new, moved_files=set(), overturns=[], today=TODAY))
    assert {c["path"] for c in bad} >= {"defendants[deltaexample].answered"}


def test_a_new_document_explains_its_value_and_a_passed_court_date_moves_on():
    prior = good()
    prior["next_court_date"]["date"] = "2026-09-01"
    new = copy.deepcopy(prior)
    new["next_court_date"] = {"date": "2026-12-01", "source": {"file_id": "fnew"}}
    cs = parity.compare(prior, new, moved_files=set(), overturns=[], today=TODAY)
    assert not parity.unexplained(cs)
    new["case_status"] = {"value": vocab.SETTLED_OWED, "source": {"file_id": "fmail"}}
    assert not parity.unexplained(parity.compare(prior, new, moved_files={"fmail"}, overturns=[], today=TODAY))
    assert parity.unexplained(parity.compare(prior, new, moved_files=set(), overturns=[], today=TODAY))


def test_a_new_matter_is_one_explained_change():
    cs = parity.compare(None, good(), moved_files=set(), overturns=[], today=TODAY)
    assert cs == [{"path": "matter", "old": None, "new": "added to the list", "explained": True, "why": "new matter"}]


def test_the_workbook_has_every_tab_and_the_rosters_name(tmp_path):
    a = good()
    b = copy.deepcopy(good())
    b.update(matter_id="m2", number="100002", responsible="Beta Example")
    b["case_status"]["value"] = vocab.SETTLED_OWED
    b["defendants"][0]["status"] = vocab.SERVED_NO_ANSWER
    b["defendants"][0]["answered"] = {"date": None, "source": None}
    b["discovery_served_on_client"] = [
        {
            "set": "Special interrogatories, set one",
            "served_by": "Delta Example",
            "date": "2026-07-01",
            "source": {"file_id": "f9", "name": "SROGS.pdf", "doc_date": None},
            "responses_served": {"value": "no", "date": None, "source": None},
        }
    ]
    c = copy.deepcopy(good())
    c.update(matter_id="m3", number="100003", responsible="Alpha Example", defendants=[])
    c["case_status"]["value"] = vocab.NOT_FILED
    c["notes"] = "Demand stage; no complaint (second read confirmed, see 4de824d4 results.json)"
    out = tmp_path / "w.xlsx"
    stats = book.build(
        [a, b, c],
        {
            "m1": [
                {
                    "path": "defendants[Delta Example].served",
                    "old": "2026-04-09",
                    "new": "2026-04-10",
                    "explained": True,
                    "why": "x",
                }
            ]
        },
        out,
        TODAY,
    )
    wb = load_workbook(out)
    assert wb.sheetnames == [
        "Summary",
        "Needs action",
        "Settled - dismissal owed",
        "All defendants",
        "Discovery",
        "Case notes",
        "Not filed",
        "What changed",
    ]
    assert stats["needs_action"] == 0  # the settled matter's defendant is owed a dismissal, not an answer
    assert [r[0].value for r in wb["Summary"].iter_rows(min_row=2)] == ["Alpha Example", "Beta Example"]
    assert wb["Settled - dismissal owed"].max_row == 2
    disc = [[c.value for c in r] for r in wb["Discovery"].iter_rows(min_row=2)]
    assert any(r[3] == "Served on the firm's client" and r[8] == "no" for r in disc)
    assert any(r[3] == "Propounded by the firm" for r in disc)
    nf = wb["Not filed"]["D2"].value
    assert "second read" not in nf and "4de824d4" not in nf and "results.json" not in nf
    assert wb["What changed"]["D2"].value == "Delta Example: served"
    assert gates.leak_scan(out) == []


def test_needs_action_lists_an_unanswered_defendant_in_an_active_case(tmp_path):
    a = good()
    a["defendants"][0].update(status=vocab.SERVED_NO_ANSWER, answered={"date": None, "source": None})
    assert book.build([a], {}, tmp_path / "w.xlsx", TODAY)["needs_action"] == 1


def test_clean_strips_internal_vocabulary():
    s = book.clean(
        "Answer filed (second read: verifier saw it) see [doc 4] id 11111111-1111-4111-8111-111111111111 — ok"
    )
    for bad in ("second read", "verifier", "[doc", "11111111-", "—"):
        assert bad not in s


# ---- the sent-out date (2026-10-24) ---------------------------------------------------
def _cells(d):
    m = good()
    m["defendants"] = [d]
    return dict(zip(book.DEF_HEADER, book.defendant_rows([m])[0]["cells"]))


def test_the_sent_out_columns_follow_served_and_leave_the_sort_keys_alone():
    h = book.DEF_HEADER
    assert h[9:13] == ["Served", "Served (source)", "Sent out for service", "Sent out for service (source)"]
    assert h[:3] == ["Responsible (Smokeball)", "Next court date", "Matter #"]
    assert len(book.DEF_WIDTHS) == len(h)


def test_a_sent_date_renders_with_its_document():
    c = _cells(good()["defendants"][0])
    assert c["Sent out for service"] == "2026-03-20"
    assert c["Sent out for service (source)"] == "Service request.pdf"
    assert c["Served"] == "2026-04-10 (personal)"


def test_an_unread_sent_date_is_not_read_yet_never_not_in_file():
    """An absent key is a value nobody looked for: printing "Not in file"
    would state a finding nobody made."""
    d = good()["defendants"][0]
    del d["sent_for_service"]
    c = _cells(d)
    assert c["Sent out for service"] == "Not read yet"
    assert c["Sent out for service (source)"] == ""


def test_a_read_that_found_no_request_says_not_in_file_and_why():
    d = good()["defendants"][0]
    d["sent_for_service"] = {"date": None, "source": None}
    c = _cells(d)
    assert c["Sent out for service"] == "Not in file"
    assert c["Sent out for service (source)"] == book.NOT_IN_FILE_WHY


def test_the_built_workbook_carries_the_sent_out_columns_and_review_checks_them(tmp_path):
    from medchron.litigation import review

    a = good()
    b = copy.deepcopy(good())
    b.update(matter_id="m2", number="100002")
    del b["defendants"][0]["sent_for_service"]
    out = tmp_path / "w.xlsx"
    book.build([a, b], {}, out, TODAY)
    ws = load_workbook(out)["All defendants"]
    col = [c.value for c in ws[1]].index("Sent out for service")
    assert [r[col].value for r in ws.iter_rows(min_row=2)] == ["2026-03-20", "Not read yet"]
    assert gates.leak_scan(out) == []
    scope = [{"number": "100001"}, {"number": "100002"}]
    assert review.findings(out.read_bytes(), [a, b], scope) == []
    # The review compares All defendants cell by cell, so the new columns too.
    a2 = copy.deepcopy(a)
    a2["defendants"][0]["sent_for_service"]["date"] = "2026-03-19"
    assert [f for f in review.findings(out.read_bytes(), [a2, b], scope) if f.check == "rows"]


def test_a_new_sent_date_with_a_source_is_explained_and_a_vanished_one_holds():
    prior = good()
    del prior["defendants"][0]["sent_for_service"]
    new = good()
    cs = parity.compare(prior, new, moved_files=set(), overturns=[], today=TODAY)
    assert [c["path"] for c in cs] == ["defendants[deltaexample].sent_for_service"] and not parity.unexplained(cs)
    gone = parity.unexplained(parity.compare(new, prior, moved_files=set(), overturns=[], today=TODAY))
    assert [c["path"] for c in gone] == ["defendants[deltaexample].sent_for_service"]
    assert book.humanize("defendants[Delta Example].sent_for_service") == "Delta Example: sent for service"


# ---- the court-only fold-in (2026-10-23) ----------------------------------------------
def test_discovery_paths_name_every_discovery_parity_point_and_nothing_else():
    m = good()
    m["discovery_served_on_client"] = [
        {
            "set": "Special interrogatories, set one",
            "served_by": "Delta Example",
            "date": "2026-07-01",
            "source": {"file_id": "f9"},
            "responses_served": {"value": "no", "date": None, "source": None},
        }
    ]
    pts = parity.points(m)
    disc = {p for p in pts if p.startswith(vocab.DISCOVERY_PATHS)}
    assert len(disc) == 3 and all(p.startswith("discovery_") for p in disc)
    assert all(not p.startswith("discovery_") for p in set(pts) - disc)


def test_clean_drops_a_correction_lead_in():
    assert book.clean("Corrected on double-check: served 4/10/2026") == "served 4/10/2026"
    assert book.clean("corrected on review: no answer filed") == "no answer filed"
