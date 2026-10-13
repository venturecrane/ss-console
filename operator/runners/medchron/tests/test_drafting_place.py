"""The attorney's own sets, attached to his request, placed in the court's
caption on the firm's pleading paper (drafting/place.py, placing.py).

Every name and sentence is invented (Example / Exampletown / CV-0001)."""

from __future__ import annotations

import base64
import hashlib
import io
import json
from pathlib import Path

import pytest

from drafting_testkit import ScriptedClient, make_inputs, make_job, seat_with, standard_docs
from medchron_testkit import PRICING
from medchron.drafting import format_check, job as job_mod, place, placing, pleading
from medchron.drafting.run import DraftingRun

CHECKER = Path(__file__).resolve().parents[3] / "templates" / "drafting" / "drafting_gate_check.py"
REQUEST = "Attached are my three sets. Please put each one into the caption on pleading paper and file them."
LEAD = {
    "email": "admin@firm.example",
    "block": [
        "Alpha Example, Esq. (SBN 000001)",
        "EXAMPLE & EXAMPLE, LLP",
        "1 Example Way",
        "Exampletown, California 90000",
    ],
}
KINDS = {
    "special": (
        "SPECIAL INTERROGATORIES",
        "SPECIAL INTERROGATORY",
        "State each fact supporting YOUR contention that the floor was dry.",
    ),
    "admission": ("REQUESTS FOR ADMISSION", "REQUEST FOR ADMISSION", "Admit that YOU mopped the floor at 10:00 a.m."),
    "production": (
        "REQUESTS FOR PRODUCTION",
        "REQUEST FOR PRODUCTION",
        "All video of the aisle on the day of the fall.",
    ),
}


@pytest.fixture(autouse=True)
def _checker(monkeypatch):
    monkeypatch.setenv("SMD_DRAFTING_GATE_CHECK", str(CHECKER))


@pytest.fixture
def pricing(tmp_path: Path) -> Path:
    p = tmp_path / "pricing.json"
    p.write_text(json.dumps(PRICING), encoding="utf-8")
    return p


def his_set(kind: str, n: int = 3, title: bool = True, signature: bool = True) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    heading, label, text = KINDS[kind]
    d = docx.Document()

    def centered(t: str) -> None:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(t).bold = True

    if title:
        centered(f"PLAINTIFF GAMMA EXAMPLE'S {heading} TO DEFENDANT DELTA EXAMPLE, SET ONE")
    d.add_paragraph("PROPOUNDING PARTY: Plaintiff GAMMA EXAMPLE")
    d.add_paragraph("RESPONDING PARTY: Defendant DELTA EXAMPLE")
    d.add_paragraph("SET NO.: ONE")
    centered("DEFINITIONS")
    d.add_paragraph('"YOU" and "YOUR" mean Defendant DELTA EXAMPLE.')
    centered(heading)
    for i in range(1, n + 1):
        p = d.add_paragraph()
        p.add_run(f"{label} NO. {i}:").bold = True
        d.add_paragraph(f"{text} ({i})")
    if signature:
        d.add_paragraph("Dated: October 12, 2026")
        d.add_paragraph("By: ____________ Alpha Example, Attorneys for Plaintiff")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def att(name: str, blob: bytes) -> dict:
    return {"filename": name, "sha256": hashlib.sha256(blob).hexdigest(), "b64": base64.b64encode(blob).decode()}


def _run(tmp_path, pricing, client, attachments, request=REQUEST):
    inputs = tmp_path / "inputs"
    if not inputs.is_dir():
        make_inputs(inputs, attorneys=[LEAD])
    jd = make_job(tmp_path / "job", cls="discovery_set")
    doc = json.loads((jd / "job.json").read_text())
    doc["request_text"], doc["attachments"] = request, attachments
    (jd / "job.json").write_text(json.dumps(doc))
    seat = seat_with(standard_docs())
    r = DraftingRun(
        jd,
        inputs_dir=str(inputs),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=client,
        log=lambda _m: None,
        readback_pause=0.0,
    )
    return r, r.run(), seat


def _text(path: Path) -> str:
    import docx

    return format_check.doc_text(docx.Document(io.BytesIO(path.read_bytes())))


def test_three_attached_sets_are_placed_verbatim_each_on_its_own_pleading_paper(tmp_path, pricing):
    from medchron.demand import deliver

    blobs = {k: his_set(k, signature=(k != "production")) for k in KINDS}
    client = ScriptedClient()
    r, v, seat = _run(tmp_path, pricing, client, [att(f"{k}.docx", b) for k, b in blobs.items()])
    assert v.outcome == "delivered", (v.stage, v.reason)
    assert [f["role"] for f in v.files] == ["draft", "draft", "draft", "attorney_notes"]
    # nothing of his was composed, audited or repaired: no paid drafting at all
    stages = json.dumps([c.get("system") for c in client.calls])
    assert "COMPOSE-PROMPT" not in stages and "AUDIT-PROMPT" not in stages and "REPAIR-PROMPT" not in stages
    spec = pleading.spec_for(r.firm.data["format"])
    for (k, blob), f in zip(blobs.items(), v.files):
        path = deliver.out_dir(r.data) / f["name"]
        text = _text(path)
        theirs = place.read(f"{k}.docx", blob)
        assert place.missing(theirs.plain[place._start(theirs) :], text) == [], f["name"]
        assert "Case No. CV-0001" in text and "EXAMPLE & EXAMPLE, LLP" in text
        assert text.upper().count("PROOF OF SERVICE") == 1
        import docx

        assert pleading.check(docx.Document(io.BytesIO(path.read_bytes())), path.read_bytes(), spec) == []
    # his own signature kept; the firm's signature lines only where he wrote none
    prod = _text(deliver.out_dir(r.data) / v.files[2]["name"])
    assert "By: ______________________________" in prod and "Attorneys for Plaintiff GAMMA EXAMPLE" in prod
    notes = (r.data / "notes-final.md").read_text()
    assert "placed production.docx word for word" in notes


def test_a_placement_that_loses_his_words_is_held_for_smd(tmp_path, pricing, monkeypatch):
    real = place.place

    def lossy(t, fields, block, pos):
        out = real(t, fields, block, pos)
        out.markdown = out.markdown.replace("the floor was dry", "the floor was clean")
        return out

    monkeypatch.setattr(placing.place, "place", lossy)
    r, v, seat = _run(tmp_path, pricing, ScriptedClient(), [att("special.docx", his_set("special"))])
    assert v.outcome == "failed" and v.reason.startswith("review_held: ") and "word for word" in v.reason
    assert seat.created == []


def test_attachments_that_are_not_sets_hold_with_a_sentence(tmp_path, pricing):
    import docx

    d = docx.Document()
    d.add_paragraph("A letter about scheduling.")
    buf = io.BytesIO()
    d.save(buf)
    r, v, _ = _run(tmp_path, pricing, ScriptedClient(), [att("letter.docx", buf.getvalue())])
    assert v.outcome == "held" and "carry no discovery requests" in v.reason


def test_the_notes_speak_to_the_attorney_not_about_the_machinery(tmp_path, pricing):
    """Rehearsal 2026-10-13: the notes said "What the job did" and carried the
    renderer's own lines ("named styles defined on the starter base")."""
    r, v, _ = _run(tmp_path, pricing, ScriptedClient(), [att("special.docx", his_set("special"))])
    assert v.outcome == "delivered", v.reason
    notes = (r.data / "notes-final.md").read_text()
    for internal in ("the job", "starter base", "page number field"):
        assert internal not in notes, internal


def test_without_an_authored_block_the_courts_attorney_block_prints_as_a_block():
    att_ = {"name": "ALPHA EXAMPLE", "state_bar_number": "000001", "firm": "EXAMPLE LAW", "email": "a@x.example"}
    assert place.court_attorney_block(att_) == ["ALPHA EXAMPLE (SBN 000001)", "EXAMPLE LAW", "Email: a@x.example"]


def test_his_title_is_kept_and_a_missing_title_is_built_from_his_identification_block():
    blob = his_set("admission", title=False)
    t = place.read("a.docx", blob)
    out = place.place(
        t,
        {"plaintiffs": [{"value": "GAMMA EXAMPLE"}], "defendants": [{"value": "DELTA EXAMPLE"}]},
        None,
        "# PROOF OF SERVICE",
    )
    assert "PLAINTIFF GAMMA EXAMPLE'S REQUESTS FOR ADMISSION TO DEFENDANT DELTA EXAMPLE, SET ONE" in out.markdown
    assert "{{NOT IN RECORD: case number" in out.markdown  # never a guessed case number


def test_an_attachment_whose_bytes_do_not_match_is_refused(tmp_path):
    a = att("x.docx", b"abc")
    a["sha256"] = "0" * 64
    with pytest.raises(job_mod.DraftingJobError, match="sha256"):
        job_mod._attachments([a])
    with pytest.raises(job_mod.DraftingJobError, match="docx"):
        job_mod._attachments([att("../x.pdf", b"abc")])
