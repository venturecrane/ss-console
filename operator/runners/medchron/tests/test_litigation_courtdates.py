"""The calendar's court dates merged into the list and the printable PDF
(``litigation/courtdates.py``, ``litigation/printout.py``). Synthetic data only."""

from __future__ import annotations

import copy
import datetime as dt
from pathlib import Path

from medchron.litigation import courtdates as cd, printout

TODAY = dt.date(2026, 10, 29)


def ev(start: str, subject: str, *, all_day: bool = False, ltb: bool = False, deleted: bool = False) -> dict:
    return {"startTime": start, "allDay": all_day, "subject": subject, "isDeleted": deleted, "ltb": ltb}


#: Shaped on a real calendar: LawToolBox enters the set events timed and its
#: rule deadlines all-day; staff enter mediations and reminders by hand.
CAL = [
    ev("2027-02-09T08:30:00", "2/9 - Example. Trial Day 1 dept 47", ltb=True),
    ev("2027-01-04T09:30:00", "1/4 - Example. MSC dept 59", ltb=True),
    ev("2026-11-17T13:00:00", "MEDIATION EXAMPLE"),
    ev("2026-10-29T13:30:00", "MEDIATION EXAMPLE", deleted=True),
    ev("2027-01-11T00:00:00", "MOTION/DISCOVERY CUT-OFF", all_day=True, ltb=True),
    ev("2026-12-11T00:00:00", "30 DAY RMDR - MOTION/DISCOVERY CUT-OFF", all_day=True, ltb=True),
    ev("2026-12-21T00:00:00", "*Exchange expert info after demand", all_day=True, ltb=True),
    ev("2026-11-25T00:00:00", "*Demand expert exchange & disclosure - BY HAND", all_day=True, ltb=True),
    ev("2027-01-08T00:00:00", "Hold summary judgment hearing ", all_day=True, ltb=True),
    ev("2026-12-22T00:00:00", "Request phone appearance at MSC (MSC dept 59)", all_day=True, ltb=True),
    ev("2027-01-08T07:30:00", "1/8 - 1 MONTH REMINDER. Example. TRIAL STARTS 2/9"),
    ev("2026-12-09T00:00:00", "HIDDEN *Possible latest file/serve papers", all_day=True, ltb=True),
    ev("2026-11-05T10:00:00", "OFF - 11/5 Example - hearing"),
    ev("2026-11-03T00:00:00", "11/3 - Deadline. Example. LD to Submit MSC and Trial Dates", all_day=True),
    ev("2025-04-11T08:30:00", "4/11 - Example. CMC"),
]


def test_each_kind_is_the_set_event_never_a_rule_deadline_reminder_or_deleted_entry():
    got = cd.upcoming(CAL, TODAY)
    assert {k: (v["date"], v["time"]) for k, v in got.items()} == {
        "trial": ("2027-02-09", "8:30 AM"),
        "msc": ("2027-01-04", "9:30 AM"),
        "mediation": ("2026-11-17", "1:00 PM"),
        "discovery_cutoff": ("2027-01-11", ""),
        "expert_exchange": ("2026-12-21", ""),
    }
    assert got["trial"]["source"] == cd.LTB and got["mediation"]["source"] == cd.STAFF


def test_the_rules_that_exclude():
    assert cd.kind(ev("2027-01-08T00:00:00", "Hold summary judgment hearing", all_day=True, ltb=True)) is None
    assert cd.kind(ev("2027-01-08T09:00:00", "Hearing on motion to compel")) == "hearing"
    assert cd.kind(ev("2026-11-25T00:00:00", "File unanimous stipulation to mediation", all_day=True, ltb=True)) is None
    assert cd.kind(ev("2026-11-25T00:00:00", "Trial setting conference", all_day=False)) == "hearing"
    assert cd.kind(ev("2026-11-25T10:00:00", "CMC", deleted=True)) is None
    assert cd.kind(ev("2026-11-25T10:00:00", "VACATED - CMC")) is None


def test_a_lawtoolbox_all_day_entry_headed_by_its_kind_is_the_event_itself():
    assert cd.kind(ev("2027-01-11T00:00:00", "TRIAL- Trial date", all_day=True, ltb=True)) == "trial"
    assert (
        cd.kind(ev("2027-08-30T00:00:00", "**MSC- Mandatory Settlement Conference (MSC)", all_day=True, ltb=True))
        == "msc"
    )
    assert (
        cd.kind(ev("2027-05-26T00:00:00", "CMC- Initial case management conference", all_day=True, ltb=True)) == "cmc"
    )
    assert cd.kind(ev("2027-05-26T00:00:00", "INITIAL CMC- *Pay CCP 631(b) jury fee", all_day=True, ltb=True)) is None
    assert cd.kind(ev("2026-11-19T00:00:00", "BINDING ARBITRATION", all_day=True)) == "hearing"
    assert cd.kind(ev("2026-11-16T00:00:00", "File Arb Docs w/Arbitrator", all_day=True)) is None


def test_the_next_court_date_comes_from_the_calendar_when_the_papers_have_none():
    m = {"matter_id": "m1", "next_court_date": {"date": None, "event": None, "source": None}}
    got = cd.merge(m, CAL, TODAY)
    assert got["next"]["date"] == "2027-01-04" and got["next"]["source"]["name"] == cd.LTB
    assert got["flag"] == "" and got["read"] is True
    assert [label for label, _r in got["rows"]] == [
        "Trial", "Settlement conference", "Mediation", "Discovery cutoff", "Expert exchange",
    ]  # fmt: skip


def test_papers_and_calendar_on_the_same_day_agree_and_cite_both():
    src = {"file_id": "f1", "name": "Notice of MSC.pdf", "doc_date": "2026-08-01"}
    m = {"matter_id": "m1", "next_court_date": {"date": "2027-01-04", "event": "MSC", "source": src}}
    got = cd.merge(m, CAL, TODAY)
    assert got["flag"] == "" and got["next"]["date"] == "2027-01-04"
    assert cd.LTB in got["next"]["source"]["name"] and "Notice of MSC.pdf" in got["next"]["source"]["name"]


def test_a_papers_date_the_calendar_lacks_is_shown_and_flagged():
    src = {"file_id": "f1", "name": "Notice of hearing.pdf", "doc_date": "2026-09-01"}
    m = {"matter_id": "m1", "next_court_date": {"date": "2026-12-01", "event": "Hearing", "source": src}}
    got = cd.merge(m, CAL, TODAY)  # 2026-12-01 holds no court entry (a rule deadline on that day is not one)
    assert got["flag"] == cd.DISAGREE and got["next"]["date"] == "2026-12-01"
    assert got["papers"]["date"] == "2026-12-01" and got["calendar"]["date"] == "2027-01-04"


def test_a_past_papers_date_gives_way_to_the_calendar_without_a_flag():
    m = {"matter_id": "m1", "next_court_date": {"date": "2025-04-11", "event": "CMC", "source": {"name": "x.pdf"}}}
    got = cd.merge(m, CAL, TODAY)
    assert got["next"]["date"] == "2027-01-04" and got["flag"] == "" and got["papers"] is None


def test_shown_never_changes_the_final_values_and_an_unread_calendar_is_said():
    finals = [_matter("m1", "1001", "A v. B", "Lee"), _matter("m2", "1002", "C v. D", "Lee")]
    before = copy.deepcopy(finals)
    out = cd.shown(finals, {"m1": CAL}, TODAY)
    assert finals == before
    assert out[0]["next_court_date"]["date"] == "2027-01-04"
    assert out[1]["court_dates"]["read"] is False and out[1]["next_court_date"]["date"] is None
    assert cd.shown(finals, None, TODAY) is finals


def _matter(mid: str, num: str, case: str, who: str, **kw) -> dict:
    return {
        "matter_id": mid,
        "number": num,
        "case_name": case,
        "responsible": who,
        "court": "Superior Court",
        "case_number": "26CV0001",
        "case_status": {"value": "Active", "detail": "", "source": None},
        "complaint_filed": {"date": "2026-01-02", "source": {"name": "Complaint.pdf"}},
        "next_court_date": {"date": None, "event": None, "source": None},
        "defendants": [{"name": "Def One", "status": "Served", "served": {"date": "2026-02-01"}, "answered": {}}],
        "matter_flags": [],
        "notes": "",
        **kw,
    }


def test_the_pdf_is_ordered_by_attorney_then_case_and_prints_every_court_date(tmp_path: Path):
    finals = [
        _matter("m1", "1001", "Zed v. Co", "Avery"),
        _matter("m2", "1002", "Able v. Co", "Avery"),
        _matter("m3", "1003", "Mid v. Co", "Blake"),
    ]
    shown = cd.shown(finals, {"m1": CAL, "m2": [], "m3": CAL}, TODAY)
    out = tmp_path / "list court dates.pdf"
    stats = printout.build(shown, out, TODAY)
    assert stats["cases"] == 3 and stats["with_court_date"] == 2 and stats["pages"] >= 4
    assert [m["number"] for m in printout.order(shown)] == ["1002", "1001", "1003"]
    assert printout.check(out, shown) == [] and printout.scan(out) == []
    text = "\n".join(printout.text_of(out))
    for d in ("2027-02-09", "2027-01-04", "2026-11-17", "2027-01-11", "2026-12-21"):
        assert d in text
    assert text.index("1002 Able v. Co") < text.index("1001 Zed v. Co") < text.index("1003 Mid v. Co")


def test_the_pdf_check_can_fail(tmp_path: Path):
    shown = cd.shown([_matter("m1", "1001", "A v. B", "Lee")], {"m1": CAL}, TODAY)
    out = tmp_path / "x.pdf"
    printout.build(shown, out, TODAY)
    planted = copy.deepcopy(shown)
    planted[0]["court_dates"]["rows"].append(("Hearing", {"date": "2027-03-03", "time": "", "what": "x",
                                                          "source": cd.STAFF}))  # fmt: skip
    assert any("not printed" in p for p in printout.check(out, planted))
    missing = shown + cd.shown([_matter("m9", "1009", "Q v. R", "Lee")], {}, TODAY)
    assert any("no page" in p for p in printout.check(out, missing))


def test_the_pdf_scan_catches_internal_text(tmp_path: Path):
    shown = cd.shown([_matter("m1", "1001", "A v. B", "Lee", notes="checked on pass 2")], {"m1": CAL}, TODAY)
    out = tmp_path / "x.pdf"
    printout.build(shown, out, TODAY)
    assert any(w == "internal process wording" for _p, w, _m in printout.scan(out))


# ---- the run: the calendar reaches the files, never the stored state ---------------------------
def test_a_run_shows_the_calendar_date_in_every_file_and_stores_only_the_papers(tmp_path: Path, pricing_path: Path):
    from litigation_testkit import M1, LitSeat, ScriptedLit, make_inputs, make_job, matter_docs
    from openpyxl import load_workbook

    from medchron.litigation import manifest
    from medchron.litigation.run import PRINTOUT_ROLE, LitigationRun

    (tmp_path / "pricing.json").write_text(pricing_path.read_text())
    seat = LitSeat({M1: matter_docs()})
    seat.events[M1] = [
        {"startTime": "2027-03-01T08:30:00", "allDay": False, "subject": "3/1 - Trial dept 9", "isDeleted": False,
         "additionalData": {"LawToolBox": {"id": "x"}}},
    ]  # fmt: skip
    today = dt.date(2026, 10, 7)
    r = LitigationRun(
        make_job(tmp_path / "job"),
        inputs_dir=str(make_inputs(tmp_path / "inputs")),
        pricing=str(tmp_path / "pricing.json"),
        state_dir=tmp_path / "state",
        seat_factory=lambda: seat,
        client=ScriptedLit(),
        log=lambda m: None,
        today=lambda: today,
        readback_pause=0.0,
    )
    v = r.run()
    assert v.verdict == "delivered", v.reason
    pdf = next(f for f in v.files if f["role"] == PRINTOUT_ROLE)
    assert pdf["name"] == printout.path_for(r._out()).name and pdf["sha256"]
    assert "2027-03-01" in "\n".join(printout.text_of(printout.path_for(r._out())))
    wb = load_workbook(r._out(), read_only=True)
    notes = [list(x) for x in wb["Case notes"].iter_rows(values_only=True)]
    wb.close()
    assert any("2027-03-01" in str(c) for row in notes[1:] for c in row)
    stored = manifest.load_prior(tmp_path / "state", M1)
    assert "2027-03-01" not in str(stored.get("next_court_date"))
