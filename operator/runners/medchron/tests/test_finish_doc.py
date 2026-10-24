"""build_doc, the two classifiers, the strip, the coverage gate, the billing
chart and the worksheet, in-process. Real PDFs; a scripted client for the
one paid stage."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace as NS

import pytest
import yaml

from medchron import config as config_mod, job as job_mod
from medchron.stages import billing_chart, billing_docx, build_doc, classify, coverage, strip
from medchron.stages.base import StageRun
from medchron_testkit import FIRM_CONFIG, FakeSeat, make_pdf

PROSE = (
    "Patient seen on 01/02/2026 for neck pain after the collision. Blood pressure 120/80. "
    "Ibuprofen 400 mg prescribed. Riverside Imaging ordered an MRI. "
) * 4
ORDER_PAGE = (
    "Order #: 12345\nSubject Name: Alpha Example\nRequested Date Range: 2025-2026\nvendor@example-retrieval.com"
)
INDEX_PAGE = "WORD INDEX\nThis list is computer generated and may contain errors.\nabdomen 4\nankle 9\n" * 3


class Usage:
    input_tokens, output_tokens, cache_read_input_tokens, cache_creation_input_tokens = 100, 20, 0, 0


class Scripted:
    def __init__(self, reply):
        self.calls: list[dict] = []
        self._reply = reply
        self.messages = NS(create=self._create, stream=None, batches=None)

    def _create(self, **params):
        self.calls.append(params)
        return self._reply(params, len(self.calls))


@pytest.fixture
def firm(tmp_path: Path) -> Path:
    cfg = json.loads(json.dumps(FIRM_CONFIG))
    cfg["nonrecord"]["page_classes"] = [
        {"name": "INDEX", "patterns": ["this list is computer generated"]},
        {"name": "ORDER", "patterns": [r"order\s*#\s*:.{0,600}example-retrieval\.com"]},
    ]
    cfg["providers"]["aliases"] = [{"match": r"example clinic|ex clinic", "label": "Example Clinic"}]
    cfg["billing"]["provider_match"] = {"Example Clinic": [r"example clinic", r"ex clinic"]}
    p = tmp_path / "firm.yaml"
    p.write_text(yaml.safe_dump(cfg))
    return p


def _sr(job_dir: Path, firm: Path, data_root: Path, client=None, log: list[str] | None = None) -> StageRun:
    job = job_mod.load(job_dir)
    cfg = config_mod.load(str(firm))
    lines = log if log is not None else []
    return StageRun(
        job=job,
        cfg=cfg,
        unit=job.units[0],
        slug_dir=data_root / "example-matter",
        decided={},
        log=lines.append,
        seat_factory=lambda: FakeSeat([], [], {}),
        client_factory=lambda: client,
    )


# ---- build_doc ---------------------------------------------------------------------
ENTRIES = (
    "Prior Medical History\n\nChronic neck pain before the incident. (Exhibit 1 - p. 1)\n\n"
    "01/02/2026\nExample Clinic (Jane Doe, MD) | Patient Complaints & Limitations\n\n"
    "Neck pain since the Subject Incident. (Exhibit 1 - p. 2)\n\n"
    "01/09/2026\nEx Clinic | Medical Diagnoses\n\nCervical strain. (Exhibit 1 - p. 3)\n"
)
MAP_INDEX = (
    "## ENTRIES\nnone in this chunk\n\n## INDEX\n2026-01-02 | Example Clinic | S13.4XXA, M54.2 (as recorded) | clinic note.pdf\n"
    "2026-01-09 | Ex Clinic | S523.3XXA | clinic note.pdf\n2027-01-01 | Other | Z00.0 | other.pdf\n\n## FILES-SEEN\n"
)


def _icd_tables(data_root: Path) -> None:
    d = data_root / "controls" / "icd"
    d.mkdir(parents=True)
    line10 = "{:<5} {:<7} {} {:<60} {}".format
    (d / "icd10cm_order.txt").write_text(
        line10(
            "00001",
            "S134XXA",
            "1",
            "Sprain of lig of cervical spine, init",
            "Sprain of ligaments of cervical spine, initial encounter",
        )
        + "\n"
        + line10("00002", "M542", "1", "Cervicalgia", "Cervicalgia")
        + "\n"
    )
    (d / "CMS32_DESC_LONG_DX.txt").write_text("7242 Lumbago\n")
    (d / "VERSION.json").write_text(json.dumps({"icd10cm": {"label": "test"}}))


def test_build_doc_composes_sections_from_the_entries_and_tables_from_code(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    sr = _sr(job_dir, firm, data_root)
    rd = sr.slug_dir / "runs" / "alpha"
    rd.mkdir(parents=True)
    (rd / "entries_scoped_final.md").write_text(ENTRIES)
    (rd / "map-01.md").write_text(MAP_INDEX)
    (rd / "preincident_note.txt").write_text("[NTD: 1 additional pre-incident encounter was reviewed.]")
    out = sr.slug_dir / "out" / "alpha"
    out.mkdir(parents=True)
    (out / "page_map.json").write_text(
        json.dumps([{"exhibit": 1, "title": "Exhibit 1 - Example Clinic - 01-02-2026 (Medical Records)", "files": []}])
    )
    _icd_tables(sr.job.data_root)
    assert build_doc.run(sr) == 0
    doc = (rd / "final-chronology.md").read_text()
    assert doc.startswith("Alpha Example - Medical Chronology\n")
    assert (
        "| Example Clinic | 01/02/2026 - 01/09/2026 | 2 | Exhibit 1 - p. 2 |" in doc
    )  # facility level, both spellings
    assert (
        "| S13.4XXA | Sprain of ligaments of cervical spine, initial encounter | 01/02/2026 | Exhibit 1 - p. 2 |" in doc
    )
    assert (
        "| M54.2 | Cervicalgia |" in doc and "S523.3XXA" not in doc and "Z00.0" not in doc
    )  # bad shape and unkept date dropped
    assert "## Medical Chronology\n\nPrior Medical History\n\nChronic neck pain" in doc and "[NTD: 1 additional" in doc
    assert "| 1 | Example Clinic - 01-02-2026 (Medical Records) |" in doc
    assert "## Records Reviewed and Limitations" in doc and "prepared from 0 documents in the matter file" in doc


def test_build_doc_refuses_unrecognised_preamble(job_dir: Path, firm: Path, data_root: Path) -> None:
    sr = _sr(job_dir, firm, data_root)
    rd = sr.slug_dir / "runs" / "alpha"
    rd.mkdir(parents=True)
    (rd / "entries_scoped_final.md").write_text(
        "Some narrative the model wrote first.\n\n" + ENTRIES.split("\n\n", 2)[2]
    )
    (sr.slug_dir / "out" / "alpha").mkdir(parents=True)
    assert build_doc.run(sr) == 1


# ---- the classifiers and the strip -----------------------------------------------------------
def _exhibit(sr: StageRun, pages: list[str]) -> Path:
    out = sr.slug_dir / "out" / "alpha"
    out.mkdir(parents=True, exist_ok=True)
    title = "Exhibit 1 - Example Clinic - 01-02-2026 (Medical Records)"
    pdf = out / f"{title}.pdf"
    pdf.write_bytes(make_pdf(pages))
    (out / "page_map.json").write_text(
        json.dumps(
            [
                {
                    "exhibit": 1,
                    "title": title,
                    "total_pages": len(pages),
                    "files": [{"file": "clinic note.pdf", "old_exhibit": 1, "start_page": 1, "pages": len(pages)}],
                }
            ]
        )
    )
    return pdf


def _doc(sr: StageRun, text: str) -> Path:
    rd = sr.slug_dir / "runs" / "alpha"
    rd.mkdir(parents=True, exist_ok=True)
    p = rd / "final-chronology.md"
    p.write_text(text)
    return p


DOC_3 = (
    "Alpha Example - Medical Chronology\n\n## Treatment Timeline\n\n| Medical Provider | Treatment Period | Visits | Reference |\n"
    "| Example Clinic | 01/02/2026 | 1 | Exhibit 1 - p. 3 |\n\n## Diagnostic Highlights\n\n| ICD Code | Description | First Diagnosed | Reference |\n\n"
    "## Medical Chronology\n\n01/02/2026\nExample Clinic | Medical Diagnoses\n\nCervical strain noted. (Exhibit 1 - p. 3)\n\n"
    "## Exhibit List\n\n| Exhibit No. | Description |\n"
)


def test_classify_nonrecord_fingerprints_classes_and_reports_collisions(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    sr = _sr(job_dir, firm, data_root)
    _exhibit(sr, [ORDER_PAGE, INDEX_PAGE, PROSE, ""])
    _doc(sr, DOC_3)
    assert classify.run_nonrecord(sr) == 0
    nr = json.loads((sr.slug_dir / "nonrecord.json").read_text())["1"]
    assert nr["drop_pages"] == [1, 2] and nr["unknown"] == [4] and nr["cited_collision"] == []
    assert [b["class"] for b in nr["blocks"]] == ["ORDER", "INDEX"]
    _doc(sr, DOC_3.replace("(Exhibit 1 - p. 3)", "(Exhibit 1 - p. 2)"))
    assert classify.run_nonrecord(sr) == 0
    assert json.loads((sr.slug_dir / "nonrecord.json").read_text())["1"]["cited_collision"] == [2]


def _scanned_setup(sr: StageRun, unknown: list[int]) -> None:
    """An exhibit whose listed pages are unclassifiable by text, with controls authored."""
    _exhibit(sr, [PROSE, PROSE, PROSE, ""])
    _doc(sr, DOC_3)
    (sr.slug_dir / "nonrecord.json").write_text(
        json.dumps({"1": {"pages": 4, "blocks": [], "drop_pages": [], "unknown": unknown, "cited_collision": []}})
    )
    ctl = sr.job.data_root / "controls"
    ctl.mkdir(parents=True, exist_ok=True)
    (ctl / "control-order.pdf").write_bytes(make_pdf([ORDER_PAGE]))
    (ctl / "control-index.pdf").write_bytes(make_pdf([INDEX_PAGE]))
    (ctl / "controls.json").write_text(
        json.dumps(
            [
                {"pdf": "controls/control-order.pdf", "page": 1, "label": "ORDER"},
                {"pdf": "controls/control-index.pdf", "page": 1, "label": "INDEX"},
            ]
        )
    )
    (sr.slug_dir / "record_control.json").write_text(json.dumps({"exhibit": 1, "page": 2}))


def _label_reply(p, n):
    labels = [b["text"].split()[1].rstrip(":") for b in p["messages"][0]["content"] if b["type"] == "text"]
    answer = {"CTL0": "ORDER", "CTL1": "INDEX", "CTL2": "RECORD"}
    return NS(
        content=[NS(type="text", text="\n".join(f"{lb} = {answer.get(lb, 'RECORD')}" for lb in labels))],
        stop_reason="end_turn",
        usage=Usage(),
    )


def test_classify_scanned_batches_are_bounded_by_bytes_and_a_too_large_request_is_split(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Live 2026-09-16: twelve scanned pages at 110 dpi passed the API's request
    limit (413) and the stage died with them. The limit is on bytes, not pages;
    and a refusal is split, never fatal."""
    import httpx
    from anthropic import RequestTooLargeError

    client = Scripted(_label_reply)
    sr = _sr(job_dir, firm, data_root, client)
    _scanned_setup(sr, unknown=[1, 2, 3])
    # every page image "weighs" 11 MB: with a 20 MB cap, no two share a request
    monkeypatch.setattr(classify, "_png", lambda doc, p, dpi=110: "x" * (11 * 2**20))
    assert classify.run_scanned(sr) == 0
    sizes = [sum(1 for b in c["messages"][0]["content"] if b["type"] == "image") for c in client.calls]
    assert sizes == [1] * 6, sizes  # 3 pages + 3 controls, one per request
    labels = json.loads((sr.slug_dir / "scanned_labels.json").read_text())
    assert labels["controls_ok"] is True and set(labels["labels"]) >= {"Ex1p1", "Ex1p2", "Ex1p3"}

    # the API refuses any request carrying more than one image: halve until it takes them
    def refuse_multi(p, n):
        if sum(1 for b in p["messages"][0]["content"] if b["type"] == "image") > 1:
            raise RequestTooLargeError(
                "too large", response=httpx.Response(413, request=httpx.Request("POST", "https://x")), body=None
            )
        return _label_reply(p, n)

    client2 = Scripted(refuse_multi)
    sr2 = _sr(job_dir, firm, data_root, client2)
    monkeypatch.setattr(classify, "_png", lambda doc, p, dpi=110: "x" * 10)
    assert classify.run_scanned(sr2) == 0
    sizes2 = [sum(1 for b in c["messages"][0]["content"] if b["type"] == "image") for c in client2.calls]
    # one refused batch of six, halves refused in turn, every page finally sent alone
    assert sizes2[0] == 6 and sizes2.count(1) == 6 and len(sizes2) < 12, sizes2
    labels2 = json.loads((sr2.slug_dir / "scanned_labels.json").read_text())
    assert labels2["controls_ok"] is True and set(labels2["labels"]) >= {"Ex1p1", "Ex1p2", "Ex1p3"}


def test_classify_scanned_labels_head_pages_and_needs_its_controls(job_dir: Path, firm: Path, data_root: Path) -> None:
    def reply(p, n):
        labels = [b["text"].split()[1].rstrip(":") for b in p["messages"][0]["content"] if b["type"] == "text"]
        answer = {"Ex1p4": "AUTH", "CTL0": "ORDER", "CTL1": "INDEX", "CTL2": "RECORD"}
        return NS(
            content=[NS(type="text", text="\n".join(f"{lb} = {answer.get(lb, 'RECORD')}" for lb in labels))],
            stop_reason="end_turn",
            usage=Usage(),
        )

    client = Scripted(reply)
    sr = _sr(job_dir, firm, data_root, client)
    _exhibit(
        sr, [PROSE, PROSE, PROSE, ""]
    )  # page 4 is blank: unclassifiable by text, within HEAD_PAGES? no (start 1: pages 1-3)
    _doc(sr, DOC_3)
    (sr.slug_dir / "nonrecord.json").write_text(
        json.dumps({"1": {"pages": 4, "blocks": [], "drop_pages": [], "unknown": [1, 4], "cited_collision": []}})
    )
    assert classify.run_scanned(sr) == 1  # controls not authored
    ctl = sr.job.data_root / "controls"
    ctl.mkdir(parents=True)
    (ctl / "control-order.pdf").write_bytes(make_pdf([ORDER_PAGE]))
    (ctl / "control-index.pdf").write_bytes(make_pdf([INDEX_PAGE]))
    (ctl / "controls.json").write_text(
        json.dumps(
            [
                {"pdf": "controls/control-order.pdf", "page": 1, "label": "ORDER"},
                {"pdf": "controls/control-index.pdf", "page": 1, "label": "INDEX"},
            ]
        )
    )
    (sr.slug_dir / "record_control.json").write_text(json.dumps({"exhibit": 1, "page": 2}))
    assert classify.run_scanned(sr) == 0
    labels = json.loads((sr.slug_dir / "scanned_labels.json").read_text())
    assert labels["controls_ok"] is True and labels["labels"]["Ex1p1"] == "RECORD" and labels["nonrecord"] == {}
    sent = [b["text"] for b in client.calls[0]["messages"][0]["content"] if b["type"] == "text"]
    assert sent == ["page Ex1p1:", "page CTL0:", "page CTL1:", "page CTL2:"]  # p.4 is outside the file head


def test_classify_scanned_with_no_record_control_or_failed_controls_keeps_every_scanned_page(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    """Rows 3 and 15 (2026-10-24). An all-fax matter has no record control
    (decide_control writes ``skip``), and a weak record control can be called
    something else by the classifier; both stopped the run. Now the stage
    calls nothing for a skipped control, records untrusted labels with no
    drops for failed ones, and exits 0: the strip keeps every scanned page."""
    client = Scripted(_label_reply)
    sr = _sr(job_dir, firm, data_root, client)
    _scanned_setup(sr, unknown=[1, 2, 3])
    (sr.slug_dir / "record_control.json").write_text(json.dumps({"skip": True, "_decided": {"reason": "no text"}}))
    assert classify.run_scanned(sr) == 0
    assert client.calls == []  # nothing to falsify the classifier against, so it is not paid for
    labels = json.loads((sr.slug_dir / "scanned_labels.json").read_text())
    assert labels["controls_ok"] is False and labels["nonrecord"] == {} and "no text" in labels["skipped"]

    def wrong_record(p, n):  # the classifier calls the record control a fax cover
        r = _label_reply(p, n)
        r.content[0].text = (
            r.content[0].text.replace("CTL2 = RECORD", "CTL2 = FAX").replace("Ex1p1 = RECORD", "Ex1p1 = AUTH")
        )
        return r

    client2 = Scripted(wrong_record)
    sr2 = _sr(job_dir, firm, data_root, client2)
    (sr2.slug_dir / "record_control.json").write_text(json.dumps({"exhibit": 1, "page": 2}))
    assert classify.run_scanned(sr2) == 0
    labels = json.loads((sr2.slug_dir / "scanned_labels.json").read_text())
    assert labels["controls_ok"] is False and labels["nonrecord"] == {}


def test_strip_falsifies_dry_runs_and_applies_with_a_verified_remap(job_dir: Path, firm: Path, data_root: Path) -> None:
    sr = _sr(job_dir, firm, data_root)
    pdf = _exhibit(sr, [ORDER_PAGE, INDEX_PAGE, PROSE, PROSE + " second"])
    _doc(sr, DOC_3.replace("(Exhibit 1 - p. 3)", "(Exhibit 1 - p. 3-4)"))
    assert classify.run_nonrecord(sr) == 0
    assert strip.run_falsify(sr) == 0 and not pdf.with_name(pdf.name + ".orig").exists()
    assert strip.run_dry(sr) == 0 and not (sr.slug_dir / "strip_result.json").exists()
    assert strip.run_apply(sr) == 0
    doc = (sr.slug_dir / "runs" / "alpha" / "final-chronology.md").read_text()
    assert "(Exhibit 1 - p. 1-2)" in doc and "| Exhibit 1 - p. 1 |" in doc  # body and table cell both remapped
    result = json.loads((sr.slug_dir / "strip_result.json").read_text())
    assert result["drops"] == {"1": [1, 2]} and result["new_page_counts"] == {"1": 2}
    assert Path(str(pdf) + ".orig").exists()
    import pymupdf

    assert len(pymupdf.open(str(pdf))) == 2
    assert strip.run_apply(sr) == 1  # never twice


def test_strip_keeps_a_cited_page_the_classifier_calls_non_record(job_dir: Path, firm: Path, data_root: Path) -> None:
    """Pamplona 2026-10-24: one cited page called non-record held a 4,418-page
    job. The citation's paper stays; an authored drop of it still refuses."""
    import pymupdf

    sr = _sr(job_dir, firm, data_root)
    _exhibit(sr, [ORDER_PAGE, PROSE])
    _doc(
        sr,
        DOC_3.replace("(Exhibit 1 - p. 3)", "(Exhibit 1 - p. 1)").replace(
            "| Exhibit 1 - p. 3 |", "| Exhibit 1 - p. 1 |"
        ),
    )
    assert classify.run_nonrecord(sr) == 0
    assert json.loads((sr.slug_dir / "nonrecord.json").read_text())["1"]["cited_collision"] == [1]
    authored = sr.slug_dir / "drops.json"
    authored.write_text(json.dumps({"manual_drops": {"1": [[1, "authored"]]}, "citation_edits": []}))
    assert strip.run_dry(sr) == strip.EXIT_CITED_DROP  # 2 in every mode since 2026-10-24
    assert strip.run_falsify(sr) == 2
    authored.unlink()
    assert strip.run_dry(sr) == 0
    assert strip.run_apply(sr) == 0
    [path] = classify.resolve_exhibit_files(sr.slug_dir / "out" / sr.unit.unit).values()
    with pymupdf.open(str(path)) as doc:
        assert len(doc) == 2


def _cite(n: str) -> str:
    """DOC_3 with its body citation and table cell on the given page spec."""
    return DOC_3.replace("(Exhibit 1 - p. 3)", f"(Exhibit 1 - p. {n})").replace(
        "| Exhibit 1 - p. 3 |", f"| Exhibit 1 - p. {n} |"
    )


def test_strip_passes_an_out_of_range_citation_through_unchanged(job_dir: Path, firm: Path, data_root: Path) -> None:
    """Row 2 (2026-10-24): a citation to a page beyond the exhibit (or p. 0)
    crashed the strip with a KeyError, and a crash is `failed`, terminal. The
    composer's page is not the strip's to judge: it passes through unchanged,
    logged, and the audit grades it PAGE_OUT_OF_RANGE."""
    log: list[str] = []
    sr = _sr(job_dir, firm, data_root, log=log)
    _exhibit(sr, [ORDER_PAGE, PROSE, "second " + PROSE])
    _doc(sr, _cite("2, 9").replace("Cervical strain noted.", "Cervical strain noted (Exhibit 1 - p. 0)."))
    assert classify.run_nonrecord(sr) == 0
    assert strip.run_falsify(sr) == 0
    assert strip.run_dry(sr) == 0
    assert strip.run_apply(sr) == 0
    doc = (sr.slug_dir / "runs" / "alpha" / "final-chronology.md").read_text()
    assert "(Exhibit 1 - p. 1, 9)" in doc and "| Exhibit 1 - p. 1, 9 |" in doc and "(Exhibit 1 - p. 0)" in doc
    assert any("outside the exhibit" in line and "9" in line for line in log), log


def test_strip_leaves_an_exhibit_unstripped_when_a_moved_cited_page_fails_the_content_check(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Row 23 (2026-10-24): on a form or annotated PDF the page delete_pages
    rewrites can differ from the original in pixels though it is the same
    paper. That refused the run. Leaving the exhibit as it was, with its
    citations unchanged, is always true to the paper: the strip only removes
    non-record pages, so keeping them costs nothing and needs no disclosure."""
    import pymupdf

    real = strip.page_match

    def form_page(a, b, dpi=60):  # the form page renders differently once its neighbours are gone
        return (False, 0.4) if "FORM PAGE" in a.get_text() else real(a, b, dpi)

    monkeypatch.setattr(strip, "page_match", form_page)
    sr = _sr(job_dir, firm, data_root)
    pdf = _exhibit(sr, [ORDER_PAGE, INDEX_PAGE, "FORM PAGE " + PROSE])
    _doc(sr, DOC_3)
    assert classify.run_nonrecord(sr) == 0
    assert strip.run_dry(sr) == 0
    assert strip.run_apply(sr) == 0
    doc = (sr.slug_dir / "runs" / "alpha" / "final-chronology.md").read_text()
    assert "(Exhibit 1 - p. 3)" in doc and "| Exhibit 1 - p. 3 |" in doc  # citations unchanged
    with pymupdf.open(str(pdf)) as d:
        assert len(d) == 3  # nothing removed
    result = json.loads((sr.slug_dir / "strip_result.json").read_text())
    assert result["drops"] == {"1": []} and result["unstripped"][0]["exhibit"] == 1


def test_strip_falsifier_ignores_a_byte_identical_neighbour_and_still_fires_on_a_blind_check(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Row 24 (2026-10-24): a fax that repeats a page gives the cited page a
    byte-identical neighbour, so the one-page offset lands on the same paper
    and the falsifier called the check blind. An identical neighbour cannot
    tell the check anything; it is skipped. A check that really cannot fail
    still stops the run (exit 1)."""
    sr = _sr(job_dir, firm, data_root)
    _exhibit(sr, [ORDER_PAGE, PROSE, PROSE])
    _doc(sr, _cite("2"))
    assert classify.run_nonrecord(sr) == 0
    assert strip.run_falsify(sr) == 0
    # the control: distinct neighbours, and an instrument that cannot fail
    _exhibit(sr, [ORDER_PAGE, PROSE, "second " + PROSE])
    assert classify.run_nonrecord(sr) == 0
    monkeypatch.setattr(strip, "page_match", lambda a, b, dpi=60: (True, 0.0))
    assert strip.run_falsify(sr) == 1


def test_strip_falsifier_does_not_count_an_offset_past_the_last_page_as_proof(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """N6 (2026-10-24): when the only moved cited page is an exhibit's last
    page, the one-page offset lands past the end. That compared no content, yet
    it counted as a mismatch, so a blind check "proved" it could fail. It is
    not comparable; with nothing comparable the falsifier is not applicable."""
    log: list[str] = []
    sr = _sr(job_dir, firm, data_root, log=log)
    _exhibit(sr, [ORDER_PAGE, PROSE])
    _doc(sr, _cite("2"))
    assert classify.run_nonrecord(sr) == 0
    monkeypatch.setattr(strip, "page_match", lambda a, b, dpi=60: (True, 0.0))  # blind
    assert strip.run_falsify(sr) == 0
    assert any("not applicable" in line for line in log), log
    assert not any("Check can fail; trustworthy" in line for line in log), log


def _two_exhibits(sr: StageRun) -> tuple[Path, Path]:
    out = sr.slug_dir / "out" / "alpha"
    out.mkdir(parents=True, exist_ok=True)
    rows, pdfs = [], []
    for n, (label, body) in enumerate([("Example Clinic", PROSE), ("Example Imaging", "imaging " + PROSE)], 1):
        title = f"Exhibit {n} - {label} - 01-02-2026 (Medical Records)"
        pages = [ORDER_PAGE, INDEX_PAGE, body, body + " second"]
        (out / f"{title}.pdf").write_bytes(make_pdf(pages))
        pdfs.append(out / f"{title}.pdf")
        rows.append(
            {
                "exhibit": n,
                "title": title,
                "total_pages": 4,
                "files": [{"file": f"{label}.pdf", "old_exhibit": n, "start_page": 1, "pages": 4}],
            }
        )
    (out / "page_map.json").write_text(json.dumps(rows))
    _doc(
        sr,
        DOC_3.replace("(Exhibit 1 - p. 3)", "(Exhibit 1 - p. 3-4) and imaging. (Exhibit 2 - p. 4)").replace(
            "| Exhibit 1 - p. 3 |", "| Exhibit 2 - p. 3 |"
        ),
    )
    return pdfs[0], pdfs[1]


def _killed_at(monkeypatch: pytest.MonkeyPatch, n: int, step: str = "replace") -> None:
    """The process dies at the n-th exhibit ``step`` (the ``.orig`` hard link,
    or the replace that follows it), as an OOM kill would."""
    real, seen = getattr(Path, step), []

    def wrapped(self, target):
        if self.name.endswith(".stripped" if step == "replace" else ".orig"):
            seen.append(self)
            if len(seen) == n:
                raise SystemExit("killed")
        return real(self, target)

    monkeypatch.setattr(Path, step, wrapped)


@pytest.mark.parametrize("step", ["hardlink_to", "replace"])
def test_an_apply_killed_between_two_exhibits_resumes_and_completes(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch, step: str
) -> None:
    """N3 (2026-10-24): the daemon requeues a signal-killed runner. An apply
    killed after replacing one of two exhibits came back REFUSED "already
    stripped once", which needs a person. It resumes: the replaced exhibit is
    kept because it IS the stripped result by content, the other is stripped,
    and the chronology is rewritten once from the pre-strip copy."""
    import pymupdf

    sr = _sr(job_dir, firm, data_root)
    ex1, ex2 = _two_exhibits(sr)
    assert classify.run_nonrecord(sr) == 0
    _killed_at(monkeypatch, 2, step)
    with pytest.raises(SystemExit):
        strip.run_apply(sr)
    monkeypatch.undo()
    assert Path(str(ex1) + ".orig").exists()
    # killed before the second link, or between the link and its replace
    assert not Path(str(ex2) + ".orig").exists() or Path(str(ex2) + ".orig").samefile(ex2)
    assert Path(str(ex2) + ".orig").exists() == (step == "replace")
    assert strip.run_apply(sr) == 0
    doc = (sr.slug_dir / "runs" / "alpha" / "final-chronology.md").read_text()
    assert "(Exhibit 1 - p. 1-2)" in doc and "(Exhibit 2 - p. 2)" in doc and "| Exhibit 2 - p. 1 |" in doc
    for pdf in (ex1, ex2):
        with pymupdf.open(str(pdf)) as d:
            assert len(d) == 2
    assert json.loads((sr.slug_dir / "strip_result.json").read_text())["new_page_counts"] == {"1": 2, "2": 2}
    assert not strip.prestrip_path(sr.slug_dir / "runs" / "alpha" / "final-chronology.md").exists()
    assert strip.run_apply(sr) == 1  # a finished apply still never runs twice


def test_an_apply_killed_after_the_chronology_rewrite_rewrites_it_the_same(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The rewrite is idempotent: a kill after the chronology was renumbered
    resumes from the pre-strip copy, never renumbering the renumbered text."""
    sr = _sr(job_dir, firm, data_root)
    _two_exhibits(sr)
    assert classify.run_nonrecord(sr) == 0
    real = Path.write_text

    def write_text(self, data, *a, **k):
        if self.name == "strip_result.json":
            raise SystemExit("killed")
        return real(self, data, *a, **k)

    monkeypatch.setattr(Path, "write_text", write_text)
    with pytest.raises(SystemExit):
        strip.run_apply(sr)
    monkeypatch.undo()
    cpath = sr.slug_dir / "runs" / "alpha" / "final-chronology.md"
    first = cpath.read_text()
    assert strip.run_apply(sr) == 0
    assert cpath.read_text() == first and "(Exhibit 1 - p. 1-2)" in first


def test_a_resumed_apply_refuses_an_exhibit_that_is_not_the_stripped_result(
    job_dir: Path, firm: Path, data_root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The marker is never trusted: an exhibit with a `.orig` whose current
    bytes are not the stripped result is a real defect, and still refuses."""
    sr = _sr(job_dir, firm, data_root)
    ex1, _ = _two_exhibits(sr)
    assert classify.run_nonrecord(sr) == 0
    _killed_at(monkeypatch, 2)
    with pytest.raises(SystemExit):
        strip.run_apply(sr)
    monkeypatch.undo()
    ex1.write_bytes(make_pdf(["something else", "entirely"]))
    assert strip.run_apply(sr) == 1


def test_strip_gives_each_refusal_its_own_exit_code(job_dir: Path, firm: Path, data_root: Path) -> None:
    """Other (2026-10-24): strip_falsify's exit 1 meant three things, so the
    reason the driver printed ("the content check cannot fail") was wrong for
    two of them. Each cause has its own code and reason in dag.py."""
    from medchron import dag

    sr = _sr(job_dir, firm, data_root)
    _exhibit(sr, [ORDER_PAGE, PROSE, "second " + PROSE])
    _doc(sr, _cite("2"))
    assert strip.run_falsify(sr) == strip.EXIT_NO_CLASSIFIER  # classify_nonrecord has not run
    assert classify.run_nonrecord(sr) == 0
    (sr.slug_dir / "drops.json").write_text(
        json.dumps({"manual_drops": {}, "citation_edits": [{"find": "no such text", "replace": "x"}]})
    )
    assert strip.run_falsify(sr) == strip.EXIT_EDIT_NOT_FOUND
    codes = {strip.EXIT_BLIND, strip.EXIT_CITED_DROP, strip.EXIT_EDIT_NOT_FOUND, strip.EXIT_NO_CLASSIFIER}
    assert len(codes) == 4
    for name in ("strip_falsify", "strip_dry", "strip_apply"):
        reasons = {c: dag.BY_NAME[name].exit_map[c][1] for c in codes - {strip.EXIT_BLIND}}
        assert "citation edit" in reasons[strip.EXIT_EDIT_NOT_FOUND]
        assert "classify_nonrecord" in reasons[strip.EXIT_NO_CLASSIFIER]
    assert "cannot fail" in dag.BY_NAME["strip_falsify"].exit_map[strip.EXIT_BLIND][1]


def test_strip_keeps_scanned_pages_when_the_classifiers_controls_failed(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    """Rows 3 and 15 (2026-10-24): failed scanned-classifier controls refused
    the strip. Its labels are simply not used: the scanned pages stay, the text
    classifier's drops still apply, the run goes on."""
    sr = _sr(job_dir, firm, data_root)
    _exhibit(sr, [ORDER_PAGE, PROSE, ""])
    _doc(sr, _cite("2"))
    assert classify.run_nonrecord(sr) == 0
    (sr.slug_dir / "scanned_labels.json").write_text(
        json.dumps({"labels": {"Ex1p3": "AUTH"}, "nonrecord": {"1": [[3, "AUTH"]]}, "controls_ok": False})
    )
    assert strip.run_dry(sr) == 0
    assert strip.run_apply(sr) == 0
    assert json.loads((sr.slug_dir / "strip_result.json").read_text())["drops"] == {"1": [1]}


# ---- coverage gate --------------------------------------------------------------------------
def test_coverage_gate_accounts_for_every_pulled_file(job_dir: Path, firm: Path, data_root: Path) -> None:
    sr = _sr(job_dir, firm, data_root)
    d = sr.slug_dir
    (d / "units").mkdir(parents=True)
    (d / "units" / "alpha.json").write_text(
        json.dumps(
            [
                {"id": "a", "name": "clinic note", "ext": ".pdf"},
                {"id": "b", "name": "Retainer signed", "ext": ".pdf"},
                {"id": "c", "name": "mystery scan", "ext": ".pdf"},
            ]
        )
    )
    (d / "raw_manifest.jsonl").write_text(
        "".join(
            json.dumps(r) + "\n"
            for r in [
                {"id": "a", "name": "clinic note", "ext": ".pdf", "ok": True},
                {"id": "b", "name": "Retainer signed", "ext": ".pdf", "ok": True},
                {"id": "c", "name": "mystery scan", "ext": ".pdf", "ok": True},
                {"id": "d", "name": "stray photo", "ext": ".png", "ok": True},
            ]
        )
    )
    rd = d / "runs" / "alpha"
    rd.mkdir(parents=True)
    (rd / "entries_final.md").write_text(
        "01/02/2026\nExample Clinic | Medical Diagnoses\n\nStrain. (Exhibit 1 - p. 1)\n"
    )
    (d / "out" / "alpha").mkdir(parents=True)
    (d / "out" / "alpha" / "page_map.json").write_text(
        json.dumps([{"exhibit": 1, "files": [{"file": "clinic note.pdf"}]}])
    )
    log: list[str] = []
    sr = _sr(job_dir, firm, data_root, log=log)
    assert coverage.run(sr) == 1  # the stray photo is owned by nobody
    assert any("COVERAGE GATE FAIL" in line and "1 pulled file" in line for line in log)
    (d / "orphans.json").write_text(
        json.dumps({"orphans": [{"name": "stray photo.png", "reason": "email signature graphic"}]})
    )
    log.clear()
    assert coverage.run(sr) == 1  # mystery scan is in the unit, uncited, unaccounted: hold
    assert json.loads((rd / "coverage_unexplained.json").read_text()) == ["mystery scan.pdf"]
    (d / "billing_docs.json").write_text(json.dumps({"docs": [{"name": "mystery scan", "path": "x", "pages": 1}]}))
    log.clear()
    assert coverage.run(sr) == 0
    assert any("engagement document" in line for line in log) and any("GATE PASS" in line for line in log)


def test_coverage_gate_takes_the_composers_own_account_of_an_uncited_file(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    """Live 2026-09-16: 97 pleadings, filings, photos and letters, each with the
    composer's `nothing extractable: <reason>` in FILES-SEEN, held the run
    because the firm's name table had never met them. The composer read every
    page; its reason is the record's own statement. `entries: 0` with no reason
    still explains nothing, and a renamed duplicate is found by its file id."""
    sr = _sr(job_dir, firm, data_root)
    d = sr.slug_dir
    (d / "units").mkdir(parents=True)
    files = [
        {"id": "a", "name": "clinic note", "ext": ".pdf"},
        {"id": "p", "name": "COMPLAINT", "ext": ".pdf"},
        {"id": "b", "name": "EXHIBIT B", "ext": ".pdf"},
        {"id": "s", "name": "silent scan", "ext": ".pdf"},
        {"id": "L", "name": "ledger scan", "ext": ".pdf"},
        {"id": "E", "name": "has entries", "ext": ".pdf"},
        {"id": "msgatt-2", "name": "image001 (2)", "ext": ".jpg"},
    ]
    (d / "units" / "alpha.json").write_text(json.dumps(files))
    (d / "raw_manifest.jsonl").write_text("".join(json.dumps({**f, "ok": True}) + "\n" for f in files))
    rd = d / "runs" / "alpha"
    rd.mkdir(parents=True)
    (rd / "entries_final.md").write_text(
        "01/02/2026\nExample Clinic | Medical Diagnoses\n\nStrain. (Exhibit 1 - p. 1)\n"
    )
    (rd / "map-01.md").write_text(
        "## ENTRIES\nnone in this chunk\n\n## INDEX\n\n## BILLING-DATES\nnone in this chunk\n\n"
        "## CONFLICTS / REFERENCED-BUT-ABSENT\nnone observed\n\n## FILES-SEEN\n"
        "=== FILE: clinic note.pdf (fileId a) === entries: 1\n"
        "=== FILE: COMPLAINT.pdf (fileId p) === | nothing extractable: pleading, no treatment record\n"
        "=== FILE: EXHIBIT B.pdf (fileId b) === billing-dates: 29\n"
        "=== FILE: ledger scan.pdf (fileId L) === entries: 0, billing-dates: 332\n"
        "=== FILE: has entries.pdf (fileId E) === entries: 4, billing-dates: 7\n"
        "=== FILE: silent scan.pdf (fileId s) === entries: 0\n"
        "=== FILE: image001.jpg (fileId msgatt-2) === nothing extractable: photo of an insurance card\n"
    )
    (d / "out" / "alpha").mkdir(parents=True)
    (d / "out" / "alpha" / "page_map.json").write_text(
        json.dumps([{"exhibit": 1, "files": [{"file": "clinic note.pdf"}]}])
    )
    log: list[str] = []
    sr = _sr(job_dir, firm, data_root, log=log)
    assert coverage.run(sr) == 1
    # "has entries.pdf" has citable content and is not cited: LOSS, an
    # integrity stop, never disclosed. The pure ledger is explained
    # by its own billing account. Since 2026-10-23 `entries: 0` IS an account:
    # the composer read the silent scan and drew nothing, so it is explained
    # by rule and recorded for the firm-facing disclosure.
    assert json.loads((rd / "coverage_unexplained.json").read_text()) == ["has entries.pdf"], log
    assert any("silent scan.pdf" in line and "no dated medical" in line for line in log), log
    from medchron import coverage_decide

    rule = coverage_decide.rows(d, "alpha")
    assert [(r["name"], r["by"]) for r in rule] == [("silent scan.pdf", "rule")]
    assert any("COMPLAINT.pdf" in line and "composer: pleading" in line for line in log), log
    assert any("EXHIBIT B.pdf" in line and "billing dates only" in line for line in log), log
    # the COMBINED form the composer actually writes: `entries: 0, billing-dates: N`.
    # A prefix test saw only the bare form and held a finished chronology on a
    # live matter (an A&P matter, two ledgers, 348 dates of service).
    assert any("ledger scan.pdf" in line and "billing dates only (332)" in line for line in log), log
    assert any("image001 (2).jpg" in line and "composer: photo" in line for line in log), "found by id after the rename"
    # the rehearsal answers from the same source
    lines = coverage.rehearse(sr)
    # "with nothing cited": the clinic note and the file with entries. Never
    # the explained five -- the pure ledger by its own billing account rather
    # than by its name, and the silent scan by its entries: 0 (2026-10-23).
    assert lines[0].endswith("2 would need a citation") and not any("silent scan" in ln for ln in lines), lines
    assert not any("COMPLAINT" in ln or "EXHIBIT B" in ln or "image001" in ln or "ledger scan" in ln for ln in lines), (
        lines
    )


# ---- billing chart and worksheet -------------------------------------------------------------------
def test_money_parses_printed_shapes_and_flags_lost_decimals() -> None:
    sus: list = []
    assert billing_chart.money("$1,800.00", sus, 500_000) == 1800.0
    assert billing_chart.money("112,547,03", sus, 500_000) == 112547.03  # decimal comma
    assert billing_chart.money("$4700 I 00", sus, 500_000) == 4700.0  # CMS-1500 adjacent boxes
    assert (
        billing_chart.money("235000", sus, 500_000) == 235000.0
        and billing_chart.money("$1,?00.00", sus, 500_000) is None
    )
    assert billing_chart.money("1,000,000.00", sus, 500_000) == 1000000.0 and sus == [("1,000,000.00", 1000000.0)]
    assert billing_chart.total_rank("28. TOTAL CHARGE") == 0 and billing_chart.total_rank("Overall - Total") == 3


def test_billing_chart_prefers_a_ledger_total_and_reports_the_rest(job_dir: Path, firm: Path, data_root: Path) -> None:
    sr = _sr(job_dir, firm, data_root)
    _doc(sr, DOC_3)
    rows = [
        {
            "file": "Example Clinic ledger",
            "pages": 2,
            "failures": [],
            "chunks": [
                {
                    "doc_type": "LEDGER",
                    "provider": "Example Clinic",
                    "patient": "Alpha Example",
                    "date_first": "01/02/2026",
                    "date_last": "01/09/2026",
                    "printed_totals": [
                        {"label": "Total Charges", "amount": "$1,800.00", "page": 1},
                        {"label": "Insurance adjustment", "amount": "-$300.00", "page": 2},
                    ],
                },
                {
                    "provider": "",
                    "printed_totals": [{"label": "Overall - Total", "amount": "$2,100.00", "page": 2}],
                    "doc_type": "LEDGER",
                },
            ],
        },
        {
            "file": "claim form",
            "pages": 1,
            "failures": [],
            "chunks": [
                {
                    "doc_type": "MEDICAL_BILL",
                    "provider": "Ex Clinic",
                    "printed_totals": [{"label": "28. TOTAL CHARGE", "amount": "$465.00", "page": 1}],
                }
            ],
        },
        {
            "file": "recovery vendor",
            "pages": 1,
            "failures": [],
            "chunks": [
                {
                    "doc_type": "LIEN_SUBROGATION",
                    "provider": "Example Health Plan",
                    "printed_totals": [{"label": "claim", "amount": "$900.00", "page": 1}],
                }
            ],
        },
        {
            "file": "unknown bill",
            "pages": 1,
            "failures": [],
            "chunks": [
                {
                    "doc_type": "MEDICAL_BILL",
                    "provider": "Somebody Else Imaging",
                    "printed_totals": [{"label": "Balance", "amount": "$50.00", "page": 1}],
                }
            ],
        },
    ]
    (sr.slug_dir / "billing_extract.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows))
    assert billing_chart.run(sr) == 0
    chart = json.loads((sr.slug_dir / "billing_chart.json").read_text())
    row = chart["rows"][0]
    assert row["provider"] == "Example Clinic" and row["total"] == 2100.0 and row["basis"] == "Example Clinic ledger"
    assert row["adjustments"] == [
        {"amount": 300.0, "label": "Insurance adjustment", "doc": "Example Clinic ledger", "page": 2}
    ]
    assert chart["grand_total"] == 2100.0 and chart["inherited_attribution"] == [
        ["Example Clinic ledger", None, "Example Clinic"]
    ]
    assert (
        chart["unmatched"] == {"Somebody Else Imaging": [50.0]} and chart["subrogation"][0][1] == "Example Health Plan"
    )
    assert billing_docx.run(sr) == 0
    out = list((sr.slug_dir / "out" / "alpha").glob("* - Medical Billing Worksheet *.docx"))
    assert len(out) == 1 and out[0].name.startswith("Alpha Example - Medical Billing Worksheet ")
    from docx import Document

    text = "\n".join(p.text for p in Document(str(out[0])).paragraphs)
    assert "Alpha Example - Medical Billing Summary" in text and "Health-Plan and Subrogation Claims" in text


def test_billing_docx_refuses_a_suspect_amount(job_dir: Path, firm: Path, data_root: Path) -> None:
    sr = _sr(job_dir, firm, data_root)
    sr.slug_dir.mkdir(parents=True, exist_ok=True)
    (sr.slug_dir / "billing_chart.json").write_text(
        json.dumps({"rows": [], "suspect_amounts": [["1,000,000", 1000000.0]]})
    )
    assert billing_docx.run(sr) == 1


# ---- stop audit rows 17 and 18: an unread page and a large amount are disclosed, never a stop ----
def _ledger_with(sr: StageRun, amount: str, page_text: str) -> None:
    """A ledger whose page 1 prints `page_text` (empty = a scan, no text
    layer) and whose transcription read `amount` as the account total."""
    raw = sr.slug_dir / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    p = raw / "Example Clinic ledger.pdf"
    p.write_bytes(make_pdf([page_text]))
    (sr.slug_dir / "billing_docs.json").write_text(
        json.dumps({"docs": [{"id": "b1", "name": "Example Clinic ledger", "path": str(p)}]})
    )
    rec = {
        "file": "Example Clinic ledger",
        "pages": 1,
        "failures": [],
        "chunks": [
            {
                "doc_type": "LEDGER",
                "provider": "Example Clinic",
                "patient": "Alpha Example",
                "date_first": "01/02/2026",
                "date_last": "01/09/2026",
                "printed_totals": [{"label": "Total Charges", "amount": amount, "page": 1}],
            }
        ],
    }
    (sr.slug_dir / "billing_extract.jsonl").write_text(json.dumps(rec) + "\n")


def _worksheet_text(sr: StageRun) -> str:
    from docx import Document

    out = list((sr.slug_dir / "out" / "alpha").glob("* - Medical Billing Worksheet *.docx"))
    assert len(out) == 1
    d = Document(str(out[0]))
    # Document order, paragraphs and table cells interleaved as they render.
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return "\n".join("".join(p.itertext()) for p in d.element.body.iter(f"{w}p"))


def test_a_large_amount_the_page_prints_with_its_decimal_is_accepted(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    """A spinal-fusion ledger can legitimately total over the $500k suspect
    threshold. When the page's own text prints that figure with its decimal
    point, it is not a lost decimal: the worksheet renders (row 18)."""
    sr = _sr(job_dir, firm, data_root)
    _doc(sr, DOC_3)
    _ledger_with(sr, "$612,345.67", "Example Clinic\nTotal Charges   $612,345.67\n")
    assert billing_chart.run(sr) == 0
    chart = json.loads((sr.slug_dir / "billing_chart.json").read_text())
    assert chart["suspect_amounts"] == [] and chart["rows"][0]["total"] == 612345.67
    assert billing_docx.run(sr) == 0
    assert json.loads((sr.slug_dir / "billing_dropped_rows.json").read_text()) == []
    assert "$612,345.67" in _worksheet_text(sr)


@pytest.mark.parametrize(
    "page_text",
    ["Example Clinic\nTotal Charges   $6,123.45\n", ""],
    ids=["page-prints-another-figure", "scan-no-text-layer"],
)
def test_a_large_amount_the_page_does_not_print_is_dropped_and_disclosed(
    job_dir: Path, firm: Path, data_root: Path, page_text: str
) -> None:
    """A transcription that lost its decimal ("612345" for $6,123.45), or a
    figure on a scan with nothing to check it against, is not put in the
    total and is not guessed at either: the row is left out, recorded in
    billing_dropped_rows.json, and named on the worksheet (row 18)."""
    sr = _sr(job_dir, firm, data_root)
    _doc(sr, DOC_3)
    _ledger_with(sr, "612345", page_text)
    assert billing_chart.run(sr) == 0
    chart = json.loads((sr.slug_dir / "billing_chart.json").read_text())
    assert chart["suspect_amounts"] == [] and chart["rows"][0]["total"] is None
    assert chart["grand_total"] == 0.0, "no amount is invented for the dropped row"
    dropped = json.loads((sr.slug_dir / "billing_dropped_rows.json").read_text())
    assert [(r["file"], r["page"], r["amount_text"]) for r in dropped] == [("Example Clinic ledger", 1, "612345")]
    assert set(dropped[0]) == {"file", "page", "amount_text", "reason"} and dropped[0]["reason"]
    assert billing_docx.run(sr) == 0
    text = _worksheet_text(sr)
    assert "Amounts Left Out" in text and "Example Clinic ledger" in text and "p. 1" in text
    assert "612345" not in text and "612,345" not in text, "an unconfirmed figure is never shown as a figure"


def test_the_worksheet_names_each_unread_billing_page_as_not_totalled(
    job_dir: Path, firm: Path, data_root: Path
) -> None:
    """Row 17's worksheet half: a FAILED_PAGE is named by file and page, and
    the worksheet says plainly its amount is not in the total."""
    sr = _sr(job_dir, firm, data_root)
    _doc(sr, DOC_3)
    _ledger_with(sr, "$1,800.00", "Total Charges $1,800.00")
    rec = json.loads((sr.slug_dir / "billing_extract.jsonl").read_text())
    rec["chunks"].append({"FAILED_PAGE": 2})
    rec["failures"] = [2]
    (sr.slug_dir / "billing_extract.jsonl").write_text(json.dumps(rec) + "\n")
    assert billing_chart.run(sr) == 0
    assert billing_docx.run(sr) == 0
    text = _worksheet_text(sr)
    assert "Pages Not Read" in text and "not totalled" in text
    assert "Example Clinic ledger" in text and "p. 2" in text


# ---- stop audit row 28: a clinician-named spine and facility-labelled bills ----
CLINICIAN_DOC = (
    "Alpha Example - Medical Chronology\n\n## Treatment Timeline\n\n"
    "| Medical Provider | Treatment Period | Visits | Reference |\n"
    "| Priya Rao, PA-C (NVRH 9W Laboratory) | 01/05/2026 | 1 | Exhibit 1 - p. 1 |\n"
    "| Mark Hale, MD (NVRH Ophthalmology) | 01/07/2026 - 01/20/2026 | 2 | Exhibit 1 - p. 2 |\n"
    "| Eva Ramos, M.D. | 02/01/2026 | 1 | Exhibit 2 - p. 1 |\n"
    "| Eva M Ramos, MD | 02/10/2026 | 1 | Exhibit 2 - p. 4 |\n"
    "| Joan Kerr, RN (Valley Care at Home) | 03/01/2026 - 03/09/2026 | 3 | Exhibit 3 - p. 1 |\n\n"
    "## Diagnostic Highlights\n\n| ICD Code | Description | First Diagnosed | Reference |\n\n"
    "## Medical Chronology\n\n01/05/2026\nPriya Rao, PA-C (NVRH 9W Laboratory) | Labs\n\nCBC drawn. "
    "(Exhibit 1 - p. 1)\n\n## Exhibit List\n\n| Exhibit No. | Description |\n"
)


def _ledger(file: str, provider: str, amount: str, label: str = "Total Charges") -> dict:
    return {
        "file": file,
        "pages": 1,
        "failures": [],
        "chunks": [
            {
                "doc_type": "LEDGER",
                "provider": provider,
                "patient": "Alpha Example",
                "date_first": "01/05/2026",
                "date_last": "03/09/2026",
                "printed_totals": [{"label": label, "amount": amount, "page": 1}],
            }
        ],
    }


def test_the_facility_rule_matches_by_words_and_initials_and_not_by_generic_words() -> None:
    from medchron.stages import billing_match as BM

    assert BM.acronym("NV Ridge Health") == "nvrh" and BM.acronym("NV Ridge Medical Center") == "nvrmc"
    assert BM.facility_hit("NV Ridge Health", "NVRH 9W Laboratory")
    assert BM.facility_hit("Valley Care at Home, Inc.", "Valley Care at Home")
    assert not BM.facility_hit("Oakmont Health", "Valley Health"), "a shared generic word is not a match"
    assert not BM.facility_hit("Medical Center", "Valley Medical Center Hospital")
    assert BM.person_key("Eva Ramos, M.D.") == BM.person_key("Eva M Ramos, MD") == ("eva", "ramos")
    assert BM.person_key("Example Clinic") is None, "an uncredentialed name pair is not read as a person"
    assert BM.person_key("Joan Kerr, RN (Valley Care at Home)") == ("joan", "kerr")
    assert BM.facilities("Joan Kerr, RN (Valley Care at Home)") == ["Valley Care at Home"]
    assert BM.facilities("Example Clinic") == ["Example Clinic"]


def test_facility_bills_total_against_a_clinician_named_spine(job_dir: Path, data_root: Path, tmp_path: Path) -> None:
    """The seat replay shape (row 28): the Treatment Timeline lists
    clinicians, the bills are labelled by facility, and the firm authored no
    provider_match patterns for this matter. The old matcher matched nothing
    and the worksheet read SUBTOTAL $0.00 with every provider a gap. Now a
    bill matches through the facility in the provider cell (its own initials:
    "NV Ridge Health" is NVRH), one clinician listed twice is one row, and a
    bill that matches no provider is shown as printed under its own subtotal,
    OUTSIDE the grand total (a first version of this worksheet once totalled
    3x the firm's lien report), and a document that already supplied a
    matched row adds nothing to it (same paper, counted once)."""
    cfg = json.loads(json.dumps(FIRM_CONFIG))
    cfg["billing"]["provider_match"] = {}
    firm = tmp_path / "bare-firm.yaml"
    firm.write_text(yaml.safe_dump(cfg))
    log: list[str] = []
    sr = _sr(job_dir, firm, data_root, log=log)
    _doc(sr, CLINICIAN_DOC)
    rows = [
        _ledger("NVRH statement", "NV Ridge Health", "$12,345.67", "Overall - Total"),
        _ledger("Valley home care ledger", "Valley Care at Home", "$2,000.00"),
        _ledger("Ramos ledger", "Eva Ramos MD", "$300.00"),
        _ledger("Ambulance invoice", "Sierra Ambulance Authority", "$1,500.00"),
    ]
    # The NVRH statement also carries a page billed under a name that matches
    # nothing: that paper already supplied a matched row, so it stays there.
    rows[0]["chunks"].append(
        {
            "doc_type": "LEDGER",
            "provider": "Ridgeline Billing Office",
            "printed_totals": [{"label": "Total Charges", "amount": "$12,345.67", "page": 2}],
        }
    )
    (sr.slug_dir / "billing_extract.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows))
    assert billing_chart.run(sr) == 0
    chart = json.loads((sr.slug_dir / "billing_chart.json").read_text())
    by = {r["provider"]: r for r in chart["rows"]}
    assert chart["grand_total"] == 14645.67, (by, log)
    nvrh = next(r for r in chart["rows"] if r.get("merged_from") and r["total"] == 12345.67)
    assert sorted(nvrh["merged_from"]) == [
        "Mark Hale, MD (NVRH Ophthalmology)",
        "Priya Rao, PA-C (NVRH 9W Laboratory)",
    ]
    ramos = [r for r in chart["rows"] if "Ramos" in r["provider"]]
    assert len(ramos) == 1 and ramos[0]["total"] == 300.0, "one clinician is one row"
    assert (ramos[0]["first"], ramos[0]["last"]) == ("2026-02-01", "2026-02-10")
    assert by["Joan Kerr, RN (Valley Care at Home)"]["total"] == 2000.0
    assert "Sierra Ambulance Authority" not in by, "an as-printed row is not in the grand total"
    printed = {r["provider"]: r for r in chart["as_printed_rows"]}
    assert list(printed) == ["Sierra Ambulance Authority"], "the same paper is counted once"
    assert printed["Sierra Ambulance Authority"]["total"] == 1500.0
    assert printed["Sierra Ambulance Authority"]["as_printed"] is True and chart["as_printed_subtotal"] == 1500.0
    assert set(chart["unmatched"]) == {"Sierra Ambulance Authority", "Ridgeline Billing Office"}, "still reported"
    assert billing_docx.run(sr) == 0
    text = _worksheet_text(sr)
    assert "$14,645.67" in text and "GRAND TOTAL" in text
    head = "Billed under a name not matched to a treating provider"
    assert head in text and text.index(head) > text.index("GRAND TOTAL")
    assert "Sierra Ambulance Authority" in text and "SUBTOTAL (not in the grand total)" in text
    assert "$1,500.00" in text and "Ridgeline Billing Office" not in text
