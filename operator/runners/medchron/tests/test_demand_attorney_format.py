"""An attorney's own demand format, chosen by the responsible attorney
(demand-firm.yaml ``formats`` + ``attorneys.<name>.format``; house_numbered.py).

The firm asked on 2026-10-14 that one attorney's demands follow his own
format. These tests hold the three things that ask turns on: the default
attorney's demand is unchanged to the byte, the mapped attorney's demand is
rendered and format-checked in his own format, and an unmapped attorney still
holds. Every fixture is synthetic (demand_numbered_testkit.py)."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from demand_numbered_testkit import (
    BANNER_LINES,
    NUMBERED_DRAFT,
    SIGNER,
    edit_config,
    make_numbered_inputs,
    make_numbered_reference,
    variant,
)
from demand_testkit import DRAFT, SIGNATURE, ScriptedClient, make_inputs, make_job, make_pdf, seat_with, standard_docs
from medchron_testkit import PRICING
from medchron.demand import firm as firm_mod, format_check, house, house_numbered
from medchron.demand.run import DemandRun

CHECKER = Path(__file__).resolve().parents[3] / "templates" / "drafting" / "drafting_gate_check.py"
NUMBERED_ATTORNEY = "Gamma Counsel"


@pytest.fixture(autouse=True)
def _checker(monkeypatch):
    monkeypatch.setenv("SMD_DRAFTING_GATE_CHECK", str(CHECKER))


@pytest.fixture
def pricing(tmp_path: Path) -> Path:
    p = tmp_path / "pricing.json"
    p.write_text(json.dumps(PRICING), encoding="utf-8")
    return p


def _run(tmp_path: Path, pricing: Path, inputs: Path, attorney: str, draft: str, docs=None):
    seat = seat_with(docs or standard_docs())
    seat.facts = {**seat.facts, "responsible_attorney": attorney}
    client = ScriptedClient(draft=draft)
    jd = make_job(tmp_path / "job")
    r = DemandRun(
        jd,
        inputs_dir=str(inputs),
        pricing=str(pricing),
        seat_factory=lambda: seat,
        client=client,
        log=lambda _m: None,
        readback_pause=0.0,
        vendor_factory=lambda: None,
    )
    return r.run(), seat, client, tmp_path / "job" / "data" / "out" / "demand"


# ---- the config ------------------------------------------------------------------
def test_an_attorney_format_validates_and_selects_by_the_attorney_row(tmp_path):
    firm = firm_mod.load(make_numbered_inputs(tmp_path / "in"))
    row = firm.attorney_for(" gamma  COUNSEL ")
    assert row and row["format"] == "numbered" and row["initials"] == ""
    assert firm.variants_for(row)["pre_suit"]["layout"] == "numbered"
    assert firm.variants_for(firm.attorney_for("Example Lawyer")) is firm.data["variants"]
    picked = firm.select("pre_suit", row)
    assert picked.input_path("house_reference").name == "house-reference-numbered.docx"
    assert picked.input_path("skeleton").name == "skeleton-numbered.md"
    default = firm.select("pre_suit", firm.attorney_for("Example Lawyer"))
    assert default.input_path("house_reference").name == "house-reference.docx"


@pytest.mark.parametrize(
    "edit, want",
    [
        (lambda d: d["formats"]["numbered"]["pre_suit"].update(lead_headng="X"), "lead_headng: unknown key"),
        (lambda d: d["formats"]["numbered"]["pre_suit"].pop("lead_heading"), "lead_heading: required for the numbered"),
        (lambda d: d["formats"]["numbered"]["pre_suit"].update(layout="fancy"), "layout: expected one of"),
        (lambda d: d["formats"]["numbered"]["pre_suit"].pop("headings"), "headings: required"),
        (lambda d: d["formats"]["numbered"].update(appeal={}), "unknown variant"),
        (lambda d: d["formats"]["numbered"].pop("pre_suit"), "formats.numbered.pre_suit: required"),
        (lambda d: d.update(fromats=d.pop("formats")), "fromats: unknown section"),
        (lambda d: d["attorneys"][NUMBERED_ATTORNEY].update(format="nope"), "names no entry in formats"),
        (lambda d: d["attorneys"][NUMBERED_ATTORNEY].update(fromat="numbered"), "expected {signer, title, initials}"),
        (lambda d: d["attorneys"][NUMBERED_ATTORNEY].update(title="Attorney for {cleint}"), "only placeholder"),
        (lambda d: d["attorneys"][NUMBERED_ATTORNEY].update(initials=None), "initials: expected a string"),
        (lambda d: d["attorneys"][NUMBERED_ATTORNEY].update(signer=""), "signer: expected a non-empty string"),
    ],
)
def test_the_attorney_format_keys_are_a_closed_set(tmp_path, edit, want):
    root = make_numbered_inputs(tmp_path / "in")
    edit_config(root, lambda d: (edit(d), d)[1])
    with pytest.raises(firm_mod.DemandConfigError) as exc:
        firm_mod.load(root)
    assert want in str(exc.value)


def test_a_config_with_no_formats_section_still_validates(tmp_path):
    firm = firm_mod.load(make_inputs(tmp_path / "in"))
    assert (
        "formats" not in firm.data and firm.variants_for(firm.attorney_for("Example Lawyer"))["pre_suit"]["mail_line"]
    )


# ---- selection, end to end -------------------------------------------------------
def _members(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_the_default_attorneys_demand_is_byte_for_byte_unchanged_by_another_attorneys_format(tmp_path, pricing):
    inputs = make_numbered_inputs(tmp_path / "in")
    v1, _s, _c, out1 = _run(tmp_path / "a", pricing, inputs, "Example Lawyer", DRAFT)

    def drop(d):
        d.pop("formats")
        d["attorneys"].pop(NUMBERED_ATTORNEY)

    edit_config(inputs, drop)
    v2, _s, _c, out2 = _run(tmp_path / "b", pricing, inputs, "Example Lawyer", DRAFT)
    assert v1.outcome == v2.outcome == "delivered", (v1.reason, v2.reason)
    name = "Demand.Alpha Example.docx"
    assert _members(out1 / name) == _members(out2 / name)


def test_the_mapped_attorneys_demand_is_rendered_and_checked_in_his_own_format(tmp_path, pricing):
    import docx

    v, seat, client, out = _run(
        tmp_path, pricing, make_numbered_inputs(tmp_path / "in"), NUMBERED_ATTORNEY, NUMBERED_DRAFT
    )
    assert v.outcome == "delivered", v.reason
    compose = next(c for c in client.calls if "COMPOSE-PROMPT" in json.dumps(c.get("system")))
    assert "## Enclosures: Index of Exhibits" in json.dumps(compose["system"])  # his skeleton, not the default
    path = out / "Demand.Alpha Example.docx"
    doc = docx.Document(str(path))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts[texts.index("Dear Pat Adjuster:") + 1] == "CASE DISTILLED"
    assert [t for t in texts if t[:3] in ("I. ", "II.", "III", "IV.", "V. ", "VI.", "VII")][:3] == [
        "I. A COLLISION ON JANUARY 15, 2026",
        "II. DAMAGES",
        "III. INTRODUCTION",
    ]
    i = texts.index("Cordially,")
    assert texts[i + 1 : i + 5] == [SIGNATURE, SIGNER, SIGNER, "Attorney for Alpha Example"]
    assert texts[i + 5] == "Enclosures: Index of Exhibits"  # no initials line
    assert all(line in texts for line in BANNER_LINES)
    assert doc.tables[0].cell(0, 0).text == "Re:" and doc.tables[0].cell(3, 1).text == "CLM-0001"
    foot = doc.sections[0].footer.paragraphs[0].text
    assert "Example  |  Claim No. CLM-0001" in foot and "{" not in foot
    assert {s["name"] for s in seat.sent} >= {"Demand.Alpha Example.docx"}


def test_an_unmapped_attorney_still_holds_when_another_attorney_has_a_format(tmp_path, pricing):
    v, seat, client, _ = _run(tmp_path, pricing, make_numbered_inputs(tmp_path / "in"), "Other Attorney", DRAFT)
    assert v.outcome == "held" and "no signature is authored" in v.reason and "Other Attorney" in v.reason
    assert client.calls == [] and seat.sent == []


def test_a_litigation_file_of_an_attorney_whose_format_has_no_litigation_variant_holds(tmp_path, pricing):
    docs = standard_docs() + [
        ("d12", "Pleading.pdf", make_pdf(["SUPERIOR COURT OF CALIFORNIA. COMPLAINT FOR DAMAGES."]), "f-corr")
    ]
    inputs = make_numbered_inputs(tmp_path / "in")
    v, _seat, client, _ = _run(tmp_path, pricing, inputs, NUMBERED_ATTORNEY, NUMBERED_DRAFT, docs)
    assert v.outcome == "held" and f"{NUMBERED_ATTORNEY}'s demands" in v.reason and client.calls == []


def test_a_redo_after_render_reviews_the_mapped_attorneys_demand_in_his_own_format(tmp_path, pricing):
    """FALSIFIER (review 2026-10-17 B3): a resume after render went straight to
    filing without ``_select_format``, so the redo's review read the firm's
    default front keys and signature against a numbered letter and held every
    redo of his demands."""
    from medchron.demand import review as rv

    seat = seat_with(standard_docs())
    seat.facts = {**seat.facts, "responsible_attorney": NUMBERED_ATTORNEY}
    client = ScriptedClient(
        draft=NUMBERED_DRAFT,
        review="| MUST-FIX | 1 wording | planted | p. 1 |\nVERDICT: 1 MUST-FIX, 0 SHOULD-FIX, 0 NOTE",
    )
    jd = make_job(tmp_path / "job")
    inputs = make_numbered_inputs(tmp_path / "in")

    def run():
        return DemandRun(
            jd,
            inputs_dir=str(inputs),
            pricing=str(pricing),
            seat_factory=lambda: seat,
            client=client,
            log=lambda _m: None,
            readback_pause=0.0,
            vendor_factory=lambda: None,
        ).run()

    v1 = run()
    assert v1.outcome == "failed" and v1.stage == "review" and seat.sent == [], v1.reason
    (jd / rv.REDO_MARKER).write_text("{}")
    client.review = "VERDICT: 0 MUST-FIX, 0 SHOULD-FIX, 0 NOTE"
    v2 = run()
    assert v2.outcome == "delivered", v2.reason
    assert {s["name"] for s in seat.sent} >= {"Demand.Alpha Example.docx"}


# ---- render and format check, directly ---------------------------------------------
def _render(tmp_path: Path, md: str = NUMBERED_DRAFT, **kw):
    tmp_path.mkdir(parents=True, exist_ok=True)
    ref = tmp_path / "ref.docx"
    make_numbered_reference(ref)
    path, notes = house_numbered.render(
        md,
        ref,
        tmp_path / "out",
        variant=kw.pop("variant", variant()),
        signature=SIGNATURE,
        signer=SIGNER,
        signer_title="Attorney for Alpha Example",
        author="Example & Example, LLP",
        client_label="Alpha Example",
    )
    return path, notes


def _check(path: Path, v: dict | None = None, **kw):
    v = v or variant()
    return format_check.check(
        path,
        signature=SIGNATURE,
        signer_title=kw.get("title", "Attorney for Alpha Example"),
        footer_markers=v["footer_markers"],
        variant=format_check.Variant.from_config(v),
        signer=SIGNER,
        initials=kw.get("initials", ""),
    )


def test_a_numbered_render_round_trips_through_its_format_check(tmp_path):
    path, notes = _render(tmp_path)
    res = _check(path)
    assert res.ok, res.fails
    assert any("open item" in n for n in res.notes) and "HELD OUT" in notes


def test_each_numbered_list_restarts_and_a_sub_bullet_stays_inside_its_list(tmp_path):
    import docx
    from docx.oxml.ns import qn

    path, _ = _render(tmp_path)
    doc = docx.Document(str(path))
    ids = []
    for p in doc.paragraphs:
        n = p._p.find(".//" + qn("w:numId"))
        if n is not None and p.text.startswith(("January 15", "Billing", "A second")):
            ids.append(n.get(qn("w:val")))
    assert len(ids) == 3 and ids[0] == ids[1] != ids[2]
    nums = doc.part.numbering_part.element.findall(qn("w:num"))
    restarted = [n for n in nums if n.get(qn("w:numId")) in ids]
    assert all(n.find(".//" + qn("w:startOverride")).get(qn("w:val")) == "1" for n in restarted)


@pytest.mark.parametrize(
    "edit, want",
    [
        (lambda md: md.replace("## Facts and Liability", "## Facts"), "'Facts and Liability' missing"),
        (lambda md: md.replace("- **Release.**", "- **Releases.**"), "lack 'Release.'"),
        (
            lambda md: (
                md.replace("## Damages", "## @@")
                .replace("## The Duty to Accept", "## Damages")
                .replace("## @@", "## The Duty to Accept")
            ),
            "out of",
        ),
        (lambda md: md.replace("{{FILL: the acceptance date}}", "{{TODO}}"), "curly-brace marker"),
    ],
)
def test_the_numbered_format_check_refuses_an_off_format_letter(tmp_path, edit, want):
    path, _ = _render(tmp_path, edit(NUMBERED_DRAFT))
    res = _check(path)
    assert not res.ok and any(want in f for f in res.fails), res.fails


def test_the_numbered_check_refuses_a_missing_banner_line_and_a_wrong_signature(tmp_path):
    path, _ = _render(tmp_path)
    v = variant(banner_lines=[*BANNER_LINES, "A line the letter does not carry"])
    assert any("banner line" in f for f in _check(path, v).fails)
    assert any("the title" in f for f in _check(path, title="Attorney at Law").fails)
    assert any("the initials" in f for f in _check(path, initials="GQC/xx").fails)


def test_each_format_refuses_the_other_formats_letter(tmp_path, pricing):
    numbered, _ = _render(tmp_path / "n")
    default_variant = firm_mod.load(make_inputs(tmp_path / "in")).data["variants"]["pre_suit"]
    house_res = format_check.check(
        numbered,
        signature=SIGNATURE,
        signer_title="Attorney for Alpha Example",
        footer_markers=["Settlement Communication"],
        variant=format_check.Variant.from_config(default_variant),
    )
    assert not house_res.ok
    v, *_rest, out = _run(tmp_path / "h", pricing, make_inputs(tmp_path / "h" / "in"), "Example Lawyer", DRAFT)
    assert v.outcome == "delivered", v.reason
    assert not _check(out / "Demand.Alpha Example.docx").ok


def test_an_unknown_table_and_a_missing_front_matter_key_are_render_errors(tmp_path):
    with pytest.raises(house.RenderError, match="no table in this format"):
        _render(tmp_path / "a", NUMBERED_DRAFT.replace("| Item | Summary |", "| Item | Summary | Extra |"))
    with pytest.raises(house.RenderError, match="footer_name"):
        _render(tmp_path / "b", NUMBERED_DRAFT.replace("footer_name: Example\n", ""))
