"""`workbook`: the caseload's ranked list, filed to the firm's library matter.

Ported from the file-review toolset's triage workbook, with the ranking now
code's (``rank.py``) and one sheet added, **Proposed changes**: every numbered
line a person may approve by replying with its number (a task closure the
second reader confirmed, a Medicals row from a bill whose text was read).

Sheets: Priorities (one row per matter, most urgent first), Proposed changes,
Needs you (the reader's time-sensitive findings and questions), Records to
order, Low limits (policies the firm's thresholds call low, decided in code),
New files (opened recently with none of the opening tasks), Scanned bills not read (bills this job did not transcribe, so no row was
proposed from them), and Not reviewed this run (matters held alone, by a
firm-safe reason, never a stack trace).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import book
from .rank import LABEL

POSTURE_WORDS = {
    "pre_demand": "Pre-demand",
    "demand_out": "Demand out",
    "offer_pending": "Offer pending",
    "settled_open_items": "Settled, items open",
    "settled_disbursed": "Settled, disbursed",
    "litigation": "Litigation",
    "declined": "Declined",
    "unknown": "Unclear from file",
}
KIND_WORDS = {"task": "Task", "medicals": "Medicals row", "event": "Calendar entry"}
FILL = {1: "F8D7DA", 2: "FFE8CC", 3: "FFF6CC", 4: "E6F2E6"}
HELD_WORDS = {
    "inventory": "Smokeball could not be read for this file",
    "fetch": "documents could not be fetched",
    "read": "the review did not finish",
    "validate": "the review did not pass its checks",
}


def _style(ws: Any, widths: list[int]) -> None:
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
        ws.cell(row=1, column=c).font = Font(bold=True)
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _priorities(wb: Any, rows: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "Priorities"
    head = ["Rank", "Priority", "Matter", "Client", "Attorney", "Where it stands", "Last activity", "Why now"]
    ws.append([*head, "Key dates", "This run"])
    for i, r in enumerate(rows, 1):
        rk, rec = r["rank"], r["review"]
        kd = "\n".join(f"{k['date']}: {k['what']}" for k in rk["key_dates"])
        ws.append(
            [
                i,
                LABEL[rk["priority"]],
                str(r["number"]),
                ", ".join(rec.get("clients") or []),
                r.get("responsible") or "",
                POSTURE_WORDS.get(str(rec.get("posture")), "Unclear from file"),
                rk.get("last_activity") or "",
                rk["reason"],
                kd,
                r["status_word"],
            ]
        )
        for c in range(1, len(head) + 3):
            ws.cell(row=i + 1, column=c).fill = PatternFill("solid", fgColor=FILL[rk["priority"]])
    _style(ws, [6, 13, 10, 24, 18, 18, 12, 48, 34, 22])


def _proposed(wb: Any, lines: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Proposed changes")
    ws.append(["Change", "Matter", "Kind", "What", "Proof in the file", "Why", "Second reading"])
    for ln in lines:
        f = ln["facts"]
        if ln["kind"] == "task":
            what = f"Mark the task complete: {f['subject']}"
            proof, second = "; ".join(f.get("proof_names") or []), f.get("audit_reason") or ""
        elif ln["kind"] == "event":
            what = f"Add to the calendar: {f['subject']}, {f['date']} ({f['how']})"
            proof, second = "; ".join(f.get("source_names") or []), "Not on the matter's calendar that day"
        else:
            what = (
                f"Add to the Medicals tab: {f['provider']}, {f['charge']}, {f['service_start']} to {f['service_end']}"
            )
            proof, second = f"{f.get('source_name')} (page {f.get('page') or 'n/a'})", "Charge found in the bill's text"
        ws.append(
            [
                ln["n"],
                f.get("number"),
                KIND_WORDS.get(ln["kind"], ln["kind"]),
                what,
                proof,
                f.get("why") or "",
                second,
            ]
        )
    if not lines:
        ws.append(["", "", "", "No changes are proposed in this run.", "", "", ""])
    _style(ws, [8, 10, 13, 52, 40, 44, 30])


def _plain(wb: Any, title: str, head: list[str], rows: list[list[Any]], widths: list[int], empty: str) -> None:
    ws = wb.create_sheet(title)
    ws.append(head)
    for r in rows:
        ws.append(r)
    if not rows:
        ws.append([empty, *[""] * (len(head) - 1)])
    _style(ws, widths)


def build(path: Path, rows: list[dict[str, Any]], lines: list[dict[str, Any]], held: list[dict[str, Any]]) -> dict:
    """Write the workbook; returns counts for the verdict."""
    wb = Workbook()
    _priorities(wb, rows)
    _proposed(wb, lines)
    needs = [
        [str(r["number"]), LABEL[r["rank"]["priority"]], str(x.get("finding") if isinstance(x, dict) else x)]
        for r in rows
        for x in [
            *(r["review"].get("time_sensitive") or []),
            *book.questions_kept(r["review"], r.get("verdicts") or {}),
        ]
    ]
    _plain(wb, "Needs you", ["Matter", "Priority", "What"], needs, [10, 13, 90], "Nothing is waiting on you.")
    orders = [
        [str(r["number"]), str(o.get("facility")), str(o.get("note") or "")]
        for r in rows
        for o in r["review"].get("records_to_order") or []
    ]
    _plain(wb, "Records to order", ["Matter", "Facility", "Note"], orders, [10, 36, 70], "No records to order.")
    lows = [
        [str(r["number"]), LABEL[r["rank"]["priority"]], p["side_label"], p["carrier"], p["limits"],
         f"{p['bills']:,.2f}", p["why"]]
        for r in rows
        for p in (r.get("flags") or {}).get("low_limits") or []
    ]  # fmt: skip
    _plain(
        wb, "Low limits", ["Matter", "Priority", "Policy", "Carrier", "Limits", "Bills on the file", "Why"],
        lows, [10, 13, 18, 28, 14, 14, 60], "No low limits on these files.",
    )  # fmt: skip
    new = [
        [str(r["number"]), nm["opened"], "; ".join(nm["missing"])]
        for r in rows
        if (nm := (r.get("flags") or {}).get("new_matter"))
    ]
    _plain(
        wb, "New files", ["Matter", "Opened", "Opening tasks not set (nothing was created)"], new,
        [10, 12, 90], "No new file is missing its opening tasks.",
    )  # fmt: skip
    scans = [[str(r["number"]), name] for r in rows for name in r.get("scanned_bills") or []]
    _plain(wb, "Scanned bills not read", ["Matter", "Bill (a scan, not read in this run)"], scans, [10, 80], "None.")
    nr = [[str(h["number"]), HELD_WORDS.get(h["why"], "the review did not finish")] for h in held]
    _plain(wb, "Not reviewed this run", ["Matter", "Why"], nr, [10, 60], "Every file in scope was reviewed.")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {
        "rows": len(rows),
        "lines": len(lines),
        "needs": len(needs),
        "orders": len(orders),
        "held": len(held),
        "low_limits": len(lows),
        "new_files": len(new),
    }


__all__ = ["POSTURE_WORDS", "build"]
