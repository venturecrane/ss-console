"""`extract`: local text extraction over the pulled `raw/` set. $0.

Per file: pypdf text with `[p.N]` page markers, docx via python-docx. A PDF
whose text layer is absent, is glyph indices, or is a constant-offset cipher
goes to the scan queue for the vision pass, and so does a text-layer PDF
holding image pages its text layer does not carry (``vision_pages``: those
pages only, decided per page); the two detectors and their
thresholds are calibrated on delivered runs (see the docstrings), and both
failure shapes produced a chronology that looked complete while omitting or
understating a record, which is worse than omitting the file.

Output (the shapes every later stage reads):
  extracted.jsonl   one record per file: {id, name, folder, ext, pages, chars,
                    text_path | scan: true}
  text/<id>.txt     the text with [p.N] markers
  scan_queue.json   files needing the vision pass
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from ..update import medlegal as ML
from .base import StageRun, read_jsonl

GLYPH_RUN = re.compile(r"(?:/\d+ ?){4,}")
GLYPH_THRESHOLD = 0.08  # junk files score 0.21-0.31; the highest clean file 0.031
ENGLISH_FLOOR = 0.02  # ciphered files 0.0000; lowest legitimate (a bill) 0.0317
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".tif", ".tiff"}
WORD_EXTS = {".docx", ".doc"}

# A short, deliberately boring list: words that appear in ANY page of English
# prose or clinical narrative, including tables and forms.
_STOP = frozenset(
    "the and of to in for with was were is are on at by a an patient this that "
    "no not his her he she had has been from as or be date name page report".split()
)


def pdf_text(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    pages = []
    for i, pg in enumerate(PdfReader(str(path)).pages, 1):
        try:
            t = pg.extract_text() or ""
        except Exception:  # noqa: BLE001 - one unreadable page is an empty page
            t = ""
        pages.append((i, t))
    return pages


def docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        for row in tb.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def not_english(txt: str) -> bool:
    """A text layer that decodes to something that is not language: a
    constant character offset ("6HH$GGHQGXP%HORZ" for "See Addendum Below")
    contains no slashes and passes every structural check, and its digits
    decode to punctuation, so a composed entry reports the study as "not
    legible" while the page states it plainly."""
    toks = re.findall(r"[A-Za-z]{2,}", txt.lower())
    if len(toks) < 50:
        return False
    return sum(1 for w in toks if w in _STOP) / len(toks) < ENGLISH_FLOOR


def glyph_junk(txt: str) -> bool:
    """A text layer of glyph INDICES ("/0/1/2/3"): a subsetted font with no
    ToUnicode map. Passes every has-a-text-layer test; unreadable to the model."""
    if not txt:
        return False
    covered = sum(len(m.group()) for m in GLYPH_RUN.finditer(txt))
    return covered / len(txt) > GLYPH_THRESHOLD


def _text_layer_absent(total: int, pages: int) -> bool:
    return total < 40 * max(1, pages) * 0.05 and total < 400


#: A page's own text layer is THIN below this many non-space characters (fax
#: and print header lines removed) or without a run of three words.
THIN_CHARS = 200
_HEADER_LINE = re.compile(r"(?im)^.*\b(?:fax|printed|print date|page \d+ of \d+|sent|received)\b.*$")
_WORD_RUN = re.compile(r"[A-Za-z]{2,}(?:\s+[A-Za-z]{2,}){2,}")


def thin(text: str) -> bool:
    body = _HEADER_LINE.sub("", text or "")
    return len(re.sub(r"\s", "", body)) < THIN_CHARS or not _WORD_RUN.search(body)


#: A date whose digits a bad OCR layer read as letters: "0l/2S/198O".
_MISREAD_DATE = re.compile(r"(?<![\w/])[\dOolIS]{1,2}/[\dOolIS]{1,2}/[\dOolIS]{2,4}(?![\w/])")
#: Letters and digits interleaved inside one word ("pat1ent", "l0mg"), the
#: OCR confusion that also turns Z94.0 into 294.0. Codes ("M54.2", "99213",
#: "S13.4XXA") and units ("10mg") have their own shapes and are not counted.
_MIXED = re.compile(r"^(?=.*[A-Za-z])(?=.*\d)[A-Za-z\d]+$")
_CODE = re.compile(r"^(?:[A-Z]\d{2}[A-Z\d]{0,5}|\d+(?:mg|mcg|ml|cc|kg|lbs?|cm|mm)|\d+(?:st|nd|rd|th))$", re.I)
_NONWORD = re.compile(r"[^aeiouy]{5,}", re.I)
GARBLE_MIXED, GARBLE_NONWORD = 0.06, 0.2


_DOB_TOKEN = re.compile(r"(?i)\b(?:D\.?O\.?B\.?|Date\s+of\s+Birth|Birth\s*Date)\s*[:#]?\s*(\d[\d/!'|\-.]{4,})")
_CLEAN_DATE = re.compile(r"^\d{1,2}[/-]\d{1,2}[/-](?:\d{2}|\d{4})[,.;]?$")
_LEAD_ICD9 = re.compile(r"(?m)^\s*\(?\d{3}\.\d{1,2}\)?\s+[A-Za-z]{3,}")
_ANY_DATE = re.compile(r"(?<![\d/])(\d{1,2})/(\d{1,2})/(\d{4})(?![\d/])")
ICD10_FROM_YEAR, FAR_YEAR = 2016, 2028


def _ocr_tells(body: str) -> bool:
    """Signs measured on a real garbled OCR layer (2026-10-14, 111 pages) that
    the length and word-shape tests all scored clean:
    * a date of birth whose separators the OCR lost or misread ("1201/1980",
      "5/15!1983"): the client's own DOB then reads as another person's;
    * a diagnosis line that opens with an ICD-9-shaped numeric code ("(294.0)
      ...") on a page dated 2016 or later, when ICD-10 had replaced ICD-9: an
      ICD-10 code whose letter the OCR read as a digit (Z94.0 as 294.0);
    * a date whose year is impossible (later than FAR_YEAR or three digits)."""
    if any(not _CLEAN_DATE.match(t) for t in _DOB_TOKEN.findall(body)):
        return True
    years = [int(y) for _, _, y in _ANY_DATE.findall(body)]
    if any(y > FAR_YEAR for y in years) or re.search(r"(?<![\d/])\d{1,2}/\d{1,2}/\d{3}(?![\d/])", body):
        return True
    return bool(_LEAD_ICD9.search(body)) and any(y >= ICD10_FROM_YEAR for y in years)


def garbled(text: str) -> bool:
    """A text layer too damaged to trust, though long enough: an OCR layer
    that confuses digits and letters. Three measurable signals, any one:
    a date carrying letters in its digit places; 6% or more of word tokens
    with letters and digits interleaved (codes and units aside); 20% or more
    of alphabetic words of four letters or more with five consonants in a
    row or no vowel at all. A clean typed page scores near zero on each."""
    body = _HEADER_LINE.sub("", text or "")
    if any(re.search(r"[OolIS]", d) for d in _MISREAD_DATE.findall(body)):
        return True
    if _ocr_tells(body):
        return True
    toks = [t.strip(".,;:()[]") for t in body.split()]
    toks = [t for t in toks if len(t) >= 3]
    if not toks:
        return False
    mixed = sum(1 for t in toks if _MIXED.match(t) and not _CODE.match(t))
    alpha = [t for t in toks if t.isalpha() and len(t) >= 4]
    odd = sum(1 for t in alpha if _NONWORD.search(t) or not re.search(r"[aeiouy]", t, re.I))
    return mixed / len(toks) >= GARBLE_MIXED or (len(alpha) >= 10 and odd / len(alpha) >= GARBLE_NONWORD)


def vision_pages(path: Path, pages: list[tuple[int, str]]) -> list[int]:
    """Pages of a text-layer PDF that are IMAGES the text layer does not
    carry, or carries garbled: a thin or garbled (``garbled``) text layer over
    an embedded image. Decided per page, because one scanned page inside a
    typed record (an 834-page file held 64) is never read when the document as
    a whole has text, and a garbled OCR layer turned Z94.0 into 294.0 and a
    client's own date of birth into "another person". A thin page with no
    image is a short typed page or a blank one: vision would read nothing more."""
    cand = [n for n, t in pages if thin(t) or garbled(t)]
    if not cand:
        return []
    import pymupdf

    with pymupdf.open(str(path)) as doc:
        return [n for n in cand if 1 <= n <= len(doc) and doc[n - 1].get_images()]


def live_rows(rows: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], int]:
    """Last write wins (retries append); a byte-duplicate row retires its id."""
    ok: dict[str, dict[str, Any]] = {}
    dupes = 0
    for r in rows:
        if r.get("ok"):
            if r.get("duplicate_of"):
                dupes += 1
                ok.pop(r["id"], None)
                continue
            ok[r["id"]] = r
    return ok, dupes


def _extract_one(d: Path, r: dict[str, Any], scans: list[dict[str, Any]]) -> dict[str, Any]:
    ext = (r.get("ext") or "").lower()
    path = Path(r["path"])
    rec: dict[str, Any] = {"id": r["id"], "name": r["name"], "folder": r["folder"], "ext": ext}
    text_path = d / "text" / f"{r['id']}.txt"
    try:
        if ext == ".pdf":
            pages = pdf_text(path)
            total = sum(len(t.strip()) for _, t in pages)
            rec["pages"] = len(pages)
            txt = "\n".join(f"[p.{n}]\n{t}" for n, t in pages)
            junk = glyph_junk(txt) or not_english(txt)
            if _text_layer_absent(total, len(pages)) or junk:
                rec["scan"] = True
                if junk:
                    rec["glyph_junk"] = True
                    # A stale text file from a previous extract satisfies the
                    # vision pass's resume check, so routing to vision means
                    # removing the text that routed it there.
                    text_path.unlink(missing_ok=True)
                scans.append(rec)
            else:
                text_path.write_text(txt, encoding="utf-8")
                rec["chars"] = len(txt)
                rec["text_path"] = str(text_path)
                imaged = vision_pages(path, pages)
                empty = [n for n, t in pages if not t.strip() and n not in imaged]
                if empty:
                    rec["empty_pages"] = empty
                if imaged:
                    # The vision pass transcribes these pages only and merges
                    # them into this text file (``stages/vision.py``).
                    rec["vision_pages"] = imaged
                    scans.append(rec)
        elif ext in WORD_EXTS:
            try:
                txt = docx_text(path)
            except Exception as exc:  # noqa: BLE001 - docx extraction failing is recorded as that document's error with empty text; the stage continues
                rec["error"] = f"docx: {exc}"[:120]
                txt = ""
            if txt:
                text_path.write_text(txt, encoding="utf-8")
                rec["chars"] = len(txt)
                rec["text_path"] = str(text_path)
            else:
                rec.setdefault("error", "no text")
        elif ext in IMAGE_EXTS:
            rec["scan"] = True
            rec["image"] = True
            scans.append(rec)
        else:
            rec["error"] = "unhandled ext"
    except Exception as exc:  # noqa: BLE001 - one broken file is one row
        rec["error"] = str(exc)[:200]
    return rec


def first_page_only(rec: dict[str, Any]) -> bool:
    """On an update, a medical-legal report named as one (``update/
    medlegal``) is reviewed and never itemized, so vision reads its first page
    only: enough to classify it, and no spend on pages nothing will compose.
    False when nothing of it is left for vision (a text-layer file whose
    first page already has text)."""
    if rec.get("image") or not ML.by_name(str(rec.get("name") or "")):
        return True
    if rec.get("scan"):
        rec["vision_pages"] = [1]
        return True
    rec["vision_pages"] = [p for p in rec.get("vision_pages") or [] if p == 1]
    return bool(rec["vision_pages"])


def run(sr: StageRun) -> int:
    d = sr.slug_dir
    (d / "text").mkdir(parents=True, exist_ok=True)
    ok, dupes = live_rows(read_jsonl(d / "raw_manifest.jsonl"))
    scans: list[dict[str, Any]] = []
    with (d / "extracted.jsonl").open("w", encoding="utf-8") as out:
        for i, r in enumerate(ok.values(), 1):
            out.write(json.dumps(_extract_one(d, r, scans)) + "\n")
            if i % 40 == 0:
                sr.log(f"{i}/{len(ok)}")
    if sr.job.update is not None:
        scans = [rec for rec in scans if first_page_only(rec)]
    (d / "scan_queue.json").write_text(json.dumps(scans, indent=1), encoding="utf-8")
    sr.log(f"DONE {len(ok)} files, {len(scans)} to scan queue, {dupes} byte-duplicates skipped")
    return 0
