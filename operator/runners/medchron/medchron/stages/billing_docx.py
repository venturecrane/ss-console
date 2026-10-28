"""`billing_docx`: the billing chart rendered as a worksheet matching the
chronology. $0.

Content rules, same as the chronology's: every figure names the document it
came from; nothing is computed that is not shown; a provider whose total is
not derivable says so instead of carrying a plausible number; the grand
total is labelled SUBTOTAL, with a warning, whenever any row is missing. A
worksheet is where a misattributed figure does damage, because it reads as
the specials number: it refuses (exit 1) a chart built without the patient
filter on a joint matter, or one still carrying a suspect (lost-decimal)
amount. The second is now an integrity falsifier only: billing_chart settles
every suspect amount by rule (accepted when the page prints it with its
decimal, otherwise left out and named), so a chart that still carries one
was built by a defect (stop audit row 18, 2026-10-24). The amounts left out
and the pages no read could transcribe are each named here by document and
page, as not in the total.
"""

from __future__ import annotations

from datetime import date
from typing import Any

from .. import billing_paid as BP
from .. import docx_style as S
from .base import StageRun, read_json
from .billing_chart import NO_BILL


def money(v: float | None) -> str:
    return f"${v:,.2f}" if v is not None else "not derivable"


def dshow(d: str | None) -> str:
    return f"{d[5:7]}/{d[8:10]}/{d[0:4]}" if d else "-"


def _cell_text(c: Any, text: str, size: float = 9.5, bold: bool = False, right: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = c.paragraphs[0]
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size, r.bold, r.font.color.rgb = Pt(size), bold, S.BLACK


def table(
    doc: Any,
    header: list[str],
    rows: list[list[str]],
    widths: list[float],
    total_row: list[str] | None = None,
    right_cols: tuple[int, ...] = (),
) -> None:
    t = doc.add_table(rows=len(rows) + 1 + (1 if total_row else 0), cols=len(header))
    S.fixed_layout(t)
    S.set_widths(t, widths)

    def fill(cells: Any, values: list[str], bold: bool, shaded: bool) -> None:
        for j, v in enumerate(values):
            c = cells[j]
            if shaded:
                S.shade(c, "f1f1f1")
            S.borders(c)
            S.cell_margins(c)
            S.vcenter(c)
            _cell_text(c, str(v), bold=bold, right=j in right_cols)

    fill(t.rows[0].cells, header, True, True)
    for i, row in enumerate(rows, 1):
        S.no_split(t.rows[i])
        fill(t.rows[i].cells, row, False, False)
    if total_row:
        S.no_split(t.rows[-1])
        fill(t.rows[-1].cells, total_row, True, True)


def para(doc: Any, text: str, size: float = 9.5, italic: bool = False, before: int = 8) -> None:
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(2)
    r = p.add_run(text)
    r.font.size, r.italic, r.font.color.rgb = Pt(size), italic, S.BLACK


def _dedup(rows: list[list[str]], key: Any) -> list[list[str]]:
    seen: set = set()
    out = []
    for r in rows:
        k = key(r)
        if k not in seen:
            seen.add(k)
            out.append(r)
    return out


def _as_printed(doc: Any, data: dict[str, Any]) -> None:
    """Row 28 (b), as revised: bills under a name matched to no treating
    provider, with their printed totals, under their OWN subtotal, after and
    outside the grand total, so the same care cannot be counted twice."""
    printed = data.get("as_printed_rows") or []
    if not printed:
        return
    doc.add_paragraph("Billed under a name not matched to a treating provider", style="Heading 1")
    para(
        doc,
        "Each bill below names a provider, as printed on the bill, that could not be matched to a provider in the "
        "medical chronology. Its amount is shown here so nothing on file is left unseen, but it is NOT in the grand "
        "total above: the same care may already be counted there under the treating provider's name, and adding it "
        "again could count it twice.",
    )
    gaps = any(r["total"] is None for r in printed)
    table(
        doc,
        ["Provider (as printed on the bill)", "Source", "Total Billed"],
        [[r["provider"], r["basis"], money(r["total"])] for r in printed],
        [2.9, 3.2, 1.1],
        total_row=["SUBTOTAL (not in the grand total)", "", money(data.get("as_printed_subtotal") or 0.0)]
        if not gaps
        else ["SUBTOTAL, INCOMPLETE (not in the grand total)", "", money(data.get("as_printed_subtotal") or 0.0)],
        right_cols=(2,),
    )


def _paid_cells(r: dict[str, Any]) -> tuple[str, str]:
    """(the figures as printed, the provider's paid total or why there is none)."""
    paid = r.get("paid") or {}
    figs = paid.get("figures") or []
    if not figs:
        return (BP.NO_PAYMENT if r.get("docs") else BP.NO_BILL_ON_FILE), "-"
    total = paid.get("total")
    return "\n".join(BP.cite(f) for f in figs), (money(total) if total is not None else str(paid.get("note")))


def _paid(doc: Any, data: dict[str, Any]) -> None:
    """Amounts paid beside the charges: every figure as printed with its
    source, summed only where billing_paid says a sum cannot count a payment
    twice. A balance or an amount due is never shown here."""
    rows = [r for r in data["rows"] if "paid" in r]
    printed = [r for r in data.get("as_printed_rows") or [] if (r.get("paid") or {}).get("figures")]
    if not rows and not printed:
        return
    doc.add_paragraph("Amounts Paid (as printed)", style="Heading 1")
    para(
        doc,
        "Each amount below is a payment exactly as a bill prints it (by an insurer, a plan or the patient), with the "
        "document and page it was read from. Balances, amounts due and adjustments are not payments and are not "
        "shown here.",
    )
    body = [[r["provider"], *_paid_cells(r)] for r in rows]
    body += [[f"{r['provider']} (name not matched; not in the total)", *_paid_cells(r)] for r in printed]
    summed = [r["paid"]["total"] for r in rows if r["paid"].get("total") is not None]
    short = [r["provider"] for r in rows if r["paid"].get("total") is None]
    table(
        doc,
        ["Provider", "Amounts Paid (as printed)", "Total Paid"],
        body,
        [2.0, 4.0, 1.2],
        total_row=[
            "TOTAL PAID" if summed and not short else "SUBTOTAL PAID (incomplete)",
            "",
            money(round(sum(summed), 2)) if summed else "not computed",
        ],
        right_cols=(2,),
    )
    unpaid = BP.unpaid_billed(rows)
    if unpaid:
        para(doc, BP.limitation_sentence(len(unpaid)) + " No payment printed: " + "; ".join(unpaid) + ".")
    unsummed = [r["provider"] for r in rows if r["paid"].get("figures") and r["paid"].get("total") is None]
    if unsummed:
        para(doc, "Printed payments shown above but not in the total paid: " + "; ".join(unsummed) + ".")
    _owed(doc, rows + printed)


_OWED_KINDS = (
    ("balances", "Balance or amount due (owed, not paid)"),
    ("combined", "Payments and adjustments combined, as printed"),
    ("adjustments", "Adjustment (not a payment)"),
)


def _owed(doc: Any, rows: list[dict[str, Any]]) -> None:
    """Balances, amounts due and adjustments, each on its own line, never
    counted as paid."""
    body = [
        [r["provider"], kind, BP.cite(f)]
        for r in rows
        for key, kind in _OWED_KINDS
        for f in (r.get("paid") or {}).get(key) or []
    ]
    if not body:
        return
    doc.add_paragraph("Balances and Adjustments (as printed)", style="Heading 1")
    para(doc, "These figures are printed on the bills but are not payments. None is in the total paid above.")
    table(doc, ["Provider", "What the bill prints", "Figure (as printed)"], body, [2.0, 1.9, 3.3])


def _disagree(doc: Any, rows: list[dict[str, Any]]) -> None:
    alts = [r for r in rows if r.get("alternates")]
    if not alts:
        return
    doc.add_paragraph("Where Documents Disagree", style="Heading 1")
    para(
        doc,
        "More than one figure appears in the file for these providers. The chart uses the one named in the Source "
        "column; the others are listed so the choice is visible.",
    )
    table(
        doc,
        ["Provider", "Figures found in the file"],
        [[r["provider"], ", ".join(f"${a:,.2f}" for a in r["alternates"])] for r in alts],
        [2.6, 4.6],
    )


def _excluded(doc: Any, data: dict[str, Any]) -> None:
    if data.get("subrogation"):
        doc.add_paragraph("Health-Plan and Subrogation Claims", style="Heading 1")
        para(
            doc,
            "These are claims by a health plan (or its recovery vendor) for benefits it paid, not charges by a treating "
            "provider. They are NOT included in the total above, and including them would double-count the underlying care.",
        )
        srows = [
            [str(pv)[:40], fl[:40], ", ".join(str(a) for a in (amts or [])[:2])] for fl, pv, amts in data["subrogation"]
        ]
        table(
            doc,
            ["Claimant", "Source Document", "Amounts Printed"],
            _dedup(srows, lambda r: (r[0], r[1]))[:14],
            [2.2, 3.0, 2.0],
        )
    if data.get("vendor_invoices"):
        doc.add_paragraph("Excluded: Vendor Invoices", style="Heading 1")
        para(
            doc, "Invoices billed to the firm by service vendors (record retrieval and similar). Not medical specials."
        )
        vrows = [[str(pv)[:36], fl[:44]] for fl, pv in data["vendor_invoices"]]
        table(doc, ["Vendor", "Source Document"], _dedup(vrows, lambda r: r[1][:40])[:14], [2.6, 4.6])
    if data.get("quarantined"):
        doc.add_paragraph("Excluded: Documents Naming Another Patient", style="Heading 1")
        para(
            doc,
            "This matter holds documents for more than one person. The following name a different patient and are "
            "excluded from every figure above.",
        )
        qrows = [[str(wh)[:26], str(pv)[:28], fl[:36]] for fl, wh, pv in data["quarantined"]]
        table(
            doc,
            ["Patient Named", "Provider", "Source Document"],
            _dedup(qrows, lambda r: (r[0], r[2][:34]))[:16],
            [1.8, 2.2, 3.2],
        )


def _not_in_total(doc: Any, data: dict[str, Any], gaps: list[dict[str, Any]]) -> None:
    """What the figures above do not include, each named by document and page:
    amounts left out (row 18), providers with no supportable total, and pages
    no read could transcribe (row 17)."""
    left_out = [r for r in data.get("dropped_rows") or [] if isinstance(r, dict)]
    if left_out:
        doc.add_paragraph("Amounts Left Out", style="Heading 1")
        para(
            doc,
            "An amount was read from each page below but could not be confirmed against the page itself (the figure "
            "is large enough that a missing decimal point would change it, and the page does not show it with its "
            "decimal point). Each was left out of every figure above rather than estimated, and an unconfirmed "
            "figure is not repeated here.",
        )
        lrows = [[str(r.get("file"))[:56], f"p. {r.get('page')}"] for r in left_out[:24]]
        table(doc, ["Source Document", "Page"], lrows, [6.2, 1.0])
    _needed(doc, [r for r in gaps if "Amounts Left Out" not in str(r.get("basis"))])
    if data.get("failed_pages"):
        doc.add_paragraph("Pages Not Read", style="Heading 1")
        para(
            doc,
            f"{len(data['failed_pages'])} billing page(s) could not be read. No amount from them is in any figure "
            f"above (not totalled), so the figures above are a floor until these pages are read.",
        )
        frows = _dedup([[str(fl)[:56], f"p. {pg}"] for fl, pg in data["failed_pages"]], lambda r: tuple(r))
        table(doc, ["Source Document", "Page"], frows[:40], [6.2, 1.0])


#: What the file holds, said only of the rows it is true of (the independent
#: review's B2, 2026-10-24: one claim-form sentence over every gap row told
#: the firm a provider with no bill at all had claim forms on file).
_NEEDED_SAYS = {
    "claim_forms": (
        "For the providers below, the file holds individual claim forms but no account ledger or statement of the "
        "total billed. Summing the claim forms is not reliable here because the same bills appear both individually "
        "and inside compiled record sets, so a sum would count them twice."
    ),
    "no_bill": "For the providers below, no bill was found in the file.",
    "other": "For the providers below, no total can be supported from the documents on file.",
}
_NEEDED_ASK = "A current ledger from each provider, or a Lienholder and Balances report for this matter, would complete the chart."


def _needed_kind(basis: str) -> str:
    if basis.strip() == NO_BILL:
        return "no_bill"
    return "claim_forms" if "claim form" in basis else "other"


def _needed(doc: Any, needed: list[dict[str, Any]]) -> None:
    """What Is Needed: one paragraph and table per basis (claim forms with no
    ledger; no bill located), each paragraph true of every row under it."""
    if not needed:
        return
    doc.add_paragraph("What Is Needed", style="Heading 1")
    for kind in _NEEDED_SAYS:
        rows = [r for r in needed if _needed_kind(str(r.get("basis") or "")) == kind]
        if rows:
            para(doc, f"{_NEEDED_SAYS[kind]} {_NEEDED_ASK}")
            table(doc, ["Provider", "What is on file"], [[r["provider"], r["basis"]] for r in rows], [2.6, 4.6])


def _refusal(sr: StageRun, data: dict[str, Any] | None) -> str | None:
    if data is None:
        return "billing_chart.json not found; billing_chart runs first"
    if sr.job.joint and not data.get("patient_filter"):
        return "chart was built without a patient filter on a joint matter; refusing to render a worksheet"
    if data.get("suspect_amounts"):
        return (
            f"chart carries {len(data['suspect_amounts'])} suspect amount(s) (likely lost decimal); refusing to render"
        )
    return None


def _sections(doc: Any, data: dict[str, Any], rows: list[dict[str, Any]], gaps: list[dict[str, Any]]) -> None:
    """Everything after the Billing by Provider table, in order."""
    _as_printed(doc, data)
    _paid(doc, data)
    _disagree(doc, rows)
    _excluded(doc, data)
    _not_in_total(doc, data, gaps)


def run(sr: StageRun) -> int:
    from docx import Document
    from docx.shared import Inches

    d = sr.slug_dir
    unit = sr.unit.unit
    sfx = f"-{unit}" if sr.job.joint else ""
    data = read_json(d / f"billing_chart{sfx}.json", None)
    refused = _refusal(sr, data)
    if refused or data is None:
        sr.log(refused or "billing_chart.json not found")
        return 1
    rows = data["rows"]
    if not rows and not any(data.get(k) for k in ("subrogation", "vendor_invoices", "as_printed_rows")):
        sr.log("no billing rows and no excluded items; no worksheet for this unit")
        return 0
    gaps = [r for r in rows if r["total"] is None]
    font = str(sr.cfg.get("format", "font") or "Calibri")
    doc = Document()
    S.build_styles(doc, font)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.7)
        s.top_margin = s.bottom_margin = Inches(0.8)
    S.page_number_footer(doc, font)
    who = sr.unit.client_name
    doc.add_paragraph(f"{who} - Medical Billing Summary", style="Title")
    para(
        doc,
        f"Prepared {date.today():%B %-d, %Y}. Providers and treatment dates are taken from the medical chronology "
        f"for this matter. Each amount names the document it was read from. No figure here was calculated by us "
        f"except the totals, each the sum of the rows shown above it.",
        italic=True,
    )
    if gaps:
        para(
            doc,
            f"INCOMPLETE: {len(gaps)} of {len(rows)} providers have no total that can be supported from the documents "
            f"on file. The subtotal below is therefore NOT this case's medical specials figure. The sections at the "
            f"end say what is missing for each.",
        )
    doc.add_paragraph("Billing by Provider", style="Heading 1")
    body: list[list[str]] = []
    grand, orphan = 0.0, False
    for r in rows:
        grand += r["total"] or 0.0
        name = r["provider"]
        if r["first"] is None and r["last"] is None and r.get("lien"):
            name, orphan = name + " *", True
        body.append([name, dshow(r["first"]), dshow(r["last"]), r["basis"], money(r["total"])])
    table(
        doc,
        ["Provider", "First DOS", "Last DOS", "Source", "Total Billed"],
        body,
        [2.3, 0.8, 0.8, 2.35, 1.05],
        total_row=["SUBTOTAL" if gaps else "GRAND TOTAL", "", "", "", money(round(grand, 2))],
        right_cols=(4,),
    )
    if orphan:
        para(
            doc,
            "* This line appears on the firm's lien report with no provider name printed, and no matching provider or "
            "treatment appears in the medical chronology. It is included so this summary reconciles to the firm's "
            "own report, but it cannot be attributed to a provider from the documents on file.",
            size=9,
        )
    _sections(doc, data, rows, gaps)
    stamp = sr.date_stamp or date.today().strftime("%m-%d-%y")
    glob = str(sr.cfg.get("delivery", "worksheet_glob") or "* - Medical Billing Worksheet *.docx")
    out = d / "out" / unit / glob.replace("*", who, 1).replace("*", stamp, 1)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    sr.log(
        f"wrote {out.name}: {len(rows)} provider(s), {'subtotal' if gaps else 'total'} {money(round(grand, 2))}"
        + (f", {len(gaps)} gap(s)" if gaps else "")
    )
    return 0
