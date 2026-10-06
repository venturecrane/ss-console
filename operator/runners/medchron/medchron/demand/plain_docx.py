"""The gap audit, the coverage report and the attorney notes as Word files.

These are internal documents for the firm's attorney, not letters, so they are
rendered plainly: headings, paragraphs, bullets, pipe tables, ``**bold**`` and
``*italic*``. The firm's house format belongs to the demand alone
(``house.py``). Firm markers are kept verbatim as text, never resolved.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


def _runs(p: Any, text: str) -> None:
    for seg in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            p.add_run(seg[2:-2]).bold = True
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            p.add_run(seg[1:-1]).italic = True
        else:
            p.add_run(seg)


def _table(doc: Any, block: list[str]) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in block]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not rows:
        return
    width = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=width)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(width):
            cell = t.cell(i, j)
            cell.text = ""
            _runs(cell.paragraphs[0], r[j] if j < len(r) else "")
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def render(md: str, out: Path, title: str, author: str) -> Path:
    import docx

    doc = docx.Document()
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        m = re.match(r"^(#{1,4}) (.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 3))
        elif line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, block)
            continue
        elif re.match(r"^\s*[-*] ", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", line))
        elif line.strip():
            _runs(doc.add_paragraph(), line.strip())
        i += 1
    cp = doc.core_properties
    cp.title, cp.author, cp.last_modified_by, cp.comments = title, author, author, ""
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
