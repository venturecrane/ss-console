"""Refuse a demand file that is not in the firm's house format, BEFORE filing.

Ported from the engagements laptop tool ``format_check.py`` (2026-09-28), which
``file_to_matter.sh`` ran before writing anything. It checks the shape of the
FILE that would be filed, not the prose:

  1. the name is ``Demand.<Client>.docx``
  2. the RE block carries Our Client / Your Insured / Claim Number / Date of Loss
  3. the first table is the two-row banner: a time-limited demand title, then
     "... Expires at 5:00 P.M. Pacific Time on <Weekday>, <Month D, YYYY>"
  4. a salutation
  5. Summary of Injuries < Liability < Damages < Demand, centered, and an ICD-10 table
  6. the footer carries the firm's settlement-communication markers
  7. "Cordially," then the firm's signature line, a signer, and the title
  8. no drafting machinery left in: skeleton ``[INSERT`` markers, ``PROTO:``
     prototypes, the drafting end-lists, or a curly marker that is not one of
     the open-item markers a requester may ask for
  9. the document properties name only this file

Open items in the firm's bracket dialect (``[TO BE SUPPLIED: ...]``,
``[ATTORNEY: ...]``, ``[CONFIRM ... BEFORE TRANSMITTAL]``) and the three curly
markers a requester can name in her brief (``{{NOT IN RECORD: ...}}``,
``{{FILL: ...}}``, ``{{ATTORNEY: ...}}``, first asked for on 2026-10-06) are
allowed and COUNTED, not refused: they are the draft telling the attorney what
the file does not establish. Any other ``{{`` is a leftover template slot.

What changed from the laptop tool: the firm's signature, title and footer
markers come from the firm config, and the curly open-item markers are allowed.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

HEADINGS = ("summary of injuries", "liability", "damages", "demand")
OPEN_ITEM = re.compile(r"\[(TO BE SUPPLIED|ATTORNEY|CONFIRM)[^\]]*\]|\{\{(NOT IN RECORD|FILL|ATTORNEY)\b[^}]*\}\}")
CURLY_ALLOWED = re.compile(r"\{\{(NOT IN RECORD|FILL|ATTORNEY)\b[^}]*\}\}")


@dataclass
class FormatResult:
    fails: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.fails


def _qn(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


def _lines(doc: Any) -> list[tuple[str, str | None, Any]]:
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag == _qn("w:p"):
            out.append(("p", "".join(t.text or "" for t in el.iter(_qn("w:t"))).strip(), el))
        elif el.tag == _qn("w:tbl"):
            out.append(("t", None, el))
    return out


def _centered(el: Any) -> bool:
    jc = el.find(_qn("w:pPr") + "/" + _qn("w:jc"))
    return jc is not None and jc.get(_qn("w:val")) == "center"


def _banner(tables: list[Any], fails: list[str]) -> None:
    if not tables:
        fails.append("no tables: the firm's banner, ICD-10 and Demand tables are missing")
        return
    b = tables[0]
    if len(b.rows) != 2 or len(b.columns) != 1:
        fails.append("first table is not the two-row banner")
    else:
        cells = [r.cells[0].text.strip() for r in b.rows]
        if not re.search(r"time[- ]limited.*demand", cells[0], re.I):
            fails.append(f"banner line 1 is not a time-limited demand title: {cells[0][:60]!r}")
        # The expiry date is the line's point, but on a draft for review it may
        # still be an open item the attorney sets (the 2026-10-06 smoke: a
        # requester's brief asked for {{FILL}} on firm-supplied values).
        date_or_open = r"((\w+day, )?\w+ \d{1,2}, \d{4}|" + OPEN_ITEM.pattern + ")"
        if not re.search(r"expires at 5:00 p\.m\. pacific time on " + date_or_open, cells[1], re.I):
            fails.append(f"banner line 2 is not the expiry line: {cells[1][:70]!r}")
    if not any(t.cell(0, 0).text.strip().lower().startswith("icd") for t in tables):
        fails.append("no ICD-10 Code table under Summary of Injuries")


def _headings(lines: list[tuple[str, str | None, Any]], fails: list[str], headings: tuple[str, ...]) -> None:
    heads = [str(t).lower() for k, t, el in lines if k == "p" and t and _centered(el)]
    pos = []
    for h in headings:
        idx = [i for i, x in enumerate(heads) if x == h]
        if not idx:
            fails.append(f"section heading {h.title()!r} missing (or not a centered heading)")
        else:
            pos.append(idx[0])
    if pos != sorted(pos):
        fails.append("section headings are out of the profile's order (" + ", ".join(h.title() for h in headings) + ")")


def _machinery(alltext: str, fails: list[str], notes: list[str]) -> None:
    for pat, why in (
        (r"\[INSERT", "an unfilled skeleton [INSERT ...] marker"),
        (r"PROTO:", "a template prototype"),
        (r"HELD OUT PENDING", "the drafting HELD OUT list"),
        (r"^TO BE SUPPLIED$", "the drafting TO BE SUPPLIED end-list"),
        (r"ATTORNEY decisions reserved", "the drafting ATTORNEY end-list"),
    ):
        if re.search(pat, alltext, re.M):
            fails.append(f"letter contains {why}")
    if "{{" in CURLY_ALLOWED.sub("", alltext):
        fails.append("letter contains a curly-brace marker that is not an open-item marker")
    opens = [m.group(1) or m.group(2) for m in OPEN_ITEM.finditer(alltext)]
    if opens:
        notes.append(f"{len(opens)} open item(s) for the attorney: " + ", ".join(sorted(set(opens))))


@dataclass(frozen=True)
class Variant:
    """One demand variant's shape, authored in demand-firm.yaml (``variants``).

    Validated 2026-10-06 against a real pre-suit demand of the firm's, read on
    the seat and never copied off it: its RE block names Your Insured, Claim
    Number and Date of Loss (no "Our Client" line), a bold "CERTIFIED MAIL" line
    and an "Attn:" line open the address block, and "Procedure" heads the body.
    A litigation demand has no mail line and no carrier block and reads "vs."
    in its RE line."""

    headings: tuple[str, ...] = HEADINGS
    re_required: tuple[str, ...] = ("Claim Number:", "Date of Loss:")
    re_one_of: tuple[str, ...] = ("Your Insured:", "Our Client:")
    mail_line: str = ""
    banner: bool = True
    forbidden: tuple[str, ...] = ()
    re_vs: bool = False

    @classmethod
    def from_config(cls, v: dict[str, Any]) -> "Variant":
        return cls(
            headings=tuple(v["headings"]),
            re_required=tuple(v["re_required"]),
            re_one_of=tuple(v["re_one_of"]),
            mail_line=str(v["mail_line"]),
            banner=bool(v["banner"]),
            forbidden=tuple(v["forbidden"]),
            re_vs=bool(v["re_vs"]),
        )


def _address_block(texts: list[str], variant: Variant, fails: list[str]) -> None:
    head = texts[: next((i for i, t in enumerate(texts) if t.startswith("Dear ")), 20)]
    if variant.mail_line and not any(t.strip().upper() == variant.mail_line.upper() for t in head):
        fails.append(f"no {variant.mail_line!r} line above the salutation")
    if variant.mail_line and not any(t.startswith("Attn:") for t in head):
        fails.append("no 'Attn:' line in the address block")
    for label in variant.re_required:
        if not any(label in t for t in head):
            fails.append(f"RE block lacks {label!r}")
    if variant.re_one_of and not any(label in t for t in head for label in variant.re_one_of):
        fails.append(f"RE block names none of {list(variant.re_one_of)}")
    if variant.re_vs and not any(t.startswith("RE:") and re.search(r"\bvs?\.\s", t) for t in head):
        fails.append("the RE line does not read '<plaintiff> vs. <defendant>'")
    for bad in variant.forbidden:
        if any(bad.upper() in t.upper() for t in head):
            fails.append(f"{bad!r} belongs to another demand variant")


def _signature(
    texts: list[str], fails: list[str], signature: str, signer: str, signer_title: str, initials: str
) -> None:
    if "Cordially," not in texts:
        fails.append("no 'Cordially,' sign-off")
        return
    i = texts.index("Cordially,")
    blk = [t.strip() for t in texts[i : i + 6]]
    for want, what in (
        (signature, "the firm"),
        (signer, "the signer"),
        (signer_title, "the title"),
        (initials, "the initials"),
    ):
        if want and want not in blk:
            fails.append(f"signature block lacks {what} ({want!r})")


def _own_render(name: str, doc: Any, path: Path, fails: list[str]) -> None:
    """Hygiene of a file THIS job rendered: its exact name and no properties
    carried over from the reference. The firm's own files are not held to it:
    the firm saves "Demand_<Client> (7).docx" and its titles go stale."""
    if not re.fullmatch(r"Demand\.[A-Za-z][^/\\]*\.docx", name):
        fails.append(f"file name {name!r} is not Demand.<Client>.docx")
    stem = name[:-5] if name.endswith(".docx") else name
    title = (doc.core_properties.title or "").strip()
    if title and title != stem:
        fails.append(f"document title property {title!r} is not this file ({stem!r})")
    with zipfile.ZipFile(str(path)) as z:
        if "docProps/custom.xml" in z.namelist():
            fails.append("custom document properties present (they carry a source file's matter ids)")


def check(
    path: Path,
    *,
    signature: str,
    signer_title: str,
    footer_markers: list[str],
    headings: tuple[str, ...] | list[str] | None = None,
    variant: Variant | None = None,
    signer: str = "",
    initials: str = "",
    own_render: bool = True,
) -> FormatResult:
    """The file's shape against one variant; ``own_render`` adds the hygiene
    checks for a file this job rendered (off when validating a firm demand)."""
    import docx

    v = variant or Variant()
    if headings is not None:
        v = Variant(**{**v.__dict__, "headings": tuple(headings)})
    res = FormatResult()
    name = Path(path).name
    if not own_render and not re.match(r"(?i)demand[._ ]", name):
        res.fails.append(f"file name {name!r} does not start with Demand")
    try:
        doc = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - not a docx at all is a refusal, named
        res.fails.append(f"cannot open as .docx: {exc}")
        return res
    lines = _lines(doc)
    texts = [str(t) for k, t, _ in lines if k == "p" and t]
    tables = doc.tables
    _address_block(texts, v, res.fails)
    if v.banner:
        _banner(tables, res.fails)
    elif not any(t.cell(0, 0).text.strip().lower().startswith("icd") for t in tables):
        res.fails.append("no ICD-10 Code table under Summary of Injuries")
    if not any(t.startswith("Dear ") for t in texts):
        res.fails.append("no salutation")
    _headings(lines, res.fails, tuple(h.lower() for h in v.headings))
    foot = " ".join(p.text for s in doc.sections for f in (s.footer, s.first_page_footer) for p in f.paragraphs)
    if not all(m in foot for m in footer_markers):
        res.fails.append("footer lacks the firm's settlement-communication line")
    _signature(texts, res.fails, signature, signer, signer_title, initials)
    table_text = " ".join(c.text for t in tables for r in t.rows for c in r.cells)
    _machinery("\n".join(texts) + "\n" + table_text, res.fails, res.notes)
    if own_render:
        _own_render(name, doc, Path(path), res.fails)
    return res
