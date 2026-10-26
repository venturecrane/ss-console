"""The firm's answer to "what are these records?": one row per file an update
was asked to read, filed beside the chronology. $0: built only from what the
code decided.

WHY (2026-10-25, the Captain). The firm's attorney asked, of 222 files a
delivered chronology had not cited, "What are the records and how many pages
are they? Are they duplicates?", then "Please read the uncovered records ...
and flag duplicates; incorporate nonduplicates into the med chron". The update
recorded a reason and a page count for every file (``outcome``) but the reply
carried counts by kind only, so the two questions he asked first went
unanswered. This document answers them, file by file.

Each row is the file's name as the matter shows it, its kind, its pages, and
what was done with it: added to the chronology as its exhibit, or the coded
reason it added nothing (a copy of a delivered exhibit, every page already in
the exhibits, billing pages only, no record text, read in full with no new
visit, left out by the firm's rule, password-protected, ...), in the outcome
record's own wording. Nothing in it is model prose.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from ..stages.base import read_json
from . import outcome as O

TEMPLATE = "{CLIENT} - Records Review {MM-DD-YY}.docx"
KIND_BY_EXT = {
    ".pdf": "PDF",
    ".msg": "email",
    ".eml": "email",
    ".rpmsg": "email",
    ".docx": "Word document",
    ".doc": "Word document",
    ".xlsx": "spreadsheet",
    ".xls": "spreadsheet",
}
ADDED = "added"


def name(client: str, stamp: str) -> str:
    return TEMPLATE.replace("{CLIENT}", client).replace("{MM-DD-YY}", stamp)


def _kind(file_name: str) -> str:
    ext = ("." + file_name.rsplit(".", 1)[-1].lower()) if "." in file_name else ""
    if ext in O.IMAGE_EXTS:
        return "image"
    return KIND_BY_EXT.get(ext, ext.lstrip(".").upper() or "file")


#: The delivered run's own labels on a finding (``prior_reading``): internal
#: words, never shown to the firm (proof 5 of the first real update, 2026-10-25:
#: 196 of 223 rows read "composer: ...").
_INTERNAL = re.compile(r"^(?:composer|compose skipped|set aside)\s*:\s*", re.I)
_BILL_DATES = re.compile(r"^evidenced billing dates only \((\d+)\); the billing chart carries them$", re.I)


def plain(f: dict[str, Any]) -> str:
    """The finding as the firm reads it: the outcome record's text, with a
    reused finding's internal label replaced by what it means."""
    text = str(f.get("text") or "")
    if f.get("reason") != O.PRIOR:
        return text
    m = _BILL_DATES.match(text)
    if m:
        return (
            f"read for the delivered chronology: billing dates only ({m.group(1)}); they are in the billing worksheet"
        )
    text = _INTERNAL.sub("", text)
    low = text.lower()
    if low.startswith("billing-only source (vendor_invoice)"):
        text = "an invoice to the firm (records or services), not a medical bill"
    elif low.startswith("billing-only source"):
        text = "billing pages only; their charges are in the billing worksheet"
    elif low.startswith("read in full;") and "nothing in it was cited" in low:
        text = "read in full; nothing in it to cite"
    return f"read for the delivered chronology: {text}"


def rows(files: list[dict[str, Any]], added: dict[str, str], requested: set[str] | None = None) -> list[dict[str, Any]]:
    """One row per outcome file; `added`: {file id: the exhibit's title} for a
    file the update drew entries from; `requested`: the ids the request named
    (a file found inside one of them, an email's attachment, is not)."""
    out = []
    for f in files:
        rid, title = str(f.get("id")), added.get(str(f.get("id")))
        finding = f"added to the chronology as {title}" if title else plain(f)
        out.append(
            {
                "id": rid,
                "file": str(f.get("name") or rid),
                "kind": _kind(str(f.get("name") or "")),
                "pages": int(f.get("pages") or 0),
                "finding": finding,
                "reason": ADDED if title else str(f.get("reason") or ""),
                "requested": requested is None or rid in requested,
            }
        )
    out.sort(key=lambda r: (r["reason"] != ADDED, r["reason"], r["file"].lower()))
    return out


def render(path: Path, client: str, matter: str, table: list[dict[str, Any]]) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph(f"{client}: records reviewed for the chronology update", style="Title")
    added = sum(1 for r in table if r["reason"] == ADDED)
    named = sum(1 for r in table if r.get("requested", True))
    inside = f", plus {len(table) - named} found inside them" if len(table) > named else ""
    doc.add_paragraph(
        f"Matter {matter}. The {named} file(s) named in the request{inside}, each with its kind, its pages and what "
        f"was done with it. {added} added to the chronology; {len(table) - added} added nothing, for the reason shown."
    )
    t = doc.add_table(rows=1, cols=4)
    for cell, head in zip(t.rows[0].cells, ("File", "Kind", "Pages", "Finding")):
        cell.text = head
    for r in table:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = (
            r["file"],
            r["kind"],
            str(r["pages"]) if r["pages"] else "",
            r["finding"],
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def build(sr: Any, files: list[dict[str, Any]], added: dict[str, str]) -> Path:
    """The review document for this unit, written under ``out/<unit>/``."""
    stamp = sr.date_stamp or __import__("datetime").date.today().strftime("%m-%d-%y")
    path = sr.slug_dir / "out" / sr.unit.unit / name(sr.unit.client_name, stamp)
    inc = read_json(sr.slug_dir / "include.json", {}) or {}
    requested = {str(x) for x in inc.get("include_file_ids") or []} or None
    return render(path, sr.unit.client_name, str(sr.job.matter_number), rows(files, added, requested))


__all__ = ["ADDED", "TEMPLATE", "build", "name", "plain", "render", "rows"]
