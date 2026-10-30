"""`review`: the built workbook re-read from its own bytes before it is filed.

WHY (2026-10-13). "Delivered" was the job's own claim, and three wrong outputs
across the lanes were marked delivered with no alert. The litigation list
already carries the strongest content checks of any lane: every value is read
twice more by independent verify and audit models (passes 2 and 3), and the
gates refuse a date or status without a source on that matter, a future date,
a Smokeball save date standing in for a court date, and an unresolved
settlement hit. So the review here is CODE ONLY, proportionate to that: it
adds what nothing upstream checks, which is that the FILE we are about to put
on the firm's matter says what the data it was built from says.

* every in-scope matter has exactly one row in Case notes (a matter dropped
  or doubled by the build is a list the attorneys would trust and should not);
* the All defendants tab, re-read from the .xlsx bytes, agrees row for row
  and cell for cell with the rows built from ``final.json`` (a writer bug, a
  sort that lost a row, a truncated save);
* a people-scoped list's "Litigated files tagged on" tab agrees row for row
  with the rows built from the data, and no other list carries the tab;
* the one-sheet-per-case copy (``percase.py``, 2026-10-27) has one tab per
  case and every All defendants row exactly once with the same values, so the
  copy filed beside the list never disagrees with it;
* no cell carries the system's own words or a sentinel (``proof.checks``; the
  gates' ``leak_scan`` covers internal process wording and ids, this covers
  the shared system vocabulary every lane now refuses).

Each finding is MUST-FIX: code has settled it. A hold is recorded ``failed``
with a ``review_held:`` reason (litigation ``held`` is final; failed is
resumable, and SMD is paged), and an SMD release (``release-review.json`` in
the job dir) files past it with the record saying so.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from ..proof import checks, review as proof_review
from . import book, percase, printout as printout_mod

CASE_NOTES = "Case notes"
ALL_DEFENDANTS = "All defendants"


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def read_sheets(xlsx: bytes, *, header: bool = False) -> dict[str, list[list[str]]]:
    """Every sheet's data rows (row 1 dropped unless ``header``), every cell
    as text, from the file's BYTES (never the in-memory workbook the build held)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    out: dict[str, list[list[str]]] = {}
    try:
        for ws in wb.worksheets:
            rows = [[_norm(c) for c in row] for row in ws.iter_rows(values_only=True)]
            out[ws.title] = rows if header else rows[1:]
    finally:
        wb.close()
    return out


def _trim(row: list[str]) -> list[str]:
    while row and row[-1] == "":
        row = row[:-1]
    return row


def _matters(notes: list[list[str]] | None, in_scope: list[dict[str, Any]]) -> list[proof_review.Finding]:
    """Every in-scope matter listed exactly once in Case notes, and nothing else."""
    if notes is None:
        return [proof_review.code(proof_review.MUST, "matters", "the Case notes tab is missing")]
    out: list[proof_review.Finding] = []
    seen: dict[str, int] = {}
    for r in notes:
        if len(r) > 1 and r[1]:
            seen[r[1]] = seen.get(r[1], 0) + 1
    want = {_norm(m.get("number")) for m in in_scope}
    for n in sorted(want):
        if seen.get(n, 0) != 1:
            out.append(
                proof_review.code(
                    proof_review.MUST, "matters", f"matter {n} has {seen.get(n, 0)} rows in Case notes, not one"
                )
            )
    for n in sorted(set(seen) - want):
        out.append(proof_review.code(proof_review.MUST, "matters", f"matter {n} is in the list and not in scope"))
    return out


def _rows(got: list[list[str]] | None, finals: list[dict[str, Any]]) -> list[proof_review.Finding]:
    """The All defendants tab equals the rows built from the data: the count,
    then the first row that differs (one finding is enough to hold)."""
    if got is None:
        return [proof_review.code(proof_review.MUST, "rows", "the All defendants tab is missing")]
    built = [_trim([_norm(book.clean(v)) for v in r["cells"]]) for r in book.defendant_rows(finals)]
    got = [_trim(r) for r in got]
    out: list[proof_review.Finding] = []
    if len(got) != len(built):
        out.append(
            proof_review.code(
                proof_review.MUST,
                "rows",
                f"All defendants has {len(got)} rows in the file and {len(built)} in the data it was built from",
            )
        )
    for i, (g, b) in enumerate(zip(got, built), start=2):
        if g != b:
            col = next((j for j in range(max(len(g), len(b))) if (g[j : j + 1] or [""]) != (b[j : j + 1] or [""])), 0)
            out.append(
                proof_review.code(
                    proof_review.MUST,
                    "rows",
                    f"All defendants row {i}, column {col + 1}, differs from the data it was built from",
                    f"{ALL_DEFENDANTS}!row {i}",
                )
            )
            break
    return out


def _vocabulary(sheets: dict[str, list[list[str]]], first: int = 2) -> list[proof_review.Finding]:
    """No system wording or sentinel in any cell of any tab; ``first`` is the
    sheet row the first given row sits on."""
    out: list[proof_review.Finding] = []
    for title, rows in sheets.items():
        for i, r in enumerate(rows, start=first):
            for cell in r:
                hits = checks.system_vocabulary(cell) + checks.sentinels(cell)
                if hits:
                    out.append(
                        proof_review.code(
                            proof_review.MUST,
                            "vocabulary",
                            f"internal wording in a cell: {', '.join(hits)[:120]}",
                            f"{title}!row {i}",
                        )
                    )
    return out


def _tagged(got: list[list[str]] | None, want: tuple[list[str], list[list[Any]]] | None) -> list[proof_review.Finding]:
    """A people-scoped list's "Litigated files tagged on" tab equals the rows
    built from the data (and no other list carries the tab)."""
    if want is None:
        if got is not None:
            return [
                proof_review.code(proof_review.MUST, "tagged", f"the {book.TAGGED} tab is in a list with no people")
            ]
        return []
    if got is None:
        return [proof_review.code(proof_review.MUST, "tagged", f"the {book.TAGGED} tab is missing")]
    built = [_trim([_norm(book.clean(v)) for v in r]) for r in want[1]]
    got = [_trim(r) for r in got]
    if got == built:
        return []
    return [
        proof_review.code(
            proof_review.MUST,
            "tagged",
            f"{book.TAGGED} has {len(got)} rows in the file and {len(built)} in the data it was built from"
            if len(got) != len(built)
            else f"a {book.TAGGED} row differs from the data it was built from",
        )
    ]


def findings(
    xlsx: bytes,
    finals: list[dict[str, Any]],
    in_scope: list[dict[str, Any]],
    tagged: tuple[list[str], list[list[Any]]] | None = None,
) -> list[proof_review.Finding]:
    """The code review of one built workbook. Empty means it may be filed."""
    try:
        sheets = read_sheets(xlsx)
    except Exception as exc:  # noqa: BLE001 - an unreadable workbook is the finding
        return [proof_review.code(proof_review.MUST, "workbook", f"the built workbook does not open: {exc}"[:300])]
    return (
        _matters(sheets.get(CASE_NOTES), in_scope)
        + _rows(sheets.get(ALL_DEFENDANTS), finals)
        + _tagged(sheets.get(book.TAGGED), tagged)
        + _vocabulary(sheets)
    )


def _per_case(xlsx: bytes, finals: list[dict[str, Any]]) -> list[proof_review.Finding]:
    """The one-sheet-per-case copy: agrees with the data, and no system
    wording in any cell (its title row included)."""
    try:
        problems = percase.check(xlsx, finals)
        sheets = {f"{t} (per case)": rows for t, rows in read_sheets(xlsx, header=True).items()}
    except Exception as exc:  # noqa: BLE001 - an unreadable workbook is the finding
        return [proof_review.code(proof_review.MUST, "per_case", f"the per-case workbook does not open: {exc}"[:300])]
    return [proof_review.code(proof_review.MUST, "per_case", p) for p in problems] + _vocabulary(sheets, first=1)


def _printout(pdf: Path, finals: list[dict[str, Any]]) -> list[proof_review.Finding]:
    """The printable PDF: a page for every case, every court date it shows
    printed, and no system wording on any page."""
    if not pdf.is_file():
        return [proof_review.code(proof_review.MUST, "printout", "the print-out is not on disk")]
    try:
        problems = printout_mod.check(pdf, finals) + [f"{c}: {w}" for c, w, _ in printout_mod.scan(pdf)]
    except Exception as exc:  # noqa: BLE001 - an unreadable PDF is the finding
        return [proof_review.code(proof_review.MUST, "printout", f"the print-out does not open: {exc}"[:300])]
    return [proof_review.code(proof_review.MUST, "printout", p) for p in problems]


def run(
    data: Path,
    job_dir: Path,
    xlsx: Path,
    finals: list[dict[str, Any]],
    in_scope: list[dict[str, Any]],
    *,
    tagged: tuple[list[str], list[list[Any]]] | None = None,
    per_case: Path | None = None,
    printout: Path | None = None,
) -> str:
    """Review, record and decide. Returns ``code_only`` or ``released``;
    raises ``LitigationFailed`` (``review_held: ...``) on a hold. ``per_case``,
    the one-sheet-per-case copy, is re-read too: one tab per case, every
    defendant row once with the main list's values (``percase.check``)."""
    from .outcome import LitigationFailed

    if not xlsx.is_file() or (per_case is not None and not per_case.is_file()):
        raise LitigationFailed("stage_unfinished: review: the built workbook is not on disk; the book stage reruns")
    fs = findings(xlsx.read_bytes(), finals, in_scope, tagged)
    if per_case is not None:
        fs += _per_case(per_case.read_bytes(), finals)
    if printout is not None:
        fs += _printout(printout, finals)
    proof_review.save(data, fs, model="", lane="litigation")
    # The code review is re-run on every attempt, so a redo needs nothing
    # forgotten; consuming it here would forget the record just saved and
    # read as "missing", never "held" (2026-10-13 integration review).
    status, rec = proof_review.gate(data, job_dir, consume_redo=False)
    if status == "held":
        raise LitigationFailed(proof_review.held_reason(rec or {}))
    return "released" if status == "released" else "code_only"


__all__ = ["findings", "read_sheets", "run"]
