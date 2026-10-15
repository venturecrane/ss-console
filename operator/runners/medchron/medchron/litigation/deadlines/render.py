"""The workbooks' layout, the hand kit's ``build.py`` sheet for sheet: a navy
header row, frozen and filtered, dates as m/d/yyyy, HIDDEN entries last in
gray. Data only arrives here; nothing is decided here."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable

from .book import HDR, STATUTE_HDR, TABS, TASK_HDR, WIDTH, Rules, stamp


def _styles() -> tuple[Any, Any, Any, Any]:
    from openpyxl.styles import Alignment, Font, PatternFill

    return PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF"), Font(color="808080"), Alignment


def sheet(
    wb: Any,
    title: str,
    hdr: list[str],
    data: list[list[Any]],
    widths: list[int],
    sort_key: Callable[[list[Any]], Any],
    hidden_col: int | None = None,
    date_cols: tuple[int, ...] = (0,),
) -> Any:
    from openpyxl.utils import get_column_letter

    head_fill, head_font, gray, alignment = _styles()
    ws = wb.create_sheet(title)
    ws.append(hdr)
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
    if hidden_col is not None:
        data = sorted(data, key=lambda r: (r[hidden_col] == "Y", sort_key(r)))
    else:
        data = sorted(data, key=sort_key)
    for r in data:
        ws.append(r)
        if hidden_col is not None and r[hidden_col] == "Y":
            for c in ws[ws.max_row]:
                c.font = gray
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    for col in date_cols:
        for cells in ws.iter_cols(min_col=col + 1, max_col=col + 1, min_row=2):
            for cell in cells:
                cell.number_format = "m/d/yyyy"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _new() -> Any:
    import openpyxl

    wb = openpyxl.Workbook()
    first = wb.active
    if first is not None:
        wb.remove(first)
    return wb


def deadlines_workbook(path: Path, rows: dict[str, list[list[Any]]], tasks: list[list[Any]],
                       cases: list[list[Any]], r: Rules) -> dict[str, int]:  # fmt: skip
    """Writes the file; returns each sheet's data row count (the review's expectation)."""
    wb = _new()
    for t in TABS:
        hdr = HDR + (["List"] if t == "Next 14 days" else [])
        sheet(wb, t, hdr, rows[t], WIDTH + ([12] if t == "Next 14 days" else []),
              lambda x: (x[0], x[1] == "", x[1], str(x[4])), hidden_col=7)  # fmt: skip
    sheet(wb, "OPEN TASKS", TASK_HDR, tasks, [11, 8, 34, 11, 80, 30, 9, 20, 20], lambda x: (x[0], str(x[3])))
    sheet(
        wb,
        "CASES",
        ["Case", "Matter #", "Status", "Responsible attorney", "Paralegal", "Case number", "Filed date",
         "Next date", "Next date: what", f"Dates through {r.horizon.month}/{r.horizon.day}/{r.horizon.strftime('%y')}"],
        cases,
        [34, 11, 9, 20, 20, 18, 11, 11, 80, 10],
        lambda x: (x[7] is None, x[7] or r.today, str(x[1])),
        date_cols=(7,),
    )  # fmt: skip
    wb.save(path)
    return {**{t: len(rows[t]) for t in TABS}, "OPEN TASKS": len(tasks), "CASES": len(cases)}


def statute_workbook(path: Path, up: list[list[Any]], past: list[list[Any]], r: Rules) -> dict[str, int]:
    wb = _new()
    widths = [12, 9, 34, 11, 9, 20, 20, 12]
    sheet(wb, f"Statute by {stamp(r.statute_horizon)}", STATUTE_HDR, up, widths, lambda x: (x[0], str(x[3])))
    sheet(
        wb,
        "Statute passed, Open",
        STATUTE_HDR[:1] + ["Days since"] + STATUTE_HDR[2:],
        [[x[0], -x[1]] + x[2:] for x in past],
        widths,
        lambda x: (x[0], str(x[3])),
    )
    wb.save(path)
    return {f"Statute by {stamp(r.statute_horizon)}": len(up), "Statute passed, Open": len(past)}
