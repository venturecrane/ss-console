"""The mechanical gates: nothing reaches the firm until every rule passes.

Every rule of the hand run's ``qa.py``, plus the ones its failures added:

* settlement scan: an open matter whose read text carries settlement words
  must have every hit resolved by pass 3 (a settlement stated plainly in an
  email was missed once);
* no future dates (only the next court date may be ahead), and every date
  in ``YYYY-MM-DD``;
* every date and every status has a source, and every source is a file on
  that matter;
* the Smokeball SAVE date is not a court date: a value whose date equals its
  source file's saved date, and which its source's text does not carry in any
  common form, is refused (a third of the second read's overturns were this);
* an unreadable or missing file is never absence: each one is named in the
  matter's flags;
* the BUILT workbook carries no internal vocabulary (``leak_scan``).

Each rule returns the items it refuses; ``passed`` is every list empty.
"""

from __future__ import annotations

import datetime as dt
import json
import re
from pathlib import Path
from typing import Any, Callable

from . import vocab

DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
EM_DASH = "—"
LEAK_PATTERNS = [
    ("hex id", re.compile(r"\b(?=[0-9a-f]{0,31}[a-f])(?=[0-9a-f]{0,31}[0-9])[0-9a-f]{8,32}\b")),
    ("uuid", re.compile(r"\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}\b")),
    (
        "internal process wording",
        re.compile(
            r"(?i)\b(second read|first read|third read|verifier|double-check|overturn\w*|re-?read|pass [123])\b"
        ),
    ),
    ("storage error", re.compile(r"(?i)nosuchkey")),
    ("internal file name", re.compile(r"(?i)\b[\w-]+\.(json|jsonl|py|sh|yaml)\b|\bread[123]\b|\bupload_manifest\b")),
    ("doc reference", re.compile(r"(?i)\[doc ?\d+\]|\bdoc_ref\b|\bfile_id\b")),
    ("encoding jargon", re.compile(r"(?i)\bcp1252\b")),
    ("status-note jargon", re.compile(r"(?i)normalized status note|source_first_read")),
    ("em dash", re.compile(EM_DASH)),
]
TextOf = Callable[[str, str], "str | None"]


def _dated(m: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    """Every value on a matter that carries a date, by a readable path."""
    out = [("complaint_filed", m.get("complaint_filed") or {}), ("next_court_date", m.get("next_court_date") or {})]
    for d in m.get("defendants") or []:
        out += [
            (f"{d.get('name')}: served", d.get("served") or {}),
            (f"{d.get('name')}: answered", d.get("answered") or {}),
        ]
    for i, r in enumerate(m.get("discovery_propounded") or []):
        out.append((f"discovery propounded #{i + 1}", r))
    for i, r in enumerate(m.get("discovery_served_on_client") or []):
        out += [(f"discovery served on client #{i + 1}", r), (f"responses #{i + 1}", r.get("responses_served") or {})]
    return [(p, v) for p, v in out if isinstance(v, dict)]


def _status(m: dict[str, Any]) -> str:
    cs = m.get("case_status")
    return str((cs or {}).get("value") or "") if isinstance(cs, dict) else str(cs or "")


def _all_uim(m: dict[str, Any]) -> bool:
    ds = m.get("defendants") or []
    return bool(ds) and all(d.get("status") == vocab.UIM for d in ds)


def date_forms(iso: str) -> list[str]:
    d = dt.date.fromisoformat(iso)
    mon, mo3 = d.strftime("%B"), d.strftime("%b")
    return [
        iso,
        f"{d.month:02d}/{d.day:02d}/{d.year}",
        f"{d.month}/{d.day}/{d.year}",
        f"{d.month:02d}/{d.day:02d}/{d.year % 100:02d}",
        f"{d.month}/{d.day}/{d.year % 100:02d}",
        f"{d.month:02d}-{d.day:02d}-{d.year}",
        f"{mon} {d.day}, {d.year}",
        f"{mo3} {d.day}, {d.year}",
        f"{mo3}. {d.day}, {d.year}",
        f"{d.day} {mon} {d.year}",
    ]


def text_carries(text: str, iso: str) -> bool:
    """The date appears in the text in a written form. Court forms often
    print a filled date as spaced characters ("1 / 2 5 / 2 0 2 7", Yolo's CMC
    notice, 2026-10-06), so the comparison is also made with all whitespace
    removed; every digit must still match."""
    flat = " ".join(text.split()).lower()
    if any(f.lower() in flat for f in date_forms(iso)):
        return True
    squeezed = "".join(text.split()).lower()
    return any("".join(f.split()).lower() in squeezed for f in date_forms(iso))


def _case_rules(m: dict[str, Any], n: str, out: dict[str, list[Any]]) -> None:
    cs = _status(m)
    if cs not in vocab.CASE_STATUSES:
        out["case status not normalized"].append((n, cs))
    if cs != vocab.UNCLEAR and not (m.get("case_status") or {}).get("source"):
        out["case status without source"].append(n)
    if not str(m.get("responsible") or "").strip():
        out["no responsible person"].append(n)
    if cs != vocab.NOT_FILED and not (m.get("complaint_filed") or {}).get("date") and not _all_uim(m):
        out["filed case with no complaint date"].append((n, cs))
    if not m.get("defendants") and cs != vocab.NOT_FILED:
        out["no defendants"].append(n)
    if EM_DASH in json.dumps(m, ensure_ascii=False):
        out["em dash in text"].append(n)


def _date_rules(m: dict[str, Any], n: str, t: str, out: dict[str, list[Any]]) -> None:
    cs = _status(m)
    for path, v in _dated(m):
        d = v.get("date")
        if d is None:
            continue
        if not (isinstance(d, str) and DATE.match(d)):
            out["date not YYYY-MM-DD"].append((n, path, d))
            continue
        if not v.get("source"):
            out["date without source"].append((n, path))
        if path == "next_court_date":
            if d < t and cs in (vocab.ACTIVE, vocab.STAYED):
                out["next court date is in the past (active case)"].append((n, d))
        elif d > t:
            out["future date"].append((n, path, d))


def _defendant_rules(m: dict[str, Any], n: str, out: dict[str, list[Any]]) -> None:
    for d in m.get("defendants") or []:
        st = d.get("status")
        if st not in vocab.DEFENDANT_STATUSES:
            out["defendant status not normalized"].append((n, d.get("name"), st))
        ans = (d.get("answered") or {}).get("date")
        if st in (vocab.ANSWERED, vocab.ANSWERED_ORIGINAL) and not ans:
            out["answered but no answer date"].append((n, d.get("name")))
        if ans and st in vocab.NO_ANSWER:
            out["answer date but status says no answer"].append((n, d.get("name"), st))
        if (d.get("served") or {}).get("date") and st in (vocab.NOT_SERVED, vocab.OUT_FOR_SERVICE):
            out["served date but status says not served"].append((n, d.get("name")))


def rule_matters(ms: list[dict[str, Any]], today: dt.date) -> dict[str, list[Any]]:
    """The qa.py rules, matter by matter."""
    t = today.isoformat()
    out: dict[str, list[Any]] = {
        k: []
        for k in (
            "duplicate matter",
            "case status not normalized",
            "no responsible person",
            "filed case with no complaint date",
            "future date",
            "date not YYYY-MM-DD",
            "next court date is in the past (active case)",
            "no defendants",
            "defendant status not normalized",
            "date without source",
            "case status without source",
            "answered but no answer date",
            "answer date but status says no answer",
            "served date but status says not served",
            "em dash in text",
        )
    }
    seen: dict[str, int] = {}
    for m in ms:
        n = str(m.get("number") or m.get("matter_id"))
        seen[n] = seen.get(n, 0) + 1
        _case_rules(m, n, out)
        _date_rules(m, n, t, out)
        _defendant_rules(m, n, out)
    out["duplicate matter"] = [n for n, k in seen.items() if k > 1]
    return out


def rule_sources(
    ms: list[dict[str, Any]], files: dict[str, set[str]], saved: dict[str, dict[str, set[str]]], text_of: TextOf
) -> dict[str, list[Any]]:
    """Every source is a file on the matter; a save date is not a court date."""
    bad, save = [], []
    for m in ms:
        mid, n = str(m.get("matter_id")), str(m.get("number"))
        vals = _dated(m) + [("case_status", m.get("case_status") or {})]
        for path, v in vals:
            src = v.get("source") if isinstance(v, dict) else None
            if not src:
                continue
            fid = str(src.get("file_id") or "")
            if fid not in files.get(mid, set()):
                bad.append((n, path))
                continue
            d = v.get("date")
            if not (isinstance(d, str) and DATE.match(d)) or d not in saved.get(mid, {}).get(fid, set()):
                continue
            text = text_of(mid, fid)
            if text is None or not text_carries(text, d):
                save.append((n, path, d))
    return {"source is not a file on the matter": bad, "date is the Smokeball save date, not in the document": save}


def rule_integrity(ms: list[dict[str, Any]]) -> dict[str, list[Any]]:
    """Every unreadable or missing file is named in the matter's flags."""
    out = []
    for m in ms:
        flags = " ".join(str(f) for f in m.get("matter_flags") or [])
        for rec in m.get("integrity") or []:
            if str(rec.get("name") or "") not in flags:
                out.append((m.get("number"), rec.get("name"), rec.get("problem")))
    return {"unreadable or missing file not named on the matter": out}


def settlement_hits(texts: dict[str, str], firm: Any) -> list[str]:
    """File ids whose text carries a settlement term."""
    rx = firm.settlement_rx
    return sorted(fid for fid, t in texts.items() if any(p.search(t) for p in rx))


def rule_settlement(ms: list[dict[str, Any]], hits: dict[str, list[str]]) -> dict[str, list[Any]]:
    out = []
    for m in ms:
        if _status(m) not in vocab.OPEN_STATUSES:
            continue
        reviewed = {str(r.get("file_id")) for r in m.get("settlement_reviewed") or []}
        out += [(m.get("number"), fid) for fid in hits.get(str(m.get("matter_id")), []) if fid not in reviewed]
    return {"settlement-scan hit not resolved": out}


def run(
    ms: list[dict[str, Any]],
    *,
    today: dt.date,
    files: dict[str, set[str]],
    saved: dict[str, dict[str, set[str]]],
    text_of: TextOf,
    hits: dict[str, list[str]],
) -> dict[str, Any]:
    issues: dict[str, list[Any]] = {}
    issues.update(rule_matters(ms, today))
    issues.update(rule_sources(ms, files, saved, text_of))
    issues.update(rule_integrity(ms))
    issues.update(rule_settlement(ms, hits))
    issues = {k: v for k, v in issues.items() if v}
    return {"passed": not issues, "issues": issues}


def leak_scan(xlsx: Path) -> list[tuple[str, str, str]]:
    """(sheet!cell, what, the matched text) for every leak in the built file."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if not isinstance(c.value, str):
                        continue
                    for what, rx in LEAK_PATTERNS:
                        m = rx.search(c.value)
                        if m:
                            out.append((f"{ws.title}!{c.coordinate}", what, m.group(0)))
    finally:
        wb.close()
    return out
