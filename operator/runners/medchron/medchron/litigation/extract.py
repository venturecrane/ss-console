"""`extract`: text from every fetched file, by kind, page by page.

The 2026-10-07 hand run lost facts at this layer four different ways, so each
kind has its own reader and nothing falls through silently:

* PDF: the text layer per page; a page with no text layer is a scan and goes
  to vision (``ocr``, the doorway's read model). Checkbox facts are never
  taken from this text; pass 3 looks at the page image (``page_image``).
* ``.msg``: demand's reader, which retries a body in cp1252 (the encoding that
  silently dropped a settlement email on 2026-10-07); headers, body, and the
  text of PDF/DOCX attachments.
* ``.eml``: the stdlib MIME parser; plain or HTML body; attachments as above.
* ``.htm``/``.html``: BeautifulSoup text. ``.rtf``: RTFDE for encapsulated
  RTF, else striprtf when installed, else a small control-word stripper.
  ``.doc``: antiword when it is on PATH. ``.docx``: python-docx. ``.txt``.
* Anything else, a reader that raises, or text that comes out EMPTY is an
  extraction FAILURE record. A failure is never absence: the gates refuse a
  negative finding on a matter until the unreadable file is named on it.
* A STOP is not a failure: a spend cap (``LimitHold``) or a vision call that
  ran out of network retries (``deadline.is_transient``) says nothing about
  the file. It propagates, so the job holds or fails as a whole and a resume
  reads the file again; a failure record is permanent for the job.
"""

from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
import threading
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any, Callable

from .. import unopened
from ..deadline import is_transient
from ..limits import LimitHold
from ..stages.base import append_jsonl, read_jsonl

PAGE_MARK = "===== page {n} ====="
SCAN_FLOOR = 25  # characters: a page with less carries no text layer
HEAD, TAIL = 3, 2
DEFERRED = "[scanned page {n}: not transcribed; call view_page with this doc and page {n} to see it]"
MIN_TEXT = 1
OCR_PROMPT = (
    "Transcribe every word on this page exactly as printed, top to bottom, including stamps, handwriting you can "
    "read, and form labels. For a checkbox write [X] when it is marked and [ ] when it is empty. Output the text only."
)
Ocr = Callable[[bytes], str]
#: Bump by hand with ANY change to how a reader turns bytes into text (a new
#: decoding, a new file kind, a fixed email body, a changed page rule): it is
#: part of the text cache's method version (``textcache.method_version``), so
#: text made the old way is read again instead of served forever.
EXTRACT_VERSION = 1
#: The file and page a vision call is for, per thread (read by the
#: rehearsal's counting OCR stub, which records them).
CURRENT = threading.local()


#: Problems older runs recorded for a STOP (before stops propagated): a
#: resume drops these records and reads the file again.
_STOP_PROBLEMS = frozenset(
    f"reader_error: {n}"
    for n in (
        "LimitHold",
        "CallDeadlineError",
        "APIConnectionError",
        "APITimeoutError",
        "RateLimitError",
        "InternalServerError",
        "OverloadedError",
    )
)


def _stop(exc: BaseException) -> bool:
    return isinstance(exc, LimitHold) or is_transient(exc)


class ExtractFailure(RuntimeError):
    def __init__(self, problem: str) -> None:
        self.problem = problem
        super().__init__(problem)


#: The failure code of a PDF locked with a password (``unopened``). Without it
#: a locked attachment raised out of its email and the whole email, body and
#: every other attachment, became one unreadable file.
PASSWORD = "pdf_password"


def _pages(texts: list[str]) -> str:
    return "\n".join(f"{PAGE_MARK.format(n=i)}\n{t.strip()}" for i, t in enumerate(texts, 1))


# MEMORY (2026-10-13: A&P litigation job 01M4V2G6 was OOM-killed in its 1 GiB
# cgroup at 2525 of 4808 documents). MuPDF keeps a PROCESS-WIDE store of
# decoded images and fonts that outlives ``doc.close()`` (it fills toward its
# 256 MB default and stays there), and pymupdf appends every MuPDF warning to a
# module-level list nothing trims. Both grow with the number of documents read,
# not with the largest one. So every MuPDF use here runs under ``_MUPDF`` and
# ends with ``_release()``, a PDF is opened from its path (its bytes are never
# all in Python at once), and no PDF stays open across a vision call: a
# document's footprint is gone before the next one starts.
_MUPDF = threading.Lock()
RENDER_LONG_SIDE = 1600


def _release() -> None:
    """Drop MuPDF's cached decoded objects and pymupdf's warning list. What an
    open document still references MuPDF keeps; only the leftovers go."""
    import pymupdf

    pymupdf.TOOLS.store_shrink(100)
    pymupdf.TOOLS.reset_mupdf_warnings()


def _open_pdf(src: bytes | Path) -> Any:
    import pymupdf

    if isinstance(src, Path):
        return pymupdf.open(str(src), filetype="pdf")
    return pymupdf.open(stream=src, filetype="pdf")


def _png(doc: Any, page: int, long_side: int) -> bytes:
    import pymupdf

    pg = doc[page - 1]
    zoom = long_side / (max(pg.rect.width, pg.rect.height) or 1.0)
    return pg.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).tobytes("png")


def render_png(pdf: bytes | Path, page: int, long_side: int = RENDER_LONG_SIDE) -> bytes:
    with _MUPDF:
        try:
            doc = _open_pdf(pdf)
        except Exception:
            _release()
            raise
        try:
            return _png(doc, page, long_side)
        finally:
            doc.close()
            _release()


def pdf_text(data: bytes | Path, ocr: Ocr | None) -> tuple[str, int]:
    """(paged text, scanned pages sent to vision). The page images vision
    needs (at most ``HEAD + TAIL``) are rendered and the document closed
    BEFORE the first vision call."""
    with _MUPDF:
        try:
            doc = _open_pdf(data)
        except Exception as exc:  # noqa: BLE001 - a PDF that does not open is one failure record
            _release()
            raise ExtractFailure(f"pdf_unreadable: {type(exc).__name__}") from None
        try:
            if unopened.locked(doc):
                raise ExtractFailure(f"{PASSWORD}: {unopened.REASON}")
            texts = [pg.get_text() or "" for pg in doc]
            picks = ocr_pages(texts)
            if picks and ocr is None:
                raise ExtractFailure("scan_without_vision: a page has no text layer and no vision reader was given")
            images = [(i, _png(doc, i + 1, RENDER_LONG_SIDE)) for i in picks]
        finally:
            doc.close()
            _release()
    scanned = 0
    while images and ocr is not None:
        i, png = images.pop(0)  # each image is let go once it is transcribed
        CURRENT.page = i + 1
        texts[i] = ocr(png)
        scanned += 1
    for i, t in enumerate(texts):
        if len(t.strip()) < SCAN_FLOOR and i not in picks:
            texts[i] = DEFERRED.format(n=i + 1)
    return _pages(texts), scanned


def ocr_pages(texts: list[str]) -> list[int]:
    """The scanned pages transcribed at extract: every one on a short
    document; on a long one only those among its first ``HEAD`` and last
    ``TAIL`` pages (the caption, the filing stamp, the signature and the
    proof of service). The rest are read on demand: ``view_page``."""
    n = len(texts)
    keep = set(range(n)) if n <= HEAD + TAIL else set(range(HEAD)) | set(range(n - TAIL, n))
    return [i for i in sorted(keep) if len(texts[i].strip()) < SCAN_FLOOR and not texts[i].startswith("[scanned page")]


def docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _attachment_text(name: str, blob: bytes, ocr: Ocr | None) -> str:
    low = name.lower()
    try:
        if low.endswith(".pdf"):
            return pdf_text(blob, ocr)[0]
        if low.endswith(".docx"):
            return docx_text(blob)
    except ExtractFailure as exc:
        return f"[attachment could not be read: {exc.problem.split(':')[0]}]"
    return ""


def _email(head: dict[str, str], body: str, atts: list[tuple[str, bytes]], ocr: Ocr | None) -> str:
    lines = [f"{k}: {v}" for k, v in head.items()]
    out = "\n".join(lines) + "\n\n" + body.strip()
    for name, blob in atts:
        out += f"\n\n[ATTACHMENT] {name}"
        text = _attachment_text(name, blob, ocr)
        if text.strip():
            out += "\n" + text
    return out


def msg_text(path: Path, ocr: Ocr | None) -> str:
    from ..demand.pull import _open_msg, head_parties

    m = _open_msg(path)
    if m is None:
        raise ExtractFailure("msg_unreadable: the Outlook container did not open in either encoding")
    sender, to = head_parties(m)
    head = {"From": sender, "To": to, "Date": m["date"], "Subject": m["subject"]}
    return _email(head, m["body"], m["attachments"], ocr)


def html_text(raw: bytes | str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())


def eml_text(data: bytes, ocr: Ocr | None) -> str:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        content = body_part.get_content()
        body = html_text(content) if body_part.get_content_type() == "text/html" else str(content)
    atts = []
    for part in msg.iter_attachments():
        payload = part.get_payload(decode=True)
        if isinstance(payload, bytes):
            atts.append((part.get_filename() or "attachment", payload))
    head = {k: str(msg.get(k) or "") for k in ("From", "To", "Date", "Subject")}
    return _email(head, body, atts, ocr)


_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "listtable", "listoverridetable"}
_RTF_TOKEN = re.compile(
    r"\\([a-z]{1,32})(-?\d{1,10})? ?|\\'([0-9a-f]{2})|\\([^a-z])|([{}])|[\r\n]+|([^\\{}\r\n]+)", re.I
)


def strip_rtf(text: str) -> str:
    """Plain text from RTF control words: groups, \\par, \\'hh, \\uN, and the
    destinations that carry no body text (font and colour tables, pictures)."""
    out: list[str] = []
    stack: list[bool] = []
    skip = False
    for m in _RTF_TOKEN.finditer(text):
        word, arg, hexc, sym, brace, plain = m.groups()
        if brace == "{":
            stack.append(skip)
        elif brace == "}":
            skip = stack.pop() if stack else False
        elif skip:
            continue
        elif word:
            w = word.lower()
            if w in _RTF_SKIP:
                skip = True
            elif w in ("par", "line", "row"):
                out.append("\n")
            elif w == "tab" or w == "cell":
                out.append("\t")
            elif w == "u" and arg:
                out.append(chr(int(arg) % 65536))
        elif hexc:
            out.append(bytes([int(hexc, 16)]).decode("cp1252", errors="replace"))
        elif sym == "*":
            skip = True
        elif sym in ("\\", "{", "}"):
            out.append(sym)
        elif plain:
            out.append(plain)
    return "".join(out)


def rtf_text(data: bytes) -> str:
    raw = data.decode("latin-1")
    try:
        from RTFDE.deencapsulate import DeEncapsulator

        d = DeEncapsulator(data)
        d.deencapsulate()
        if d.content_type == "html":
            return html_text(d.html)
        return str(d.text or "")
    except Exception:  # noqa: BLE001 - not encapsulated RTF (the common case): the plain readers below
        pass
    try:
        from striprtf.striprtf import rtf_to_text  # type: ignore[import-not-found]

        return str(rtf_to_text(raw))
    except ImportError:
        return strip_rtf(raw)


def doc_text(path: Path) -> str:
    exe = shutil.which("antiword")
    if not exe:
        raise ExtractFailure("doc_unreadable: antiword is not installed on this seat")
    r = subprocess.run([exe, str(path)], capture_output=True, text=True, timeout=120, check=False)  # noqa: S603 - fixed binary, a path we wrote
    if r.returncode != 0:
        raise ExtractFailure("doc_unreadable: antiword could not read the file")
    return r.stdout


def text_of(path: Path, ext: str, ocr: Ocr | None) -> tuple[str, int]:
    """(text, scanned pages); raises ExtractFailure."""
    ext = ext.lower()
    if path.stat().st_size == 0:
        raise ExtractFailure("empty_file: the file has no bytes")
    if ext == ".pdf":
        return pdf_text(path, ocr)  # from disk: MuPDF reads what it needs, Python never holds the file
    if ext == ".msg":
        return msg_text(path, ocr), 0
    data = path.read_bytes()
    if ext == ".eml":
        return eml_text(data, ocr), 0
    if ext in (".htm", ".html"):
        return html_text(data), 0
    if ext == ".rtf":
        return rtf_text(data), 0
    if ext == ".doc":
        return doc_text(path), 0
    if ext == ".docx":
        return docx_text(data), 0
    if ext == ".txt":
        return data.decode("utf-8", errors="replace"), 0
    if ext in (".jpg", ".jpeg", ".png", ".tif", ".tiff") and ocr is not None:
        return _pages([ocr(data)]), 1
    raise ExtractFailure(f"unsupported: no reader for {ext or 'a file without an extension'}")


def extract_file(path: Path, ext: str, out: Path, ocr: Ocr | None) -> dict[str, Any]:
    """One record; on success the text is at ``out``."""
    CURRENT.file, CURRENT.page = out.stem, None
    try:
        text, scanned = text_of(path, ext, ocr)
    except ExtractFailure as exc:
        return {"ok": False, "problem": exc.problem}
    except Exception as exc:  # a reader crash is a failure record, never an absence
        if _stop(exc):
            raise  # a spend cap or a network give-up stops the job; a resume re-reads the file
        return {"ok": False, "problem": f"reader_error: {type(exc).__name__}"}
    body = re.sub(r"===== page \d+ =====", "", text)
    if len(body.strip()) < MIN_TEXT:
        return {"ok": False, "problem": "empty_text: the reader produced no text"}
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    return {"ok": True, "chars": len(text), "scanned_pages": scanned, "text_path": str(out)}


def extract_matter(
    mdir: Path,
    rows: list[dict[str, Any]],
    ocr: Ocr | None,
    progress: Any = None,
    concurrency: int = 1,
    cache: Any = None,
) -> list[dict[str, Any]]:
    """Every ok pulled row of one matter; resumable through extracted.jsonl.
    Documents are read ``concurrency`` at a time (each one's vision calls in
    turn), records written by this thread only, each the moment its document
    finishes. At most ``concurrency`` documents are in flight: a slow scan at
    the head never lets finished work pile up behind it (``pool.map`` submits
    every document at once and holds each result until the head is done).
    ``progress`` (a ``progress.Progress``) is stepped once per document.
    A row filled from the text cache (``cached``, no ``path``) is recorded
    from its ``txt/`` without extracting. ``cache`` (a
    ``textcache.TextCache``) is read before a document is extracted and
    written after it is (``extract_cached``)."""
    log = mdir / "extracted.jsonl"
    done = _recorded(log)
    todo = [r for r in rows if str(r["id"]) not in done and _extractable(r)]
    for r in rows:
        if str(r["id"]) in done and progress is not None:
            progress.step()

    def one(r: dict[str, Any]) -> dict[str, Any]:
        fid = str(r["id"])
        return {"file_id": fid, "name": r.get("name"), **extract_cached(mdir, r, ocr, cache)}

    def record(rec: dict[str, Any]) -> None:
        append_jsonl(log, rec)
        done[rec["file_id"]] = rec
        if progress is not None:
            if not rec.get("ok"):
                progress.add("failures")
            progress.step()

    _in_flight(todo, one, record, max(1, concurrency))
    return [done[str(r["id"])] for r in rows if str(r["id"]) in done]


def _extractable(r: dict[str, Any]) -> bool:
    """A pulled row extract records: downloaded bytes, or a cache fill (a
    content duplicate of either is read through its original)."""
    return bool(r.get("ok") and (r.get("path") or r.get("cached")) and not r.get("duplicate_of"))


class CacheTextMissing(RuntimeError):
    """A row filled from the text cache whose ``txt/`` copy is gone. A STOP,
    never a failure record: the run's extract stage heals such rows first
    (``fetch.heal_cached``: filled again or downloaded)."""


def extract_cached(mdir: Path, r: dict[str, Any], ocr: Ocr | None, cache: Any = None) -> dict[str, Any]:
    """One pulled row's record, reusing paid text where the cache has it.

    * a row filled from the cache (``cached``, no raw bytes): its text is
      already in ``txt/``; it is recorded, never extracted (a missing copy
      raises ``CacheTextMissing``, never a failure record);
    * a downloaded row whose content (sha256) the cache holds: the text is
      copied, no vision call is made, and this matter joins the entry;
    * anything else is extracted, and a success is written to the cache (a
      failure never is: the next job tries the file again)."""
    fid = str(r["id"])
    out = mdir / "txt" / f"{fid}.txt"
    sha = r.get("sha256")
    hit = None if r.get("cached") or cache is None else cache.get_text(sha)
    if hit is not None:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(hit, encoding="utf-8")
    if r.get("cached") or hit is not None:
        if not out.is_file():
            raise CacheTextMissing(f"cached text for {fid} is not on disk")
        if cache is not None:
            cache.add_matter(sha, mdir.name)
        scanned = int((cache.meta(sha) if cache is not None else {}).get("scanned_pages") or 0)
        chars = len(out.read_text(encoding="utf-8", errors="replace"))
        return {"ok": True, "chars": chars, "scanned_pages": scanned, "text_path": str(out), "cached": True}
    rec = extract_file(Path(r["path"]), r.get("ext") or "", out, ocr)
    if rec.get("ok") and cache is not None:
        text = out.read_text(encoding="utf-8")
        cache.put_text(sha, text, scanned_pages=int(rec["scanned_pages"]), matter_id=mdir.name)
    return rec


def _recorded(log: Path) -> dict[str, dict[str, Any]]:
    """{file id: its last record}; a recorded STOP is no record at all."""
    done: dict[str, dict[str, Any]] = {}
    for r in read_jsonl(log):
        if r.get("problem") in _STOP_PROBLEMS:
            done.pop(r["file_id"], None)
        else:
            done[r["file_id"]] = r
    return done


def _in_flight(
    todo: list[dict[str, Any]],
    one: Callable[[dict[str, Any]], dict[str, Any]],
    record: Callable[[dict[str, Any]], None],
    width: int,
) -> None:
    """Run `one` over `todo`, at most `width` at a time, `record`ing each
    result on this thread. A STOP from one document is held until its
    siblings in flight have finished and been recorded (what was already paid
    for is kept), nothing new is started, and then it is raised."""
    from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait

    stop: list[BaseException] = []

    def settle(finished: set[Future[dict[str, Any]]]) -> None:
        for f in finished:
            exc = f.exception()
            if exc is None:
                record(f.result())
            else:
                stop.append(exc)

    with ThreadPoolExecutor(max_workers=width) as pool:
        live: set[Future[dict[str, Any]]] = set()
        for r in todo:
            live.add(pool.submit(one, r))
            if len(live) < width:
                continue
            finished, live = wait(live, return_when=FIRST_COMPLETED)
            settle(finished)
            if stop:
                break
        settle(set(wait(live).done))
    if stop:
        raise stop[0]


def failures(mdir: Path) -> list[dict[str, Any]]:
    return [r for r in read_jsonl(mdir / "extracted.jsonl") if not r.get("ok")]


def write_summary(data: Path, summary: dict[str, Any]) -> None:
    (data / "extract.json").write_text(json.dumps(summary, indent=1), encoding="utf-8")
