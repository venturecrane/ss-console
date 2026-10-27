"""`book`: the workbook the firm reads, styled the way the 2026-10-07 hand-run
workbooks were (bold grey header, wrapped top-aligned cells, frozen header).

Tabs: Summary, (a people-scoped list only) Litigated files tagged on, Needs
action, Settled - dismissal owed, All defendants, Discovery, Case notes, Not
filed, What changed.

The Responsible column is the Smokeball staff roster's name for the matter's
responsible staff id (inventory), never a name a model wrote. Every string
passes ``clean()`` on the way in, and the gates scan the BUILT file again
(``gates.leak_scan``) before it is filed: cleaning is the attempt, the scan is
the proof.
"""

from __future__ import annotations

import collections
import datetime as dt
import re
from pathlib import Path
from typing import Any

from . import vocab
from .people import role_labels

HEADER_FILL = "DDDDDD"
DEF_HEADER = [
    "Responsible (Smokeball)",
    "Next court date",
    "Matter #",
    "Case",
    "Court / case no.",
    "Complaint filed",
    "Case status",
    "Defendant",
    "Defendant status",
    "Sent out for service",
    "Sent out for service (source)",
    "Served",
    "Served (source)",
    "Answered",
    "Answered (source)",
    "Notes",
]
DEF_WIDTHS = [18, 24, 10, 32, 28, 14, 24, 28, 24, 16, 40, 22, 40, 14, 40, 60]
#: The sent-out cell for a defendant no pass has read for it, and for one a
#: pass read and found no request to a process server (with the reason in its
#: source cell). An unread value is never shown as a finding.
NOT_READ_YET = "Not read yet"
NOT_IN_FILE = "Not in file"
NOT_IN_FILE_WHY = "No request to a process server was found in the files read"
_CLEAN = [
    (r"\s*\((?:second read|double-check|first read said)[^()]*\)", ""),
    (r"(?i)\bcorrected on (?:double-check|second read|review):\s*", ""),  # 2026-10-23: a note's lead-in
    (r"(?i)\b(second|first|third) read\b", "review"),
    (r"(?i)\bverifier\b", "reviewer"),
    (r"(?i)\boverturn(ed|s)?\b", "corrected"),
    (r"(?i)\bre-?read\b", "reviewed"),
    (r"\s*\(download returns XML NoSuchKey\)", ""),
    (r"(?i)\bNoSuchKey\b", "missing in Smokeball"),
    (r"(?i)\bcp1252\b", "legacy encoding"),
    (r"(?i)normalized status note(?: \(audit\))?:\s*", ""),
    (r"(?i)\[doc ?\d+\]", ""),
    (r"(?i)\bdoc(ument)? (number|no\.?|#) ?\d+\b", "a document in the file"),
    (r"\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}\b", ""),
    (r"\b(?=[0-9a-f]{0,31}[a-f])(?=[0-9a-f]{0,31}[0-9])[0-9a-f]{8,32}\b", ""),
    (r"(?i)\b[\w-]+\.(json|jsonl|py|sh|yaml)\b", ""),
    ("—", ", "),
    (r"\(\s*[,;]?\s*\)", ""),
    (r"  +", " "),
]


def clean(x: Any) -> Any:
    if not isinstance(x, str):
        return x
    for pat, rep in _CLEAN:
        x = re.sub(pat, rep, x)
    return x.strip(" |;,")


def _date(v: Any) -> str:
    if not isinstance(v, dict):
        return ""
    d = v.get("date") or ""
    extra = v.get("method") or v.get("event")
    return f"{d} ({extra})".strip() if extra and d else str(d or "")


def _src(v: Any) -> str:
    s = v.get("source") if isinstance(v, dict) else None
    if not isinstance(s, dict):
        return ""
    return ", ".join(x for x in (s.get("name"), s.get("doc_date")) if x)


def _status(m: dict[str, Any]) -> str:
    return str((m.get("case_status") or {}).get("value") or "")


def _flags(x: dict[str, Any], key: str) -> str:
    return " | ".join(str(f) for f in x.get(key) or [])


def _sort_date(s: str) -> str:
    m = re.search(r"\d{4}-\d{2}-\d{2}", s or "")
    return m.group(0) if m else "9999"


def sent_cells(d: dict[str, Any]) -> tuple[str, str]:
    """(Sent out for service, its source) for one defendant: the date and its
    document; "Not in file" and why, when a pass read and found none; "Not
    read yet" when no pass has read it (the key is absent)."""
    if "sent_for_service" not in d:
        return NOT_READ_YET, ""
    v = d.get("sent_for_service")
    if isinstance(v, dict) and v.get("date"):
        return _date(v), _src(v)
    return NOT_IN_FILE, NOT_IN_FILE_WHY


def defendant_rows(ms: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for m in ms:
        for d in m.get("defendants") or []:
            rows.append(
                {
                    "m": m,
                    "d": d,
                    "cells": [
                        m.get("responsible") or "",
                        _date(m.get("next_court_date")),
                        m.get("number"),
                        m.get("case_name") or m.get("title") or "",
                        " ".join(x for x in (m.get("court"), m.get("case_number")) if x),
                        _date(m.get("complaint_filed")),
                        _status(m),
                        d.get("name"),
                        d.get("status"),
                        *sent_cells(d),
                        _date(d.get("served")),
                        _src(d.get("served")),
                        _date(d.get("answered")),
                        _src(d.get("answered")),
                        _flags(d, "flags"),
                    ],
                }
            )
    rows.sort(key=lambda r: (str(r["cells"][0]), _sort_date(r["cells"][1]), str(r["cells"][2])))
    return rows


def needs_action(r: dict[str, Any]) -> bool:
    m, d = r["m"], r["d"]
    if _status(m) not in vocab.OPEN_STATUSES or _status(m) == vocab.NOT_FILED:
        return False
    if d.get("status") == vocab.FIRM_CLIENT:
        return not (d.get("answered") or {}).get("date")
    return d.get("status") not in vocab.DEFENDANT_DONE


def _fmt(ws: Any, widths: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _sheet(wb: Any, title: str, header: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append([clean(v) for v in r])
    _fmt(ws, widths)


def _summary(wb: Any, ms: list[dict[str, Any]], need: list[dict[str, Any]], as_of: dt.date) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.append(
        [
            "Responsible (Smokeball)",
            "Matters",
            "Active",
            "Settled, dismissal not yet filed",
            "Dismissed / closed / judgment",
            "Stayed / arbitration",
            "Not filed",
            "Status unclear",
            "Defendants needing action",
            "As of",
        ]
    )
    by = collections.defaultdict(list)
    for m in ms:
        by[m.get("responsible") or ""].append(m)
    for who, group in sorted(by.items()):
        c = collections.Counter(_status(m) for m in group)
        ws.append(
            [
                clean(who),
                len(group),
                c[vocab.ACTIVE],
                c[vocab.SETTLED_OWED],
                c[vocab.DISMISSED] + c[vocab.JUDGMENT],
                c[vocab.STAYED],
                c[vocab.NOT_FILED],
                c[vocab.UNCLEAR],
                sum(1 for r in need if r["m"] is not None and (r["m"].get("responsible") or "") == who),
                as_of.isoformat(),
            ]
        )
    _fmt(ws, [22, 10, 10, 18, 18, 14, 10, 12, 16, 12])


def _discovery_rows(ms: list[dict[str, Any]]) -> list[list[Any]]:
    out = []
    for m in sorted(ms, key=lambda m: (m.get("responsible") or "", str(m.get("number")))):
        base = [m.get("responsible") or "", m.get("number"), m.get("case_name") or m.get("title") or ""]
        for r in m.get("discovery_propounded") or []:
            out.append(
                [*base, "Propounded by the firm", r.get("set"), r.get("served_on"), r.get("date"), _src(r), "", "", ""]
            )
        for r in m.get("discovery_served_on_client") or []:
            rs = r.get("responses_served") or {}
            out.append(
                [
                    *base,
                    "Served on the firm's client",
                    r.get("set"),
                    r.get("served_by"),
                    r.get("date"),
                    _src(r),
                    rs.get("value") or "",
                    rs.get("date") or "",
                    _src(rs),
                ]
            )
    return out


def humanize(path: str) -> str:
    m = re.match(r"defendants\[(.*)\]\.(\w+)$", path)
    if m:
        return f"{m.group(1)}: {m.group(2).replace('_', ' ')}"
    m = re.match(r"(discovery_\w+)\[(.*)\]\.(\w+)$", path)
    if m:
        what = "Discovery propounded" if m.group(1) == "discovery_propounded" else "Discovery served on client"
        return f"{what} ({m.group(2).replace('|', ', ')}): {m.group(3)}"
    return path.replace("_", " ").capitalize()


def _show(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, tuple):
        return " ".join(str(x) for x in v if x)
    return str(v)


def case_notes_order(ms: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The Case notes tab's order (the per-case list's tab order too)."""
    return sorted(ms, key=lambda m: (m.get("responsible") or "", str(m.get("number"))))


def case_notes_text(m: dict[str, Any]) -> str:
    """One matter's Case notes cell, before ``clean``."""
    detail = (m.get("case_status") or {}).get("detail") or ""
    role = "The firm represents the defense in this matter." if m.get("firm_role") == "defense" else ""
    return " | ".join(x for x in (role, detail, _flags(m, "matter_flags"), m.get("notes") or "") if x)


def not_filed_text(m: dict[str, Any]) -> str:
    """One not-filed matter's "Where it stands" cell, before ``clean``."""
    return " | ".join(x for x in (m.get("notes") or "", _flags(m, "matter_flags")) if x)


def _notes_sheets(wb: Any, ms: list[dict[str, Any]]) -> None:
    """Case notes and Not filed."""
    notes = []
    for m in case_notes_order(ms):
        text = case_notes_text(m)
        notes.append(
            [
                m.get("responsible") or "",
                m.get("number"),
                m.get("case_name") or m.get("title") or "",
                _status(m),
                _date(m.get("next_court_date")),
                text,
            ]
        )
    _sheet(
        wb,
        "Case notes",
        ["Responsible (Smokeball)", "Matter #", "Case", "Case status", "Next court date", "Case notes"],
        notes,
        [18, 10, 34, 24, 22, 120],
    )
    nf = [
        [
            m.get("responsible") or "",
            m.get("number"),
            m.get("case_name") or m.get("title") or "",
            not_filed_text(m),
        ]
        for m in ms
        if _status(m) == vocab.NOT_FILED
    ]
    _sheet(wb, "Not filed", ["Responsible (Smokeball)", "Matter #", "Case", "Where it stands"], nf, [18, 10, 34, 120])


def _changed_rows(ms: list[dict[str, Any]], changes: dict[str, list[dict[str, Any]]]) -> list[list[Any]]:
    by_num = {str(m.get("matter_id")): m for m in ms}
    changed = []
    for mid, cs in changes.items():
        m = by_num.get(mid) or {}
        for c in cs:
            changed.append(
                [
                    m.get("responsible") or "",
                    m.get("number"),
                    m.get("case_name") or m.get("title") or "",
                    humanize(c.get("label") or c["path"]) if c["path"] != "matter" else "Matter",
                    _show(c["old"]),
                    _show(c["new"]) if not isinstance(c["new"], dict) else "added",
                ]
            )
    return changed


#: A people-scoped list's own tab (2026-10-24): every litigated file each person
#: holds a staff role on, whether or not a complaint is filed yet.
TAGGED = "Litigated files tagged on"


def tagged_header(several: bool) -> list[str]:
    return (["Person"] if several else []) + ["Matter #", "Case", "Role", "Case status", "Complaint filed"]


def tagged_rows(ms: list[dict[str, Any]], matters: list[dict[str, Any]], who: list[dict[str, Any]]) -> list[list[Any]]:
    """One row per (person, litigated matter they hold a role on): the person
    when several, matter #, case, their role(s), case status, complaint filed.
    ``matters`` are the inventory's (they carry ``roles``); ``who`` its people."""
    roles = {str(m.get("id")): m.get("roles") or {} for m in matters}
    several = len(who) > 1
    rows: list[list[Any]] = []
    for p in who:
        for m in sorted(ms, key=lambda m: str(m.get("number"))):
            held = role_labels(roles.get(str(m.get("matter_id"))), p.get("ids") or ())
            if held:
                rows.append(
                    ([p.get("name") or ""] if several else [])
                    + [
                        m.get("number"),
                        m.get("case_name") or m.get("title") or "",
                        ", ".join(held),
                        _status(m),
                        _date(m.get("complaint_filed")),
                    ]
                )
    return rows


def build(
    ms: list[dict[str, Any]],
    changes: dict[str, list[dict[str, Any]]],
    out: Path,
    as_of: dt.date,
    *,
    discovery: bool = True,
    tagged: tuple[list[str], list[list[Any]]] | None = None,
) -> dict[str, Any]:
    """``discovery=False`` (a court-only request) leaves the Discovery tab out:
    its values were not read this run, and a list never shows a stale column
    as current. ``tagged`` (header, rows), for a people scope, adds the
    "Litigated files tagged on" tab after the Summary."""
    from openpyxl import Workbook

    rows = defendant_rows(ms)
    need = [r for r in rows if needs_action(r)]
    wb = Workbook()
    _summary(wb, ms, need, as_of)
    if tagged is not None:
        header, trows = tagged
        _sheet(wb, TAGGED, header, trows, ([22] if len(header) == 6 else []) + [10, 34, 28, 24, 18])
    _sheet(wb, "Needs action", DEF_HEADER, [r["cells"] for r in need], DEF_WIDTHS)
    _sheet(
        wb,
        "Settled - dismissal owed",
        DEF_HEADER,
        [r["cells"] for r in rows if _status(r["m"]) == vocab.SETTLED_OWED],
        DEF_WIDTHS,
    )
    _sheet(wb, "All defendants", DEF_HEADER, [r["cells"] for r in rows], DEF_WIDTHS)
    if discovery:
        _sheet(
            wb,
            "Discovery",
            [
                "Responsible (Smokeball)",
                "Matter #",
                "Case",
                "Direction",
                "Set",
                "Party",
                "Date served",
                "Source",
                "Responses served",
                "Responses date",
                "Responses (source)",
            ],
            _discovery_rows(ms),
            [18, 10, 30, 22, 30, 26, 12, 40, 12, 12, 40],
        )
    _notes_sheets(wb, ms)
    changed = _changed_rows(ms, changes)
    _sheet(
        wb,
        "What changed",
        ["Responsible (Smokeball)", "Matter #", "Case", "What", "Was", "Now"],
        changed,
        [18, 10, 30, 40, 30, 30],
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {
        "matters": len(ms),
        "defendant_rows": len(rows),
        "needs_action": len(need),
        "changes": len(changed),
        **({"tagged_rows": len(tagged[1])} if tagged is not None else {}),
    }
