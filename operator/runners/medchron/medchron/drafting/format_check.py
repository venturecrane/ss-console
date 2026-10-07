"""`format_check`: the rendered Word file against the firm's authored format.

The rules come from ``drafting-firm.yaml``'s ``format`` block and nowhere
else: the specimen briefs are exemplars and review aids, never measured. Each
assertion reads the rendered .docx (python-docx's model for what it models,
the raw XML for the footer field and the section properties):

* every run with text is the authored font at the authored size;
* plain US Letter paper, no line numbering (never pleading paper);
* mediation brief: level-1 headings centered, bold, NOT underlined; level-2
  indented, bold and underlined; body double-spaced; the authored sections in
  the authored order; settlement authority, target and bracket only as
  ``{{ATTORNEY}}`` markers;
* discovery: item labels in the authored label style (all caps, bold,
  underlined); labels and requests double-spaced with no space before or
  after; requests first-line indented;
* discovery set: a Definitions section (no preliminary statement is
  required: the Code bars one); a proof of service; the section 2030.050
  declaration present if and only if the CUMULATIVE count to the responding
  party is over 35 or a prior count is not readable (``render.decl_decision``,
  the same rule the attach applied); discovery response: a proof of service and
  no declaration;
* deposition outline and mediation brief: a PAGE field in the footer;
  mediation level-3 headings indented 1.0", bold and underlined, body
  paragraphs first-line indented 0.5".

A failure here is OUR render's fault (the renderer and the settlements before
it are code), so the job ends ``failed`` (resumable, SMD alerted), never held.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .render import DECL_LINE, FIGURE, RESERVED, _outside_markers, decl_decision

LETTER = (12240, 15840)
_ROMAN = re.compile(r"^\s*[IVXLC]+\.\s*", re.I)


@dataclass
class Result:
    fails: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.fails


def _underlined(run: Any, para: Any) -> bool:
    if run.underline is not None:
        return bool(run.underline)
    style = para.style
    while style is not None:
        if style.font.underline is not None:
            return bool(style.font.underline)
        style = style.base_style
    return False


def _bold(run: Any, para: Any) -> bool:
    if run.bold is not None:
        return bool(run.bold)
    style = para.style
    while style is not None:
        if style.font.bold is not None:
            return bool(style.font.bold)
        style = style.base_style
    return False


def _eff(p: Any, attr: str) -> Any:
    """A paragraph-format value as Word resolves it: the paragraph's own, else
    its style chain's."""
    v = getattr(p.paragraph_format, attr)
    style = p.style
    while v is None and style is not None:
        v = getattr(style.paragraph_format, attr)
        style = style.base_style
    return v


def _text_runs(p: Any) -> list[Any]:
    return [r for r in p.runs if r.text.strip()]


def _paras(doc: Any) -> list[Any]:
    return [p for p in doc.paragraphs if p.text.strip()]


def _level(p: Any) -> int:
    name = p.style.name if p.style is not None else ""
    m = re.search(r"Heading (\d)$", name)
    return int(m.group(1)) if m else 0


def check_font(doc: Any, fmt: dict[str, Any], res: Result) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    runs = [r for p in doc.paragraphs for r in _text_runs(p)]
    runs += [r for t in doc.tables for c in t._cells for p in c.paragraphs for r in _text_runs(p)]
    bad = 0
    for r in runs:
        rpr = r._r.rPr
        fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if fonts is None or fonts.get(qn("w:ascii")) != fmt["font"] or r.font.size != Pt(float(fmt["size_pt"])):
            bad += 1
    if bad:
        res.fails.append(f"font: {bad} run(s) are not {fmt['font']} {fmt['size_pt']}")


def check_paper(doc: Any, res: Result) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Twips

    for s in doc.sections:
        if (s.page_width, s.page_height) != (Twips(LETTER[0]), Twips(LETTER[1])):
            res.fails.append("paper: not US Letter")
        if s._sectPr.find(qn("w:lnNumType")) is not None:
            res.fails.append("paper: line numbering present (pleading paper)")


def check_mediation(doc: Any, fmt: dict[str, Any], res: Result) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    paras = _paras(doc)
    for p in (p for p in paras if _level(p) == 1):
        if p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            res.fails.append(f"heading: level 1 not centered: {p.text[:40]}")
        if not all(_bold(r, p) for r in _text_runs(p)):
            res.fails.append(f"heading: level 1 not bold: {p.text[:40]}")
        if any(_underlined(r, p) for r in _text_runs(p)):
            res.fails.append(f"heading: level 1 underlined: {p.text[:40]}")
    for p in (p for p in paras if _level(p) == 2):
        if not (p.paragraph_format.left_indent and p.paragraph_format.left_indent > 0):
            res.fails.append(f"heading: level 2 not indented: {p.text[:40]}")
        if not all(_bold(r, p) and _underlined(r, p) for r in _text_runs(p)):
            res.fails.append(f"heading: level 2 not bold and underlined: {p.text[:40]}")
    for p in (p for p in paras if _level(p) == 3):
        if _eff(p, "left_indent") != Inches(1.0):
            res.fails.append(f"heading: level 3 not indented 1.0 inch: {p.text[:40]}")
        if not all(_bold(r, p) and _underlined(r, p) for r in _text_runs(p)):
            res.fails.append(f"heading: level 3 not bold and underlined: {p.text[:40]}")
    body = [p for p in paras if p.style is not None and p.style.name == "SMD Body"]
    single = [p for p in body if _eff(p, "line_spacing") != 2.0]
    if single:
        res.fails.append(f"spacing: {len(single)} body paragraph(s) not double-spaced")
    prose = [p for p in body if not re.match(r"^\s*\S{1,6}\t", p.text)]  # a numbered item hangs instead
    unindented = [p for p in prose if _eff(p, "first_line_indent") != Inches(0.5)]
    if unindented:
        res.fails.append(f"indent: {len(unindented)} body paragraph(s) without the 0.5 inch first-line indent")
    heads = [_ROMAN.sub("", p.text).strip().casefold() for p in paras if _level(p) == 1]
    pos = -1
    for title in fmt["mediation_brief_sections"]:
        hit = next((i for i, h in enumerate(heads) if i > pos and _ROMAN.sub("", title).strip().casefold() in h), None)
        if hit is None:
            res.fails.append(f"sections: {title!r} missing or out of the authored order")
        else:
            pos = hit
    for p in paras:
        t = _outside_markers(p.text)
        if _level(p) or not RESERVED.search(t):
            continue
        if FIGURE.search(t):
            res.fails.append(f"reserved: a settlement figure outside an {{{{ATTORNEY}}}} marker: {p.text[:60]}")
        elif "{{ATTORNEY" not in p.text:
            res.fails.append(
                f"reserved: settlement authority/target/bracket without an {{{{ATTORNEY}}}} marker: {p.text[:60]}"
            )


def check_discovery(doc: Any, cls: str, fmt: dict[str, Any], res: Result, digest: str = "") -> None:
    from docx.shared import Inches, Pt

    paras = _paras(doc)
    labels = [p for p in paras if p.style is not None and p.style.name == "SMD Item Label"]
    items = [p for p in paras if p.style is not None and p.style.name == "SMD Item Text"]
    if not labels:
        res.fails.append("discovery: no item labels")
    if fmt["discovery_label_style"] in ("caps_bold_underline", "all caps, bold, underlined"):
        for p in labels:
            if not all(
                _bold(r, p) and _underlined(r, p) and (r.font.all_caps or r.text == r.text.upper())
                for r in _text_runs(p)
            ):
                res.fails.append(f"label: not all-caps bold underlined: {p.text[:40]}")
    for p in labels + items:
        spacing = _eff(p, "line_spacing")
        extra = [a for a in ("space_before", "space_after") if (_eff(p, a) or Pt(0)) != Pt(0)]
        if spacing != 2.0 or extra:
            res.fails.append(
                f"spacing: not double-spaced between items only ({', '.join(extra) or 'line spacing'}): {p.text[:40]}"
            )
    for p in items:
        if p.paragraph_format.first_line_indent != Inches(0.5):
            res.fails.append(f"request: not first-line indented: {p.text[:40]}")
    texts = [p.text for p in paras]
    if not any("PROOF OF SERVICE" in t.upper() for t in texts):
        res.fails.append("attachments: proof of service missing")
    has_decl = any(DECL_LINE.search(t) for t in texts)
    pos_at = next((i for i, t in enumerate(texts) if "PROOF OF SERVICE" in t.upper()), None)
    decl_at = next((i for i, t in enumerate(texts) if DECL_LINE.search(t)), None)
    if pos_at is not None and decl_at is not None and decl_at > pos_at:
        res.fails.append("attachments: the declaration must precede the proof of service")
    if cls == "discovery_set":
        if not any(t.strip().strip(":").upper() == "DEFINITIONS" for t in texts):
            res.fails.append("definitions: no Definitions section")
        d = decl_decision(doc_text(doc), digest)
        if has_decl != d.attach:
            want = "required" if d.attach else "not allowed"
            res.fails.append(
                f"attachments: section 2030.050 declaration {want} ({d.total} special interrogatories to this "
                f"party, {d.this_set} in this set; {len(d.unreadable)} prior count(s) not readable)"
            )
    elif has_decl:
        res.fails.append("attachments: a response carries no section 2030.050 declaration")


def check_page_field(blob: bytes, res: Result) -> None:
    z = zipfile.ZipFile(io.BytesIO(blob))
    footers = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.startswith("word/footer"))
    if not re.search(r"<w:instrText[^>]*>\s*PAGE\b", footers) and 'w:instr=" PAGE' not in footers:
        res.fails.append("footer: no PAGE field (page numbers at the bottom)")


def doc_text(doc: Any) -> str:
    """The rendered document as text the shared rules read: paragraphs, and
    table rows as ``| a | b |`` (the identification block is a table)."""
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        out += ["| " + " | ".join(c.text for c in row.cells) + " |" for row in t.rows]
    return "\n".join(out)


def check(path: Path, cls: str, fmt: dict[str, Any], digest: str = "") -> Result:
    """``digest`` carries the DISCOVERY-SETS index the cumulative section
    2030.050 rule reads (``render.decl_decision``, the rule the attach used)."""
    import docx

    blob = Path(path).read_bytes()
    doc = docx.Document(io.BytesIO(blob))
    res = Result()
    check_font(doc, fmt, res)
    check_paper(doc, res)
    if cls == "mediation_brief":
        check_mediation(doc, fmt, res)
    if cls in ("discovery_set", "discovery_response"):
        check_discovery(doc, cls, fmt, res, digest)
    if cls == "mediation_brief" or (cls == "depo_outline" and fmt["depo_outline_page_numbers"]):
        check_page_field(blob, res)
    return res
