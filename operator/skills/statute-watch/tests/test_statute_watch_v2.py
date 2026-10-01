"""Tests for statute-watch v2: client names from titles, urgency sections, the
"since last month" comparison and its state file, the attached workbook, and
the whole message (body AND workbook) through the send gate's own filters with
a register seeded only from the handoff file.

Run from operator/:  python -m pytest skills/statute-watch -q
(the workbook tests need openpyxl in the interpreter running pytest, as the
connector venv has it on the seat).

Every fixture name, number and date below is invented.
"""

from __future__ import annotations

import base64
import hashlib
import importlib.util
import io
import json
import sys
import zipfile
from contextlib import redirect_stdout
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest

#: The workbook is built with openpyxl (the connector venv has it on the seat).
#: Tests that build one skip without it; the body, handoff and state tests do not.
HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
needs_openpyxl = pytest.mark.skipif(not HAS_OPENPYXL, reason="openpyxl not installed")

_DIR = Path(__file__).resolve().parents[1]
_OPERATOR = _DIR.parents[1]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


pre_run = _load("statute_watch_pre_run_v2_test", _DIR / "pre_run.py")
render = pre_run._RENDER
clients = pre_run._CLIENTS
changes = pre_run._CHANGES
workbook = pre_run._WORKBOOK
pull = pre_run._PULL
report = pre_run._REPORT
citation_filter = _load("sw2_citation_filter", _OPERATOR / "safety_substrate" / "citation_filter.py")
identifier_filter = _load("sw2_identifier_filter", _OPERATOR / "safety_substrate" / "identifier_filter.py")

TODAY = date(2026, 10, 1)
RECIPIENT = "office@firm.test"
CFG = {
    "business_hours": {"timezone": "America/Los_Angeles"},
    "personas": [{"skills": [{"name": "statute-watch", "settings": {"recipient": RECIPIENT}}]}],
}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _day(offset: int) -> str:
    return (TODAY + timedelta(days=offset)).isoformat()


# ---------------------------------------------------------------------------
# Client names from the matter title
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("title", "workbook_name", "body_name"),
    [
        ("913353 - Marchetti, Rocco - Auto - 2024", "Marchetti, Rocco", "Marchetti"),
        ("913354 - Ruiz, Maria V. - Premises", "Ruiz, Maria", "Ruiz"),
        ("913355 - Ruiz, Maria V - Premises", "Ruiz, Maria", "Ruiz"),
        ("913356 - Okafor, J. Robert - Auto", "Okafor, Robert", "Okafor"),
        ("913357 - Lind, Ann | Lind, Bo - Auto", "Lind, Ann; Lind, Bo", "Lind"),
        ("913358 - Lee, Ann | Park, Ben - Auto", "Lee, Ann; Park, Ben", "Lee and Park"),
        ("913359 - Lee, Ann | Park, Ben | Cho, Dee - Dog bite", "Lee, Ann; Park, Ben; Cho, Dee", "Lee, Park and Cho"),
        ("913360 - Moss, Ann & Vale, Tom - Auto", "Moss, Ann; Vale, Tom", "Moss and Vale"),
        ("913361 - Van Der Berg, Pia - Auto", "Van Der Berg, Pia", "Van Der Berg"),
        ("913362 - O'Neil-Smith, Kay - Auto", "O'Neil-Smith, Kay", "O'Neil-Smith"),
        ("913363 - Doe (Minor), Sam - Auto", "Doe, Sam", "Doe"),
        ("913364 - Acme Holdings - Property", "Acme Holdings", "Acme Holdings"),
        ("2026-PI-107 - Okafor, Pat - Auto", "Okafor, Pat", "Okafor"),
    ],
)
def test_client_names_come_from_the_title(title, workbook_name, body_name) -> None:
    assert clients.names_for(title, "Contactname") == (workbook_name, body_name)
    line = (
        "3 | 2026-10-04 | 913353 | " + workbook_name + " | Pat Quill | Open | 1\n- **File 913353, " + body_name + "**"
    )
    assert not citation_filter.contains_citation(line)


@pytest.mark.parametrize(
    "title", [None, "", "Marchetti, Rocco - Auto", "Estate matter", "913353", "Smith-Jones - Lee, Ann - Auto"]
)
def test_a_title_without_the_shape_falls_back_to_the_contact(title) -> None:
    assert clients.names_for(title, "Lindqvist") == ("Lindqvist", "Lindqvist")
    assert clients.names_for(title, None) == (None, None)


def test_a_caption_shaped_title_never_renders_as_one() -> None:
    workbook_name, body_name = clients.names_for("913365 - Garcia v. Allstate, Inc - Auto", None)
    assert body_name == "Garcia"
    assert not citation_filter.contains_citation("- **File 913365, " + body_name + "** | " + str(workbook_name))


def test_attorney_full_name_drops_initials() -> None:
    assert clients.person_name("Pat J.", "Quill") == "Pat Quill"
    assert clients.person_name(None, None) is None


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------


def _row(number: str, days: int, client: str = "Quill") -> report.Row:
    return report.Row(
        matter_number=number,
        statute=TODAY + timedelta(days=days),
        days_left=days,
        client=client,
        attorney="Whitlock",
        attorney_assigned=True,
        court_documents=1,
        matter_id="m-" + number,
        client_workbook=client + ", Pat",
        attorney_full="Dana Whitlock",
        status="Open",
    )


def test_sections_split_by_days_left() -> None:
    rows = [_row("301", 0), _row("302", 7), _row("303", 8), _row("304", 30), _row("305", 31)]
    body = render.render_full(rows, total=5, unreadable=0, week=render.due_this_week(rows))
    week, rest = body.split("## Due this week\n\n")[1].split("## Due this month\n\n")
    month, later = rest.split("## Later\n\n")
    assert "File 301" in week and "File 302" in week and "File 303" not in week
    assert "File 303" in month and "File 304" in month and "File 305" not in month
    assert "File 305" in later
    assert body.startswith("5 open cases have a statute date in the next three months")
    assert "in Smokeball, 2 due within 7 days." in body.splitlines()[0]
    assert (
        "- **File 301, Quill**\n   Statute date October 1, 2026, due today. Attorney Whitlock. Court-named documents: 1."
        in body
    )
    assert body.rstrip().endswith(render.ATTACHED_LINE)


def test_an_empty_section_is_left_out() -> None:
    rows = [_row("305", 45)]
    body = render.render_full(rows, total=1, unreadable=0, week=0, changes=[])
    assert "## Due this week" not in body and "## Due this month" not in body and "## Later" in body
    assert body.startswith("1 open case has a statute date")
    assert "; since last month, 0 new and 0 passed with no filing recorded." in body
    assert render.NO_CHANGES_LINE in body


# ---------------------------------------------------------------------------
# Since last month: one fixture row per classification
# ---------------------------------------------------------------------------

PREV_DATE = "2026-10-20"


@pytest.mark.parametrize(
    ("record", "kind", "statute"),
    [
        ({"status": "Open", "statute": PREV_DATE, "filed": True, "caseNumber": False}, changes.FILED, PREV_DATE),
        ({"status": "Open", "statute": PREV_DATE, "filed": False, "caseNumber": True}, changes.FILED, PREV_DATE),
        ({"status": "Closed", "statute": PREV_DATE, "filed": False, "caseNumber": False}, changes.CLOSED, PREV_DATE),
        (
            {"status": "Archived", "statute": PREV_DATE, "filed": False, "caseNumber": False},
            changes.NOT_OPEN,
            PREV_DATE,
        ),
        (
            {"status": "Open", "statute": "2027-03-01", "filed": False, "caseNumber": False},
            changes.CHANGED,
            "2027-03-01",
        ),
        ({"status": "Open", "statute": None, "filed": False, "caseNumber": False}, changes.REMOVED, None),
        ({"status": "Open", "statute": "soon", "filed": False, "caseNumber": False}, changes.UNCHECKED, None),
        ({"error": True}, changes.UNCHECKED, None),
        (None, changes.UNCHECKED, None),
    ],
)
def test_each_departure_is_worded_as_the_record_shows(record, kind, statute) -> None:
    entry = {"matter_id": "m-1", "number": "913353", "statute_date": PREV_DATE}
    got_kind, got_statute = changes.classify(entry, record, TODAY)
    assert got_kind == kind
    assert (got_statute.isoformat() if got_statute else None) == statute


def test_a_passed_statute_with_nothing_filed() -> None:
    entry = {"matter_id": "m-1", "number": "913353", "statute_date": "2026-09-28"}
    record = {"status": "Open", "statute": "2026-09-28T00:00:00", "filed": False, "caseNumber": False}
    assert changes.classify(entry, record, TODAY) == (changes.PASSED, date(2026, 9, 28))


def test_the_sentences() -> None:
    def change(kind, statute=None, previous=None):
        return changes.Change(kind, "m", "913353", "Marchetti", "Marchetti, Rocco", statute, previous)

    d1, d2 = date(2026, 9, 28), date(2027, 3, 1)
    assert (
        render.change_item(change("new", d1))
        == "- **File 913353, Marchetti**: New on the list. Statute date September 28, 2026."
    )
    assert render.change_detail(change("passed", d1, d1)) == (
        "Statute date September 28, 2026 has passed. Smokeball shows no Filed date and no Case number."
    )
    assert (
        render.change_detail(change("changed", d2, d1))
        == "Statute date changed from September 28, 2026 to March 1, 2027."
    )
    assert render.change_detail(change("removed", None, d1)) == "Statute date removed in Smokeball."
    assert (
        render.change_detail(change("filed", d1, d1)) == "Filed per Smokeball (Filed date or Case number now recorded)."
    )
    assert render.change_detail(change("closed", d1, d1)) == "Closed in Smokeball."
    assert render.change_detail(change("unchecked", None, d1)) == "Could not be checked this month."


# ---------------------------------------------------------------------------
# The whole message through the gate: one row of every kind
# ---------------------------------------------------------------------------


def _matter(mid, number, statute, last, **extra) -> dict:
    row = {
        "id": mid,
        "number": number,
        "clientIds": ["c-" + mid],
        "staffId": "s-1",
        "statute": statute,
        "title": number + " - " + last + ", Pat - Auto - 2024",
        "status": "Open",
        "caseNumber": False,
        "filed": False,
    }
    row.update(extra)
    return row


#: This month: m-new (new), m-stay (on both lists), m-later; m-filed is still
#: open but now filed. Every other previous case is gone from the open list.
RAW = {
    "matters": [
        _matter("m-new", "913301", _day(3), "Marchetti"),
        _matter("m-stay", "913302", _day(20), "Lindqvist"),
        _matter("m-later", "2026-PI-107", _day(60), "Okafor"),
        _matter("m-filed", "913304", _day(10), "Ruiz", filed=True),
    ]
}
PREVIOUS = [
    {"matter_id": "m-stay", "number": "913302", "statute_date": _day(20)},
    {"matter_id": "m-later", "number": "2026-PI-107", "statute_date": _day(60)},
    {"matter_id": "m-filed", "number": "913304", "statute_date": _day(10)},
    {"matter_id": "m-closed", "number": "913305", "statute_date": _day(15)},
    {"matter_id": "m-passed", "number": "913306", "statute_date": _day(-3)},
    {"matter_id": "m-moved", "number": "913307", "statute_date": _day(19)},
    {"matter_id": "m-removed", "number": "913308", "statute_date": _day(25)},
    {"matter_id": "m-error", "number": "913309", "statute_date": _day(40)},
]


def _departed(title_last, **facts) -> dict:
    return {"title": "913300 - " + title_last + ", Pat - Auto", "filed": False, "caseNumber": False, **facts}


DETAILS = {
    "staff": {"s-1": {"firstName": "Dana", "lastName": "Whitlock"}},
    "matters": {"m-new": {"courtDocuments": 0}, "m-stay": {"courtDocuments": 2}, "m-later": {"documentsError": True}},
    "departed": {
        "m-filed": _departed("Ruiz", status="Open", statute=_day(10), filed=True),
        "m-closed": _departed("Moss", status="Closed", statute=_day(15)),
        "m-passed": _departed("Vale", status="Open", statute=_day(-3)),
        "m-moved": _departed("Cho", status="Open", statute=_day(150)),
        "m-removed": _departed("Park", status="Open", statute=None),
        "m-error": {"error": True},
    },
}


class _Clock:
    def __init__(self) -> None:
        self.t0 = datetime(2026, 10, 1, 19, 7, tzinfo=timezone.utc)
        self.calls = 0

    def __call__(self) -> datetime:
        self.calls += 1
        return self.t0 if self.calls == 1 else self.t0 + timedelta(minutes=12)


def _real_workbook(spec: dict) -> dict | None:
    return pull.build_workbook(workbook.SNIPPET, spec)


def _seed_state(tmp_path: Path, cases=PREVIOUS) -> None:
    path = tmp_path / ".smd" / "statute-watch" / "last-run.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"version": 1, "run_local_date": "2026-09-03", "cases": cases}))


DEFAULT_BUILD = _real_workbook if HAS_OPENPYXL else None


def _run(tmp_path, monkeypatch, *, dry=False, previous=PREVIOUS, build=DEFAULT_BUILD, no_state=False):
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    monkeypatch.setenv("SMD_CONNECTOR_VENV_PYTHON", sys.executable)
    monkeypatch.delenv("HERMES_TIMEZONE", raising=False)
    if no_state:
        monkeypatch.setenv("SMD_STATUTE_WATCH_NO_STATE", "1")
    else:
        monkeypatch.delenv("SMD_STATUTE_WATCH_NO_STATE", raising=False)
    if previous is not None:
        _seed_state(tmp_path, previous)
    monkeypatch.setattr(pre_run, "_heartbeat", lambda *a: True)
    seen: dict = {}

    def details(cases, departed):
        seen["departed"] = sorted(departed)
        return DETAILS

    out = io.StringIO()
    with redirect_stdout(out):
        assert (
            pre_run.run(
                CFG, pull_matters=lambda: RAW, pull_details=details, build_workbook=build, clock=_Clock(), dry=dry
            )
            == 0
        )
    return json.loads(out.getvalue().strip().splitlines()[-1]), out.getvalue(), seen


def _envelope(tmp_path: Path) -> dict:
    return json.loads((tmp_path / ".smd" / "pre_run" / "statute-watch.dispatch.json").read_text())


def _register_from_handoff(tmp_path: Path) -> identifier_filter.ProvenanceRegister:
    """Seeded ONLY from the handoff file, the way the overlay seeds it: date
    atoms, then at most 100 (matterNumber, dates) records."""
    record = json.loads((tmp_path / ".smd" / "pre_run" / "statute-watch.json").read_text())
    reg = identifier_filter.ProvenanceRegister()
    for day in record["dates"][:200]:
        reg.add_read_text(day)
    for entry in record["records"][:100]:
        reg.add_record(entry["matterNumber"], entry["dates"])
    return reg


def workbook_lines(data: bytes) -> list[str]:
    """The workbook as the overlay's scanner reads it: one line per row, cells
    joined " | ", date cells as ISO dates, empty cells empty."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data))
    lines = []
    for ws in book.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = []
            for value in row:
                if isinstance(value, (datetime, date)):
                    cells.append(value.isoformat()[:10])
                else:
                    cells.append("" if value is None else str(value))
            lines.append(" | ".join(cells))
    return lines


def _attachment_bytes(dispatch: dict) -> bytes:
    (attachment,) = dispatch["attachments"]
    return base64.b64decode(attachment["content_b64"])


def test_every_kind_renders_and_the_departures_were_read(tmp_path, monkeypatch) -> None:
    verdict, _stdout, seen = _run(tmp_path, monkeypatch)
    assert verdict["status"] == "report_ready" and verdict["cases"] == 3 and verdict["changes"] == 7
    assert verdict["attached"] is HAS_OPENPYXL
    assert seen["departed"] == sorted(["m-filed", "m-closed", "m-passed", "m-moved", "m-removed", "m-error"])
    body = _envelope(tmp_path)["dispatches"][0]["full_body"]
    since = body.split("## Since last month\n\n")[1]
    expected = [
        "- **File 913306, Vale**: Statute date September 28, 2026 has passed. Smokeball shows no Filed date and no Case number.",
        "- **File 913307, Cho**: Statute date changed from October 20, 2026 to February 28, 2027.",
        "- **File 913308, Park**: Statute date removed in Smokeball.",
        "- **File 913309, client name not available**: Could not be checked this month.",
        "- **File 913301, Marchetti**: New on the list. Statute date October 4, 2026.",
        "- **File 913304, Ruiz**: Filed per Smokeball (Filed date or Case number now recorded).",
        "- **File 913305, Moss**: Closed in Smokeball.",
    ]
    assert since.split("\n\n")[0].splitlines() == expected
    assert body.splitlines()[0] == (
        "3 open cases have a statute date in the next three months with no Filed date and no Case number in "
        "Smokeball, 1 due within 7 days; since last month, 1 new and 1 passed with no filing recorded."
    )
    assert "September 3" not in body and "2026-09-03" not in body  # last run's date is never printed


@needs_openpyxl
def test_body_and_workbook_pass_both_filters_with_a_handoff_only_register(tmp_path, monkeypatch) -> None:
    _run(tmp_path, monkeypatch)
    (dispatch,) = _envelope(tmp_path)["dispatches"]
    reg = _register_from_handoff(tmp_path)
    lines = workbook_lines(_attachment_bytes(dispatch))
    assert any(line.startswith("Statute date changed | 913307 | Cho, Pat | 2027-02-28 | ") for line in lines)
    for text in (
        dispatch["subject"] + "\n" + dispatch["full_body"],
        dispatch["body_without_attachment"],
        dispatch["skeleton_body"],
        "\n".join(lines),
    ):
        assert identifier_filter.check(text, reg).unverified == ()
        assert citation_filter.scan(text) == [] and not citation_filter.contains_citation(text)
        assert "—" not in text and "–" not in text
    for line in lines:  # the overlay scans row by row too
        assert identifier_filter.check(line, reg).unverified == (), line


def test_the_body_passes_both_filters_without_the_workbook(tmp_path, monkeypatch) -> None:
    """The same check with no workbook built: runs where openpyxl is absent."""
    _run(tmp_path, monkeypatch, build=None)
    (dispatch,) = _envelope(tmp_path)["dispatches"]
    reg = _register_from_handoff(tmp_path)
    for text in (dispatch["subject"] + "\n" + dispatch["full_body"], dispatch["skeleton_body"]):
        assert identifier_filter.check(text, reg).unverified == ()
        assert citation_filter.scan(text) == [] and not citation_filter.contains_citation(text)
    # Another case's (seeded) date on the changed-date line: only the per-line
    # pair check can see it.
    mispaired = dispatch["full_body"].replace("to February 28, 2027", "to October 4, 2026")
    assert identifier_filter.check(mispaired, reg).unverified


@needs_openpyxl
def test_the_workbook_check_can_fail(tmp_path, monkeypatch) -> None:
    """Falsifiers: an unseeded date in a cell, and two rows' dates swapped."""
    _run(tmp_path, monkeypatch)
    (dispatch,) = _envelope(tmp_path)["dispatches"]
    reg = _register_from_handoff(tmp_path)
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(_attachment_bytes(dispatch)))
    ws = book["Statute watch"]
    ws.cell(row=2, column=2).value = datetime(2026, 12, 24)
    buffer = io.BytesIO()
    book.save(buffer)
    assert identifier_filter.check("\n".join(workbook_lines(buffer.getvalue())), reg).unverified
    lines = workbook_lines(_attachment_bytes(dispatch))
    rows = [line for line in lines if line.split(" | ")[2:3] in (["913301"], ["913302"])]
    first, second = rows[0].split(" | "), rows[1].split(" | ")
    first[1], second[1] = second[1], first[1]
    assert identifier_filter.check(" | ".join(first) + "\n" + " | ".join(second), reg).unverified


@needs_openpyxl
def test_the_workbook_layout(tmp_path, monkeypatch) -> None:
    _run(tmp_path, monkeypatch)
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(_attachment_bytes(_envelope(tmp_path)["dispatches"][0])))
    assert book.sheetnames == ["Statute watch", "Since last month", "About"]
    ws = book["Statute watch"]
    assert [c.value for c in ws[1]] == list(workbook.CASE_HEADER)
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:G4"
    assert all(c.font.bold for c in ws[1])
    assert [c.value for c in ws[2]] == [
        3,
        datetime(2026, 10, 4),
        "913301",
        "Marchetti, Pat",
        "Dana Whitlock",
        "Open",
        0,
    ]
    assert ws.cell(row=2, column=2).number_format == "mm/dd/yyyy"
    assert ws.cell(row=2, column=1).fill.fgColor.rgb.endswith("F8D7D5")  # due within 7 days
    assert ws.cell(row=3, column=1).fill.fgColor.rgb.endswith("FFF2CC")  # within 30
    assert ws.cell(row=4, column=1).fill.fill_type is None
    assert ws.cell(row=4, column=7).value == workbook.NOT_CHECKED
    since = book["Since last month"]
    assert [c.value for c in since[1]] == list(workbook.CHANGE_HEADER)
    assert since.max_row == 8
    about = [r[0] for r in book["About"].iter_rows(values_only=True)]
    assert about == list(workbook.ABOUT_LINES)


@needs_openpyxl
def test_workbook_zip_parts_stay_inside_the_overlay_allowlist(tmp_path, monkeypatch) -> None:
    """The overlay refuses formulas, hyperlinks, worksheet rels, tables,
    comments, drawings and any part beyond the plain set (W1 constraint)."""
    _run(tmp_path, monkeypatch)
    data = _attachment_bytes(_envelope(tmp_path)["dispatches"][0])
    allowed_exact = {
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "xl/sharedStrings.xml",
        "xl/calcChain.xml",
        "docProps/app.xml",
        "docProps/core.xml",
    }
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        for name in names:
            sheet = name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name
            assert name in allowed_exact or sheet, name
        for name in names:
            if name.startswith("xl/worksheets/"):
                xml = archive.read(name).decode("utf-8")
                for marker in ("<f>", "<f ", "<hyperlink", "<tablePart", "<drawing", "<legacyDrawing", "r:id="):
                    assert marker not in xml, (name, marker)


@needs_openpyxl
def test_envelope_carries_the_pinned_attachment_and_the_middle_rung(tmp_path, monkeypatch) -> None:
    _run(tmp_path, monkeypatch)
    (dispatch,) = _envelope(tmp_path)["dispatches"]
    (attachment,) = dispatch["attachments"]
    assert set(attachment) == {"name", "content_type", "content_b64", "sha256"}
    assert attachment["name"] == "Statute watch October 2026.xlsx"
    assert attachment["content_type"] == XLSX
    data = base64.b64decode(attachment["content_b64"], validate=True)
    assert hashlib.sha256(data).hexdigest() == attachment["sha256"]
    assert len(data) <= 512 * 1024
    import re

    assert re.fullmatch(r"[\w .()-]{1,120}\.xlsx", attachment["name"])
    full, without = dispatch["full_body"], dispatch["body_without_attachment"]
    assert full.rstrip().endswith(render.ATTACHED_LINE)
    assert without == full.replace(render.ATTACHED_LINE, render.NOT_ATTACHED_LINE)
    assert dispatch["body_sha256_without_attachment"] == pre_run._H.canonical_body_sha256(without)
    assert dispatch["body_sha256_full"] == pre_run._H.canonical_body_sha256(full)
    assert dispatch["appends"] == [] and dispatch["subject"] == "Statute watch, October 2026: 3 cases, 1 due this week"


@needs_openpyxl
def test_the_wake_row_names_all_three_bodies(tmp_path, monkeypatch) -> None:
    """The send verifier grades a middle-rung send degraded only when the
    EMITTED_WAKE row carries the no-attachment hash."""
    _run(tmp_path, monkeypatch)
    (dispatch,) = _envelope(tmp_path)["dispatches"]
    assert pre_run._ENVELOPE.wake_stamps(_envelope(tmp_path)) == [
        {
            "body_sha256_full": dispatch["body_sha256_full"],
            "body_sha256_skeleton": dispatch["body_sha256_skeleton"],
            "body_sha256_without_attachment": dispatch["body_sha256_without_attachment"],
        }
    ]
    beats: list = []
    monkeypatch.setattr(pre_run, "_heartbeat", lambda *a: beats.append(a) or True)
    with redirect_stdout(io.StringIO()):
        pre_run.run(
            CFG,
            pull_matters=lambda: RAW,
            pull_details=lambda c, d: DETAILS,
            build_workbook=_real_workbook,
            clock=_Clock(),
        )
    (stamp,) = beats[0][3]["body_sha256"]
    assert set(stamp) == {"body_sha256_full", "body_sha256_skeleton", "body_sha256_without_attachment"}


def test_a_workbook_that_cannot_be_built_costs_only_the_attachment(tmp_path, monkeypatch) -> None:
    verdict, _stdout, _seen = _run(tmp_path, monkeypatch, build=lambda spec: None)
    assert verdict["status"] == "report_ready" and verdict["attached"] is False
    (dispatch,) = _envelope(tmp_path)["dispatches"]
    assert "attachments" not in dispatch and "body_without_attachment" not in dispatch
    assert dispatch["full_body"].rstrip().endswith(render.NOT_ATTACHED_LINE)


@needs_openpyxl
def test_no_last_run_says_changes_start_next_month(tmp_path, monkeypatch) -> None:
    verdict, _stdout, seen = _run(tmp_path, monkeypatch, previous=None)
    assert verdict["changes"] is None and seen["departed"] == []
    body = _envelope(tmp_path)["dispatches"][0]["full_body"]
    assert "## Since last month\n\n" + render.NO_PREVIOUS_LINE in body
    assert "since last month," not in body.splitlines()[0]
    lines = workbook_lines(_attachment_bytes(_envelope(tmp_path)["dispatches"][0]))
    assert render.NO_PREVIOUS_LINE + " |  |  |  | " in lines


# ---------------------------------------------------------------------------
# The state file
# ---------------------------------------------------------------------------


def _state(tmp_path: Path) -> Path:
    return tmp_path / ".smd" / "statute-watch" / "last-run.json"


def test_state_is_written_after_a_successful_envelope_write(tmp_path, monkeypatch) -> None:
    _run(tmp_path, monkeypatch, previous=None)
    path = _state(tmp_path)
    record = json.loads(path.read_text())
    assert record == {
        "version": 1,
        "run_local_date": "2026-10-01",
        "cases": [
            {"matter_id": "m-new", "number": "913301", "statute_date": "2026-10-04"},
            {"matter_id": "m-stay", "number": "913302", "statute_date": "2026-10-21"},
            {"matter_id": "m-later", "number": "2026-PI-107", "statute_date": "2026-11-30"},
        ],
    }
    assert path.stat().st_mode & 0o777 == 0o600 and path.parent.stat().st_mode & 0o777 == 0o700


def test_state_is_not_written_in_a_dry_run(tmp_path, monkeypatch) -> None:
    verdict, _stdout, _seen = _run(tmp_path, monkeypatch, dry=True, previous=None)
    assert verdict["status"] == "dry_run" and verdict["attached"] is HAS_OPENPYXL
    assert not (tmp_path / ".smd").exists()


def test_state_is_not_written_with_no_state_set(tmp_path, monkeypatch) -> None:
    _run(tmp_path, monkeypatch, no_state=True)
    assert json.loads(_state(tmp_path).read_text())["cases"] == PREVIOUS  # the seeded file, untouched


def test_state_is_not_written_when_the_envelope_is_not(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(pre_run._H, "write_dispatch_envelope", lambda skill, payload: False)
    verdict, _stdout, _seen = _run(tmp_path, monkeypatch, previous=None)
    assert verdict["status"] == "envelope_write_failed"
    assert not _state(tmp_path).exists()


def test_an_unreadable_state_reads_as_no_last_run(tmp_path, monkeypatch, capsys) -> None:
    monkeypatch.setenv("HERMES_HOME", str(tmp_path))
    _state(tmp_path).parent.mkdir(parents=True)
    _state(tmp_path).write_text("{not json")
    assert changes.load_state() is None
    assert "last-run state unreadable" in capsys.readouterr().err


def test_the_handoff_seeds_previous_dates_and_changes_come_first(tmp_path, monkeypatch) -> None:
    _run(tmp_path, monkeypatch)
    record = json.loads((tmp_path / ".smd" / "pre_run" / "statute-watch.json").read_text())
    by_number = {r["matterNumber"]: r["dates"] for r in record["records"]}
    assert sorted(by_number["913307"]) == ["2026-10-20", "2027-02-28"]  # changed: old and new
    assert by_number["913306"] == ["2026-09-28"]
    assert record["records"][0]["matterNumber"] == "913306"  # the first change, ahead of the case list
    assert "2026-09-03" not in record["dates"]
