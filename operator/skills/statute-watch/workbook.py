"""The statute report's spreadsheet: what goes in it, and the snippet that builds it.

openpyxl is installed in the Smokeball connector's venv (3.1.5) and NOT in the
gateway's python, so the workbook is built the way the reads are: ``pull.py``
runs ``SNIPPET`` under the connector-venv interpreter, the rows ride stdin as
JSON, and the snippet answers one JSON line ``{"content_b64", "sha256"}``.
Nothing about a case rides argv or is printed anywhere else.

Three sheets:

* "Statute watch": every listed case (the body caps at 100; the workbook
  carries them all). Days left, Statute date (a real date cell, mm/dd/yyyy),
  File number, Client ("Last, First"), Attorney (full name), Status,
  Court-named documents. Bold frozen header, column widths, an autofilter,
  and a fill on rows due within 7 days and within 30 days.
* "Since last month": Change, File number, Client, Statute date, Detail. The
  Detail is the same sentence the email body prints.
* "About": what is counted, in the words of the first report (letter 157).

THE SEND GATE READS THIS FILE. The overlay extracts every row as one line,
cells joined " | ", date cells as ISO dates, and runs the same identifier and
citation filters the body passes. So the only dates in the workbook are
statute dates, each on its own case's row (the handoff seeds those pairs); the
About sheet prints no date at all; names carry no initials; no em dashes.
"""

from __future__ import annotations

ABOUT_LINES = (
    "Statute watch: open cases nearing their statute of limitations with nothing filed (from Smokeball)",
    "What is listed: every Open or Pending case whose Statute of Limitation date in Smokeball falls from the day "
    "of this report through the next three months, and that has no Filed date and no Case number in Smokeball. "
    "Sorted by statute date, soonest first.",
    "Days left counts from the day of this report; 0 means the statute date is that day. Rows due within 7 days "
    "are shaded red, and rows due within 30 days are shaded yellow.",
    '"Court-named documents" counts documents on the case whose name reads like a court filing (complaint, '
    "summons, proof of service, answer, case management). A case can be filed while Smokeball's Filed date and "
    "Case number are still blank; a high count is the sign to look.",
    "Client: from the case title in Smokeball, with middle initials left off. Attorney: the case's responsible "
    "attorney in Smokeball.",
    '"Since last month" lists the cases new to the list, and every case that left it with what Smokeball shows '
    "for it now: a Filed date or Case number recorded, the case closed, the statute date passed with nothing "
    "filed, or the statute date changed or removed. A case that could not be read this month says so.",
)

CASE_HEADER = ("Days left", "Statute date", "File number", "Client", "Attorney", "Status", "Court-named documents")
CHANGE_HEADER = ("Change", "File number", "Client", "Statute date", "Detail")

NO_NUMBER = "No number on file"
NO_CLIENT = "Client name not available"
NOT_CHECKED = "not checked"

SNIPPET = """\
import base64
import hashlib
import io
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

spec = json.loads(sys.stdin.read())
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE4EE")
WEEK = PatternFill("solid", fgColor="F8D7D5")
MONTH = PatternFill("solid", fgColor="FFF2CC")


def day(value):
    return date.fromisoformat(value) if isinstance(value, str) and value else None


def sheet(wb, title, header, rows, widths, date_col):
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for cell in ws[1]:
        cell.font = BOLD
        cell.fill = HEAD
    for row in rows:
        values = list(row)
        values[date_col] = day(values[date_col])
        ws.append(values)
        cell = ws.cell(row=ws.max_row, column=date_col + 1)
        if cell.value is not None:
            cell.number_format = "mm/dd/yyyy"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:" + get_column_letter(len(header)) + str(max(ws.max_row, 1))
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == len(header) and title != "Statute watch")
    return ws


wb = Workbook()
wb.remove(wb.active)
ws = sheet(wb, "Statute watch", spec["case_header"], spec["cases"], [10, 13, 13, 34, 24, 10, 22], 1)
for row in ws.iter_rows(min_row=2):
    left = row[0].value
    fill = WEEK if isinstance(left, int) and left <= 7 else MONTH if isinstance(left, int) and left <= 30 else None
    if fill is not None:
        for cell in row:
            cell.fill = fill
sheet(wb, "Since last month", spec["change_header"], spec["changes"], [22, 13, 34, 13, 80], 3)
if spec.get("changes_note"):
    wb["Since last month"].cell(row=2, column=1, value=spec["changes_note"])
about = wb.create_sheet("About")
about.column_dimensions["A"].width = 110
for index, line in enumerate(spec["about"], start=1):
    about.cell(row=index, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
about["A1"].font = Font(bold=True, size=13)
wb.properties.creator = "Statute watch"
wb.active = 0
buffer = io.BytesIO()
wb.save(buffer)
data = buffer.getvalue()
print(json.dumps({"content_b64": base64.b64encode(data).decode("ascii"), "sha256": hashlib.sha256(data).hexdigest()}))
"""


def _docs(value) -> object:
    return value if isinstance(value, int) and not isinstance(value, bool) else NOT_CHECKED


def case_cells(row, attorney_text: str) -> list:
    """One "Statute watch" row: values read, absences explicit."""
    return [
        row.days_left,
        row.statute.isoformat(),
        row.matter_number or NO_NUMBER,
        row.client_workbook or NO_CLIENT,
        attorney_text,
        row.status or "",
        _docs(row.court_documents),
    ]


def spec(case_rows: list[list], change_rows: list[list], changes_note: str | None) -> dict:
    """The JSON the snippet reads on stdin."""
    return {
        "case_header": list(CASE_HEADER),
        "change_header": list(CHANGE_HEADER),
        "cases": case_rows,
        "changes": change_rows,
        "changes_note": changes_note,
        "about": list(ABOUT_LINES),
    }
